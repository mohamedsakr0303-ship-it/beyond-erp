"""Read the original BEYOND Accounting System workbook (v0.6 – v0.9 layout) and
convert it into interchange rows, so it goes through the normal import path.

Reads cached values (data_only=True): the workbook must have been saved by
Excel/LibreOffice after its last edit.

Accounting corrections applied on the way in (reported back to the user):
  * "Advance Payment" cost lines are supplier advances (an asset, 1310), not
    expenses — they become journal vouchers Dr Supplier Advances / Cr Bank.
  * When a row's Total Cost disagrees with Qty × Unit Cost, Total Cost wins.
  * Payments above a line's cost are booked as a supplier advance.
"""
import datetime as _dt
import re

from . import accounting as acc

COST_ACCOUNTS = {"materials": "5010", "subcontractor": "5020", "direct labor": "5030", "direct labour": "5030",
                 "site preparation": "5040", "equipment rental": "5050", "transport": "5060"}
OPEX_ACCOUNTS = {"rent": "6010", "software": "6020", "service": "6020", "subscription": "6020", "marketing": "6030",
                 "utilities": "6040", "salaries": "6050", "professional": "6060", "insurance": "6085",
                 "travel": "6075", "office": "6080", "bank": "6065"}
SUPPLIER_ADVANCES = "1310"
# V1.3 workbook codes whose meaning differs in the app's chart of accounts: sheet code -> app role (or code)
V13_ACCOUNTS = {"1150": "1180",            # accrued revenue receivable (unbilled) — 1150 is retention receivable here
                "1320": "role:custody",    # custody (عهدة) — 1320 is input VAT here
                "3150": "role:capital",    # partner capital (funding partner)
                "1400": "1410", "1450": "role:accum_depreciation", "2300": "role:income_tax_payable",
                "6500": "role:depreciation", "7010": "role:interest", "7500": "role:income_tax"}
V13_KINDS = {"project cost": "Direct Cost", "collection": "Receipt Voucher", "payroll": "Payroll Accrual",
             "distribution": "Owner Drawing", "supplier payment": "Payment Voucher", "overhead": "Operating Expense",
             "custody": "Journal Voucher", "adjustment": "Journal Voucher", "accrual (auto)": "Journal Voucher",
             "partner funding (auto)": "Capital Contribution", "client advance": "Receipt Voucher",
             "invoice (auto)": "Journal Voucher"}


def _v(x):
    if isinstance(x, str):
        x = x.strip()
        return x or None
    return x


def _f(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def _date(x):
    if isinstance(x, _dt.datetime):
        return x.date().isoformat()
    if isinstance(x, _dt.date):
        return x.isoformat()
    s = str(x or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return _dt.datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def _headers(ws, needles, scan=40, start=1):
    """Yield (row_index, {header_lower: col}) for every header row containing all needles."""
    for r in range(start, min(ws.max_row, start + scan) + 1):
        vals = [str(_v(c.value) or "").lower() for c in ws[r]]
        if all(any(n in h for h in vals) for n in needles):
            yield r, {h: i for i, h in enumerate(vals) if h}


def _col(h, *needles):
    for n in needles:
        for k, i in h.items():
            if n in k:
                return i
    return None


def _get(row, i):
    return row[i] if i is not None and i < len(row) else None


def looks_legacy(path):
    """True for the original BEYOND workbook layout (and not the new interchange workbook)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True)
        names = set(wb.sheetnames)
        wb.close()
    except Exception:  # noqa: BLE001 - unreadable files are reported by the normal reader
        return False
    return "Master Project Register" in names and "Business Units" not in names


def read_legacy_workbook(path):
    try:
        import openpyxl
    except ImportError:
        raise acc.ValidationError("Excel import needs openpyxl:  python -m pip install openpyxl") from None
    wb = openpyxl.load_workbook(str(path), data_only=True)
    names = {ws.title.lower(): ws for ws in wb.worksheets}

    def sheet(key):
        return next((ws for k, ws in names.items() if key in k), None)

    gj = sheet("general journal")
    journal_mode = gj is not None and any(h for _, h in _headers(gj, ["je #", "account", "debit", "credit"], scan=10))
    data = {k: [] for k in ("Business Units", "Price List", "Partners", "Partner Shares", "Receipt Vouchers",
                            "Direct Costs", "Operating Expenses", "Payroll", "Partner Transactions",
                            "Journal Vouchers", "Settings")}
    notes = []
    jv = [0]

    def journal(date, memo, party, lines, bu=""):
        jv[0] += 1
        for acct, dr, cr in lines:
            data["Journal Vouchers"].append({"voucher": f"LEG-{jv[0]}", "date": date, "kind": "Journal Voucher",
                                             "bu_code": bu, "party": party, "memo": memo, "account": acct,
                                             "debit": dr or "", "credit": cr or ""})

    # ---- business units
    codes = set()
    ws = sheet("project register")
    if ws is not None:
        for hr, h in _headers(ws, ["project id", "client"], scan=10):
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                code, name = _v(_get(row, _col(h, "project id"))), _v(_get(row, _col(h, "client")))
                if not code or not name or not str(code).upper().startswith("PRJ"):
                    continue
                basis = _v(_get(row, _col(h, "billing basis")))
                status = _v(_get(row, _col(h, "status")))
                contract = _get(row, _col(h, "contract value"))
                codes.add(code)
                if str(basis or "").lower().startswith("measured"):
                    notes.append(f"{code}: billing basis '{basis}' (V1.4) imported as Fixed Price — revenue comes from its "
                                 "measured BOQ invoices, exactly like the workbook.")
                data["Business Units"].append({
                    "code": code, "name": name, "type": "" if _v(_get(row, _col(h, "project type"))) == "TBD"
                    else _v(_get(row, _col(h, "project type"))) or "", "location": _v(_get(row, _col(h, "location"))) or "",
                    "area": _f(_get(row, _col(h, "area"))), "start_date": _date(_get(row, _col(h, "start date"))),
                    "status": "In Progress" if status in (None, "TBD") else status,
                    "basis": basis if basis in acc.BILLING_BASES else
                    "Fixed Price" if str(basis or "").lower().startswith("measured") else "Cost-Plus %",
                    "fee_pct": _f(_get(row, _col(h, "fee %"))),
                    "contract_value": contract if isinstance(contract, (int, float)) else ""})
            break

    # V1.3+: the Cost Allocation sheet's first rule is what untagged shared costs follow — use its split exactly
    rule = _allocation_rule(sheet("cost allocation"), codes)
    if rule:
        for bu in data["Business Units"]:
            bu["overhead_pct"] = rule["split"].get(bu["code"], 0)
        notes.append(f"Shared overhead split taken from the Cost Allocation sheet's '{rule['name']}' rule: "
                     + ", ".join(f"{c} {v:.2%}" for c, v in rule["split"].items()) + ".")
    # the original workbook shared company overhead evenly across active projects — start from the same split
    active = [b for b in data["Business Units"] if b["status"] not in ("Completed", "Closed")]
    if active and not rule:
        base = int(10000 / len(active)) / 10000
        for b in active:
            b["overhead_pct"] = base
        active[0]["overhead_pct"] = round(1 - base * (len(active) - 1), 6)
        notes.append(f"Shared overhead is split evenly across the {len(active)} active business units, as in the "
                     "original workbook — change the percentages any time (Business Units → Shared overhead allocation).")

    # ---- price list
    ws = sheet("rate matrix")
    if ws is not None:
        for hr, h in _headers(ws, ["category", "unit"], scan=10):
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                cat, desc = _v(_get(row, _col(h, "category"))), _v(_get(row, _col(h, "subcategory", "specification")))
                if not cat or not desc:
                    continue
                cost, mk = _f(_get(row, _col(h, "historical cost", "avg"))), _f(_get(row, _col(h, "markup")))
                price = _get(row, _col(h, "final client price"))
                note = _v(_get(row, _col(h, "notes")))
                data["Price List"].append({"category": cat, "item_no": str(_v(_get(row, _col(h, "item no"))) or ""),
                                           "description": desc, "unit": _v(_get(row, _col(h, "unit"))) or "",
                                           "cost": cost, "markup": mk,
                                           "price": price if isinstance(price, (int, float)) else cost * (1 + mk),
                                           "notes": "" if not note or str(note).startswith("=") else note})
            break

    # ---- partners & shares
    ws = sheet("partner equity")
    if ws is not None:
        for r in range(1, ws.max_row + 1):
            if _v(ws.cell(r, 1).value) == "Partner Name" and "/" in str(_v(ws.cell(r, 2).value) or ""):
                cols = {}
                for ci in range(2, ws.max_column + 1):
                    m = re.match(r"(PRJ-[\w-]+)", str(_v(ws.cell(r, ci).value) or ""))
                    if m and m.group(1) in codes:
                        cols[ci] = m.group(1)
                rr = r + 1
                while True:
                    name = _v(ws.cell(rr, 1).value)
                    if not name or str(name).lower().startswith("check"):
                        break
                    shares = {code: _f(ws.cell(rr, ci).value) for ci, code in cols.items()}
                    data["Partners"].append({"name": name,
                                             "default_share": round(sum(shares.values()) / len(shares), 6) if shares else 0})
                    for code, s in shares.items():
                        data["Partner Shares"].append({"partner": name, "bu_code": code, "share": s})
                    rr += 1
                break
        for r in range(1, ws.max_row + 1):
            if str(_v(ws.cell(r, 1).value) or "").upper().startswith("DISTRIBUTION LOG") and not journal_mode:
                hdr = {str(_v(ws.cell(r + 1, ci).value) or "").lower(): ci for ci in range(1, ws.max_column + 1)}
                ca, cd, cn = hdr.get("amount"), hdr.get("description"), hdr.get("notes")
                rr = r + 2
                while rr <= ws.max_row:
                    name, amt = _v(ws.cell(rr, 2).value), _f(ws.cell(rr, ca).value) if ca else 0
                    if not name and not amt:
                        break
                    if name and amt > 0:
                        bu = _v(ws.cell(rr, cn).value) if cn else None
                        data["Partner Transactions"].append({
                            "date": _date(ws.cell(rr, 1).value), "partner": name, "type": "Owner Drawing",
                            "amount": amt, "bu_code": bu if bu in codes else "",
                            "memo": (_v(ws.cell(rr, cd).value) if cd else None) or "Interim dividend"})
                    rr += 1
                break
            if str(_v(ws.cell(r, 1).value) or "").lower().startswith("reserve %"):
                data["Settings"].append({"key": "reserve_pct", "value": _f(ws.cell(r, 2).value)})

    # ---- receipts (income tracker may contain several stacked tables)
    ws = sheet("income tracker") if not journal_mode else None
    if ws is not None:
        seen = set()
        for hr, h in _headers(ws, ["date", "project id", "amount paid"], scan=ws.max_row):
            if hr in seen:
                continue
            seen.add(hr)
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                if str(_v(_get(row, 0)) or "").lower() == "date":
                    break
                code, paid = _v(_get(row, _col(h, "project id"))), _f(_get(row, _col(h, "amount paid")))
                if code not in codes or paid <= 0:
                    continue
                data["Receipt Vouchers"].append({
                    "date": _date(_get(row, _col(h, "date"))), "bu_code": code, "ref": _v(_get(row, _col(h, "invoice"))) or "",
                    "memo": _v(_get(row, _col(h, "milestone", "description"))) or "Client payment",
                    "party": _v(_get(row, _col(h, "client"))) or "", "amount": paid,
                    "method": _v(_get(row, _col(h, "payment method"))) or "Bank Transfer"})

    # ---- direct costs
    ws = sheet("project costing") if not journal_mode else None
    if ws is not None:
        for hr, h in _headers(ws, ["project id", "cost category", "total cost"], scan=10):
            for i, row in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), start=hr + 1):
                code = _v(_get(row, _col(h, "project id")))
                if code not in codes:
                    continue
                qty = _f(_get(row, _col(h, "quantity"))) or 1
                rate = _f(_get(row, _col(h, "unit cost")))
                total = _get(row, _col(h, "total cost"))
                total = _f(total) if isinstance(total, (int, float)) else qty * rate
                if total <= 0:
                    continue
                savings = max(_f(_get(row, _col(h, "negotiation savings"))), 0)
                paid = max(_f(_get(row, _col(h, "amount paid"))), 0)
                cat = str(_v(_get(row, _col(h, "cost category"))) or "Materials")
                vendor = _v(_get(row, _col(h, "vendor"))) or ""
                desc = _v(_get(row, _col(h, "item description"))) or cat
                date = _date(_get(row, _col(h, "date")))
                if cat.lower().startswith("advance"):
                    party = vendor or f"On-account: {desc}"[:60]
                    journal(date, f"Supplier advance — {desc}", party,
                            [(SUPPLIER_ADVANCES, paid or total, 0), ("1020", 0, paid or total)], code)
                    notes.append(f"Row {i}: '{desc}' ({total:,.2f}) booked as a supplier advance (asset 1310), "
                                 "not an expense. Reclassify to cost with a journal voucher once the work is done.")
                    continue
                gross = total + savings
                if abs(qty * rate - gross) > 0.01:
                    rate = gross / qty
                cost = total
                extra = max(paid - cost, 0)
                if paid < cost and not vendor:
                    vendor = "Unnamed supplier"
                data["Direct Costs"].append({
                    "date": date, "bu_code": code, "account": COST_ACCOUNTS.get(cat.lower(), "5010"),
                    "description": desc, "vendor": vendor, "deal_ref": _v(_get(row, _col(h, "deal reference"))) or "",
                    "deal_status": _v(_get(row, _col(h, "deal status"))) or "", "qty": qty,
                    "unit": _v(_get(row, _col(h, "unit"))) or "", "rate": round(rate, 6), "savings": savings,
                    "paid": min(paid, cost), "method": "Bank Transfer"})
                if extra > 0.005:
                    journal(date, f"Payment above cost — {desc}", vendor or "Unnamed supplier",
                            [(SUPPLIER_ADVANCES, extra, 0), ("1020", 0, extra)], code)
                    notes.append(f"Row {i}: paid {extra:,.2f} more than the cost — kept as a supplier advance.")
            break

    # ---- operating expenses
    ws = sheet("operating expenses") if not journal_mode else None
    if ws is not None:
        for hr, h in _headers(ws, ["date", "expense category", "amount"], scan=10):
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                amt = _f(_get(row, _col(h, "amount")))
                if amt <= 0:
                    continue
                cat = str(_v(_get(row, _col(h, "expense category"))) or "Other").lower()
                acct = next((a for k, a in OPEX_ACCOUNTS.items() if k in cat), "6090")
                data["Operating Expenses"].append({
                    "date": _date(_get(row, _col(h, "date"))), "account": acct,
                    "description": _v(_get(row, _col(h, "description"))) or cat,
                    "vendor": _v(_get(row, _col(h, "vendor"))) or "", "amount": amt, "paid": amt,
                    "method": _v(_get(row, _col(h, "payment method"))) or "Bank Transfer"})
            break

    # ---- payroll
    ws = sheet("payroll") if not journal_mode else None
    if ws is not None:
        for hr, h in _headers(ws, ["worker name", "rate"], scan=10):
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                worker = _v(_get(row, _col(h, "worker name")))
                qty, rate = _f(_get(row, _col(h, "days/hours"))), _f(_get(row, _col(h, "rate")))
                if not worker or qty <= 0 or rate <= 0:
                    continue
                ct = str(_v(_get(row, _col(h, "cost type"))) or "").lower()
                code = _v(_get(row, _col(h, "project id")))
                if "company" in ct or code not in codes:
                    ctype, code = "Administrative Salaries (Shared Overhead)", ""
                elif "non-billable" in ct or "project opex" in ct:
                    ctype = "Indirect Labour (Cost-Centre Overhead)"
                else:
                    ctype = "Direct Labour (Cost of Revenue)"
                note = _v(_get(row, _col(h, "notes")))
                data["Payroll"].append({
                    "date": _date(_get(row, _col(h, "date"))), "bu_code": code or "", "worker": worker,
                    "role": (_v(_get(row, _col(h, "role"))) or "") + (f" — {note}" if note else ""),
                    "cost_type": ctype, "period": str(_v(_get(row, _col(h, "period"))) or ""),
                    "basis": _v(_get(row, _col(h, "rate basis"))) or "", "qty": qty, "rate": rate,
                    "paid": min(_f(_get(row, _col(h, "total paid"))), qty * rate), "method": "Bank Transfer"})
            break

    if journal_mode:
        _read_new_accounts(sheet("chart of accounts"), data, notes)
        _read_journal(wb, gj, data, notes, codes)
        _read_invoice_register(wb, data, notes)
    _read_boq_and_invoices(wb, data, notes, register=journal_mode)
    if journal_mode:
        _read_partner_loans(sheet("balances tracker"), data, notes)
        _read_negotiations(sheet("negotiation register"), data, notes)
        _read_design_offer(sheet("design offer"), data, notes)
    data.pop("_remap", None)
    wb.close()
    # BEYOND's own letterhead identity, as the first app version set it; only fills settings that are still empty
    for k, v in (("company_name", "BEYOND DESIGN"), ("company_legal_name", "By SWAG For Contracting"),
                 ("company_tagline", "Interior Design & Finishing · More than spaces"),
                 ("bank_name", "CIB PHAROUS UNIVERSITY BRANCH"), ("bank_account", "100047359559"),
                 ("bank_holder", "MOHAMED ABDEL MOHSEN ABDELKADER"), ("bank_swift", "CIBEEGCX87"),
                 ("bank_iban", "EG790010008700000100018405541")):
        data["Settings"].append({"key": "default:" + k, "value": v})
    data["_notes"] = notes
    return data


_ITEM = re.compile(r"^\d+(\.\d+)+$")
_SECTION = re.compile(r"^-?\s*\d+\s*-\s*(.+)$")


def _meta_value(ws, label_words, rows=range(1, 9)):
    """Value next to a bilingual label such as 'Date / التاريخ' (label on the right, value on the left)."""
    for r in rows:
        cells = [(c.column, _v(c.value)) for c in ws[r] if _v(c.value) not in (None, "")]
        for i, (col, v) in enumerate(cells):
            if any(w in str(v) for w in label_words):
                left = [val for cc, val in cells if cc < col]
                if left:
                    return left[-1] if len(left) == 1 else left[-1]
    return None


def _item_rows(ws, cols):
    """Yield (category, row values) for every priced item row after the first 'م' header."""
    category, started = "", False
    for row in ws.iter_rows(values_only=True):
        first = str(_v(row[0]) or "")
        if first == "م":
            started = True
            continue
        m = _SECTION.match(first)
        if m and not _v(row[1] if len(row) > 1 else None):
            category = m.group(1).strip()
            continue
        if not started:
            continue
        if _ITEM.match(first.replace(" ", "")) and _v(row[1]):
            yield category, row


def _find_bu(data, *hints):
    hints = [str(h).strip().lower() for h in hints if h]
    for b in data["Business Units"]:
        loc, name = str(b.get("location") or "").lower(), str(b.get("name") or "").lower()
        if any(h and (h == loc or h in name) for h in hints):
            return b["code"]
    return ""


def _account_v13(code, roles, remap=None):
    code = str(code or "").strip()
    if remap and code in remap:
        return remap[code]
    m = V13_ACCOUNTS.get(code)
    if not m:
        return code
    return roles.get(m[5:], code) if m.startswith("role:") else m


_COA_SUBTYPE = {"Asset": "other_current_asset", "Liability": "other_current_liability", "Equity": "other_equity",
                "Revenue": "operating_revenue", "COGS": "cogs", "Expense": "other_operating"}


def _allocation_rule(ws, codes):
    """The first rule on the Cost Allocation sheet: {name, split: {project: share}} (shares as fractions)."""
    if ws is None:
        return None
    for hr, h in _headers(ws, ["allocation rule", "billable"], scan=10):
        row_h = [_v(c.value) for c in ws[hr]]
        cols = {i: str(v).strip() for i, v in enumerate(row_h) if str(v or "").strip() in codes}
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            name = _v(row[0]) if row else None
            if not name:
                continue
            split = {code: round(_f(row[i]), 10) for i, code in cols.items() if i < len(row)}
            if split and 0.999 <= sum(split.values()) <= 1.0000001:
                return {"name": name, "split": split}
            return None
    return None


def _read_new_accounts(ws, data, notes):
    """Accounts the workbook added (e.g. V1.4's 5045 Site Survey) that the app's chart does not have yet."""
    if ws is None:
        return
    from .db import DEFAULT_ACCOUNTS
    app = {c: (n, t, r) for c, n, t, _s, r in DEFAULT_ACCOUNTS}
    known = set(app) | set(V13_ACCOUNTS)
    taken = set(known)
    added, moved = [], []
    remap = data.setdefault("_remap", {})
    for hr, h in _headers(ws, ["code", "account name", "type"], scan=10):
        rows = [r for r in ws.iter_rows(min_row=hr + 1, values_only=True)]
        taken |= {str(_v(_get(r, _col(h, "code"))) or "").strip() for r in rows}
        for row in rows:
            code, name, typ = (_v(_get(row, _col(h, k))) for k in ("code", "account name", "type"))
            code = str(code or "").strip()
            if not code.isdigit() or not name or typ not in _COA_SUBTYPE:
                continue
            if code in app and code not in V13_ACCOUNTS and _clashes(name, typ, *app[code]):
                # the same number means something else in the app (e.g. 1330 partner loans vs withholding tax
                # receivable): keep the workbook's account under the next free number instead of mixing them up
                new = next(str(n) for n in range(int(code) + 1, int(code) + 1000) if str(n) not in taken)
                taken.add(new)
                remap[code] = new
                data.setdefault("Chart of Accounts", []).append({"code": new, "name": name, "type": typ, "subtype": _COA_SUBTYPE[typ]})
                moved.append(f"{code} {name} → {new} ({code} is “{app[code][0]}” in the app)")
                continue
            if code in known:
                continue
            line = str(_v(_get(row, _col(h, "income statement"))) or "").lower()
            sub = _COA_SUBTYPE[typ]
            if typ == "Expense":
                sub = ("employee" if "employee" in line else "general_admin" if "g&a" in line else
                       "depreciation" if "depreciation" in line else "interest" if "interest" in line else sub)
            data.setdefault("Chart of Accounts", []).append({"code": code, "name": name, "type": typ, "subtype": sub})
            added.append(f"{code} {name}")
        break
    if added:
        notes.append("Accounts added from the workbook's chart of accounts: " + ", ".join(added) + ".")
    if moved:
        notes.append("Account numbers that mean something else in the app, kept under a new number: " + "; ".join(moved) + ".")


_TAX_WORDS = ("tax", "vat", "ضريب", "withholding", "خصم وتحصيل", "stamp", "دمغة")
def _clashes(name, typ, app_name, app_type, app_role):
    """True when a workbook account and the app account with the same number are clearly different things: another
    type (asset vs liability…), or a tax account on one side only (e.g. 1330 partner loans vs withholding tax)."""
    if typ != app_type:
        return True
    name_l, app_l = str(name).lower(), str(app_name).lower()
    return any(w in name_l for w in _TAX_WORDS) != any(w in app_l for w in _TAX_WORDS)


def _read_design_offer(ws, data, notes):
    """The Design Offer sheet's inputs (client, project, scope lines, cost, discount, payment split, bank, terms) become an
    offer in the app on the company's Design Package layout, so it prints and exports from Sales → Offers."""
    if ws is None:
        return
    inp = {}
    for row in ws.iter_rows(values_only=True):
        for i, v in enumerate(row):
            k = str(v or "").strip().lower()
            if k and i + 1 < len(row) and row[i + 1] not in (None, "") and not k.startswith(("beyond", "dear")):
                inp.setdefault(k, row[i + 1])
    client, project = str(inp.get("client") or "").strip(), str(inp.get("project / unit") or inp.get("project") or "").strip()
    total = _f(inp.get("total cost"))
    if not client or not total:
        return
    scope = [str(inp[k]).strip() for k in sorted(x for x in inp if x.startswith("scope line")) if str(inp[k]).strip()]
    pct = lambda k: round(_f(inp.get(k)) * (100 if _f(inp.get(k)) <= 1 else 1), 2)  # noqa: E731
    days, fee, rev = _f(inp.get("working days")), pct("execution fee %"), _f(inp.get("minor revisions"))
    terms = [f"1. Time duration for the drawings of the unit is {days:g} working days." if days else "",
             f"2. Execution fees is {fee:g}% on total project cost according to the quotation (Optional)." if fee else "",
             f"3. Design is subject to {rev:g} minor Revision." if rev else ""]
    bank = {"bank": inp.get("bank"), "acct": inp.get("account number"), "holder": inp.get("account name"),
            "swift": inp.get("swift code"), "iban": inp.get("iban")}
    content = {"lang": "en", "title": "DESIGN PACKAGE :{project} ({client})",
               "intro": "Dear {client},\nKindly find this quotation (design package) for\n{project}",
               "sections": [{"kind": "text", "heading": "INTERIOR DESIGN", "style": "lines", "bold": False,
                             "body": "\n".join(f"{i}. {x}" for i, x in enumerate(scope, 1))},
                            {"kind": "pricing", "heading": "PROJECT COST", "style": "lines", "words": True,
                             "items": [{"desc": "Total Cost", "amount": total}], "discount": _f(inp.get("discount")), "vat": 0},
                            {"kind": "payment", "heading": "PAYMENT TERMS", "style": "lines",
                             "milestones": [{"label": "Advance Payment", "pct": pct("advance %")},
                                            {"label": "after Renders submission", "pct": pct("after renders %")},
                                            {"label": "after Final approval", "pct": pct("after approval %")}]},
                            {"kind": "bank", "heading": "", "style": "lines",
                             **({"bank_id": "custom", "bank_details": {k: str(v) for k, v in bank.items() if v}} if any(bank.values()) else {})},
                            {"kind": "text", "heading": "TIME DURATION", "style": "lines", "bold": True,
                             "body": "\n".join(t for t in terms if t)}],
               "closing": ""}
    import json as _json
    date = _date(inp.get("date")) or _dt.date.today().isoformat()
    data.setdefault("_extras", {}).setdefault("offers", []).append(
        {"number": f"DESIGN-{date}-{client}"[:60], "type_name": "Design Package", "date": date, "client": client, "project": project,
         "status": "Draft", "content": _json.dumps(content, ensure_ascii=False), "replace_same": True})
    notes.append(f"Design Offer: {client} — {project}, {total:,.2f} less {_f(inp.get('discount')):,.2f} — added to Sales → Offers as a draft.")


def _read_negotiations(ws, data, notes):
    """The Negotiation Register: each deal settled below its quote is matched to its journal entry by Deal Ref, so the
    app's negotiation report shows it (quoted, saving, kept or passed on). The amounts are already in the journal."""
    if ws is None:
        return
    found = []
    for hr, h in _headers(ws, ["deal ref", "quoted", "saving"], scan=10):
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            ref = _v(_get(row, _col(h, "deal ref")))
            saving = round(_f(_get(row, _col(h, "saving"))), 2)
            if not ref or not saving:
                continue
            treat = str(_v(_get(row, _col(h, "treatment"))) or "").lower()
            mode = "keep" if treat.startswith("keep") else "pass"
            rows = [r for r in data.get("Journal Vouchers", []) if str(r.get("deal_ref") or "").strip() == str(ref).strip()]
            if not rows:
                notes.append(f"Negotiation Register: deal {ref} has no journal entry with that Deal Ref — not linked.")
                continue
            je = rows[-1]["voucher"]      # the settling entry (the last one carrying this deal)
            for r in data["Journal Vouchers"]:
                if r["voucher"] == je:
                    r.update(quoted=round(_f(_get(row, _col(h, "quoted"))), 2), savings=saving, negotiation=mode)
            found.append(f"{ref} ({saving:,.2f}, {'kept' if mode == 'keep' else 'passed to the client'})")
        break
    if found:
        notes.append("Negotiation Register: " + "; ".join(found) + " — linked to their journal entries (Purchasing → Negotiations).")


def _read_partner_loans(ws, data, notes):
    """The Balances Tracker's partner-loan section (سلف الشركاء): each loan still owed goes into the Loans register as
    'already in the books' — the General Journal posted the money, so nothing is posted twice; repayments recorded in
    the app then reduce it and follow its monthly plan."""
    if ws is None:
        return
    remap = data.get("_remap") or {}
    acct = next((r["account"] for r in data.get("Journal Vouchers", []) if r.get("party") and r["account"] == remap.get("1330", "1330")), None)
    loans, start_row = [], None
    for r in range(1, ws.max_row + 1):
        head = str(ws.cell(r, 1).value or "")
        m = re.search(r"FUNDING PARTNER.*\(([^)]+)\)", head, re.I)
        if m and not any(x.get("key") == "funding_partner" for x in data.get("Settings", [])):
            # the partner whose own account is the company's bank (section 5): his money in is worked out live
            data.setdefault("Settings", []).append({"key": "funding_partner", "value": m.group(1).strip()})
            notes.append(f"Funding partner: {m.group(1).strip()} — the company's money runs through his account "
                         "(Equity → If a partner leaves shows his money in and out).")
        if not any(x.get("key") == "funding_limit" for x in data.get("Settings", [])):
            cells = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 16) + 1)]
            if any("agreed funding limit" in str(v or "").lower() for v in cells):
                nums = [v for v in cells if isinstance(v, (int, float)) and not isinstance(v, bool)]
                if nums:
                    data.setdefault("Settings", []).append({"key": "funding_limit", "value": str(nums[-1])})
                    notes.append(f"Funding limit {nums[-1]:,.2f} (Balances Tracker) — shown against his money in the company.")
        if "partner loans" in head.lower() and not start_row:
            start_row = r
    if not start_row:
        return
    for hr, h in _headers(ws, ["partner", "lent", "still owed"], scan=6, start=start_row):
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            who = _v(_get(row, _col(h, "partner")))
            if not who or str(who).upper().startswith("TOTAL") or str(who).startswith("("):
                if who and str(who).upper().startswith("TOTAL"):
                    break
                continue
            lent, owed = _f(_get(row, _col(h, "lent"))), _f(_get(row, _col(h, "still owed")))
            if lent <= 0 or owed <= 0:
                continue
            start = _date(_get(row, _col(h, "plan start", "start date")))
            months = int(_f(_get(row, _col(h, "plan length", "months)"))) or 0)
            loans.append({"code": f"LN-PL-{len(loans) + 1:02d}", "direction": "lent", "party": str(who),
                          "purpose": "Partner loan (سلفة شريك) — from the workbook's Balances Tracker",
                          "account": acct or remap.get("1330", "1330"), "interest_account": "4095", "accrued_account": "1185",
                          "start_date": start, "principal": round(lent * 100), "opening_balance": round(owed * 100),
                          "opening_date": start if abs(owed - lent) < 0.005 else _dt.date.today().isoformat(),
                          "rate": 0, "interest_type": "none", "fixed_interest": 0, "term_months": months,
                          "frequency": 1, "first_due": "", "schedule_type": "equal" if months else "none", "schedule": "[]",
                          "status": "open", "closed_on": "", "notes": "Interest-free. The workbook's plan only advises; "
                          "record repayments here (cash, or settled from his profit share by a decision).", "created_at": ""})
        break
    if loans:
        data.setdefault("_extras", {}).setdefault("loans", []).extend(loans)
        notes.append("Partner loans (Balances Tracker): " + "; ".join(
            f"{x['party']} {x['opening_balance'] / 100:,.2f} still owed, {x['term_months']} months from {x['start_date']}" for x in loans) +
            " — added to Loans & custody → Loans as already in the books (the journal entry is not posted twice).")


def _read_journal(wb, ws, data, notes, codes):
    """V1.3+: the General Journal is the single source of truth — every JE becomes one voucher, line for line."""
    from .db import DEFAULT_ACCOUNTS
    roles = {r: c for c, _n, _t, _s, r in DEFAULT_ACCOUNTS if r}
    hr, h = next(iter(_headers(ws, ["je #", "account", "debit", "credit"], scan=10)))
    groups, order = {}, []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        je = _v(_get(row, _col(h, "je #")))
        acct = _v(_get(row, _col(h, "account")))
        if not je or not acct or str(acct).lower() == "account":
            continue
        dr, cr = round(_f(_get(row, _col(h, "debit"))), 2), round(_f(_get(row, _col(h, "credit"))), 2)
        if not dr and not cr:
            continue
        if je not in groups:
            order.append(je)
        groups.setdefault(je, []).append(row)
    auto, mapped = [], set()
    n = 0
    for je in order:
        rows = groups[je]
        first = rows[0]
        src = str(_v(_get(first, _col(h, "source"))) or "").strip()
        kind = V13_KINDS.get(src.lower(), "Journal Voucher")
        if kind == "Payroll Accrual" and not any(str(_v(_get(r, _col(h, "account")))).startswith(("5", "6")) for r in rows):
            kind = "Payment Voucher"            # paying a salary already accrued
        proj = next((_v(_get(r, _col(h, "project"))) for r in rows if _v(_get(r, _col(h, "project"))) in codes), "")
        party = next((_v(_get(r, _col(h, "vendor", "party"))) for r in rows if _v(_get(r, _col(h, "vendor", "party")))), "") or ""
        deal = next((_v(_get(r, _col(h, "deal ref"))) for r in rows if _v(_get(r, _col(h, "deal ref")))), "") or ""
        memo = str(_v(_get(first, _col(h, "description"))) or src or je)
        date = _date(_get(first, _col(h, "date")))
        if "auto" in src.lower():
            memo = f"{memo} — snapshot from the Excel workbook on {date} (it updated itself there)"
            auto.append(je)
        for r in rows:
            code = str(_v(_get(r, _col(h, "account"))))
            new = _account_v13(code, roles, data.get("_remap"))
            if new != code:
                mapped.add(f"{code}→{new}")
            data["Journal Vouchers"].append({
                "voucher": je, "date": date, "kind": kind, "bu_code": proj, "party": party, "ref": je, "memo": memo,
                "deal_ref": deal, "account": new, "debit": round(_f(_get(r, _col(h, "debit"))), 2) or "",
                "credit": round(_f(_get(r, _col(h, "credit"))), 2) or "",
                "line_party": (_v(_get(r, _col(h, "vendor", "party"))) or "") if (_v(_get(r, _col(h, "vendor", "party"))) or "") != party else ""})
        n += 1
    notes.append(f"General Journal: {n} entries imported line for line (the workbook's own source of truth); the hidden "
                 "archive sheets (Project Costing, OpEx, Payroll, Income Tracker) were not imported, so nothing is counted twice.")
    if mapped:
        notes.append("Account codes that mean something else in the app were moved: " + ", ".join(sorted(mapped)) + ".")
    if auto:
        notes.append(f"{', '.join(auto)} were formulas that recalculated themselves in Excel (cost-plus accrued revenue, "
                     "funding-partner top-up). They were imported with their values as of their date — the app will not "
                     "update them; reverse and re-enter them when the figures change.")


def _read_invoice_register(wb, data, notes):
    """V1.3+: invoices live in 'Invoice Register' + 'Invoice Lines' (the print sheets only display them)."""
    data.setdefault("Sales Invoices", [])
    data.setdefault("Invoice Lines", [])
    reg = next((ws for ws in wb.worksheets if ws.title.lower() == "invoice register"), None)
    lines = next((ws for ws in wb.worksheets if ws.title.lower() == "invoice lines"), None)
    if reg is None:
        return
    codes = {b["code"]: b for b in data["Business Units"]}
    for hr, h in _headers(reg, ["invoice no", "project"], scan=10):
        for row in reg.iter_rows(min_row=hr + 1, values_only=True):
            num, bu = _v(_get(row, _col(h, "invoice no"))), _v(_get(row, _col(h, "project")))
            if not num or bu not in codes:
                continue
            status = str(_v(_get(row, _col(h, "status"))) or "Draft")
            data["Sales Invoices"].append({"number": num, "date": _date(_get(row, _col(h, "date"))), "bu_code": bu,
                                           "status": "Draft", "customer": _v(_get(row, _col(h, "client"))) or ""})
            adv = _f(_get(row, _col(h, "advance recovered")))
            ret = _f(_get(row, _col(h, "retention held")))
            extra = (f" Its advance recovery ({abs(adv):,.2f}) and retention ({abs(ret):,.2f}) are not on the app invoice — "
                     "settle them against client advances / retention with a receipt or journal voucher.") if adv or ret else ""
            posted = str(_v(_get(row, _col(h, "posted to journal"))) or "").lower().startswith("yes")
            tail = (" It is already in the General Journal (the workbook posted it), so do NOT post it again — keep it as "
                    "the invoice document only.") if posted else " Check it under Sales Invoices and press 'Save & post' when it is issued."
            notes.append(f"Invoice '{num}' ({bu}, {status} in the workbook) imported as a DRAFT with all its lines.{tail}{extra}")
        break
    if lines is None:   # V1.4: no Invoice Lines sheet — each invoice's lines are on its own 'Invoice N' sheet
        _invoice_sheet_lines(wb, data, notes)
        return
    for hr, h in _headers(lines, ["invoice no", "description", "qty this invoice"], scan=10):
        for row in lines.iter_rows(min_row=hr + 1, values_only=True):
            num = _v(_get(row, _col(h, "invoice no")))
            qty = _f(_get(row, _col(h, "qty this invoice")))
            if not num or not qty or not any(i["number"] == num for i in data["Sales Invoices"]):
                continue
            sub = _v(_get(row, _col(h, "sub-category")))
            data["Invoice Lines"].append({"number": num, "category": _v(_get(row, _col(h, "category"))) or "",
                                          "item_no": str(_v(_get(row, _col(h, "item no"))) or ""),
                                          "description": (f"{sub} — " if sub else "") + str(_v(_get(row, _col(h, "description"))) or ""),
                                          "unit": _v(_get(row, _col(h, "unit"))) or "", "qty": qty,
                                          "price": _f(_get(row, _col(h, "rate")))})
        break


def _norm_invoice(v):
    """'Invoice 01', 'Invoice  1', sheet 'Invoice 1' -> 'invoice1'."""
    t = re.sub(r"\s+", "", str(v or "")).lower()
    return re.sub(r"0+(\d)", r"\1", t)


def _invoice_sheet_lines(wb, data, notes):
    """V1.4: the register's invoice No. points at its sheet ('Invoice 01' = the 'Invoice 1' sheet); the executed
    quantities on that sheet are the invoice's lines."""
    sheets = {}
    for ws in wb.worksheets:
        if ws.title.lower().startswith("invoice") and ws.title.lower() not in ("invoice register", "invoice lines", "invoice print"):
            sheets[_norm_invoice(ws.title)] = ws
            sheets.setdefault(_norm_invoice(_meta_value(ws, ["الاصدار", "الإصدار", "Revision"])), ws)
    for inv in data["Sales Invoices"]:
        ws = sheets.get(_norm_invoice(inv["number"]))
        if ws is None:
            notes.append(f"Invoice '{inv['number']}': no sheet with its lines was found — add the lines in the app.")
            continue
        inv["boq_title"] = f"Contract BOQ — {inv['number']}"
        n, var, category, started = 0, 0, "", False
        for row in ws.iter_rows(values_only=True):
            first = str(_v(row[0]) or "")
            if first == "م":
                started = True
                continue
            m = _SECTION.match(first)
            if m and not _v(row[1] if len(row) > 1 else None):
                category = m.group(1).strip()
                continue
            if not started or len(row) < 6 or not _v(row[1]):
                continue
            ex, price = _f(row[4]), _f(row[5])
            if not ex or not price:
                continue
            item = first.replace(" ", "")
            if not _ITEM.match(item):          # a variation: work that was never in the BOQ (no item number)
                var += 1
                item = f"V-{var:02d}"
            data["Invoice Lines"].append({"number": inv["number"], "category": category, "item_no": item,
                                          "description": str(_v(row[1])).strip(), "unit": _v(row[2]) or "", "qty": ex,
                                          "price": price})
            n += 1
        notes.append(f"Invoice '{inv['number']}': {n} executed lines read from the '{ws.title}' sheet"
                     + (f", {var} of them variations outside the BOQ (numbered V-01…)." if var else "."))


def _read_boq_and_invoices(wb, data, notes, register=False):
    data.setdefault("BOQ Lines", [])
    data.setdefault("Sales Invoices", [])
    data.setdefault("Invoice Lines", [])
    for ws in wb.worksheets:
        title = ws.title.lower()
        if title in ("invoice register", "invoice lines") or (register and title == "invoice print"):
            continue
        if title.startswith("boq"):
            client = _meta_value(ws, ["Client", "العميل", "المشروع"]) or ""
            location = _meta_value(ws, ["Location", "الموقع"]) or ""
            rev = _meta_value(ws, ["Revision", "الاصدار", "الإصدار"]) or "Rev 01"
            date = _date(_meta_value(ws, ["Date", "التاريخ"]))
            bu = _find_bu(data, location, client)
            n = 0
            for cat, row in _item_rows(ws, 8):
                qty, price = _f(row[3]), _f(row[4])
                cost_total = _f(row[7]) if len(row) > 7 else 0
                data["BOQ Lines"].append({
                    "title": f"BOQ — {client}".strip(" —") or ws.title, "bu_code": bu, "rev": rev, "date": date,
                    "client": client, "location": location, "category": cat, "item_no": str(_v(row[0])),
                    "description": _v(row[1]), "unit": _v(row[2]) or "", "qty": qty, "price": price,
                    "cost": round(cost_total / qty, 6) if qty and cost_total else ""})
                n += 1
            if n:
                notes.append(f"Sheet '{ws.title}': {n} BOQ lines imported as 'BOQ — {client}'"
                             + ("" if bu else " (no matching business unit — link it from the BOQ screen)") + ".")
        elif title.startswith("invoice"):
            client = _meta_value(ws, ["المشروع", "Project", "Client"]) or ""
            location = _meta_value(ws, ["الموقع", "Location"]) or ""
            number = re.sub(r"\s+", " ", str(_meta_value(ws, ["الاصدار", "الإصدار", "Revision"]) or ws.title)).strip()
            date = _date(_meta_value(ws, ["التاريخ", "Date"]))
            bu = _find_bu(data, location, client)
            if not bu:
                notes.append(f"Sheet '{ws.title}': no business unit matches '{client}' / '{location}' — skipped.")
                continue
            boq_title = f"Contract BOQ — {number}"
            contract, executed = 0, 0
            for cat, row in _item_rows(ws, 9):
                cq, ex, price = _f(row[3]), _f(row[4]), _f(row[5])
                base = {"category": cat, "item_no": str(_v(row[0])), "description": _v(row[1]),
                        "unit": _v(row[2]) or "", "price": price}
                data["BOQ Lines"].append({**base, "title": boq_title, "bu_code": bu, "rev": "Contract", "date": date,
                                          "client": client, "location": location, "qty": cq})
                contract += 1
                if ex and not register:
                    data["Invoice Lines"].append({**base, "number": number, "qty": ex})
                    executed += 1
            if register:          # V1.3: the invoice itself comes from the Invoice Register; keep the contract BOQ
                notes.append(f"Sheet '{ws.title}': contract BOQ ({contract} lines) imported as '{boq_title}' for {bu}.")
                continue
            if executed:
                data["Sales Invoices"].append({"number": number, "date": date, "bu_code": bu,
                                               "boq_title": boq_title, "status": "Draft", "customer": client})
            notes.append(f"Sheet '{ws.title}': contract BOQ ({contract} lines) and progress billing '{number}' "
                         f"({executed} executed lines) imported for {bu} as a DRAFT — open it under Sales Invoices, "
                         "check it and press 'Save & post' to book the receivable and revenue.")
