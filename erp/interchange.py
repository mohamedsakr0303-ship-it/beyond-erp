"""The interchange format shared by Excel, JSON and CSV import/export.

One spec (SHEETS) defines every data sheet: its columns, their Arabic labels,
which calculated helper columns it has, and which accounts each row debits and
credits ("legs"). The Excel workbook uses the legs to build its Trial Balance
with SUMIFS formulas, so the workbook computes the same statements as the app.

Rows are plain dicts keyed by column key; Excel/CSV headers read like
"Date / التاريخ" and are matched on either the key or the English part.
"""
import csv
import datetime as _dt
import io
import json
import re

from . import accounting as acc
from . import bankrec, custody, loans, payroll, recurring, rentals, sales
from .db import ETA_SECRETS, SUBTYPES, to_major, to_minor

FORMAT = "beyond-erp"
FORMAT_VERSION = 1
FIRST_DATA_ROW = 5
HEADER_ROW = 4


def C(key, en, ar, kind="text", width=14, required=False, role=None, choices=None, calc=None, note=""):
    """Column spec. kind: text|date|num|money|pct|acct|bu|calc. role: default account role for blanks."""
    return {"key": key, "en": en, "ar": ar, "kind": kind, "width": width, "required": required, "role": role,
            "choices": choices, "calc": calc, "note": note}


def ROLE(role):
    """Excel formula returning the account code that holds a system role."""
    return (f"INDEX('Chart of Accounts'!$A${FIRST_DATA_ROW}:$A$3000,"
            f"MATCH(\"{role}\",'Chart of Accounts'!$F${FIRST_DATA_ROW}:$F$3000,0))")


# payroll parameters, looked up once on the Lists sheet from the Payroll Settings sheet (key, default)
PAYROLL_PARAMS = ([("si_emp_rate", 0.11), ("si_er_rate", 0.1875), ("si_min", 2700), ("si_max", 16700), ("exemption", 20000),
                   ("martyrs_rate", 0.0005), ("ot_day", 1.35), ("ot_night", 1.7), ("ot_rest", 2), ("days_month", 30), ("hours_day", 8)]
                  + [(f"band_{i}_upto", u) for i, u in enumerate((40000, 55000, 70000, 200000, 400000, 1200000), start=1)]
                  + [(f"band_{i}_rate", r) for i, r in enumerate((0, 0.10, 0.15, 0.20, 0.225, 0.25, 0.275), start=1)]
                  + [(f"start_band_{i}_above", 1e15) for i in range(2, 8)])
PAYROLL_FLAGS = ["housing", "transport", "meal", "phone", "site", "other_allow", "bonus", "commission", "incentive", "other_earn",
                 "penalty", "advance", "other_ded"]


def PRM(key):
    """The Lists cell holding a payroll parameter (see workbook: Lists!ZX…)."""
    keys = [k for k, _ in PAYROLL_PARAMS] + ["prorate_allowances"] + [f"{c}_{f}" for c in PAYROLL_FLAGS for f in ("tax", "ins", "on")]
    return f"Lists!$ZX${keys.index(key) + 1}"


def WH_ACCT(cell):
    return (f"IFERROR(INDEX(Warehouses!$C${FIRST_DATA_ROW}:$C$1000,MATCH({cell},Warehouses!$A${FIRST_DATA_ROW}:"
            f"$A$1000,0))&\"\",\"\")")


def TAXRATE(cell):
    return f"IFERROR(VLOOKUP({cell},'Tax Codes'!$A${FIRST_DATA_ROW}:$E$300,5,FALSE),0)"


def ROLECOL(key, role):
    """Helper column holding the account code of a system role (e.g. the input-VAT account)."""
    return C(key, f"{role} account", "الحساب", kind="calc", width=9, calc="=" + ROLE(role))


def eff(key, role):
    """Helper column: the entered account, or the role's default account when blank."""
    return C(f"{key}__eff", f"{key} (effective)", "الحساب الفعلي", kind="calc", width=10,
             calc=f'=IF({{{key}}}="",{ROLE(role)},{{{key}}}&"")')


def eff_money(key, method="method"):
    """Helper column: the entered cash / bank account, else cash for the Cash method and the bank for any other."""
    return C(f"{key}__eff", f"{key} (effective)", "الحساب الفعلي", kind="calc", width=10,
             calc=f'=IF({{{key}}}="",IF(TRIM({{{method}}})="Cash",{ROLE("cash")},{ROLE("bank")}),{{{key}}}&"")')


YES_NO = ["Posted", "Draft"]
METHODS = acc.PAYMENT_METHODS
PAYROLL = list(acc.PAYROLL_TYPES)

SHEETS = [
    {"name": "Settings", "ar": "الإعدادات", "module": "settings",
     "title": "Company settings & report period",
     "cols": [C("key", "Key", "المفتاح", width=26, required=True), C("value", "Value", "القيمة", width=40)]},
    {"name": "Chart of Accounts", "ar": "دليل الحسابات", "module": "accounting",
     "title": "Chart of accounts — edit codes and names freely; Role is the system tag (do not duplicate)",
     "cols": [C("code", "Code", "رقم الحساب", width=10, required=True),
              C("name", "Account Name", "اسم الحساب", width=36, required=True),
              C("type", "Type", "النوع", width=11, required=True, choices=list(SUBTYPES)),
              C("subtype", "Statement Line", "بند القوائم المالية", width=24, required=True,
                choices=sorted({s for v in SUBTYPES.values() for s in v})),
              C("normal", "Normal Balance", "الطبيعة", kind="calc", width=10,
                calc='=IF(OR({type}="Asset",{type}="COGS",{type}="Expense"),IF(OR({subtype}="accumulated_depreciation",'
                     '{subtype}="accumulated_amortization",{subtype}="allowance_doubtful"),"Credit","Debit"),IF(OR({subtype}="drawings",'
                     '{subtype}="sales_discount"),"Debit","Credit"))'),
              C("role", "System Role", "الدور في النظام", width=18),
              C("name_ar", "Arabic Name", "الاسم بالعربية", width=34),
              C("active", "Active", "نشط", width=7, choices=["Yes", "No"])]},
    {"name": "Tax Codes", "ar": "أكواد الضرائب", "module": "tax",
     "title": "VAT and withholding-tax codes with their e-invoice (ETA) tax type / sub-type. Rates are editable",
     "cols": [C("code", "Code", "الكود", width=12, required=True), C("name", "Name", "الاسم", width=30, required=True),
              C("name_ar", "Arabic Name", "الاسم بالعربية", width=30),
              C("kind", "Kind", "النوع", width=7, required=True, choices=["vat", "wht"]),
              C("rate", "Rate", "النسبة", kind="pct", width=8), C("eta_type", "ETA Tax Type", "نوع الضريبة (المنظومة)", width=10),
              C("eta_subtype", "ETA Sub-type", "النوع الفرعي", width=10), C("active", "Active", "نشط", width=7, choices=["Yes", "No"])]},
    {"name": "Business Units", "ar": "وحدات الأعمال", "module": "business_units",
     "title": "Business units (cost centres / projects). Overhead % = manual share of shared indirect costs",
     "cols": [C("code", "Code", "الكود", width=14, required=True), C("name", "Name / Client", "الاسم / العميل", width=34, required=True),
              C("type", "Type", "النوع", width=12), C("location", "Location", "الموقع", width=16),
              C("area", "Area (m²)", "المساحة", kind="num", width=9), C("start_date", "Start Date", "تاريخ البدء", kind="date"),
              C("status", "Status", "الحالة", width=12, choices=["Tendering", "In Progress", "On Hold", "Completed", "Closed"]),
              C("basis", "Billing Basis", "أساس الفوترة", width=18, choices=acc.BILLING_BASES),
              C("fee_pct", "Fee %", "نسبة الأتعاب", kind="pct", width=8),
              C("contract_value", "Contract Value", "قيمة العقد", kind="money"),
              C("overhead_pct", "Overhead Allocation %", "نسبة تحميل المصروفات المشتركة", kind="pct", width=10),
              C("negotiation", "Negotiation Earnings", "مكاسب التفاوض", width=12, choices=["", "keep", "pass"]),
              C("wip_mode", "Work in Progress From", "أعمال تحت التنفيذ من", width=14, choices=["", "cost_plus", "drafts"]),
              C("customer", "Customer", "العميل", width=24), C("notes", "Notes", "ملاحظات", width=24)]},
    {"name": "Parties", "ar": "العملاء والموردون", "module": "sales",
     "title": "Customers & suppliers master — tax registration numbers and addresses are needed for e-invoices",
     "cols": [C("name", "Name", "الاسم", width=28, required=True),
              C("type", "Type", "النوع", width=10, choices=["customer", "supplier", "both", "employee", "other"]),
              C("receiver_type", "Receiver Type", "نوع المستلم", width=8, choices=["B", "P", "F"],
                note="B = business (tax ID), P = person (national ID), F = foreign"),
              C("tax_id", "Tax Registration No.", "رقم التسجيل الضريبي", width=14),
              C("national_id", "National ID", "الرقم القومي", width=16), C("country", "Country", "الدولة", width=7),
              C("governate", "Governorate", "المحافظة", width=12), C("region", "City / Region", "المدينة", width=12),
              C("street", "Street", "الشارع", width=20), C("building", "Building No.", "رقم المبنى", width=8),
              C("postal_code", "Postal Code", "الرمز البريدي", width=8), C("email", "Email", "البريد", width=20),
              C("phone", "Phone", "الهاتف", width=14), C("default_vat", "Default VAT", "ض.ق.م الافتراضية", kind="tax", width=10),
              C("default_wht", "Default WHT", "خصم وتحصيل افتراضي", kind="tax", width=10), C("notes", "Notes", "ملاحظات", width=20),
              C("category", "Category", "التصنيف", width=14,
                note="supplier, subcontractor, driver, transport, technician, consultant, customer, employee, other — or your own"),
              C("contact_person", "Contact Person", "الشخص المسؤول", width=18), C("phone2", "Other Phone", "هاتف آخر", width=14),
              C("whatsapp", "WhatsApp", "واتساب", width=14), C("specialty", "Trade / Specialty", "النشاط / التخصص", width=16),
              C("vehicle", "Vehicle / Plate", "المركبة / اللوحة", width=12),
              C("bank_details", "Bank Details", "بيانات البنك", width=26)]},
    {"name": "Partners", "ar": "الشركاء", "module": "partners", "title": "Partners and their default profit share",
     "cols": [C("name", "Partner", "الشريك", width=30, required=True),
              C("default_share", "Default Share %", "نسبة الملكية الافتراضية", kind="pct", width=12),
              C("active", "Still a Partner", "شريك حالي", width=8, choices=["Yes", "No"]),
              C("joined", "Joined", "تاريخ الانضمام", kind="date"), C("left_on", "Left", "تاريخ الخروج", kind="date"),
              C("capital_interest", "Interest on Capital %", "فائدة رأس المال", kind="pct", width=10),
              C("loan_rate", "Loan Interest %", "فائدة القرض", kind="pct", width=10),
              C("profit_share", "Profit Share % (blank = ownership)", "نسبة الأرباح والخسائر (فارغ = الملكية)", kind="pct", width=12),
              C("salary", "Monthly Salary", "المرتب الشهري", kind="money"),
              C("drawings_interest", "Interest on Drawings %", "فائدة المسحوبات", kind="pct", width=10),
              C("notes", "Notes", "ملاحظات", width=24)]},
    {"name": "Partner Shares", "ar": "حصص الشركاء", "module": "partners",
     "title": "Ownership per business unit (overrides the default share)",
     "cols": [C("partner", "Partner", "الشريك", width=30, required=True), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu", required=True),
              C("share", "Share %", "النسبة", kind="pct", width=10)]},
    {"name": "Fixed Assets", "ar": "الأصول الثابتة", "module": "assets",
     "title": "Fixed asset register (its purchase, depreciation and disposal are journal entries with reference FA:<code>)",
     "cols": [C("code", "Asset Code", "كود الأصل", width=10, required=True), C("name", "Asset", "الأصل", width=28, required=True),
              C("account", "Asset Account", "حساب الأصل", kind="acct", required=True),
              C("accum_account", "Accumulated Depreciation Account", "حساب مجمع الإهلاك", kind="acct", role="accum_depreciation"),
              C("expense_account", "Depreciation Expense Account", "حساب مصروف الإهلاك", kind="acct", role="depreciation"),
              C("acquired", "Date Acquired", "تاريخ الشراء", kind="date", required=True),
              C("cost", "Cost", "التكلفة", kind="money", required=True), C("salvage", "Residual Value", "القيمة التخريدية", kind="money"),
              C("life_months", "Useful Life (months)", "العمر الإنتاجي (شهور)", kind="num", width=9, note="0 = not depreciated (land)"),
              C("depr_start", "Depreciate From (YYYY-MM)", "بداية الإهلاك", width=10),
              C("opening_depr", "Depreciation Before Registering", "إهلاك سابق", kind="money"),
              C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"), C("location", "Location", "الموقع", width=14),
              C("serial", "Serial / Plate No.", "الرقم التسلسلي / اللوحة", width=14), C("notes", "Notes", "ملاحظات", width=20),
              C("depr_method", "Depreciation Method", "طريقة الإهلاك", width=12, choices=["", "straight_line", "declining"],
                note="blank = the company's depreciation policy"),
              C("depr_factor", "Declining Factor", "معامل التناقص", kind="num", width=8, note="2 = double declining"),
              C("book_depr", "Depreciate in this workbook", "احتساب الإهلاك في هذا الملف", width=10, choices=["Yes", "No"],
                note="Yes (blank) = the statements here include depreciation up to the report end, like the app's depreciation "
                     "run. No = it is already in the Journal Vouchers (books exported from the app)"),
              C("dep_monthly", "Monthly Depreciation (straight line)", "الإهلاك الشهري (قسط ثابت)", kind="calc", width=14,
                calc='=IF(OR({code}="",N({life_months})=0),0,ROUND((N({cost})-N({salvage}))/N({life_months}),2))'),
              C("x_run", "depreciated up to", "الإهلاك حتى", kind="calc", width=10,
                note="The last full month by the report end (Settings → report_to)",
                calc='=IF(ReportTo=EOMONTH(ReportTo,0),ReportTo,EOMONTH(ReportTo,-1))'),
              C("x_mb", "months before", "شهور سابقة", kind="calc", width=8,
                calc='=IF({dep_monthly}>0,ROUND(N({opening_depr})/{dep_monthly},0),0)'),
              C("x_left", "months left", "الشهور المتبقية", kind="calc", width=8, calc='=MAX(0,N({life_months})-{x_mb})'),
              C("x_start", "first month depreciated", "أول شهر إهلاك", kind="calc", width=10,
                calc='=IF({acquired}="","",IF({depr_start}<>"",DATE(LEFT({depr_start},4),MID({depr_start},6,2),1),DATE(YEAR({acquired}),MONTH({acquired}),1)))'),
              C("x_k", "months now", "شهور الإهلاك", kind="calc", width=8,
                calc='=IF(OR({code}="",{acquired}="",N({life_months})=0),0,IFERROR(MAX(0,MIN({x_left},(YEAR({x_run})-YEAR({x_start}))*12'
                     '+MONTH({x_run})-MONTH({x_start})+1-{x_mb})),0))'),
              C("x_f", "declining factor", "معامل التناقص", kind="calc", width=8,
                calc='=IF({depr_method}="declining",IF(N({depr_factor})>0,N({depr_factor}),2),0)'),
              C("x_dd", "declining months", "شهور التناقص", kind="calc", width=8,
                calc='=IF({x_f}=0,0,MIN({x_k},MAX(0,INT({x_left}-N({life_months})/{x_f})+1)))'),
              C("x_dep", "depreciation booked here", "الإهلاك المحتسب", kind="calc", width=10,
                calc='=IF(OR({x_k}=0,{x_left}=0),0,ROUND(IF({x_f}=0,IF({x_k}>={x_left},1,{x_k}/{x_left})*(N({cost})-N({salvage})-N({opening_depr})),'
                     '(N({cost})-N({salvage})-N({opening_depr}))*(1-(1-{x_f}/N({life_months}))^{x_dd}'
                     '*IF({x_k}>={x_left},0,1-({x_k}-{x_dd})/({x_left}-{x_dd})))),2))'),
              C("dep_to_date", "Depreciation to Report End", "الإهلاك حتى نهاية الفترة", kind="calc", width=14,
                note="Accumulated: before registering + every full month up to the report end (the app's method: straight line "
                     "on what is left, or declining balance switching to straight line); never more than cost − residual",
                calc='=IF({code}="","",N({opening_depr})+{x_dep})'),
              C("dep_book_value", "Book Value at Report End", "القيمة الدفترية في نهاية الفترة", kind="calc", width=14,
                calc='=IF({code}="","",IFERROR(N({cost})-N({dep_to_date}),""))'),
              C("x_book", "booked here", "المحتسب هنا", kind="calc", width=10, calc='=IF({book_depr}="No",0,{x_dep})'),
              C("date", "depreciation date", "تاريخ الإهلاك", kind="calc", width=10, calc='=IF({x_book}=0,"",{x_run})'),
              eff("expense_account", "depreciation"), eff("accum_account", "accum_depreciation")],
     "legs": [("x_book", "expense_account__eff", "D"), ("x_book", "accum_account__eff", "C")]},
    {"name": "Warehouses", "ar": "المخازن", "module": "stock", "title": "Warehouses / stores and their inventory account",
     "cols": [C("name", "Warehouse", "المخزن", width=22, required=True), C("location", "Location", "الموقع", width=20),
              C("account", "Inventory Account", "حساب المخزون", kind="acct", width=12)]},
    {"name": "Stock Items", "ar": "الأصناف", "module": "stock", "title": "Item master",
     "cols": [C("name", "Item", "الصنف", width=30, required=True), C("unit", "Unit", "الوحدة", width=9),
              C("sku", "SKU", "كود الصنف", width=10), C("reorder_level", "Reorder Level", "حد إعادة الطلب", kind="num", width=10),
              C("category", "Category", "التصنيف", width=16), C("std_cost", "Standard Cost", "التكلفة المعيارية", kind="money", width=12,
                note="the usual buying cost per unit — kept even with no stock"),
              C("price", "Selling Price", "سعر البيع", kind="money", width=12), C("notes", "Notes", "ملاحظات", width=24)]},
    {"name": "Products", "ar": "المنتجات والوحدات", "module": "shop",
     "title": "Showroom catalog — items, units (apartments, villas, cars…) and services; photos are added in the app",
     "cols": [C("code", "Code", "الكود", width=10, note="blank = automatic; an existing code is updated"),
              C("name", "Name", "الاسم", width=28, required=True), C("name_ar", "Arabic Name", "الاسم بالعربية", width=24),
              C("kind", "Type", "النوع", width=9, choices=["item", "unit", "service"]),
              C("category", "Category", "التصنيف", width=14), C("unit", "Unit", "الوحدة", width=8),
              C("price", "Price", "السعر", kind="money", required=True), C("cost", "Cost per Unit", "تكلفة الوحدة", kind="money"),
              C("vat_code", "VAT Code", "كود الضريبة", kind="tax", width=9),
              C("status", "Status", "الحالة", width=10, choices=["available", "reserved", "sold", "delivered", "hidden"]),
              C("bu_code", "Project / Business Unit", "المشروع", kind="bu"),
              C("stock_item", "Stock Item (sold from stock)", "صنف المخزون", width=18),
              C("inventory_account", "Held in Account (cost of sales from)", "الحساب المحتفظ به", kind="acct",
                note="e.g. 1210 Goods & Units for Sale — when sold, its cost moves to cost of sales"),
              C("specs", "Specifications", "المواصفات", width=30, note="Label: value; Label: value — e.g. Area: 150 m²; Floor: 3"),
              C("description", "Description", "الوصف", width=30)]},
    {"name": "Price List", "ar": "قائمة الأسعار", "module": "estimating",
     "title": "Unit-rate price list (historical cost × (1 + markup) = client price)",
     "cols": [C("category", "Category", "الباب", width=24, required=True), C("item_no", "Item No.", "رقم البند", width=8),
              C("sub", "Sub-category", "البند الفرعي", width=16),
              C("description", "Description", "الوصف", width=60, required=True), C("unit", "Unit", "الوحدة", width=8),
              C("cost", "Unit Cost", "تكلفة الوحدة", kind="money"), C("markup", "Markup %", "هامش الربح", kind="pct", width=8),
              C("price", "Client Price", "سعر البيع", kind="money", note="Blank = cost × (1 + markup)"),
              C("notes", "Notes", "ملاحظات", width=20),
              C("eta_code", "ETA Item Code", "كود الصنف بالمنظومة", width=18, note="EGS code registered on the e-invoice portal")]},
    {"name": "BOQ Lines", "ar": "بنود المقايسات", "module": "estimating",
     "title": "Bills of quantities — one row per line; rows with the same BOQ Title + Business Unit form one BOQ",
     "cols": [C("title", "BOQ Title", "عنوان المقايسة", width=20, required=True), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"),
              C("rev", "Revision", "الإصدار", width=8), C("date", "Date", "التاريخ", kind="date"),
              C("client", "Client", "العميل", width=18), C("location", "Location", "الموقع", width=14),
              C("category", "Category", "الباب", width=22), C("sub", "Sub-category", "البند الفرعي", width=14),
              C("item_no", "Item No.", "رقم البند", width=8), C("description", "Description", "الوصف", width=50, required=True),
              C("unit", "Unit", "الوحدة", width=8), C("qty", "Qty", "الكمية", kind="num", width=9),
              C("price", "Unit Price", "سعر الوحدة", kind="money"), C("cost", "Unit Cost", "تكلفة الوحدة", kind="money"),
              C("amount", "Amount", "الإجمالي", kind="calc", calc="=ROUND(N({qty})*N({price}),2)")]},
    {"name": "Direct Costs", "ar": "التكاليف المباشرة", "module": "purchasing", "kind": "Direct Cost",
     "title": "Direct costs (cost of revenue) per business unit — materials, subcontractors, site work",
     "cols": [C("date", "Date", "التاريخ", kind="date", required=True), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu", required=True),
              C("account", "Cost Account", "حساب التكلفة", kind="acct", role="materials"),
              C("description", "Description", "البيان", width=34, required=True), C("vendor", "Supplier", "المورد", width=18),
              C("deal_ref", "Deal Ref.", "مرجع الاتفاق", width=10), C("deal_status", "Deal Status", "حالة الاتفاق", width=10, choices=["Open", "Settled"]),
              C("qty", "Qty", "الكمية", kind="num", width=8), C("unit", "Unit", "الوحدة", width=8),
              C("rate", "Unit Cost", "تكلفة الوحدة", kind="money"),
              C("savings", "Negotiation Discount", "خصم التفاوض", kind="money"),
              C("negotiation", "Discount Kept As", "معالجة الخصم", width=10, choices=["", "keep", "pass"],
                note="keep = company earnings (cost stays at full price) · pass = lower cost · blank = the business unit's rule"),
              C("paid", "Paid Now", "المدفوع", kind="money"), C("method", "Payment Method", "طريقة الدفع", choices=METHODS),
              C("paid_from", "Paid From", "دفع من حساب", kind="acct", role="bank"),
              C("payable_account", "Payable Account", "حساب الدائنين", kind="acct", role="ap"),
              C("vat_code", "VAT Code", "كود ض.ق.م", kind="tax", width=10), C("wht_code", "WHT Code", "كود خصم وتحصيل", kind="tax", width=10),
              C("total", "Total Cost", "صافي التكلفة", kind="calc", calc="=ROUND(N({qty})*N({rate})-IF({negotiation}=\"keep\",0,N({savings})),2)"),
              C("vat_amt", "Input VAT", "ض.ق.م المدخلات", kind="calc", calc="=ROUND({total}*" + TAXRATE("{vat_code}") + ",2)"),
              C("wht_amt", "WHT Withheld", "الخصم والتحصيل", kind="calc", calc="=ROUND({total}*" + TAXRATE("{wht_code}") + ",2)"),
              C("unpaid", "Unpaid", "المتبقي", kind="calc", calc="={total}+{vat_amt}-{wht_amt}-N({paid})"),
              eff("account", "materials"), eff_money("paid_from"), eff("payable_account", "ap"),
              ROLECOL("vat_acct", "vat_input"), ROLECOL("wht_acct", "wht_payable")],
     "legs": [("total", "account__eff", "D"), ("vat_amt", "vat_acct", "D"), ("paid", "paid_from__eff", "C"),
              ("unpaid", "payable_account__eff", "C"), ("wht_amt", "wht_acct", "C")]},
    {"name": "Operating Expenses", "ar": "المصروفات التشغيلية", "module": "purchasing", "kind": "Operating Expense",
     "title": "Operating expenses. Leave Business Unit blank for shared overhead (allocated by Overhead %). "
              "Paid Now blank = paid in full",
     "cols": [C("date", "Date", "التاريخ", kind="date", required=True), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"),
              C("account", "Expense Account", "حساب المصروف", kind="acct", role="other_overhead"),
              C("description", "Description", "البيان", width=34, required=True), C("vendor", "Supplier", "المورد", width=18),
              C("amount", "Amount", "المبلغ", kind="money", required=True), C("paid", "Paid Now", "المدفوع", kind="money"),
              C("method", "Payment Method", "طريقة الدفع", choices=METHODS),
              C("paid_from", "Paid From", "دفع من حساب", kind="acct", role="bank"),
              C("payable_account", "Payable Account", "حساب الدائنين", kind="acct", role="ap"),
              C("vat_code", "VAT Code", "كود ض.ق.م", kind="tax", width=10), C("wht_code", "WHT Code", "كود خصم وتحصيل", kind="tax", width=10),
              C("vat_amt", "Input VAT", "ض.ق.م المدخلات", kind="calc", calc="=ROUND(N({amount})*" + TAXRATE("{vat_code}") + ",2)"),
              C("wht_amt", "WHT Withheld", "الخصم والتحصيل", kind="calc", calc="=ROUND(N({amount})*" + TAXRATE("{wht_code}") + ",2)"),
              C("paid_eff", "Paid (effective)", "المدفوع الفعلي", kind="calc",
                calc='=IF({paid}="",N({amount})+{vat_amt}-{wht_amt},N({paid}))'),
              C("unpaid", "Unpaid", "المتبقي", kind="calc", calc="=N({amount})+{vat_amt}-{wht_amt}-{paid_eff}"),
              eff("account", "other_overhead"), eff_money("paid_from"), eff("payable_account", "ap"),
              ROLECOL("vat_acct", "vat_input"), ROLECOL("wht_acct", "wht_payable")],
     "legs": [("amount", "account__eff", "D"), ("vat_amt", "vat_acct", "D"), ("paid_eff", "paid_from__eff", "C"),
              ("unpaid", "payable_account__eff", "C"), ("wht_amt", "wht_acct", "C")]},
    {"name": "Payroll", "ar": "الرواتب والأجور", "module": "payroll", "kind": "Payroll Accrual",
     "title": "Payroll accruals — wages earned (and paid now) per worker",
     "cols": [C("date", "Date", "التاريخ", kind="date", required=True), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"),
              C("worker", "Employee / Worker", "الموظف / العامل", width=22, required=True), C("role", "Position", "الوظيفة", width=14),
              C("cost_type", "Cost Classification", "تصنيف التكلفة", width=30, choices=PAYROLL),
              C("account", "Expense Account", "حساب المصروف", kind="acct"),
              C("period", "Period", "الفترة", width=9), C("basis", "Rate Basis", "أساس الأجر", width=10,
                                                            choices=["Daily", "Monthly", "Lump Sum", "Hourly"]),
              C("qty", "Days / Months", "العدد", kind="num", width=8), C("rate", "Rate", "الأجر", kind="money"),
              C("paid", "Paid Now", "المدفوع", kind="money"), C("method", "Payment Method", "طريقة الدفع", choices=METHODS),
              C("paid_from", "Paid From", "دفع من حساب", kind="acct", role="bank"),
              C("payable_account", "Payroll Payable Account", "حساب الأجور المستحقة", kind="acct", role="accrued_payroll"),
              C("total", "Gross Pay", "إجمالي الأجر", kind="calc", calc="=ROUND(N({qty})*N({rate}),2)"),
              C("unpaid", "Unpaid", "المتبقي", kind="calc", calc="={total}-N({paid})"),
              C("account__eff", "account (effective)", "الحساب الفعلي", kind="calc", width=10,
                calc='=IF({account}<>"",{account}&"",IF(LEFT({cost_type},6)="Indire",' + ROLE("project_overhead") +
                     ',IF(LEFT({cost_type},6)="Admini",' + ROLE("admin_salaries") + ',' + ROLE("direct_labour") + ')))'),
              eff_money("paid_from"), eff("payable_account", "accrued_payroll")],
     "legs": [("total", "account__eff", "D"), ("paid", "paid_from__eff", "C"), ("unpaid", "payable_account__eff", "C")]},
    {"name": "Receipt Vouchers", "ar": "سندات القبض", "module": "sales", "kind": "Receipt Voucher",
     "title": "Money received from clients. Credit Account: revenue (cash basis), receivables (against an invoice) "
              "or client advances",
     "cols": [C("date", "Date", "التاريخ", kind="date", required=True), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu", required=True),
              C("ref", "Receipt No.", "رقم السند", width=10), C("memo", "Description", "البيان", width=28),
              C("party", "Received From", "المستلم منه", width=20),
              C("amount", "Amount", "المبلغ", kind="money", required=True), C("method", "Payment Method", "طريقة الدفع", choices=METHODS),
              C("received_into", "Received Into", "إيداع في حساب", kind="acct", role="bank"),
              C("credit_account", "Credit Account", "الحساب الدائن", kind="acct", role="revenue"),
              C("withheld", "Tax Withheld by Client", "خصم وتحصيل بواسطة العميل", kind="money"),
              C("vat_code", "VAT Code (cash revenue)", "كود ض.ق.م (إيراد نقدي)", kind="tax", width=10,
                note="Only when the receipt is booked straight to revenue: the amount is VAT-inclusive"),
              C("gross", "Gross", "الإجمالي", kind="calc", calc="=N({amount})+N({withheld})"),
              C("vat_amt", "Output VAT", "ض.ق.م المخرجات", kind="calc",
                calc='=IF({vat_code}="",0,{gross}-ROUND({gross}/(1+' + TAXRATE("{vat_code}") + '),2))'),
              C("net_credit", "Net Credit", "صافي الدائن", kind="calc", calc="={gross}-{vat_amt}"),
              eff_money("received_into"), eff("credit_account", "revenue"),
              ROLECOL("whtr_acct", "wht_receivable"), ROLECOL("vato_acct", "vat_output")],
     "legs": [("amount", "received_into__eff", "D"), ("withheld", "whtr_acct", "D"),
              ("net_credit", "credit_account__eff", "C"), ("vat_amt", "vato_acct", "C")]},
    {"name": "Sales Invoices", "ar": "فواتير المبيعات", "module": "sales", "kind": "Sales Invoice",
     "title": "Sales invoices / progress billings. Amount is the sum of the invoice's lines",
     "cols": [C("number", "Invoice No.", "رقم الفاتورة", width=14, required=True), C("date", "Date", "التاريخ", kind="date", required=True),
              C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu", required=True),
              C("boq_title", "Based on BOQ", "بناءً على مقايسة", width=16),
              C("status", "Status", "الحالة", width=8, choices=YES_NO),
              C("ar_account", "Receivable Account", "حساب المدينين", kind="acct", role="ar"),
              C("revenue_account", "Revenue Account", "حساب الإيراد", kind="acct", role="revenue"),
              C("customer", "Customer", "العميل", width=22), C("vat_code", "VAT Code", "كود ض.ق.م", kind="tax", width=10), C("wht_code", "WHT Code", "كود خصم وتحصيل", kind="tax", width=10),
              # extract (مستخلص) — how it prints for the client; none of these change the books
              C("extract_no", "Extract No. (on its site)", "رقم المستخلص", kind="num", width=8, note="blank = the next number on the business unit"),
              C("project", "Project (extract header)", "المشروع", width=16), C("work", "Work (extract header)", "الأعمال", width=24),
              C("prev_paid", "Paid Before (extract)", "ما سبق صرفه", kind="money", note="blank = what the client paid before the date (receipts)"),
              C("deductions", "Deductions (extract)", "خصومات", kind="money"),
              C("final", "Final Extract", "مستخلص ختامي", width=7, choices=["Yes", "No"], note="Yes = final (ختامي); blank = running (جاري)"),
              C("notes", "Extract Notes", "ملاحظات المستخلص", width=30),
              C("amount", "Net Amount", "صافي المبلغ", kind="calc",
                calc="=SUMIFS([[Invoice Lines.amount]],[[Invoice Lines.number]],{number})"),
              C("posted_amount", "Posted Amount", "المبلغ المرحّل", kind="calc", calc='=IF({status}="Draft",0,{amount})'),
              C("x_rest", "to the invoice's revenue account", "لحساب إيراد الفاتورة", kind="calc", width=10,
                calc="={posted_amount}-SUMIFS([[Invoice Lines.x_own]],[[Invoice Lines.number]],{number})"),
              C("vat_amt", "Output VAT", "ض.ق.م المخرجات", kind="calc", calc="=ROUND({posted_amount}*" + TAXRATE("{vat_code}") + ",2)"),
              C("wht_amt", "WHT by Client", "خصم وتحصيل العميل", kind="calc", calc="=ROUND({posted_amount}*" + TAXRATE("{wht_code}") + ",2)"),
              C("ar_amt", "Receivable", "المستحق من العميل", kind="calc", calc="={posted_amount}+{vat_amt}-{wht_amt}"),
              eff("ar_account", "ar"), eff("revenue_account", "revenue"),
              ROLECOL("whtr_acct", "wht_receivable"), ROLECOL("vato_acct", "vat_output")],
     "legs": [("ar_amt", "ar_account__eff", "D"), ("wht_amt", "whtr_acct", "D"),
              ("x_rest", "revenue_account__eff", "C"), ("vat_amt", "vato_acct", "C")]},
    {"name": "Invoice Lines", "ar": "بنود الفواتير", "module": "sales", "title": "Lines of each sales invoice",
     "cols": [C("number", "Invoice No.", "رقم الفاتورة", width=14, required=True), C("category", "Category", "الباب", width=20),
              C("sub", "Sub-category", "البند الفرعي", width=12), C("item_no", "Item No.", "رقم البند", width=8),
              C("description", "Description", "الوصف", width=50, required=True), C("unit", "Unit", "الوحدة", width=8),
              C("qty", "Qty Executed", "الكمية المنفذة", kind="num", width=9), C("price", "Unit Price", "سعر الوحدة", kind="money"),
              C("pct", "% Executed", "النسبة", kind="pct", width=8, note="blank = 100 %; e.g. half a month of rent = 50 %"),
              C("amount", "Amount", "الإجمالي", kind="calc", calc='=ROUND(N({qty})*N({price})*IF({pct}="",1,N({pct})),2)'),
              C("account", "Line Revenue Account", "حساب إيراد البند", kind="acct", note="blank = the invoice's revenue account"),
              # a line with its own revenue account is credited there (the invoice credits the rest)
              C("date", "invoice date", "تاريخ الفاتورة", kind="calc", width=10, calc='=IFERROR(INDEX([[Sales Invoices.date]],MATCH({number},[[Sales Invoices.number]],0)),"")'),
              C("bu_code", "invoice business unit", "وحدة الأعمال", kind="calc", width=10, calc='=IFERROR(INDEX([[Sales Invoices.bu_code]],MATCH({number},[[Sales Invoices.number]],0))&"","")'),
              C("x_acct", "own revenue account", "حساب البند", kind="calc", width=8, calc='={account}&""'),
              C("x_own", "booked to its own account", "المرحل لحساب البند", kind="calc", width=10,
                calc='=IF(OR({account}="",{date}=""),0,IF(IFERROR(INDEX([[Sales Invoices.status]],MATCH({number},[[Sales Invoices.number]],0))&"","")="Draft",0,{amount}))'),
              # the Extract Print sheet: 1 = this line is on the extract chosen there (running: the site's earlier issued ones too)
              C("x_xp", "on the printed extract", "في المستخلص المطبوع", kind="calc", width=8,
                calc='=IF({number}&""="",0,IF({number}&""=\'Extract Print\'!$J$2&"",1,IF(AND(\'Extract Print\'!$J$3="running",'
                     '{bu_code}&""=\'Extract Print\'!$J$4&"",N({date})>0,N({date})<N(\'Extract Print\'!$J$5),'
                     'IFERROR(INDEX([[Sales Invoices.status]],MATCH({number},[[Sales Invoices.number]],0))&"","")<>"Draft"),1,0)))')],
     "legs": [("x_own", "x_acct", "C")]},
    {"name": "Payment Vouchers", "ar": "سندات الصرف", "module": "purchasing", "kind": "Payment Voucher",
     "title": "Payments settling what we owe suppliers, subcontractors and workers",
     "cols": [C("date", "Date", "التاريخ", kind="date", required=True), C("party", "Payee", "المستفيد", width=22, required=True),
              C("against", "Pay Against Account", "الحساب المدين", kind="acct", role="ap"),
              C("amount", "Amount", "المبلغ", kind="money", required=True), C("method", "Payment Method", "طريقة الدفع", choices=METHODS),
              C("paid_from", "Paid From", "دفع من حساب", kind="acct", role="bank"), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"),
              C("ref", "Voucher No.", "رقم السند", width=10), C("memo", "Description", "البيان", width=26),
              eff("against", "ap"), eff_money("paid_from")],
     "legs": [("amount", "against__eff", "D"), ("amount", "paid_from__eff", "C")]},
    {"name": "Partner Transactions", "ar": "حركات الشركاء", "module": "partners",
     "title": "Capital contributions and owner drawings / dividends",
     "cols": [C("date", "Date", "التاريخ", kind="date", required=True), C("partner", "Partner", "الشريك", width=26, required=True),
              C("type", "Type", "النوع", width=20, required=True, choices=["Capital Contribution", "Owner Drawing"]),
              C("amount", "Amount", "المبلغ", kind="money", required=True), C("method", "Payment Method", "طريقة الدفع", choices=METHODS),
              C("cash_account", "Cash / Bank Account", "حساب النقدية", kind="acct", role="bank"),
              C("equity_account", "Equity Account", "حساب حقوق الملكية", kind="acct"),
              C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"), C("memo", "Description", "البيان", width=22),
              eff_money("cash_account"),
              C("equity__eff", "equity (effective)", "الحساب الفعلي", kind="calc", width=10,
                calc='=IF({equity_account}<>"",{equity_account}&"",IF(LEFT({type},7)="Capital",' + ROLE("capital") +
                     "," + ROLE("drawings") + "))"),
              C("dr_acct", "Debit Account", "الحساب المدين", kind="calc", width=10,
                calc='=IF(LEFT({type},7)="Capital",{cash_account__eff},{equity__eff})'),
              C("cr_acct", "Credit Account", "الحساب الدائن", kind="calc", width=10,
                calc='=IF(LEFT({type},7)="Capital",{equity__eff},{cash_account__eff})')],
     "legs": [("amount", "dr_acct", "D"), ("amount", "cr_acct", "C")]},
    {"name": "Stock Moves", "ar": "حركات المخزون", "module": "stock",
     "title": "Goods receipts, material issues to business units, transfers and count adjustments "
              "(Adjustment Qty = counted − system, may be negative)",
     "cols": [C("date", "Date", "التاريخ", kind="date", required=True),
              C("type", "Type", "نوع الحركة", width=11, required=True, choices=acc.STOCK_MOVE_TYPES),
              C("warehouse", "Warehouse", "المخزن", width=16, required=True), C("to_warehouse", "To Warehouse", "إلى مخزن", width=14),
              C("item", "Item", "الصنف", width=24, required=True), C("unit", "Unit", "الوحدة", width=8),
              C("qty", "Qty", "الكمية", kind="num", width=8, required=True), C("rate", "Unit Cost", "تكلفة الوحدة", kind="money"),
              C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"), C("vendor", "Supplier", "المورد", width=16),
              C("paid", "Paid Now", "المدفوع", kind="money"), C("method", "Payment Method", "طريقة الدفع", choices=METHODS),
              C("paid_from", "Paid From", "دفع من حساب", kind="acct", role="bank"),
              C("payable_account", "Payable Account", "حساب الدائنين", kind="acct", role="ap"),
              C("account", "Cost Account (issues)", "حساب التكلفة", kind="acct", role="materials"),
              C("note", "Note", "ملاحظة", width=16), C("vat_code", "VAT Code", "كود ض.ق.م", kind="tax", width=10), C("wht_code", "WHT Code", "كود خصم وتحصيل", kind="tax", width=10),
              C("x_rqty", "received qty", "الكمية المستلمة", kind="calc", width=8, calc='=IF({type}="Receipt",N({qty}),0)'),
              C("x_rval", "received value", "قيمة الاستلام", kind="calc", width=8, calc='=IF({type}="Receipt",ROUND(N({qty})*N({rate}),2),0)'),
              C("value", "Value", "القيمة", kind="calc",
                note="Qty × unit cost; with no unit cost (issues, transfers, counts) the item's average cost of receipts up to this date",
                calc='=ROUND(ABS(N({qty}))*IF({rate}<>"",N({rate}),IFERROR(SUMIFS([[Stock Moves.x_rval]],[[Stock Moves.item]],{item},'
                     '[[Stock Moves.date]],"<="&{date})/SUMIFS([[Stock Moves.x_rqty]],[[Stock Moves.item]],{item},[[Stock Moves.date]],"<="&{date}),0)),2)'),
              C("vat_amt", "Input VAT", "ض.ق.م المدخلات", kind="calc",
                calc='=IF({type}="Receipt",ROUND({value}*' + TAXRATE("{vat_code}") + ',2),0)'),
              C("wht_amt", "WHT Withheld", "الخصم والتحصيل", kind="calc",
                calc='=IF({type}="Receipt",ROUND({value}*' + TAXRATE("{wht_code}") + ',2),0)'),
              ROLECOL("vat_acct", "vat_input"), ROLECOL("wht_acct", "wht_payable"),
              eff_money("paid_from"), eff("payable_account", "ap"), eff("account", "materials"),
              C("wh_acct", "Warehouse Account", "حساب المخزن", kind="calc", width=10,
                calc='=IF(' + WH_ACCT("{warehouse}") + '="",' + ROLE("inventory") + ',' + WH_ACCT("{warehouse}") + ')'),
              C("to_acct", "To-Warehouse Account", "حساب المخزن المستلم", kind="calc", width=10,
                calc='=IF(' + WH_ACCT("{to_warehouse}") + '="",' + ROLE("inventory") + ',' + WH_ACCT("{to_warehouse}") + ')'),
              C("paid_eff", "Paid (effective)", "المدفوع الفعلي", kind="calc", calc='=IF({type}="Receipt",N({paid}),0)'),
              C("dr_acct", "Debit Account", "الحساب المدين", kind="calc", width=10,
                calc='=IF({type}="Receipt",{wh_acct},IF({type}="Issue",{account__eff},IF({type}="Transfer",{to_acct},'
                     'IF(N({qty})>=0,{wh_acct},' + ROLE("stock_adjust") + '))))'),
              C("cr_acct", "Credit Account", "الحساب الدائن", kind="calc", width=10,
                calc='=IF({type}="Receipt",{payable_account__eff},IF({type}="Adjustment",IF(N({qty})>=0,' +
                     ROLE("stock_adjust") + ',{wh_acct}),{wh_acct}))'),
              C("cr_amt", "Credit Amount", "المبلغ الدائن", kind="calc", calc="={value}+{vat_amt}-{wht_amt}-{paid_eff}")],
     "legs": [("value", "dr_acct", "D"), ("vat_amt", "vat_acct", "D"), ("cr_amt", "cr_acct", "C"),
              ("paid_eff", "paid_from__eff", "C"), ("wht_amt", "wht_acct", "C")]},
    {"name": "Journal Vouchers", "ar": "قيود اليومية", "module": "accounting", "kind": "Journal Voucher",
     "title": "Journal vouchers — opening balances, depreciation, amortization, tax, loans, accruals. "
              "Rows sharing a Voucher No. form one balanced entry",
     "cols": [C("voucher", "Voucher No.", "رقم القيد", width=10, required=True), C("date", "Date", "التاريخ", kind="date", required=True),
              C("kind", "Voucher Type", "نوع القيد", width=16), C("bu_code", "Business Unit", "وحدة الأعمال", kind="bu"),
              C("party", "Party", "الطرف", width=16), C("ref", "Reference", "المرجع", width=10),
              C("memo", "Description", "البيان", width=30), C("account", "Account", "الحساب", kind="acct", required=True),
              C("debit", "Debit", "مدين", kind="money"), C("credit", "Credit", "دائن", kind="money"),
              C("line_party", "Line Party", "طرف السطر", width=14,
                note="Only when a line belongs to someone other than the voucher's party (e.g. each partner in a share transfer)"),
              C("item", "Item", "الصنف", width=18,
                note="Item line: the item or material (debit = qty × unit price, made by the app). Leave empty on normal lines"),
              C("item_unit", "Item Unit", "وحدة الصنف", width=8), C("item_qty", "Item Qty", "كمية الصنف", kind="num"),
              C("item_rate", "Item Unit Price", "سعر وحدة الصنف", kind="money"),
              C("item_to", "Item Goes To", "وجهة الصنف", width=16,
                note="A warehouse name (received into stock) — or empty: used on the voucher's business unit")],
     "legs": [("debit", "account", "D"), ("credit", "account", "C")]},
]
from . import setup_sheets as _setup  # noqa: E402
SHEETS += _setup.specs(C)
SHEET_BY_NAME = {s["name"]: s for s in SHEETS}
POSTING_SHEETS = [s for s in SHEETS if s.get("legs")]


def input_cols(sheet):
    return [c for c in sheet["cols"] if c["kind"] != "calc"]


# =================================================================== EXPORT
def _d(v):
    return v or ""


def _num(v):
    return round(v, 4) if isinstance(v, float) else v


def export_data(db):
    """Everything in the books as {sheet name: [row dicts]} (input columns only)."""
    out = {s["name"]: [] for s in SHEETS}
    period_from = db.setting("report_from") or f"{_dt.date.today().year}-01-01"
    period_to = db.setting("report_to") or f"{_dt.date.today().year}-12-31"
    out["Settings"] = [{"key": "report_from", "value": period_from}, {"key": "report_to", "value": period_to}]
    for r in db.q("SELECT key, value FROM settings ORDER BY key"):
        if r["key"] not in ("schema_version", "report_from", "report_to", "docs_seeded", *ETA_SECRETS):
            out["Settings"].append({"key": r["key"], "value": r["value"]})
    from . import payroll as _pr   # employees and payroll rules carry no money: re-import updates them
    out["Employees"] = _pr.employee_sheet_rows(db)
    out["Payroll Settings"] = _pr.settings_sheet_rows(db)
    _site_sheets(db, out)
    from . import budgets   # budgets carry no money, so the export can include them (re-import replaces, never doubles)
    out["Budgets"] = budgets.export_rows(db)
    for t in db.q("SELECT * FROM tax_codes ORDER BY kind, code"):
        out["Tax Codes"].append({k: t[k] for k in ("code", "name", "name_ar", "kind", "rate", "eta_type", "eta_subtype")}
                                | {"active": "Yes" if t["active"] else "No"})
    for p in db.q("SELECT * FROM parties ORDER BY name"):
        out["Parties"].append({k: v for k, v in p.items() if k != "id"})
    for a in db.q("SELECT * FROM accounts ORDER BY code"):
        out["Chart of Accounts"].append({"code": a["code"], "name": a["name"], "name_ar": a["name_ar"], "type": a["type"],
                                         "subtype": a["subtype"], "role": a["role"] or "",
                                         "active": "Yes" if a["active"] else "No"})
    bus = {b["id"]: b for b in db.q("SELECT * FROM business_units ORDER BY code")}
    code = lambda bid: bus[bid]["code"] if bid in bus else ""  # noqa: E731
    for b in bus.values():
        out["Business Units"].append({k: b[k] for k in ("code", "name", "type", "location", "area", "start_date",
                                                         "status", "basis", "fee_pct", "overhead_pct", "customer",
                                                         "notes", "negotiation", "wip_mode")}
                                     | {"contract_value": to_major(b["contract_value"])})
    for p in db.q("SELECT * FROM partners ORDER BY id"):
        out["Partners"].append({"name": p["name"], "default_share": p["default_share"],
                                "active": "Yes" if p["active"] else "No", "joined": p["joined"] or "",
                                "left_on": p["left_on"] or "", "capital_interest": p["capital_interest"] or "",
                                "loan_rate": p["loan_rate"] or "", "notes": p["notes"] or "",
                                "profit_share": "" if p["profit_share"] is None else p["profit_share"],
                                "salary": to_major(p["salary"]) if p["salary"] else "", "drawings_interest": p["drawings_interest"] or ""})
    for s in db.q("SELECT p.name, s.bu_id, s.share FROM partner_shares s JOIN partners p ON p.id=s.partner_id "
                  "ORDER BY p.id, s.bu_id"):
        out["Partner Shares"].append({"partner": s["name"], "bu_code": code(s["bu_id"]), "share": s["share"]})
    for a in db.q("SELECT * FROM fixed_assets ORDER BY code"):
        out["Fixed Assets"].append({k: a[k] for k in ("code", "name", "account", "accum_account", "expense_account", "acquired",
                                                      "life_months", "depr_start", "location", "serial", "notes")}
                                   | {"depr_method": a["method"] or "", "depr_factor": a["factor"] or ""}
                                   | {"cost": to_major(a["cost"]), "salvage": to_major(a["salvage"]) or "",
                                      "opening_depr": to_major(a["opening_depr"]) or "", "bu_code": code(a["bu_id"]),
                                      "book_depr": "No"})
    whs = {w["id"]: w for w in db.q("SELECT * FROM warehouses ORDER BY id")}
    for w in whs.values():
        out["Warehouses"].append({"name": w["name"], "location": w["location"], "account": w["account"] or ""})
    items = {i["id"]: i for i in db.q("SELECT * FROM stock_items ORDER BY name")}
    for i in items.values():
        out["Stock Items"].append({"name": i["name"], "unit": i["unit"], "sku": i["sku"], "reorder_level": i["reorder_level"],
                                   "category": i.get("category") or "", "std_cost": to_major(i.get("std_cost") or 0),
                                   "price": to_major(i.get("price") or 0), "notes": i.get("notes") or ""})
    bus_ = {b["id"]: b["code"] for b in db.q("SELECT id, code FROM business_units")}
    for p in db.q("SELECT * FROM products ORDER BY code"):
        out["Products"].append({k: p[k] for k in ("code", "name", "name_ar", "kind", "category", "unit", "vat_code", "status",
                                                  "stock_item", "inventory_account", "description")}
                               | {"price": to_major(p["price"]), "cost": to_major(p["cost"]) or "", "bu_code": bus_.get(p["bu_id"], ""),
                                  "specs": "; ".join(f"{a}: {b}" for a, b in json.loads(p["specs"] or "[]"))})
    for r in db.q("SELECT * FROM rates ORDER BY category, id"):
        out["Price List"].append({"category": r["category"], "item_no": r["item_no"], "sub": r["sub"],
                                  "description": r["description"], "unit": r["unit"], "cost": to_major(r["cost"]),
                                  "markup": r["markup"], "price": to_major(r["price"]), "notes": r["notes"],
                                  "eta_code": r["eta_code"]})
    for b in db.q("SELECT * FROM boqs ORDER BY id"):
        for ln in db.q("SELECT * FROM boq_lines WHERE boq_id=? ORDER BY seq, id", (b["id"],)):
            out["BOQ Lines"].append({"title": b["title"], "bu_code": code(b["bu_id"]), "rev": b["rev"], "date": b["date"],
                                     "client": b["client"], "location": b["location"], "category": ln["category"],
                                     "sub": ln["sub"], "item_no": ln["item_no"], "description": ln["description"],
                                     "unit": ln["unit"], "qty": ln["qty"], "price": to_major(ln["price"]),
                                     "cost": to_major(ln["cost"])})
    # sales invoices from extracts
    invoice_entries = set()
    boq_titles = {b["id"]: b["title"] for b in db.q("SELECT id, title FROM boqs")}
    for x in db.q("SELECT * FROM extracts ORDER BY date, id"):
        e = acc.get_entry(db, x["entry_id"]) if x["entry_id"] else None
        if e and e["status"] == "posted" and not e["reversed_by"]:
            invoice_entries.add(e["id"])
            roles_ = {a["code"]: a["role"] for a in db.q("SELECT code, role FROM accounts")}
            dr = next(ln for ln in e["lines"] if ln["debit"] and roles_.get(ln["account"]) != "wht_receivable")
            types_ = {a["code"]: a["type"] for a in db.q("SELECT code, type FROM accounts")}
            crs = [ln for ln in e["lines"] if ln["credit"] and roles_.get(ln["account"]) != "vat_output"]
            cr = next((ln for ln in crs if types_.get(ln["account"]) == "Revenue"), None)   # none when all billed in advance
            status, ar_acc, rev_acc = "Posted", dr["account"], cr["account"] if cr else ""
        else:
            status, ar_acc, rev_acc = "Draft", "", ""
        out["Sales Invoices"].append({"number": x["number"], "date": x["date"], "bu_code": code(x["bu_id"]),
                                      "boq_title": boq_titles.get(x["boq_id"], ""), "status": status,
                                      "ar_account": ar_acc, "revenue_account": rev_acc, "customer": x["customer"],
                                      "vat_code": x["vat_code"], "wht_code": x["wht_code"], "extract_no": x.get("seq_no") or "",
                                      "project": x.get("project") or "", "work": x.get("work") or "",
                                      "prev_paid": "" if x.get("prev_paid") is None else to_major(x["prev_paid"]),
                                      "deductions": to_major(x.get("deductions") or 0) or "", "final": "Yes" if x.get("final") else "",
                                      "notes": x.get("notes") or ""})
        for ln in db.q("SELECT * FROM extract_lines WHERE extract_id=? ORDER BY seq, id", (x["id"],)):
            pct = ln.get("pct")
            out["Invoice Lines"].append({"number": x["number"], "category": ln["category"], "sub": ln["sub"],
                                         "item_no": ln["item_no"], "description": ln["description"], "unit": ln["unit"],
                                         "qty": ln["qty"], "price": to_major(ln["price"]),
                                         "pct": "" if pct in (None, 1, 1.0) else pct, "account": ln["account"] or ""})
    # stock moves (only those still attached to a live entry, or entry-less transfers/counts)
    stock_entries = set()
    moves = db.q("SELECT m.*, i.name AS item, i.unit, w.name AS warehouse FROM stock_moves m "
                 "JOIN stock_items i ON i.id=m.item_id JOIN warehouses w ON w.id=m.warehouse_id ORDER BY m.id")
    skip = set()
    jitems = {r["entry_id"] for r in db.q("SELECT DISTINCT entry_id FROM material_lines")}   # exported with their voucher
    for idx, m in enumerate(moves):
        if m["id"] in skip or m["entry_id"] in jitems:
            continue
        e = acc.get_entry(db, m["entry_id"]) if m["entry_id"] else None
        if e and (e["status"] != "posted" or e["reversed_by"]):
            continue
        row = {"date": m["date"], "type": m["kind"], "warehouse": m["warehouse"], "item": m["item"],
               "unit": m["unit"], "qty": _num(m["qty"]), "rate": to_major(m["rate"]),
               "bu_code": code(m["bu_id"]), "note": m["note"]}
        if m["kind"] == "Receipt" and e:
            row["vendor"] = e["party"]
            inv_line = next(ln for ln in e["lines"] if ln["debit"])
            for ln in e["lines"]:
                if ln["credit"]:
                    arow = db.one("SELECT subtype, role FROM accounts WHERE code=?", (ln["account"],))
                    if arow["role"] == "wht_payable":
                        continue
                    if arow["subtype"] == "cash":
                        row["paid"], row["paid_from"] = ln["credit"], ln["account"]
                    else:
                        row["payable_account"] = ln["account"]
            row["vat_code"], row["wht_code"] = e["meta"].get("vat_code", ""), e["meta"].get("wht_code", "")
            row["rate"] = round(inv_line["debit"] / m["qty"], 6) if m["qty"] else row["rate"]
        elif m["kind"] == "Issue" and e:
            row["qty"] = -m["qty"]
            row["account"] = next(ln["account"] for ln in e["lines"] if ln["debit"])
            row["rate"] = round(e["amount"] / row["qty"], 6)
        elif m["kind"] == "Transfer":
            pair = next((n for n in moves[idx + 1: idx + 3] if n["kind"] == "Transfer" and n["qty"] > 0
                         and n["item_id"] == m["item_id"] and n["entry_id"] == m["entry_id"]), None)
            if pair:
                skip.add(pair["id"])
                row["to_warehouse"] = pair["warehouse"]
            row["qty"] = -m["qty"]
        elif m["kind"] == "Adjustment" and e:
            row["rate"] = round(e["amount"] / abs(m["qty"]), 6) if m["qty"] else 0
        if e:
            stock_entries.add(e["id"])
        out["Stock Moves"].append(row)

    # journal-based documents
    accounts = {a["code"]: a for a in db.q("SELECT * FROM accounts")}
    # keep the company's own voucher numbers (e.g. JE-0001) when each belongs to one entry only
    refs = {}
    for r in db.q("SELECT ref, COUNT(*) n FROM journal_entries WHERE status='posted' AND ref<>'' GROUP BY ref"):
        refs[r["ref"]] = r["n"]
    for e in (acc.get_entry(db, r["id"]) for r in
              db.q("SELECT id FROM journal_entries WHERE status='posted' ORDER BY date, id")):
        if e["id"] in invoice_entries or e["id"] in stock_entries:
            continue
        row = None if e["id"] in jitems else _entry_to_row(e, accounts)   # item lines only fit the Journal Vouchers sheet
        if row:
            sheet, data = row
            out[sheet].append(data)
        else:
            vno = e["ref"] if e["ref"] and refs.get(e["ref"]) == 1 and len(e["ref"]) <= 30 else f"JV-{e['id']}"
            mls = db.q("SELECT m.*, w.name AS wh FROM material_lines m LEFT JOIN warehouses w ON w.id=m.warehouse_id "
                       "WHERE m.entry_id=? ORDER BY m.id", (e["id"],))
            for ml in mls:   # item lines: the app makes their debit again from qty × unit price
                out["Journal Vouchers"].append({"voucher": vno, "date": e["date"], "kind": e["kind"], "bu_code": e["bu_code"] or "",
                                                "party": e["party"], "ref": e["ref"], "memo": e["memo"], "account": ml["account"],
                                                "debit": to_major(ml["amount"]), "credit": "", "line_party": "", "item": ml["description"],
                                                "item_unit": ml["unit"], "item_qty": _num(ml["qty"]), "item_rate": to_major(ml["rate"]),
                                                "item_to": ml["wh"] or ""})
            for ln in e["lines"][len(mls):]:
                out["Journal Vouchers"].append({"voucher": vno, "date": e["date"], "kind": e["kind"],
                                                "bu_code": e["bu_code"] or "", "party": e["party"], "ref": e["ref"],
                                                "memo": e["memo"], "account": ln["account"],
                                                "debit": ln["debit"] or "", "credit": ln["credit"] or "",
                                                "line_party": ln.get("party", "")})
    return out


def _entry_to_row(e, accounts):
    """Map a posted entry back onto its document sheet; None -> goes to Journal Vouchers."""
    if e["reversed_by"] or e["reversal_of"]:
        return None
    if any(ln.get("party") and ln["party"] != e["party"] for ln in e["lines"]):
        return None  # lines for several parties (partners' capital moves) only fit the Journal Vouchers sheet
    role = lambda code: accounts[code]["role"]  # noqa: E731
    dr = [ln for ln in e["lines"] if ln["debit"] and role(ln["account"]) not in ("vat_input", "wht_receivable")]
    cr = [ln for ln in e["lines"] if ln["credit"] and role(ln["account"]) not in ("wht_payable", "vat_output")]
    taxes = {"vat_code": e["meta"].get("vat_code", ""), "wht_code": e["meta"].get("wht_code", "")}
    sub = lambda code: accounts[code]["subtype"]  # noqa: E731
    typ = lambda code: accounts[code]["type"]  # noqa: E731
    m, base = e["meta"], {"date": e["date"], "bu_code": e["bu_code"] or ""}

    def split(credits, payable_types=("Liability",)):
        paid = [ln for ln in credits if sub(ln["account"]) == "cash"]
        owed = [ln for ln in credits if sub(ln["account"]) != "cash" and typ(ln["account"]) in payable_types]
        if len(paid) > 1 or len(owed) > 1 or len(paid) + len(owed) != len(credits):
            return None
        return {"paid": paid[0]["credit"] if paid else 0, "paid_from": paid[0]["account"] if paid else "",
                "payable_account": owed[0]["account"] if owed else ""}

    k = e["kind"]
    if k == "Direct Cost" and len(dr) == 1 and typ(dr[0]["account"]) == "COGS":
        s = split(cr)
        if s is None or abs(dr[0]["debit"] - (m.get("quoted", 0) - m.get("savings", 0))) > 0.005:
            return None
        return "Direct Costs", base | s | taxes | {"account": dr[0]["account"], "description": e["memo"], "vendor": e["party"],
                                           "deal_ref": m.get("deal_ref", ""), "deal_status": m.get("deal_status", ""),
                                           "qty": m.get("qty"), "unit": m.get("unit", ""), "rate": m.get("rate"),
                                           "savings": m.get("savings") or "", "method": m.get("method", ""),
                                           "negotiation": "pass" if m.get("savings") else ""}   # booked at the lower cost
    if k == "Operating Expense" and len(dr) == 1 and typ(dr[0]["account"]) == "Expense":
        s = split(cr)
        if s is None:
            return None
        return "Operating Expenses", base | s | taxes | {"account": dr[0]["account"], "description": e["memo"],
                                                 "vendor": e["party"], "amount": dr[0]["debit"],
                                                 "method": m.get("method", "")}
    if k == "Payroll Accrual" and len(dr) == 1:
        s = split(cr)
        if s is None or abs(dr[0]["debit"] - round(m.get("qty", 0) * m.get("rate", 0), 2)) > 0.005:
            return None
        return "Payroll", base | s | {"worker": e["party"], "role": m.get("role", ""), "cost_type": m.get("cost_type", ""),
                                      "account": dr[0]["account"], "period": m.get("period", ""),
                                      "basis": m.get("basis", ""), "qty": m.get("qty"), "rate": m.get("rate"),
                                      "method": m.get("method", "")}
    if k == "Receipt Voucher" and len(dr) == 1 and len(cr) == 1 and sub(dr[0]["account"]) == "cash":
        return "Receipt Vouchers", base | {"ref": e["ref"], "memo": e["memo"], "party": e["party"],
                                           "amount": dr[0]["debit"], "method": m.get("method", ""),
                                           "received_into": dr[0]["account"], "credit_account": cr[0]["account"],
                                           "withheld": m.get("withheld") or "", "vat_code": m.get("vat_code", "")}
    if k == "Payment Voucher" and len(dr) == 1 and len(cr) == 1 and sub(cr[0]["account"]) == "cash" and \
            (accounts.get(dr[0]["account"]) or {}).get("type") == "Liability":   # e.g. an advance to a supplier stays a journal
        return "Payment Vouchers", base | {"party": e["party"], "against": dr[0]["account"], "amount": dr[0]["debit"],
                                           "method": m.get("method", ""), "paid_from": cr[0]["account"],
                                           "ref": e["ref"], "memo": e["memo"]}
    if k in ("Capital Contribution", "Owner Drawing") and len(dr) == 1 and len(cr) == 1 and \
            sub((dr[0] if k == "Capital Contribution" else cr[0])["account"]) == "cash":
        cash, eq = (dr[0], cr[0]) if k == "Capital Contribution" else (cr[0], dr[0])
        return "Partner Transactions", {"date": e["date"], "partner": e["party"], "type": k,
                                        "amount": dr[0]["debit"], "cash_account": cash["account"],
                                        "equity_account": eq["account"], "bu_code": e["bu_code"] or "",
                                        "memo": e["memo"]}
    return None


# =================================================================== IMPORT
class ImportErrors(acc.ValidationError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__(f"{len(errors)} problem(s) found — nothing was imported")


def _norm_header(h):
    h = ("" if h is None else str(h)).strip()
    return h.split(" / ")[0].strip().lower()


def header_map(sheet, headers):
    """Map incoming header cells to column keys (accepts key, English label, or 'English / Arabic')."""
    lookup = {}
    for c in sheet["cols"]:   # an input column wins over a calculated helper column with the same label
        for k in (f"{c['en']} / {c['ar']}".lower(), c["key"].lower(), c["en"].lower(), c["ar"]):
            lookup.setdefault(k, c["key"])
    out, used = {}, set()
    for i, h in enumerate(headers):
        k = lookup.get(str(h or "").strip().lower()) or lookup.get(_norm_header(h)) or lookup.get(str(h or "").strip())
        out[i] = k if k not in used else None   # the same column twice: the first one counts
        if k:
            used.add(k)
    return out


def _cell(v, col):
    if v is None:
        return ""
    if col["kind"] == "date":
        if isinstance(v, (_dt.datetime, _dt.date)):
            return acc.clean_date(v)
        s = str(v).strip()
        if not s:
            return ""
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"):
            try:
                return _dt.datetime.strptime(s[:10], fmt).date().isoformat()
            except ValueError:
                pass
        if s.replace(".", "", 1).isdigit():  # Excel serial number
            return (_dt.date(1899, 12, 30) + _dt.timedelta(days=int(float(s)))).isoformat()
        return s
    if isinstance(v, (_dt.datetime, _dt.date)):
        return acc.clean_date(v)
    if isinstance(v, float) and v.is_integer() and col["kind"] in ("text", "acct", "bu"):
        v = int(v)
    if isinstance(v, str):
        v = v.strip()
    return v


def rows_from_table(sheet, headers, raw_rows):
    hmap = header_map(sheet, headers)
    cols = {c["key"]: c for c in sheet["cols"]}
    out = []
    for raw in raw_rows:
        row = {}
        for i, v in enumerate(raw):
            k = hmap.get(i)
            if k and cols[k]["kind"] != "calc":
                row[k] = _cell(v, cols[k])
        if any(v not in ("", None) for v in row.values()):
            out.append(row)
    return out


def read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        raise acc.ValidationError("Excel import needs openpyxl:  python -m pip install openpyxl") from None
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    data = {}
    by_ar = {s["ar"]: s for s in SHEETS}
    for ws in wb.worksheets:
        if ws.title == "_BeyondData":   # everything the sheets cannot hold (see workbook._embed_extras)
            cells = [r[0] for r in ws.iter_rows(min_col=1, max_col=1, values_only=True)]
            if cells and cells[0] == "beyond-erp-extras-1":
                try:
                    data["_extras"] = json.loads("".join(str(c) for c in cells[1:] if c))
                except ValueError:
                    pass
            continue
        sheet = SHEET_BY_NAME.get(ws.title) or by_ar.get(ws.title.strip())   # English or Arabic workbook
        if not sheet:
            continue
        rows = list(ws.iter_rows(values_only=True))
        # header row: first row within the top 10 whose cells match at least 2 known columns
        scores = [(sum(1 for k in header_map(sheet, r).values() if k), -i) for i, r in enumerate(rows[:10])]
        best = max(scores, default=(0, 0))
        if best[0] < min(2, len(sheet["cols"])):
            continue
        hi = -best[1]
        data[sheet["name"]] = rows_from_table(sheet, rows[hi], rows[hi + 1:])
    wb.close()
    return data


def read_csv(sheet_name, content: bytes):
    sheet = SHEET_BY_NAME.get(sheet_name)
    if not sheet:
        raise acc.ValidationError(f"Unknown sheet {sheet_name!r}")
    textc = content.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(textc)))
    if not rows:
        return {sheet_name: []}
    return {sheet_name: rows_from_table(sheet, rows[0], rows[1:])}


def read_json(content: bytes):
    try:
        obj = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as e:
        raise acc.ValidationError(f"Not valid JSON: {e}") from None
    sheets = obj.get("sheets") if isinstance(obj, dict) and "sheets" in obj else obj
    if not isinstance(sheets, dict):
        raise acc.ValidationError('JSON must be {"sheets": {"Sheet name": [ {column: value}, ... ]}}')
    data = {}
    for name, rows in sheets.items():
        sheet = SHEET_BY_NAME.get(name)
        if not sheet or not isinstance(rows, list):
            continue
        cols = {c["key"]: c for c in sheet["cols"]}
        clean = []
        for r in rows:
            if isinstance(r, dict):
                hm = header_map(sheet, list(r.keys()))
                clean.append({hm[i]: _cell(v, cols[hm[i]]) for i, v in enumerate(r.values())
                              if hm.get(i) and cols[hm[i]]["kind"] != "calc"})
        data[name] = clean
    if isinstance(obj, dict) and isinstance(obj.get("extras"), dict):
        data["_extras"] = obj["extras"]
    return data


def _site_sheets(db, out):
    """Projects, contracts, requests, equipment and attendance carry no money of their own: export them so a re-import
    brings them back. Certificates, equipment usage / maintenance, stock counts and e-receipts post entries — their money
    is in the journal, and their records travel in the JSON backup (extras), never twice."""
    from . import contracts, equipment, labour, procurement, projects
    pj = projects.export_rows(db)
    out["Project Phases"] = [{**p, "progress": p["progress"]} for p in pj["phases"]]
    out["Project Tasks"] = [{k: t[k] for k in ("bu_code", "phase", "title", "assignee", "due", "priority", "status", "notes", "created_by", "done_at")}
                            for t in pj["tasks"]]
    out["Contracts"], out["Variation Orders"] = [], []
    for c in contracts.export_rows(db):
        head = {k: c[k] for k in ("code", "kind", "party", "bu_code", "title", "value", "signed_on", "start", "finish", "retention_pct",
                                  "retention_release", "advance", "advance_recovery_pct", "vat_code", "wht_code", "payment_terms", "status")}
        head["bu_code"] = head["bu_code"] or ""
        for i, ln in enumerate(c["lines"] or [{}]):
            out["Contracts"].append({**(head if i == 0 else {"code": c["code"], "kind": c["kind"], "party": c["party"]}), **ln})
        out["Variation Orders"] += [{"contract": c["code"], **{k: v[k] for k in ("code", "date", "description", "amount", "days", "status", "decided_by", "decided_at")}}
                                    for v in c["variations"]]
    out["Purchase Requests"], out["Supplier Quotes"] = [], []
    for r in procurement.export_rows(db):
        head = {k: r[k] or "" for k in ("code", "date", "requested_by", "bu_code", "warehouse", "needed_by", "status", "notes")}
        for i, ln in enumerate(r["lines"]):
            out["Purchase Requests"].append({**(head if i == 0 else {"code": r["code"], "date": r["date"]}), **ln})
        for q in r["quotes"]:
            for j, (desc, rate) in enumerate(q["rates"].items()):
                out["Supplier Quotes"].append({"request": r["code"], "supplier": q["supplier"], **({k: q[k] for k in ("date", "valid_until", "delivery_days",
                                               "vat_code", "terms")} if j == 0 else {}), "description": desc, "rate": rate,
                                               "awarded": "Yes" if desc in q["awarded"] else ""})
    eq = equipment.export_rows(db)
    out["Equipment"] = [{**e, "asset_code": e["asset_code"] or "", "bu_code": e["bu_code"] or ""} for e in eq["equipment"]]
    out["Maintenance Plans"] = eq["plans"]
    out["Attendance"] = labour.export_rows(db)


def site_extras(db):
    """Records that post entries (their money is restored from the journal; these link back to it by reference)."""
    from . import contracts, equipment, ereceipt, labour, stocktake
    eq = equipment.export_rows(db)
    return {"certificates": {c["code"]: c["certificates"] for c in contracts.export_rows(db) if c["certificates"]},
            "equipment_usage": eq["usage"], "maintenance_logs": eq["logs"], "stock_counts": stocktake.export_rows(db),
            "ereceipts": ereceipt.export_rows(db), "attendance": labour.export_rows(db)}


def site_extras_import(db, ex, counts):
    from . import contracts, equipment, ereceipt, stocktake
    for code, certs in (ex.get("certificates") or {}).items():
        c = db.one("SELECT id FROM contracts WHERE code=?", (code,))
        if c and not db.one("SELECT 1 FROM sub_certificates WHERE contract_id=?", (c["id"],)):
            contracts.import_certificates(db, c["id"], code, certs)
            counts["Subcontractor certificates"] = counts.get("Subcontractor certificates", 0) + len(certs)
    if (ex.get("equipment_usage") or ex.get("maintenance_logs")) and not db.one("SELECT 1 FROM equipment_usage UNION SELECT 1 FROM maintenance_logs"):
        equipment.import_rows(db, {"usage": ex.get("equipment_usage"), "logs": ex.get("maintenance_logs")}, counts)
    if ex.get("stock_counts"):
        stocktake.import_rows(db, ex["stock_counts"], counts)
    if ex.get("ereceipts"):
        ereceipt.import_rows(db, ex["ereceipts"], counts)
    if ex.get("attendance") and not db.one("SELECT 1 FROM attendance"):   # after payroll re-created the employees
        from . import labour
        labour.import_rows(db, ex["attendance"], counts)

def export_extras(db, files=True, secrets=False):
    """What the sheets cannot hold: letterhead images, offer types and offers (JSON backups only). files=False: the
    attached files are left out (a zipped export carries them in its Attachments folders instead)."""
    return {"assets": db.q("SELECT key, mime, w, h, data FROM assets ORDER BY key"),
            "offer_types": db.q("SELECT id, name, content FROM offer_types ORDER BY id"),
            "offers": db.q("SELECT o.*, u.code AS bu_code FROM offers o LEFT JOIN business_units u ON u.id=o.bu_id ORDER BY o.id"),
            "ownership_changes": _ownership_export(db), "loans": loans.export_rows(db),
            "custodies": custody.export_rows(db), "rentals": rentals.export_rows(db), "payroll": payroll.export_rows(db),
            "sales": sales.export_rows(db), "bank": bankrec.export_rows(db), "recurring": recurring.export_rows(db),
            "entry_meta": _entry_meta_export(db), "attachments": _attachments_export(db) if files else [],
            "purchase_orders": _po_export(db), "services": _svc_export(db), "chat": _chat_export(db),
            "dashboard_groups": db.q("SELECT key, name, name_ar, layout, sort, look FROM user_groups ORDER BY sort, id"),
            "line_styles": _line_styles_export(db), "crm": _crm_export(db), "care": _care_export(db), "desk": _desk_export(db), "automations": {t: db.q(f"SELECT * FROM {t} ORDER BY id") for t in ("automations", "automation_log")}, "workdays": db.q("SELECT * FROM workdays ORDER BY id"), "kb_pages": db.q("SELECT * FROM kb_pages ORDER BY id"), "entry_tags": _tags_export(db),
            "bu_groups": {r["code"]: r["grp"] for r in db.q("SELECT code, grp FROM business_units WHERE grp<>''")}, "site": site_extras(db),
            "people": _people_export(db), "approvals": _approvals_export(db), "history": _history_export(db),
            "bills": __import__("erp.billsubs", fromlist=["x"]).export_rows(db),
            "year_closes": __import__("erp.yearend", fromlist=["x"]).export_rows(db),
            **({"secrets": _secrets_export(db)} if secrets else {})}


def _secrets_export(db):
    """Keys of connected tools — only in a super administrator's backup: the API keys people made for outside tools (each
    kept as its fingerprint, so the same key keeps working after a restore), the e-invoice / e-receipt (ETA) credentials
    and, on the desktop app, the AI provider key. (On the hosted app the AI provider key lives in the workspace server,
    not in these books.)"""
    from . import cloud
    from .db import ETA_SECRETS
    names = {u["id"]: u["username"] for u in db.q("SELECT id, username FROM users")}
    return {"api_keys": [{**{k: r[k] for k in r.keys() if k not in ("id", "user_id")}, "user": names.get(r["user_id"], "")}
                         for r in db.q("SELECT * FROM api_keys ORDER BY id")],
            "settings": {k: db.setting(k) for k in ETA_SECRETS if db.setting(k)},
            "device": {k: cloud.local_get(db, k) for k in ("ai_key",) if cloud.local_get(db, k)}}


def _secrets_import(db, sec, mode, counts):
    """Only from a super administrator's import (the import route drops this part for anyone else)."""
    from . import cloud
    from .db import ETA_SECRETS
    ids = {u["username"]: u["id"] for u in db.q("SELECT id, username FROM users")}
    cols = {r["name"] for r in db.q("PRAGMA table_info(api_keys)")} - {"id"}
    n = 0
    for r in sec.get("api_keys") or []:
        if not isinstance(r, dict) or not ids.get(r.get("user")) or not r.get("key_hash") \
                or db.one("SELECT 1 FROM api_keys WHERE key_hash=?", (r["key_hash"],)):
            continue
        vals = {k: r[k] for k in r if k in cols}
        vals["user_id"] = ids[r["user"]]
        db.conn.execute(f"INSERT INTO api_keys({','.join(vals)}) VALUES ({','.join('?' * len(vals))})", list(vals.values()))
        n += 1
    if n:
        counts["API keys"] = n
    m = 0
    for k, v in (sec.get("settings") or {}).items():
        if k in ETA_SECRETS and v and (mode == "replace" or not db.setting(k)):
            db.set_setting(k, str(v))
            m += 1
    for k, v in (sec.get("device") or {}).items():
        if k == "ai_key" and v and (mode == "replace" or not cloud.local_get(db, k)):
            cloud.local_set(db, k, str(v))
            m += 1
    if m:
        counts["Connected tools' keys"] = m


def read_sqlite(raw):
    """A database backup file (.sqlite3, "Database backup file") → the same data a JSON backup holds, so it imports
    through the normal path: checked first, and in the shared workspace it reaches every device. An older backup is
    upgraded on a temporary copy first."""
    import tempfile
    from pathlib import Path
    from . import db as _dbm
    if not raw[:16] == b"SQLite format 3\x00":
        raise acc.ValidationError("That is not a database backup file (.sqlite3)")
    tmpdir = None
    if _dbm.sqlite3 is None:                     # browser: open the image with sql.js
        from . import jsdb
        src = _dbm.Database(":memory:", connection=jsdb.image(raw))
    else:
        tmpdir = tempfile.TemporaryDirectory()
        p = Path(tmpdir.name) / "restore.sqlite3"
        p.write_bytes(raw)
        try:
            pre = _dbm.sqlite3.connect(str(p))
            ok = pre.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='journal_entries'").fetchone()
            pre.close()
        except _dbm.sqlite3.DatabaseError:
            ok = None
        if not ok:
            tmpdir.cleanup()
            raise acc.ValidationError("That database file is not a BeyondERP backup")
        src = _dbm.Database(str(p))
    try:
        if not src.one("SELECT 1 FROM users"):
            raise acc.ValidationError("That database file is not a BeyondERP backup")
        return read_json(to_json(src, secrets=True).encode("utf-8"))
    finally:
        try:
            src.conn.close()
        except Exception:   # noqa: BLE001 — already closed
            pass
        if tmpdir:
            tmpdir.cleanup()


def _history_export(db):
    """The activity log, support reports and each person's AI assistant conversations (by username)."""
    names = {u["id"]: u["username"] for u in db.q("SELECT id, username FROM users")}
    return {"audit_log": db.q("SELECT ts, username, action, detail FROM audit_log ORDER BY id"),
            "support_tickets": [{k: r[k] for k in r.keys() if k != "id"} for r in db.q("SELECT * FROM support_tickets ORDER BY id")],
            "alert_feed": [{k: r[k] for k in r.keys() if k != "id"} for r in db.q("SELECT * FROM alert_feed ORDER BY id")],
            "ai_threads": [{**{k: r[k] for k in r.keys() if k not in ("id", "user_id")}, "user": names.get(r["user_id"], "")}
                           for r in db.q("SELECT * FROM ai_threads ORDER BY id") if names.get(r["user_id"])]}


def _history_import(db, h, mode, counts):
    def put(table, rows):
        cols = {r["name"] for r in db.q(f"PRAGMA table_info({table})")} - {"id"}
        n = 0
        for r in rows:
            if isinstance(r, dict):
                vals = {k: r[k] for k in r if k in cols}
                if vals:
                    db.conn.execute(f"INSERT INTO {table}({','.join(vals)}) VALUES ({','.join('?' * len(vals))})", list(vals.values()))
                    n += 1
        return n
    if isinstance(h.get("audit_log"), list) and h["audit_log"] and mode == "replace":
        have = {(r["ts"], r["username"], r["action"], r["detail"]) for r in db.q("SELECT ts, username, action, detail FROM audit_log")}
        n = put("audit_log", [r for r in h["audit_log"] if isinstance(r, dict)        # lines this log already has stay single
                              and (r.get("ts"), r.get("username"), r.get("action"), r.get("detail")) not in have])
        if n:
            counts["Activity log lines"] = n
    if isinstance(h.get("support_tickets"), list) and h["support_tickets"] and (mode == "replace" or not db.one("SELECT 1 FROM support_tickets")):
        db.conn.execute("DELETE FROM support_tickets")
        put("support_tickets", h["support_tickets"])
    if isinstance(h.get("alert_feed"), list) and h["alert_feed"] and (mode == "replace" or not db.one("SELECT 1 FROM alert_feed")):
        db.conn.execute("DELETE FROM alert_feed")
        put("alert_feed", h["alert_feed"])
    if isinstance(h.get("ai_threads"), list) and h["ai_threads"] and (mode == "replace" or not db.one("SELECT 1 FROM ai_threads")):
        ids = {u["username"]: u["id"] for u in db.q("SELECT id, username FROM users")}
        db.conn.execute("DELETE FROM ai_threads")
        n = put("ai_threads", [{**{k: v for k, v in r.items() if k != "user"}, "user_id": ids[r["user"]]}
                               for r in h["ai_threads"] if isinstance(r, dict) and ids.get(r.get("user"))])
        if n:
            counts["AI assistant conversations"] = n


PEOPLE_COLS = ("username", "full_name", "email", "is_admin", "is_super", "perms", "group_key", "needs_approval", "grants", "active")


def _people_export(db):
    """Who works in these books, with their roles and permissions — never their passwords (API keys stay out too)."""
    have = {r["name"] for r in db.q("PRAGMA table_info(users)")}
    return [{k: u[k] for k in PEOPLE_COLS if k in have} for u in db.q("SELECT * FROM users ORDER BY id")]


def _people_import(db, rows):
    """A full restore brings back the people who are not in these books yet, without a password: an administrator
    sets one (Users → edit), or they sign in through the shared workspace. People already here are left as they are."""
    have = {r["name"] for r in db.q("PRAGMA table_info(users)")}
    n = 0
    for r in rows:
        if not isinstance(r, dict) or not acc.text(r.get("username")):
            continue
        name = acc.text(r["username"], maxlen=60)
        if db.one("SELECT 1 FROM users WHERE username=?", (name,)):
            continue
        email = acc.text(r.get("email"), maxlen=200).lower()
        if email and db.one("SELECT 1 FROM users WHERE lower(email)=?", (email,)):
            email = ""
        perms = r.get("perms") if isinstance(r.get("perms"), str) else json.dumps(r.get("perms") or {})
        grants = r.get("grants") if isinstance(r.get("grants"), str) else json.dumps(r.get("grants") or [])
        vals = {"username": name, "full_name": acc.text(r.get("full_name"), maxlen=120), "pw_hash": "!", "email": email,
                "is_admin": 1 if r.get("is_admin") else 0, "is_super": 1 if r.get("is_super") else 0, "perms": perms,
                "group_key": acc.text(r.get("group_key"), maxlen=60), "needs_approval": 1 if r.get("needs_approval") else 0,
                "grants": grants, "active": 0 if r.get("active") in (0, False, "0") else 1}
        vals = {k: v for k, v in vals.items() if k in have}
        db.conn.execute(f"INSERT INTO users({','.join(vals)}) VALUES ({','.join('?' * len(vals))})", list(vals.values()))
        n += 1
    return n


def _approvals_export(db):
    """Entries waiting for approval (and the decided ones, for their history). The entry an edit changes, and the
    entry an approval posted, are named by their fingerprint: restored entries get new numbers."""
    from . import attachments
    out = []
    for a in db.q("SELECT * FROM approvals ORDER BY id"):
        row = {k: a[k] for k in a.keys() if k not in ("id", "target_id", "entry_id")}
        row["target"] = attachments._fp(db, a["target_id"]) if a["target_id"] else None
        row["entry"] = attachments._fp(db, a["entry_id"]) if a["entry_id"] else None
        out.append(row)
    return out


def _approvals_import(db, rows):
    from . import attachments
    pool = {}
    for e in db.q("SELECT id FROM journal_entries ORDER BY id"):
        fp = attachments._fp(db, e["id"])
        if fp:
            pool.setdefault(tuple(fp), e["id"])
    cols = {r["name"] for r in db.q("PRAGMA table_info(approvals)")}
    n = 0
    for r in rows:
        if not isinstance(r, dict) or not r.get("dtype") or not r.get("created_by"):
            continue
        target = pool.get(tuple(r["target"])) if isinstance(r.get("target"), list) else None
        if r.get("target") and not target and r.get("status") == "pending":
            continue          # an edit of an entry that is not in these books: approving it would post a new one
        vals = {k: r[k] for k in r if k in cols and k not in ("id", "target_id", "entry_id")}
        vals["target_id"] = target
        vals["entry_id"] = pool.get(tuple(r["entry"])) if isinstance(r.get("entry"), list) else None
        db.conn.execute(f"INSERT INTO approvals({','.join(vals)}) VALUES ({','.join('?' * len(vals))})", list(vals.values()))
        n += 1
    return n


def _tags_export(db):
    """Tags by the entry's fingerprint (restored entries get new numbers)."""
    from . import attachments
    out = {}
    for r in db.q("SELECT entry_id, tag FROM entry_tags ORDER BY entry_id"):
        fp = attachments._fp(db, r["entry_id"])
        if fp:
            out.setdefault(json.dumps(fp), []).append(r["tag"])
    return [{"key": json.loads(k), "tags": v} for k, v in out.items()]


def _desk_export(db):
    from . import desk
    return desk.export_rows(db)


def _care_export(db):
    from . import care
    return care.export_rows(db)


def _crm_export(db):
    from . import crm
    return crm.export_rows(db)


def _line_styles_export(db):
    """Each styled BOQ / invoice line's look, by its document and its place in it (the sheets hold the lines)."""
    out = {"boq": [], "invoice": []}
    for b in db.q("SELECT b.id, b.title, u.code AS bu FROM boqs b LEFT JOIN business_units u ON u.id=b.bu_id"):
        for i, ln in enumerate(db.q("SELECT style FROM boq_lines WHERE boq_id=? ORDER BY seq, id", (b["id"],))):
            if ln["style"]:
                out["boq"].append([b["title"], b["bu"] or "", i, ln["style"]])
    for x in db.q("SELECT id, number FROM extracts"):
        for i, ln in enumerate(db.q("SELECT style FROM extract_lines WHERE extract_id=? ORDER BY seq, id", (x["id"],))):
            if ln["style"]:
                out["invoice"].append([x["number"], i, ln["style"]])
    return out


def _line_styles_import(db, ls):
    n = 0
    for title, bu, i, st in ls.get("boq") or []:
        b = db.one("SELECT b.id FROM boqs b LEFT JOIN business_units u ON u.id=b.bu_id WHERE b.title=? AND COALESCE(u.code,'')=? "
                   "ORDER BY b.id DESC", (acc.text(title), acc.text(bu)))
        rows = db.q("SELECT id FROM boq_lines WHERE boq_id=? ORDER BY seq, id", (b["id"],)) if b else []
        if isinstance(i, int) and 0 <= i < len(rows):
            db.conn.execute("UPDATE boq_lines SET style=? WHERE id=?", (acc.line_style(st), rows[i]["id"]))
            n += 1
    for number, i, st in ls.get("invoice") or []:
        x = db.one("SELECT id FROM extracts WHERE number=?", (acc.text(number),))
        rows = db.q("SELECT id FROM extract_lines WHERE extract_id=? ORDER BY seq, id", (x["id"],)) if x else []
        if isinstance(i, int) and 0 <= i < len(rows):
            db.conn.execute("UPDATE extract_lines SET style=? WHERE id=?", (acc.line_style(st), rows[i]["id"]))
            n += 1
    return n


def _chat_export(db):
    from . import chat
    return chat.export_rows(db)


def _svc_export(db):
    from . import services
    return services.export_rows(db)


def _po_export(db):
    from . import purchasing
    return purchasing.export_rows(db)


MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
MONTHS_AR = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]


def _safe(name, n=90):
    return (re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", str(name or "")).strip() or "file")[:n].rstrip(" .")


def workbook_zip(db, lang="en", secrets=False):
    """The whole books as an Excel workbook plus every attached file, in one .zip:
      BEYOND_ERP_<date>.xlsx
      Attachments/2026-03 March 2026/<record> — <file name>      (by the month of the entry, or when the file was added)
      Attachments/links.csv                                      (files kept as links)
      Attachments/manifest.json                                  (which record each file belongs to — importing the
                                                                   zip back puts every file back on its record)"""
    import base64
    import zipfile
    from . import attachments, workbook
    names = MONTHS_AR if lang == "ar" else MONTHS_EN
    buf = io.BytesIO()
    rows, used, links = attachments.export_rows(db), set(), []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"BEYOND_ERP_{acc.today()}{'_AR' if lang == 'ar' else ''}.xlsx", workbook.build_workbook(db, template=False, lang=lang, files=False, secrets=secrets))
        manifest = []
        for r in rows:
            day = (r["key"][0] if r["owner"] == "entry" and isinstance(r.get("key"), list) else "") or str(r.get("created_at") or "")[:10] or acc.today()
            try:
                y, m = int(day[:4]), int(day[5:7])
                folder = f"{y}-{m:02d} {names[m - 1]} {y}"
            except (ValueError, IndexError):
                folder = "Undated"
            what = (f"{r['key'][1]} {r['key'][2]}".strip() if r["owner"] == "entry" else f"{r['owner']} {r['key']}")
            item = {k: r.get(k) for k in ("owner", "key", "title", "mime", "thumb", "url", "size", "created_by", "created_at")}
            if r.get("url") and not r.get("data"):
                links.append([folder, what, r.get("title") or "", r["url"]])
                manifest.append(item)
                continue
            title = _safe(r.get("title") or "file")
            path = f"Attachments/{folder}/{_safe(day + ' ' + what, 70)} — {title}"
            base, k = path, 2
            while path.lower() in used:
                stem, dot, ext = base.rpartition(".")
                path = f"{stem} ({k}).{ext}" if dot and "/" not in ext else f"{base} ({k})"
                k += 1
            used.add(path.lower())
            try:
                z.writestr(path, base64.b64decode(r.get("data") or ""))
            except (ValueError, TypeError):
                continue
            manifest.append({**item, "path": path})
        if links:
            out = io.StringIO()
            w = csv.writer(out)
            w.writerow(["month", "record", "title", "link"])
            w.writerows(links)
            z.writestr("Attachments/links.csv", "\ufeff" + out.getvalue())
        z.writestr("Attachments/manifest.json", json.dumps({"format": "beyond-erp-files-1", "files": manifest}, ensure_ascii=False, indent=1))
    return buf.getvalue()


def is_files_zip(raw):
    """A zip made by workbook_zip (and not an .xlsx, which is a zip too)."""
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            n = z.namelist()
            return "[Content_Types].xml" not in n and any(x.lower().endswith(".xlsx") for x in n)
    except zipfile.BadZipFile:
        return False


def read_files_zip(raw, reader):
    """(the workbook's data, with the files from the Attachments folders put back in its extras). reader(path) reads
    the workbook found inside."""
    import base64
    import tempfile
    import zipfile
    from pathlib import Path
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        books = [x for x in z.namelist() if x.lower().endswith(".xlsx") and "/" not in x.strip("/") and not Path(x).name.startswith("~$")]
        if not books:
            raise acc.ValidationError("The zip has no Excel workbook in it")
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            f.write(z.read(books[0]))
            tmp = f.name
        try:
            data = reader(tmp)
        finally:
            Path(tmp).unlink(missing_ok=True)
        try:
            man = json.loads(z.read("Attachments/manifest.json").decode("utf-8"))
        except (KeyError, ValueError):
            man = {}
        files = []
        for it in man.get("files") or []:
            if not isinstance(it, dict):
                continue
            row = {k: it.get(k) for k in ("owner", "key", "title", "mime", "thumb", "url", "size", "created_by", "created_at")}
            if it.get("path"):
                try:
                    row["data"] = base64.b64encode(z.read(it["path"])).decode("ascii")
                except KeyError:
                    continue
            files.append(row)
    if files:
        data.setdefault("_extras", {})
        if not isinstance(data["_extras"], dict):
            data["_extras"] = {}
        data["_extras"]["attachments"] = files
    return data


def _attachments_export(db):
    from . import attachments
    return attachments.export_rows(db)


def _entry_fp(e):
    return [e["date"], e["kind"], e["memo"] or "", e["party"] or "", int(e["total"] or 0)]


def _entry_meta_export(db):
    """What a voucher knows beyond its lines (bank matches, the sale / recurring entry it belongs to, its form, its
    reference). The sheets re-create vouchers with new numbers, so these are matched back by date, type, memo,
    party and amount."""
    out = []
    for e in db.q("SELECT e.id, e.date, e.kind, e.memo, e.party, e.ref, e.meta, e.created_by, e.created_at, e.history, "
                  "(SELECT SUM(debit) FROM journal_lines WHERE entry_id=e.id) total FROM journal_entries e WHERE e.status='posted' ORDER BY e.id"):
        out.append({"fp": _entry_fp(e), "ref": e["ref"] or "", "meta": json.loads(e["meta"] or "{}"),
                    "by": e["created_by"] or "", "at": e["created_at"] or "", "history": e["history"] or ""})   # who entered it, when, its edits
    return out


def _entry_meta_import(db, rows):
    pool = {}
    for e in db.q("SELECT e.id, e.date, e.kind, e.memo, e.party, e.ref, e.meta, (SELECT SUM(debit) FROM journal_lines WHERE entry_id=e.id) total "
                  "FROM journal_entries e WHERE e.status='posted' ORDER BY e.id"):
        pool.setdefault(json.dumps(_entry_fp(e)), []).append(e)
    loose = {}   # the same entry re-created by a sheet that does not carry its party (e.g. a stock issue for a sale)
    for k, q in pool.items():
        f = json.loads(k)
        loose.setdefault(json.dumps(f[:3] + f[4:]), []).extend(q)
    n = 0
    for r in rows or []:
        if not isinstance(r, dict) or not isinstance(r.get("fp"), list):
            continue
        q = pool.get(json.dumps(r["fp"])) or [e for e in loose.get(json.dumps(r["fp"][:3] + r["fp"][4:]), []) if not e.get("_used")]
        q = [e for e in q if not e.get("_used")]
        if not q:
            continue
        e = q[0]
        e["_used"] = True
        cur = json.loads(e["meta"] or "{}")
        meta = {**(r.get("meta") or {}), **{k: v for k, v in cur.items() if v not in (None, "", {}, [])}}
        db.conn.execute("UPDATE journal_entries SET meta=?, ref=?, party=CASE WHEN party='' OR party IS NULL THEN ? ELSE party END WHERE id=?",
                        (json.dumps(meta, ensure_ascii=False), r.get("ref") or e["ref"] or "", r["fp"][3] or "", e["id"]))
        if r.get("by") or r.get("at"):
            db.conn.execute("UPDATE journal_entries SET created_by=COALESCE(NULLIF(?,''),created_by), created_at=COALESCE(NULLIF(?,''),created_at), "
                            "history=COALESCE(NULLIF(?,''),history) WHERE id=?", (acc.text(r.get("by"), maxlen=120), acc.text(r.get("at"), maxlen=40),
                                                                                    r.get("history") if isinstance(r.get("history"), str) else "", e["id"]))
        n += 1
    return n


def _ownership_export(db):
    """Ownership history with partners by name (ids differ after an import)."""
    names = {str(p["id"]): p["name"] for p in db.q("SELECT id, name FROM partners")}
    out = []
    for ch in db.q("SELECT * FROM ownership_changes ORDER BY id"):
        conv = {}
        for side in ("before", "after"):
            st = json.loads(ch[side])
            conv[side] = {"shares": {names.get(k, k): v for k, v in st["shares"].items()},
                          "active": {names.get(k, k): v for k, v in st.get("active", {}).items()},
                          "dates": {names.get(k, k): v for k, v in st.get("dates", {}).items()},
                          "bu": [[names.get(p, p), b, v] for p, b, v in st.get("bu", [])]}
        out.append({"date": ch["date"], "kind": ch["kind"], "memo": ch["memo"], "created_by": ch["created_by"],
                    "created_at": ch["created_at"], **conv})
    return out


def _imp_extras(db, ex, mode, counts):
    from . import branding
    if ex.get("entry_meta"):
        _entry_meta_import(db, ex["entry_meta"])
    if mode == "replace":
        for t in ("offers", "offer_types", "loans", "custodies"):
            db.conn.execute(f"DELETE FROM {t}")
    if ex.get("loans"):
        counts["Loans"] = loans.import_rows(db, ex["loans"])
    if ex.get("custodies"):
        counts["Cash custody"] = custody.import_rows(db, ex["custodies"])
    if isinstance(ex.get("payroll"), dict) and (mode == "replace" or not db.one("SELECT 1 FROM employees")):
        payroll.import_rows(db, ex["payroll"], counts)
    if isinstance(ex.get("rentals"), dict) and (mode == "replace" or not db.one("SELECT 1 FROM rentals")):
        rentals.import_rows(db, ex["rentals"], counts)
    if isinstance(ex.get("sales"), dict) and (mode == "replace" or not db.one("SELECT 1 FROM sales_orders")):
        sales.import_rows(db, ex["sales"], counts)
    if isinstance(ex.get("bank"), dict) and (mode == "replace" or not db.one("SELECT 1 FROM bank_statements")):
        bankrec.import_rows(db, ex["bank"], counts)
    if isinstance(ex.get("recurring"), dict) and (mode == "replace" or not db.one("SELECT 1 FROM recurring")):
        recurring.import_rows(db, ex["recurring"], counts)
    for a in ex.get("assets") or []:
        if isinstance(a, dict) and a.get("key") in branding.ASSET_KEYS and a.get("data"):
            db.conn.execute("INSERT OR REPLACE INTO assets(key,mime,w,h,data) VALUES (?,?,?,?,?)",
                            (a["key"], a.get("mime") or "image/png", int(a.get("w") or 0), int(a.get("h") or 0), str(a["data"])))
    ids = {}
    for t in ex.get("offer_types") or []:
        if isinstance(t, dict) and t.get("name"):
            ids[t.get("id")] = db.conn.execute("INSERT INTO offer_types(name, content) VALUES (?,?)",
                                               (acc.text(t["name"]), str(t.get("content") or "{}"))).lastrowid
    n = 0
    for o in ex.get("offers") or []:
        if isinstance(o, dict) and o.get("number"):
            bu = db.one("SELECT id FROM business_units WHERE code=?", (o.get("bu_code") or "",))
            tid = ids.get(o.get("type_id"))
            if not tid and o.get("type_name"):     # e.g. a workbook's design offer: the company's own offer type, by name
                t = db.one("SELECT id FROM offer_types WHERE name=?", (acc.text(o["type_name"]),))
                tid = t["id"] if t else None
            if o.get("replace_same") and db.one("SELECT 1 FROM offers WHERE number=?", (acc.text(o["number"]),)):
                db.conn.execute("DELETE FROM offers WHERE number=?", (acc.text(o["number"]),))
            db.conn.execute("INSERT INTO offers(number,type_id,date,client,project,bu_id,status,content) VALUES (?,?,?,?,?,?,?,?)",
                            (acc.text(o["number"]), tid, o.get("date"), acc.text(o.get("client")),
                             acc.text(o.get("project")), bu["id"] if bu else None, o.get("status") or "Draft",
                             str(o.get("content") or "{}")))
            n += 1
    if ex.get("offer_types") or ex.get("offers"):
        counts["Offers"] = n
    if isinstance(ex.get("services"), dict) and (mode == "replace" or not db.one("SELECT 1 FROM subscriptions")):
        from . import services
        services.import_rows(db, ex["services"], counts)
    if ex.get("purchase_orders"):
        from . import purchasing
        if mode == "replace":
            db.conn.execute("DELETE FROM purchase_orders")
        counts["Purchase orders"] = purchasing.import_rows(db, ex["purchase_orders"])
    if isinstance(ex.get("site"), dict):
        site_extras_import(db, ex["site"], counts)
    if ex.get("attachments"):
        from . import attachments
        if mode == "replace":
            db.conn.execute("DELETE FROM attachments")
        counts["Attached files"] = attachments.import_rows(db, ex["attachments"])
    if isinstance(ex.get("dashboard_groups"), list) and ex["dashboard_groups"]:   # each group's dashboard
        for g in ex["dashboard_groups"]:
            if isinstance(g, dict) and g.get("key"):
                from . import dashboards
                look = g.get("look")
                look = json.loads(look) if isinstance(look, str) and look.strip() else look if isinstance(look, dict) else {}
                db.conn.execute("INSERT INTO user_groups(key,name,name_ar,layout,sort,look) VALUES (?,?,?,?,?,?) ON CONFLICT(key) DO UPDATE SET "
                                "name=excluded.name, name_ar=excluded.name_ar, layout=excluded.layout, sort=excluded.sort, look=excluded.look",
                                (acc.text(g["key"]), acc.text(g.get("name")) or g["key"], acc.text(g.get("name_ar")),
                                 json.dumps(dashboards._clean_layout(g.get("layout"))), int(g.get("sort") or 0),
                                 json.dumps(dashboards.clean_page(look))))
        counts["Dashboard groups"] = len(ex["dashboard_groups"])
    if isinstance(ex.get("crm"), dict) and any(ex["crm"].values()) and (mode == "replace" or not db.one("SELECT 1 FROM crm_deals")
                                                                        and not db.one("SELECT 1 FROM crm_activities")):
        from . import crm
        counts["Customers & deals"] = crm.import_rows(db, ex["crm"], replace=mode == "replace")
    if isinstance(ex.get("kb_pages"), list) and ex["kb_pages"] and (mode == "replace" or not db.one("SELECT 1 FROM kb_pages")):
        db.conn.execute("DELETE FROM kb_pages")
        cols = [r["name"] for r in db.q("PRAGMA table_info(kb_pages)")]
        for row in ex["kb_pages"]:
            keys = [k for k in cols if k in row]
            db.conn.execute(f"INSERT INTO kb_pages({','.join(keys)}) VALUES ({','.join('?' * len(keys))})", [row[k] for k in keys])
        counts["Knowledge base"] = len(ex["kb_pages"])
    if isinstance(ex.get("workdays"), list) and ex["workdays"] and (mode == "replace" or not db.one("SELECT 1 FROM workdays")):
        db.conn.execute("DELETE FROM workdays")
        cols = [r["name"] for r in db.q("PRAGMA table_info(workdays)")]
        for row in ex["workdays"]:
            keys = [k for k in cols if k in row]
            db.conn.execute(f"INSERT INTO workdays({','.join(keys)}) VALUES ({','.join('?' * len(keys))})", [row[k] for k in keys])
        counts["Working days"] = len(ex["workdays"])
    if isinstance(ex.get("automations"), dict) and ex["automations"].get("automations") and (mode == "replace" or not db.one("SELECT 1 FROM automations")):
        db.conn.execute("DELETE FROM automation_log")
        db.conn.execute("DELETE FROM automations")
        for t in ("automations", "automation_log"):
            cols = [r["name"] for r in db.q(f"PRAGMA table_info({t})")]
            for row in ex["automations"].get(t) or []:
                keys = [k for k in cols if k in row]
                db.conn.execute(f"INSERT INTO {t}({','.join(keys)}) VALUES ({','.join('?' * len(keys))})", [row[k] for k in keys])
        counts["Automations"] = len(ex["automations"]["automations"])
    if isinstance(ex.get("desk"), dict) and any(ex["desk"].values()):
        from . import desk
        n = desk.import_backup(db, ex["desk"], mode)
        if n:
            counts["Online desk"] = n
    if isinstance(ex.get("care"), dict) and any(ex["care"].values()):
        from . import care
        n = care.import_rows(db, ex["care"], mode)
        if n:
            counts["Healthcare (patients, visits, claims)"] = n
    if isinstance(ex.get("bu_groups"), dict):
        for code, grp in ex["bu_groups"].items():
            db.conn.execute("UPDATE business_units SET grp=? WHERE code=?", (acc.text(grp, maxlen=80), acc.text(code)))
    if isinstance(ex.get("entry_tags"), list) and ex["entry_tags"]:
        from . import attachments
        pool = {}
        for e in db.q("SELECT id FROM journal_entries WHERE status='posted' ORDER BY id"):
            fp = attachments._fp(db, e["id"])
            if fp:
                pool.setdefault(tuple(fp), []).append(e["id"])
        n = 0
        for r in ex["entry_tags"]:
            ids = pool.get(tuple(r.get("key") or ()))
            if ids and isinstance(r.get("tags"), list):
                for t in r["tags"][:12]:
                    db.conn.execute("INSERT OR IGNORE INTO entry_tags(entry_id, tag) VALUES (?,?)", (ids[0], acc.text(t, maxlen=40)))
                n += 1
        if n:
            counts["Tagged entries"] = n
    if isinstance(ex.get("line_styles"), dict):            # the look of BOQ / invoice lines
        n = _line_styles_import(db, ex["line_styles"])
        if n:
            counts["Styled lines"] = n
    if mode == "replace" and isinstance(ex.get("people"), list):    # before the chat, so everyone's messages find them
        n = _people_import(db, ex["people"])
        if n:
            counts["People (set their passwords in Users)"] = n
    if isinstance(ex.get("approvals"), list) and ex["approvals"] and (mode == "replace" or not db.one("SELECT 1 FROM approvals")):
        db.conn.execute("DELETE FROM approvals")
        n = _approvals_import(db, ex["approvals"])
        if n:
            counts["Approval requests"] = n
    if isinstance(ex.get("chat"), dict):
        from . import chat
        chat.import_rows(db, ex["chat"], mode, counts)
    if isinstance(ex.get("history"), dict):
        _history_import(db, ex["history"], mode, counts)
    if isinstance(ex.get("bills"), list) and ex["bills"] and (mode == "replace" or not db.one("SELECT 1 FROM cost_subs")):
        from . import billsubs
        n = billsubs.import_rows(db, ex["bills"], replace=mode == "replace")
        if n:
            counts["Subscriptions & bills"] = n
    if isinstance(ex.get("year_closes"), list) and (mode == "replace" or ex["year_closes"]):
        from . import yearend
        n = yearend.import_rows(db, ex["year_closes"], replace=mode == "replace")
        if n:
            counts["Closed financial years"] = n
    if isinstance(ex.get("secrets"), dict):
        _secrets_import(db, ex["secrets"], mode, counts)
    if ex.get("ownership_changes"):
        if mode == "replace":
            db.conn.execute("DELETE FROM ownership_changes")
        ids = {p["name"]: str(p["id"]) for p in db.q("SELECT id, name FROM partners")}
        for ch in ex["ownership_changes"]:
            if not isinstance(ch, dict):
                continue
            conv = {}
            for side in ("before", "after"):
                st = ch.get(side) or {}
                conv[side] = json.dumps({"shares": {ids[k]: v for k, v in (st.get("shares") or {}).items() if k in ids},
                                         "active": {ids[k]: v for k, v in (st.get("active") or {}).items() if k in ids},
                                         "dates": {ids[k]: v for k, v in (st.get("dates") or {}).items() if k in ids},
                                         "bu": [[ids[p], b, v] for p, b, v in (st.get("bu") or []) if p in ids]})
            db.conn.execute("INSERT INTO ownership_changes(date,kind,memo,before,after,created_by,created_at) "
                            "VALUES (?,?,?,?,?,?,?)", (acc.clean_date(ch.get("date")), acc.text(ch.get("kind")),
                                                      acc.text(ch.get("memo")), conv["before"], conv["after"],
                                                      acc.text(ch.get("created_by")), acc.text(ch.get("created_at"))))


def to_json(db, secrets=False):
    return json.dumps({"format": FORMAT, "version": FORMAT_VERSION,
                       "exported_at": _dt.datetime.now().isoformat(timespec="seconds"),
                       "sheets": export_data(db), "extras": export_extras(db, secrets=secrets)}, ensure_ascii=False, indent=1, default=str)


def to_csv(db, sheet_name, template=False):
    sheet = SHEET_BY_NAME[sheet_name]
    cols = input_cols(sheet)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"{c['en']} / {c['ar']}" for c in cols])
    if not template:
        for r in export_data(db)[sheet_name]:
            w.writerow([r.get(c["key"], "") for c in cols])
    return ("﻿" + buf.getvalue()).encode("utf-8")


WIPE_ORDER = ("year_closes", "cost_subs", "entry_tags", "kb_pages", "workdays", "automation_log", "automations", "desk_rows", "desk_threads", "care_visit_lines", "care_visits", "care_claims", "care_packages", "care_services", "care_patients", "crm_followups", "crm_activities", "crm_deals", "sub_cert_lines", "sub_certificates", "variation_orders", "contract_lines", "contracts", "pr_quote_lines", "pr_quotes",
              "pr_lines", "purchase_requests", "stock_count_lines", "stock_counts", "equipment_usage", "maintenance_logs", "maintenance_plans",
              "equipment", "attendance", "ereceipts", "project_tasks", "project_phases", "payroll_lines", "payroll_runs", "crew_sheets", "employees", "sub_recognitions", "sub_visits", "sub_invoices", "subscriptions", "service_plans", "attachments", "po_receipts", "po_lines", "purchase_orders", "budgets", "bank_lines", "bank_statements", "bank_rules", "recurring_runs", "recurring", "sale_cheques", "sale_installments", "sale_lines", "sales_orders", "product_media", "products", "rent_moves", "rental_invoices", "rental_charges", "rental_handovers", "rental_lines", "rentals", "rent_lots",
              "rent_kit_lines", "rent_kits", "rent_models", "rent_items", "loans", "custodies", "parties", "ownership_changes", "fixed_assets", "stock_moves", "stock_levels", "stock_items", "extract_lines", "extracts", "boq_lines", "boqs",
              "journal_lines", "journal_entries", "partner_shares", "partners", "rates", "business_units")


# "New rows only": what identifies a record that is already in the books. Rows of these sheets are skipped when their
# key is there (so changes made in the app are never overwritten); several rows sharing a key (an invoice and its
# lines, an order and its items) come in or stay out together. Other sheets compare the whole row.
ONLY_NEW_KEYS = {
    "Chart of Accounts": ("code",), "Tax Codes": ("code",), "Business Units": ("code",), "Parties": ("name",), "Partners": ("name",),
    "Fixed Assets": ("code",), "Warehouses": ("name",), "Stock Items": ("name",), "Products": ("code",), "Employees": ("code",),
    "Loans": ("code",), "Cash Custody": ("code",), "Rental Equipment": ("code",), "Rental Contracts": ("ref",), "Sales Contracts": ("code",),
    "Service Plans": ("code",), "Service Subscriptions": ("code",), "Purchase Orders": ("code",), "Contracts": ("code",),
    "Purchase Requests": ("code",), "Sales Invoices": ("number",), "Invoice Lines": ("number",), "E-Receipts": ("number",),
    "Stock Counts": ("code",), "BOQ Lines": ("title",), "Subcontractor Certificates": ("contract", "number"),
    "Equipment": ("code",), "Variation Orders": ("contract", "code"), "Recurring Entries": ("name",), "Settings": ("key",),
}


def _norm_cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return f"{round(float(v), 2):.2f}"
    if isinstance(v, (_dt.date, _dt.datetime)):
        return v.isoformat()[:10]
    t = re.sub(r"\s+", " ", str(v)).strip().lower()
    try:
        return f"{round(float(t.replace(',', '')), 2):.2f}"
    except ValueError:
        pass
    return t[:10] if re.match(r"^\d{4}-\d{2}-\d{2}", t) else t


WHOLE_DOCS = ("BOQ Lines", "Sales Invoices", "Invoice Lines")


def only_new(db, data):
    """Keep only what is not in the books yet; returns (data, {sheet: {"new": n, "skipped": n}})."""
    have, out, report = export_data(db), {}, {}
    for name, rows in data.items():
        if name.startswith("_") or name not in SHEET_BY_NAME or not rows or name in WHOLE_DOCS:
            out[name] = rows             # BOQs and invoices: compared as whole documents when imported
            continue
        cols = [c["key"] for c in input_cols(SHEET_BY_NAME[name])]
        existing = have.get(name) or []
        if name in ONLY_NEW_KEYS:
            kc = ONLY_NEW_KEYS[name]
            key = lambda r: tuple(_norm_cell(r.get(k)) for k in kc)  # noqa: E731
            if name == "Settings":
                key = lambda r: (_norm_cell(str(r.get("key") or "").replace("default:", "")),)  # noqa: E731
            known = {key(r) for r in existing}
            keep = [r for r in rows if key(r) not in known]
        elif name == "Journal Vouchers":     # a voucher is known when an entry has the same date and the same lines
            def groups(rs):
                g = {}
                for r in rs:
                    g.setdefault(_norm_cell(r.get("voucher")) or id(r), []).append(r)
                return g
            fp = lambda rs: (_norm_cell(rs[0].get("date")), tuple(sorted((_norm_cell(r.get("account")), _norm_cell(r.get("debit")),  # noqa: E731
                                                                          _norm_cell(r.get("credit"))) for r in rs)))
            pool = {}
            for rs in groups(existing).values():
                pool[fp(rs)] = pool.get(fp(rs), 0) + 1
            keep = []
            for rs in groups(rows).values():
                f = fp(rs)
                if pool.get(f):
                    pool[f] -= 1
                else:
                    keep += rs
        else:                                # same content, counted: two equal rows in the file against one in the books → one new
            pool = {}
            for r in existing:
                f = tuple(_norm_cell(r.get(k)) for k in cols)
                pool[f] = pool.get(f, 0) + 1
            keep = []
            for r in rows:
                f = tuple(_norm_cell(r.get(k)) for k in cols)
                if pool.get(f):
                    pool[f] -= 1
                else:
                    keep.append(r)
        out[name] = keep
        report[name] = {"new": len(keep), "skipped": len(rows) - len(keep)}
    return out, report


def import_data(db, data, mode="append", dry_run=False, user=""):
    """Import sheets. All-or-nothing: any row error rolls everything back and is reported.
    mode "new": only rows that are not in the books yet (see only_new), then as "append"."""
    if mode not in ("append", "replace", "new"):
        raise acc.ValidationError("mode must be append, new or replace")
    errors, counts, stats = [], {}, {}
    skipped = None
    if mode == "new":
        data, skipped = only_new(db, data)
        mode = "append"

    def run(name, fn):
        rows = data.get(name) or []
        n = 0
        for i, row in enumerate(rows, start=FIRST_DATA_ROW):
            sheet = SHEET_BY_NAME[name]
            missing = [c["en"] for c in sheet["cols"] if c["required"] and row.get(c["key"]) in ("", None)]
            if missing:
                errors.append({"sheet": name, "row": i, "error": "missing " + ", ".join(missing)})
                continue
            try:
                fn(row)
                n += 1
            except (acc.ValidationError, LookupError) as e:
                errors.append({"sheet": name, "row": i, "error": str(e)})
        if rows:
            counts[name] = n

    class _Rollback(Exception):
        pass

    try:
        with db.tx():
            if mode == "replace":
                for t in WIPE_ORDER:
                    db.conn.execute(f"DELETE FROM {t}")
                db.conn.execute("DELETE FROM warehouses WHERE id NOT IN (SELECT MIN(id) FROM warehouses)")
            # the lock date is applied last, so the entries it protects can be restored first
            lock = [r for r in (data.get("Settings") or []) if str(r.get("key")) in ("lock_date", "default:lock_date")]
            data["Settings"] = [r for r in (data.get("Settings") or []) if r not in lock]
            saved_lock = db.setting("lock_date") or ""
            if mode == "replace":  # a restore; appending into closed periods stays refused
                db.conn.execute("UPDATE settings SET value='' WHERE key='lock_date'")
            run("Settings", lambda r: _imp_setting(db, r))
            run("Chart of Accounts", lambda r: _imp_account(db, r))
            run("Tax Codes", lambda r: _imp_tax_code(db, r))
            run("Business Units", lambda r: _imp_bu(db, r))
            run("Parties", lambda r: _imp_party(db, r))
            run("Partners", lambda r: _imp_partner(db, r))
            run("Partner Shares", lambda r: _imp_share(db, r))
            run("Fixed Assets", lambda r: _imp_asset(db, r))
            run("Warehouses", lambda r: _imp_wh(db, r))
            run("Stock Items", lambda r: _imp_item(db, r))
            run("Price List", lambda r: _imp_rate(db, r))
            run("Products", lambda r: _imp_product(db, r))
            _imp_boqs(db, data.get("BOQ Lines") or [], errors, counts, stats)
            run("Direct Costs", lambda r: acc.record_project_cost(db, r, user))
            run("Operating Expenses", lambda r: acc.record_expense(db, r, user))
            run("Payroll", lambda r: acc.record_payroll(db, r, user))
            run("Receipt Vouchers", lambda r: acc.record_collection(db, r, user))
            _imp_invoices(db, data, errors, counts, user, stats)
            run("Partner Transactions", lambda r: acc.record_partner_tx(db, {**r, "type": "Capital" if str(
                r.get("type", "")).startswith("Capital") else "Distribution"}, user))
            data["Stock Moves"] = sorted(data.get("Stock Moves") or [],
                                         key=lambda r: (str(r.get("date") or ""), r.get("type") != "Receipt"))
            run("Stock Moves", lambda r: _imp_stock(db, r, user))
            run("Payment Vouchers", lambda r: acc.record_payment(db, {**r, "allow_overpay": True}, user))
            _imp_journal(db, data.get("Journal Vouchers") or [], errors, counts, user)
            # modules built on their own screens: rows create new records through the module itself
            rows_fn = dict(_setup.ROW_IMPORTERS)
            groups = {n: (k, f) for n, k, f in _setup.GROUP_IMPORTERS}
            pays = data.get("Subcontract Payments") or []
            for name in _setup.ORDER:
                if name == "Subcontractor Certificates":   # advances first: certificates recover them as they are posted
                    for i, r in enumerate(pays, start=FIRST_DATA_ROW):
                        if str(r.get("kind") or "").lower() == "advance":
                            try:
                                r["_done"] = True
                                _setup.sub_payment(db, r, user)
                            except (acc.ValidationError, LookupError) as e:
                                errors.append({"sheet": "Subcontract Payments", "row": i, "error": str(e)})
                if name == "Subcontract Payments":
                    fn = rows_fn[name]
                    run(name, lambda r, fn=fn: None if r.get("_done") else fn(db, r, user))
                    continue
                if name in rows_fn:
                    fn = rows_fn[name]
                    run(name, lambda r, fn=fn: fn(db, r, user))
                    continue
                key, fn = groups[name]
                rows = data.get(name) or []
                sheet = SHEET_BY_NAME[name]
                ok_rows = []
                for i, row in enumerate(rows, start=FIRST_DATA_ROW):
                    missing = [c["en"] for c in sheet["cols"] if c["required"] and row.get(c["key"]) in ("", None)]
                    if missing:
                        errors.append({"sheet": name, "row": i, "error": "missing " + ", ".join(missing)})
                    else:
                        ok_rows.append((i, row))
                n = 0
                for _k, items in _setup._groups([r for _, r in ok_rows], key):
                    first = ok_rows[[r for _, r in ok_rows].index(items[0][1])][0]
                    try:
                        fn(db, items, user)
                        n += 1
                    except (acc.ValidationError, LookupError) as e:
                        errors.append({"sheet": name, "row": first, "error": str(e)})
                if rows:
                    counts[name] = n
            if isinstance(data.get("_extras"), dict):
                _imp_extras(db, data["_extras"], mode, counts)
            from . import wip      # every project's work in progress, from what was just imported
            wip.sync_all(db)
            db.conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES ('lock_date', ?)",
                            (acc.text(lock[-1].get("value")) if lock else saved_lock,))
            if errors or dry_run:
                raise _Rollback()
    except _Rollback:
        pass
    if errors:
        raise ImportErrors(errors[:200])
    out = {"counts": counts, "dry_run": dry_run, "mode": "new" if skipped is not None else mode}
    if skipped is not None:
        out["skipped"] = {k: v["skipped"] for k, v in skipped.items() if v["skipped"]}
    for k, v in stats.get("skipped", {}).items():          # whole documents found unchanged
        out.setdefault("skipped", {})[k] = out.get("skipped", {}).get(k, 0) + v
    if stats.get("updated"):
        out["updated"] = stats["updated"]
    return out


def _imp_setting(db, r):
    key = acc.text(r.get("key"))
    if key in ("schema_version", "docs_seeded", *ETA_SECRETS):
        return
    if (key, acc.text(r.get("value"))) in (("company_name", "Company Name"), ("company_tagline", "Company Slogan")):
        return                      # a universal template's placeholders never replace the company's own details
    if key.startswith("default:"):  # fill only when empty
        key = key[8:]
        if (db.setting(key) or "").strip():
            return
    db.conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES (?,?)", (key, acc.text(r.get("value"), maxlen=20000)))


def _imp_account(db, r):
    code = acc.text(r.get("code"))
    existing = db.one("SELECT * FROM accounts WHERE code=?", (code,))
    role = acc.text(r.get("role")) or None
    by_role = db.one("SELECT * FROM accounts WHERE role=?", (role,)) if role else None
    d = {"code": code, "name": r.get("name"), "name_ar": r.get("name_ar"), "type": r.get("type"), "subtype": r.get("subtype"),
         "active": 0 if str(r.get("active", "Yes")).lower() in ("no", "0", "false") else 1}
    if by_role and by_role["code"] != code:
        acc.save_account(db, d, old_code=by_role["code"])  # renumbered system account
    elif existing:
        acc.save_account(db, d, old_code=code)
    else:
        acc.save_account(db, d)
        if role:
            db.conn.execute("UPDATE accounts SET role=? WHERE code=?", (role, code))


def _imp_tax_code(db, r):
    code = acc.text(r.get("code")).upper()
    old = code if db.one("SELECT 1 FROM tax_codes WHERE code=?", (code,)) else None
    acc.save_tax_code(db, {**r, "code": code, "rate": _pct(r.get("rate"))}, old_code=old)


def _imp_party(db, r):
    p = db.one("SELECT id FROM parties WHERE name=?", (acc.text(r.get("name")),))
    acc.save_party(db, r, p["id"] if p else None)


def _pct(v):
    v = acc.num(v, "percentage")
    return v / 100 if v > 1 else v


def _imp_bu(db, r):
    vals = (acc.text(r.get("name"), "name", required=True), acc.text(r.get("type")), acc.text(r.get("location")),
            acc.num(r.get("area"), "area"), acc.text(r.get("start_date")), acc.text(r.get("status")) or "In Progress",
            r.get("basis") if r.get("basis") in acc.BILLING_BASES else "Cost-Plus %", _pct(r.get("fee_pct")),
            to_minor(acc.num(r.get("contract_value"), "contract value")), _pct(r.get("overhead_pct")),
            acc.text(r.get("notes"), maxlen=4000), acc.text(r.get("customer")),
            acc.text(r.get("negotiation")).lower() if acc.text(r.get("negotiation")).lower() in acc.NEGOTIATION_MODES else "",
            acc.text(r.get("wip_mode")).lower() if acc.text(r.get("wip_mode")).lower() in ("cost_plus", "drafts") else "")
    code = acc.text(r.get("code"), "code", required=True)
    if db.one("SELECT 1 FROM business_units WHERE code=?", (code,)):
        db.conn.execute("UPDATE business_units SET name=?,type=?,location=?,area=?,start_date=?,status=?,basis=?,"
                        "fee_pct=?,contract_value=?,overhead_pct=?,notes=?,customer=?,negotiation=?,wip_mode=? WHERE code=?", vals + (code,))
    else:
        db.conn.execute("INSERT INTO business_units(name,type,location,area,start_date,status,basis,fee_pct,"
                        "contract_value,overhead_pct,notes,customer,negotiation,wip_mode,code) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        vals + (code,))


def _imp_partner(db, r):
    active = 0 if str(r.get("active") or "").strip().lower() in ("no", "0", "false") else 1
    db.conn.execute("INSERT INTO partners(name,default_share,active,joined,left_on,capital_interest,loan_rate,notes) "
                    "VALUES (?,?,?,?,?,?,?,?) ON CONFLICT(name) DO UPDATE SET default_share=excluded.default_share, "
                    "active=excluded.active, joined=excluded.joined, left_on=excluded.left_on, "
                    "capital_interest=excluded.capital_interest, loan_rate=excluded.loan_rate, notes=excluded.notes",
                    (acc.text(r.get("name")), _pct(r.get("default_share")), active,
                     acc.clean_date(r["joined"]) if r.get("joined") else "",
                     acc.clean_date(r["left_on"]) if r.get("left_on") else "", _pct(r.get("capital_interest")),
                     _pct(r.get("loan_rate")), acc.text(r.get("notes"), maxlen=2000)))
    if any(k in r for k in ("profit_share", "salary", "drawings_interest")):
        db.conn.execute("UPDATE partners SET profit_share=?, salary=?, drawings_interest=? WHERE name=?",
                        (None if str(r.get("profit_share") if r.get("profit_share") is not None else "").strip() == "" else _pct(r.get("profit_share")),
                         to_minor(acc.num(r.get("salary"), "salary") if str(r.get("salary") or "").strip() else 0), _pct(r.get("drawings_interest")), acc.text(r.get("name"))))


def _imp_share(db, r):
    p = db.one("SELECT id FROM partners WHERE name=?", (acc.text(r.get("partner")),))
    if not p:
        raise acc.ValidationError(f"Unknown partner {r.get('partner')}")
    db.conn.execute("INSERT OR REPLACE INTO partner_shares(partner_id,bu_id,share) VALUES (?,?,?)",
                    (p["id"], acc.bu_id_or_none(db, r.get("bu_code"), required=True), _pct(r.get("share"))))


def _imp_asset(db, r):
    from . import assets
    old = db.one("SELECT id FROM fixed_assets WHERE code=?", (acc.text(r.get("code")),))
    assets.save(db, {**r, "pay": "none"}, old["id"] if old else None)


def _imp_product(db, r):
    p = db.one("SELECT id, status FROM products WHERE code=?", (acc.text(r.get("code")),)) if r.get("code") else None
    specs = [[a.strip(), b.strip()] for a, _, b in (x.partition(":") for x in str(r.get("specs") or "").split(";")) if a.strip()]
    d = {**r, "specs": specs, "kind": (str(r.get("kind") or "").lower() or "item")}
    d["status"] = str(r.get("status") or "").lower() or (p["status"] if p else "available")
    sales.save_product(db, d, p["id"] if p else None)


def _imp_wh(db, r):
    account = acc.account_code(db, r["account"], types=("Asset",)) if r.get("account") else None
    db.conn.execute("INSERT INTO warehouses(name,location,account) VALUES (?,?,?) ON CONFLICT(name) DO UPDATE SET "
                    "location=excluded.location, account=excluded.account",
                    (acc.text(r.get("name")), acc.text(r.get("location")), account))


def _imp_item(db, r):
    db.conn.execute("INSERT INTO stock_items(name,unit,sku,reorder_level,category,std_cost,price,notes) VALUES (?,?,?,?,?,?,?,?) "
                    "ON CONFLICT(name) DO UPDATE SET unit=excluded.unit, sku=excluded.sku, reorder_level=excluded.reorder_level, "
                    "category=excluded.category, std_cost=excluded.std_cost, price=excluded.price, notes=excluded.notes",
                    (acc.text(r.get("name")), acc.text(r.get("unit")) or "pcs", acc.text(r.get("sku")),
                     acc.num(r.get("reorder_level"), "reorder level"), acc.text(r.get("category")),
                     to_minor(acc.num(r.get("std_cost"), "standard cost")), to_minor(acc.num(r.get("price"), "selling price")),
                     acc.text(r.get("notes"), maxlen=1000)))


def _imp_rate(db, r):
    cost, mk = acc.num(r.get("cost"), "cost"), _pct(r.get("markup"))
    price = acc.num(r.get("price"), "price") or cost * (1 + mk)
    db.conn.execute("INSERT INTO rates(category,item_no,sub,description,unit,cost,markup,price,notes,eta_code) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (acc.text(r.get("category")), acc.text(r.get("item_no")), acc.text(r.get("sub")),
                     acc.text(r.get("description"), maxlen=2000), acc.text(r.get("unit")), to_minor(cost), mk,
                     to_minor(price), acc.text(r.get("notes")), acc.text(r.get("eta_code"))))


# BOQs and invoices are whole documents: one already in the books (a BOQ with the same title and business unit, an
# invoice with the same number) is compared as a whole — unchanged it is skipped, changed it is updated in place
# (a posted invoice is reopened: its entry voided and kept in the audit trail, then posted again with the new lines)
_LINE_COLS = ("category", "sub", "item_no", "description", "unit", "qty", "price")


def _same_lines(file_lines, db_rows, cols=_LINE_COLS):
    a = [tuple(ln[c] for c in cols) for ln in acc._lines_in(file_lines)]
    b = [tuple(r[c] for c in cols) for r in db_rows]
    norm = lambda rows: [tuple(round(v, 6) if isinstance(v, float) else (v or "") for v in t) for t in rows]  # noqa: E731
    return norm(a) == norm(b)


def _imp_boqs(db, rows, errors, counts, stats):
    groups = {}
    for i, r in enumerate(rows, start=FIRST_DATA_ROW):
        groups.setdefault((acc.text(r.get("title")), acc.text(r.get("bu_code"))), []).append((i, r))
    n = 0
    for (title, bu), items in groups.items():
        first = items[0][1]
        d = {"title": title, "bu_code": bu, "rev": first.get("rev"), "date": first.get("date"), "client": first.get("client"),
             "location": first.get("location"), "lines": [r for _, r in items]}
        try:
            bu_id = acc.bu_id_or_none(db, bu) if bu else None
            old = db.one("SELECT * FROM boqs WHERE title=? AND " + ("bu_id=?" if bu_id else "bu_id IS NULL") + " ORDER BY id DESC",
                         (title, bu_id) if bu_id else (title,))
            if old:
                same_head = all(acc.text(d.get(k)) in ("", acc.text(old.get(k))) for k in ("rev", "client", "location")) and \
                    (not d.get("date") or acc.clean_date(d["date"]) == old["date"])
                if same_head and _same_lines(d["lines"], db.q("SELECT * FROM boq_lines WHERE boq_id=? ORDER BY seq, id", (old["id"],))):
                    stats.setdefault("skipped", {}).setdefault("BOQ Lines", 0)
                    stats["skipped"]["BOQ Lines"] += 1
                    continue
                for k in ("rev", "client", "location", "date"):     # what the file leaves blank stays as it was
                    if d.get(k) in ("", None):
                        d[k] = old.get(k)
                acc.save_boq(db, d, boq_id=old["id"])
                stats.setdefault("updated", {}).setdefault("BOQ Lines", 0)
                stats["updated"]["BOQ Lines"] += 1
            else:
                acc.save_boq(db, d)
            n += 1
        except (acc.ValidationError, LookupError) as e:
            errors.append({"sheet": "BOQ Lines", "row": items[0][0], "error": str(e)})
    if rows:
        counts["BOQ Lines"] = n


def _imp_invoices(db, data, errors, counts, user, stats):
    lines = {}
    for r in data.get("Invoice Lines") or []:
        lines.setdefault(acc.text(r.get("number")), []).append(r)
    n = 0
    heads = {acc.text(r.get("number")) for r in data.get("Sales Invoices") or []}
    for num_ in [k for k in lines if k and k not in heads]:   # only the lines of an invoice already in the books
        if not db.one("SELECT 1 FROM extracts WHERE number=?", (num_,)):
            errors.append({"sheet": "Invoice Lines", "row": FIRST_DATA_ROW, "error": f"Invoice {num_} is not in 'Sales Invoices' or in the books"})
    rows = list(data.get("Sales Invoices") or []) + [{"number": k, "_lines_only": True} for k in lines if k and k not in heads
                                                      and db.one("SELECT 1 FROM extracts WHERE number=?", (k,))]
    for i, r in enumerate(rows, start=FIRST_DATA_ROW):
        sheet = "Invoice Lines" if r.get("_lines_only") else "Sales Invoices"
        try:
            number = acc.text(r.get("number"), "invoice number", required=True)
            old = db.one("SELECT * FROM extracts WHERE number=?", (number,))
            if r.get("_lines_only"):      # take everything else from the invoice as it is
                r = {"number": number, "date": old["date"], "bu_code": (db.one("SELECT code FROM business_units WHERE id=?", (old["bu_id"],)) or {}).get("code"),
                     "customer": old["customer"], "vat_code": old["vat_code"], "wht_code": old["wht_code"], "extract_no": old["seq_no"],
                     "project": old["project"], "work": old["work"], "notes": old["notes"], "final": old["final"],
                     "status": "Posted" if old["status"] == "posted" else "Draft", "_boq_id": old["boq_id"],
                     "prev_paid": None if old["prev_paid"] is None else acc.to_major(old["prev_paid"]),
                     "deductions": acc.to_major(old["deductions"] or 0)}
            boq_id = r.get("_boq_id")
            if r.get("boq_title"):
                b = db.one("SELECT id FROM boqs WHERE title=? ORDER BY id DESC", (acc.text(r.get("boq_title")),))
                boq_id = b["id"] if b else None
            if not lines.get(number):
                if old:                    # header only: the lines stay as they are
                    lines[number] = [dict(x, price=acc.to_major(x["price"])) for x in db.q("SELECT * FROM extract_lines WHERE extract_id=? ORDER BY seq, id", (old["id"],))]
                else:
                    raise acc.ValidationError(f"Invoice {number} has no rows in 'Invoice Lines'")
            d = {"number": number, "date": r.get("date") or (old or {}).get("date"), "bu_code": r.get("bu_code"),
                 "boq_id": boq_id if boq_id or not old else old["boq_id"], "lines": lines[number],
                 "customer": r.get("customer"), "vat_code": r.get("vat_code"), "wht_code": r.get("wht_code"), "seq_no": r.get("extract_no"),
                 "project": r.get("project"), "work": r.get("work"), "prev_paid": r.get("prev_paid"), "deductions": r.get("deductions"),
                 "final": r.get("final"), "notes": r.get("notes")}
            want_posted = str(r.get("status") or "Posted").lower() != "draft"
            if old:                        # what the file leaves blank stays as it was
                for k, ok_ in (("customer", "customer"), ("vat_code", "vat_code"), ("wht_code", "wht_code"), ("seq_no", "seq_no"),
                               ("project", "project"), ("work", "work"), ("notes", "notes")):
                    if d.get(k) in ("", None):
                        d[k] = old[ok_]
                if not d.get("bu_code"):
                    d["bu_id"] = old["bu_id"]
            if old:
                cols = _LINE_COLS + ("account", "pct")
                same = _same_lines(d["lines"], db.q("SELECT * FROM extract_lines WHERE extract_id=? ORDER BY seq, id", (old["id"],)), cols) \
                    and acc.clean_date(d["date"]) == old["date"] and acc.text(d.get("customer")) in ("", old["customer"] or "") \
                    and (old["status"] == "posted") == want_posted
                if same:
                    stats.setdefault("skipped", {}).setdefault(sheet, 0)
                    stats["skipped"][sheet] += 1
                    continue
                if old["status"] == "posted":
                    st = str(old.get("eta_status") or "").lower()
                    if old.get("eta_uuid") and st and not any(w in st for w in ("cancel", "reject", "invalid")):
                        raise acc.ValidationError(f"Invoice {number} was sent to the tax authority — cancel the e-invoice first, then import the change")
                    if old["entry_id"]:
                        acc.void_entry(db, old["entry_id"], "Updated from an imported file", user)
                xid = acc.save_extract(db, d, extract_id=old["id"])
                stats.setdefault("updated", {}).setdefault(sheet, 0)
                stats["updated"][sheet] += 1
            else:
                xid = acc.save_extract(db, d)
            if want_posted:
                acc.post_extract(db, xid, user=user, ar_account=r.get("ar_account") or None,
                                 revenue_account=r.get("revenue_account") or None)
            n += 1
        except (acc.ValidationError, LookupError) as e:
            errors.append({"sheet": sheet, "row": i, "error": str(e)})
    if rows:
        counts["Sales Invoices"] = n


def _imp_stock(db, r, user):
    typ = r.get("type")
    d = dict(r)
    if typ == "Adjustment":
        wh = acc._warehouse(db, r.get("warehouse"))
        item = acc._item(db, r)
        d["counted"] = acc._level(db, item["id"], wh["id"])["qty"] + acc.num(r.get("qty"), "qty", allow_negative=True)
    acc.stock_move(db, d, user)


def _imp_journal(db, rows, errors, counts, user):
    groups = {}
    for i, r in enumerate(rows, start=FIRST_DATA_ROW):
        groups.setdefault(acc.text(r.get("voucher")), []).append((i, r))
    n = 0
    for v, items in groups.items():
        first = items[0][1]
        try:
            acc.record_journal(db, {"date": first.get("date"), "kind": first.get("kind") or "Journal Voucher",
                                    "bu_code": first.get("bu_code"), "party": first.get("party"),
                                    "ref": first.get("ref") or ("" if re.fullmatch(r"JV-\d+", v) else v), "memo": first.get("memo"),
                                    "deal_ref": first.get("deal_ref"), "quoted": first.get("quoted"), "savings": first.get("savings"),
                                    "negotiation": first.get("negotiation"),
                                    "items": [{"item": r.get("item"), "unit": r.get("item_unit"), "qty": r.get("item_qty"),
                                               "rate": r.get("item_rate"), "warehouse": r.get("item_to") or None,
                                               "account": r.get("account") if not r.get("item_to") else None}
                                              for _, r in items if acc.text(r.get("item"))],
                                    "lines": [{"account": r.get("account"), "debit": r.get("debit"),
                                               "credit": r.get("credit"), "party": r.get("line_party")}
                                              for _, r in items if not acc.text(r.get("item"))]}, user)
            n += 1
        except acc.ValidationError as e:
            errors.append({"sheet": "Journal Vouchers", "row": items[0][0], "error": f"voucher {v}: {e}"})
    if rows:
        counts["Journal Vouchers"] = n
