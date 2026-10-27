"""Financial statements as printable tables + renderers (Excel, Word, HTML→PDF, CSV).

Every report is first built as a neutral "table" dict:
    {"title", "subtitle", "columns": [{"label", "align", "fmt"}],
     "rows": [{"cells": [...], "style": "item|header|subtotal|total|note", "indent": 0}]}
and then rendered by one of the writers. The UI shows the same table, so what
you see on screen is exactly what gets exported.
"""
import datetime as _dt
import html
import io
import zipfile

from . import accounting as acc, xlsafe

# ----------------------------------------------------------------- i18n labels
T = {
    "income_statement": ("Income Statement (Profit & Loss)", "قائمة الدخل (الأرباح والخسائر)"),
    "balance_sheet": ("Balance Sheet (Statement of Financial Position)", "الميزانية العمومية (قائمة المركز المالي)"),
    "cash_flow": ("Cash Flow Statement", "قائمة التدفقات النقدية"),
    "trial_balance": ("Trial Balance", "ميزان المراجعة"),
    "general_ledger": ("General Ledger", "دفتر الأستاذ العام"),
    "bu_pnl": ("Profitability by Business Unit", "ربحية وحدات الأعمال"),
    "project_cash": ("Cash In & Out by Project", "النقدية الواردة والمنصرفة حسب المشروع"),
    "payables": ("Accounts Payable by Party", "أرصدة الدائنين"),
    "receivables": ("Accounts Receivable by Business Unit", "أرصدة المدينين"),
    "stock": ("Inventory Valuation", "تقييم المخزون"),
    "vouchers": ("Vouchers", "السندات والقيود"),
    "journal": ("Journal", "دفتر اليومية"),
    "revenue": ("Revenue", "الإيرادات"),
    "gross_sales": ("Gross sales", "إجمالي المبيعات"),
    "discounts": ("Sales discounts & allowances", "خصومات ومسموحات المبيعات"),
    "net_sales": ("Net sales", "صافي المبيعات"),
    "cost_of_sales": ("Cost of sales", "تكلفة المبيعات"),
    "gross_profit": ("Gross profit", "مجمل الربح"),
    "operating_expenses": ("Operating expenses", "المصروفات التشغيلية"),
    "Employee expenses": ("Employee expenses", "مصروفات العاملين"),
    "General & administrative expenses": ("General & administrative expenses", "مصروفات عمومية وإدارية"),
    "Marketing & advertising expenses": ("Marketing & advertising expenses", "مصروفات تسويق وإعلان"),
    "Utilities (electricity / telephone)": ("Utilities (electricity / telephone)", "مرافق (كهرباء / اتصالات)"),
    "Other operating expenses": ("Other operating expenses", "مصروفات تشغيلية أخرى"),
    "Allocated shared overhead": ("Allocated shared overhead", "حصة المصروفات المشتركة المحملة"),
    "ebitda": ("EBITDA", "الأرباح قبل الفوائد والضرائب والإهلاك والاستهلاك (EBITDA)"),
    "depreciation": ("Depreciation", "الإهلاك"),
    "amortization": ("Amortization", "الاستهلاك (الأصول غير الملموسة)"),
    "d_and_a": ("Depreciation & amortization", "الإهلاك والاستهلاك"),
    "ebit": ("EBIT / Operating profit", "الربح التشغيلي (EBIT)"),
    "interest": ("Interest expense", "مصروفات الفوائد"),
    "other_nonoperating": ("Other non-operating expenses", "مصروفات غير تشغيلية أخرى"),
    "other_income": ("Other income", "إيرادات أخرى"),
    "ebt": ("EBT / Profit before tax", "الربح قبل الضرائب"),
    "income_tax": ("Income tax", "ضريبة الدخل"),
    "net_profit": ("Net profit / (loss)", "صافي الربح / (الخسارة)"),
    "margin": ("margin", "هامش"),
    "assets": ("ASSETS", "الأصول"),
    "current_assets": ("Current assets", "الأصول المتداولة"),
    "noncurrent_assets": ("Non-current assets", "الأصول غير المتداولة"),
    "total_assets": ("TOTAL ASSETS", "إجمالي الأصول"),
    "liabilities": ("LIABILITIES", "الالتزامات"),
    "current_liabilities": ("Current liabilities", "الالتزامات المتداولة"),
    "noncurrent_liabilities": ("Non-current liabilities", "الالتزامات غير المتداولة"),
    "total_liabilities": ("TOTAL LIABILITIES", "إجمالي الالتزامات"),
    "equity": ("EQUITY", "حقوق الملكية"),
    "total_equity": ("TOTAL EQUITY", "إجمالي حقوق الملكية"),
    "total_le": ("TOTAL LIABILITIES & EQUITY", "إجمالي الالتزامات وحقوق الملكية"),
    "cf_ops": ("CASH FLOWS FROM OPERATING ACTIVITIES", "التدفقات النقدية من الأنشطة التشغيلية"),
    "net_income": ("Net income", "صافي الدخل"),
    "adjustments": ("Adjustments for non-cash charges", "تسويات البنود غير النقدية"),
    "wc_changes": ("Changes in operating assets and liabilities", "التغير في الأصول والالتزامات التشغيلية"),
    "net_ops": ("Net cash provided by operating activities", "صافي النقد من الأنشطة التشغيلية"),
    "cf_inv": ("CASH FLOWS FROM INVESTING ACTIVITIES", "التدفقات النقدية من الأنشطة الاستثمارية"),
    "net_inv": ("Net cash used in investing activities", "صافي النقد من الأنشطة الاستثمارية"),
    "cf_fin": ("CASH FLOWS FROM FINANCING ACTIVITIES", "التدفقات النقدية من الأنشطة التمويلية"),
    "net_fin": ("Net cash provided by financing activities", "صافي النقد من الأنشطة التمويلية"),
    "beginning_cash": ("Beginning cash", "النقدية أول المدة"),
    "change_cash": ("(+/-) Change in cash & cash equivalents", "صافي التغير في النقدية"),
    "ending_cash": ("Ending cash", "النقدية آخر المدة"),
    "code": ("Code", "الكود"), "account": ("Account", "الحساب"), "debit": ("Debit", "مدين"),
    "credit": ("Credit", "دائن"), "total": ("TOTAL", "الإجمالي"), "opening": ("Opening", "أول المدة"),
    "movement": ("Movement", "الحركة"), "closing": ("Closing", "آخر المدة"), "balance": ("Balance", "الرصيد"),
    "date": ("Date", "التاريخ"), "voucher": ("Voucher", "القيد"), "description": ("Description", "البيان"),
    "party": ("Party", "الطرف"), "bu": ("Business unit", "وحدة الأعمال"),
    "as_of": ("As of", "في"), "period": ("Period", "الفترة"), "to": ("to", "إلى"),
    "all_units": ("All business units", "كل وحدات الأعمال"), "shared_only": ("Shared / company level", "على مستوى الشركة"),
    "balanced_ok": ("Balanced ✓", "متوازن ✓"), "unbalanced": ("OUT OF BALANCE", "غير متوازن"),
}
CF_AR = {"Depreciation": "الإهلاك", "Amortization": "الاستهلاك", "Accounts receivable": "المدينون",
         "Inventory": "المخزون", "Prepaid expenses & advances": "المصروفات المقدمة والدفعات",
         "Other current assets": "أصول متداولة أخرى", "Accounts payable": "الدائنون",
         "Accrued liabilities & payroll": "المستحقات والأجور المستحقة",
         "Deferred revenue (client advances)": "إيرادات مؤجلة (دفعات العملاء المقدمة)", "Taxes payable": "ضرائب مستحقة",
         "Other current liabilities": "التزامات متداولة أخرى",
         "Capital expenditures (fixed assets)": "النفقات الرأسمالية (أصول ثابتة)",
         "Additions to intangibles": "إضافات أصول غير ملموسة", "Other long-term assets": "أصول طويلة الأجل أخرى",
         "Investments bought / (sold)": "استثمارات (مشتراة / مباعة)", "Partners' current accounts": "جاري الشركاء",
         "Short-term debt issued / (repaid)": "قروض قصيرة الأجل (مستلمة / مسددة)",
         "Long-term debt issued / (repaid)": "قروض طويلة الأجل (مستلمة / مسددة)",
         "Other long-term liabilities": "التزامات طويلة الأجل أخرى", "Capital contributed": "رأس المال المدفوع",
         "Other equity movements": "حركات حقوق ملكية أخرى",
         "Dividends / partner drawings paid": "توزيعات الأرباح / مسحوبات الشركاء",
         "Retained earnings adjustments": "تسويات الأرباح المحتجزة"}


def L(key, lang):
    v = T.get(key)
    if not v:
        return key
    return v[1] if lang == "ar" else v[0]


# -------------------------------------------------------------------- periods
def periods(date_from, date_to, compare="none"):
    """[(label, from, to)] for a range, optionally split by month/quarter/year."""
    today = _dt.date.today()
    d1 = _dt.date.fromisoformat(acc.clean_date(date_from)) if date_from else today.replace(month=1, day=1)
    d2 = _dt.date.fromisoformat(acc.clean_date(date_to)) if date_to else today
    if d2 < d1:
        raise acc.ValidationError("'To' date is before 'From' date")
    if compare in (None, "", "none"):
        return [(f"{d1.isoformat()} → {d2.isoformat()}", d1.isoformat(), d2.isoformat())]
    out, cur = [], d1
    while cur <= d2 and len(out) < 36:
        if compare == "month":
            nxt = (cur.replace(day=1) + _dt.timedelta(days=32)).replace(day=1)
            label = cur.strftime("%b %Y")
        elif compare == "quarter":
            q = (cur.month - 1) // 3
            nxt = _dt.date(cur.year + (1 if q == 3 else 0), 1 if q == 3 else q * 3 + 4, 1)
            label = f"Q{q + 1} {cur.year}"
        elif compare == "year":
            nxt = _dt.date(cur.year + 1, 1, 1)
            label = f"FY{cur.year}"
        else:
            raise acc.ValidationError("compare must be none, month, quarter or year")
        end = min(nxt - _dt.timedelta(days=1), d2)
        out.append((label, cur.isoformat(), end.isoformat()))
        cur = nxt
    return out


def _row(label, values, style="item", indent=0, code=""):
    return {"cells": [label] + list(values), "style": style, "indent": indent, "code": code}


_NAMES = {}


def _names(db, lang):
    """code -> display name in the report language (Arabic name when set)."""
    if lang != "ar":
        return {}
    return {r["code"]: r["name_ar"] for r in db.q("SELECT code, name_ar FROM accounts WHERE name_ar<>''")}


def _merge_lines(results, getter):
    """Union of account lines across periods, preserving order: [(code, name, [amount per period])]."""
    order, names = [], {}
    for r in results:
        for ln in getter(r):
            k = ln["code"] or ln["name"]
            if k not in names:
                order.append(k)
                names[k] = ln["name"]
    out = []
    for k in order:
        vals = []
        for r in results:
            vals.append(next((ln["amount"] for ln in getter(r) if (ln["code"] or ln["name"]) == k), 0.0))
        out.append((k if k != names[k] else "", names[k], vals))
    return out


def _subtitle(db, lang, pers, bu_id, as_of=False):
    unit = L("all_units", lang)
    if bu_id == "none":
        unit = L("shared_only", lang)
    elif bu_id:
        b = db.one("SELECT code, name FROM business_units WHERE id=?", (int(bu_id),))
        unit = f"{b['code']} — {b['name']}" if b else unit
    if as_of:
        when = f"{L('as_of', lang)} {pers[-1][2]}"
    else:
        when = f"{L('period', lang)}: {pers[0][1]} {L('to', lang)} {pers[-1][2]}"
    return f"{when} · {unit} · {db.setting('currency', 'EGP')}"


def _cols(pers, lang, first=None):
    return [{"label": first or "", "align": "left"}] + [{"label": p[0], "align": "right", "fmt": "money"} for p in pers]


# ------------------------------------------------------------------ statements
def income_statement_table(db, date_from=None, date_to=None, compare="none", bu_id=None, detail=True,
                           lang="en", allocate=True):
    pers = periods(date_from, date_to, compare)
    res = [acc.income_statement(db, f, t, bu_id, allocate_overhead=allocate) for _, f, t in pers]
    nm = _names(db, lang)
    rows = []
    v = lambda key: [r[key] for r in res]  # noqa: E731

    def block(getter, indent=1):
        if detail:
            for code, name, vals in _merge_lines(res, getter):
                rows.append(_row(f"{code}  {nm.get(code, name)}".strip(), vals, "item", indent, code))

    rows.append(_row(L("revenue", lang), ["" for _ in pers], "header"))
    block(lambda r: r["revenue"])
    rows.append(_row(L("gross_sales", lang), v("gross_sales"), "subtotal"))
    if any(r["discounts"] for r in res):
        block(lambda r: r["discounts"])
        rows.append(_row(L("discounts", lang), [-sum(x["amount"] for x in r["discounts"]) for r in res], "item"))
    rows.append(_row(L("net_sales", lang), v("net_sales"), "subtotal"))
    block(lambda r: r["cogs"])
    rows.append(_row(L("cost_of_sales", lang), [-x for x in v("cost_of_sales")], "subtotal"))
    rows.append(_row(L("gross_profit", lang), v("gross_profit"), "total"))
    rows.append(_row(L("operating_expenses", lang), ["" for _ in pers], "header"))
    labels = [g["label"] for g in res[0]["opex_groups"]]
    for i, lbl in enumerate(labels):
        vals = [r["opex_groups"][i]["total"] if i < len(r["opex_groups"]) else 0 for r in res]
        if not any(vals):
            continue
        if detail:
            for code, name, lv in _merge_lines(res, lambda r, i=i: r["opex_groups"][i]["lines"]
                                                if i < len(r["opex_groups"]) else []):
                rows.append(_row(f"{code}  {nm.get(code, name)}".strip(), [-x for x in lv], "item", 2, code))
        rows.append(_row(L(lbl, lang), [-x for x in vals], "item" if not detail else "subtotal", 1))
    rows.append(_row(L("operating_expenses", lang), [-x for x in v("operating_expenses")], "subtotal"))
    rows.append(_row(L("ebitda", lang), v("ebitda"), "total"))
    rows.append(_row(L("margin", lang) + " %", [r["margins"]["ebitda"] for r in res], "note"))
    rows.append(_row(L("depreciation", lang), [-sum(x["amount"] for x in r["depreciation"]) for r in res], "item", 1))
    rows.append(_row(L("amortization", lang), [-sum(x["amount"] for x in r["amortization"]) for r in res], "item", 1))
    rows.append(_row(L("ebit", lang), v("ebit"), "total"))
    rows.append(_row(L("interest", lang), [-sum(x["amount"] for x in r["interest"]) for r in res], "item", 1))
    if any(r["other_nonoperating"] for r in res):
        rows.append(_row(L("other_nonoperating", lang), [-sum(x["amount"] for x in r["other_nonoperating"])
                                                         for r in res], "item", 1))
    if any(r["other_income"] for r in res):
        rows.append(_row(L("other_income", lang), [sum(x["amount"] for x in r["other_income"]) for r in res], "item", 1))
    rows.append(_row(L("ebt", lang), v("ebt"), "total"))
    rows.append(_row(L("income_tax", lang), [-x for x in v("income_tax")], "item", 1))
    rows.append(_row(L("net_profit", lang), v("net_profit"), "grand"))
    rows.append(_row(L("margin", lang) + " %", [r["margins"]["net"] for r in res], "note"))
    if not bu_id and len(pers) == 1:
        # management memo: money received that running projects still need for their remaining costs is not profit yet
        from . import partner_pack
        pk = partner_pack.compute(db, pers[0][1], pers[0][2])
        if pk["held_total"]:
            ar = lang == "ar"
            rows.append(_row("منه: ربح أعمال تحت التنفيذ لم يُفوتر" if ar else "Of which: profit on work in progress, not yet billed", [pk["wip_profit"]], "item", 1))
            rows.append(_row("منه: محتجز لتكاليف المشروعات المتبقية" if ar else "Of which: held for remaining project costs", [pk["held_total"]], "item", 1))
            rows.append(_row("صافي الربح بعد التكاليف المتبقية (أساس التوزيع)" if ar else "Clean net profit (distributable base)", [pk["clean_total"]], "total"))
    return {"title": L("income_statement", lang), "subtitle": _subtitle(db, lang, pers, bu_id),
            "columns": _cols(pers, lang), "rows": rows, "lang": lang}


def balance_sheet_table(db, date_to=None, date_from=None, compare="none", bu_id=None, lang="en"):
    pers = periods(date_from or date_to or None, date_to, compare) if compare not in (None, "", "none") \
        else [(f"{L('as_of', lang)} {acc.clean_date(date_to)}", None, acc.clean_date(date_to))]
    res = [acc.balance_sheet(db, t, bu_id) for _, _, t in pers]
    nm = _names(db, lang)
    rows = []

    def sect(title, key, total_key, sign_label):
        rows.append(_row(L(title, lang), ["" for _ in pers], "header"))
        for code, name, vals in _merge_lines(res, lambda r: r[key]):
            rows.append(_row(f"{code}  {nm.get(code, name)}".strip(), vals, "item", 1, code))
        rows.append(_row(L(sign_label, lang), [r[total_key] for r in res], "subtotal"))

    rows.append(_row(L("assets", lang), ["" for _ in pers], "section"))
    sect("current_assets", "current_assets", "total_current_assets", "current_assets")
    sect("noncurrent_assets", "noncurrent_assets", "total_noncurrent_assets", "noncurrent_assets")
    rows.append(_row(L("total_assets", lang), [r["total_assets"] for r in res], "grand"))
    rows.append(_row(L("liabilities", lang), ["" for _ in pers], "section"))
    sect("current_liabilities", "current_liabilities", "total_current_liabilities", "current_liabilities")
    sect("noncurrent_liabilities", "noncurrent_liabilities", "total_noncurrent_liabilities", "noncurrent_liabilities")
    rows.append(_row(L("total_liabilities", lang), [r["total_liabilities"] for r in res], "total"))
    rows.append(_row(L("equity", lang), ["" for _ in pers], "section"))
    eq_ar = {"Retained earnings — closed years": "أرباح محتجزة — سنوات مُقفلة", "Current earnings (profit to date)": "أرباح الفترة الحالية"}
    for code, name, vals in _merge_lines(res, lambda r: r["equity"]):
        name = eq_ar.get(name, name) if lang == "ar" and not code else name
        rows.append(_row(f"{code}  {nm.get(code, name)}".strip(), vals, "item", 1, code))
    rows.append(_row(L("total_equity", lang), [r["total_equity"] for r in res], "total"))
    rows.append(_row(L("total_le", lang), [r["total_liabilities_equity"] for r in res], "grand"))
    ok = all(r["balanced"] for r in res)
    rows.append(_row(L("balanced_ok" if ok else "unbalanced", lang), ["" for _ in pers], "note"))
    return {"title": L("balance_sheet", lang), "subtitle": _subtitle(db, lang, pers, bu_id, as_of=True),
            "columns": _cols(pers, lang), "rows": rows, "lang": lang, "ok": ok}


def cash_flow_table(db, date_from=None, date_to=None, compare="none", bu_id=None, lang="en"):
    pers = periods(date_from, date_to, compare)
    res = [acc.cash_flow(db, f, t, bu_id) for _, f, t in pers]
    lab = lambda s: s if lang != "ar" else (s.split(") ", 1)[0] + ") " + CF_AR.get(s.split(") ", 1)[1], s.split(") ", 1)[1])
                                            if ") " in s else CF_AR.get(s, s))  # noqa: E731
    rows = [_row(L("cf_ops", lang), ["" for _ in pers], "section"),
            _row(L("net_income", lang), [r["net_income"] for r in res], "subtotal"),
            _row(L("adjustments", lang), ["" for _ in pers], "header")]
    n_adj = 2
    for i in range(len(res[0]["operating"])):
        if i == n_adj:
            rows.append(_row(L("wc_changes", lang), ["" for _ in pers], "header"))
        rows.append(_row(lab(res[0]["operating"][i]["label"]), [r["operating"][i]["amount"] for r in res], "item", 1))
    rows.append(_row(L("net_ops", lang), [r["net_operating"] for r in res], "total"))
    rows.append(_row(L("cf_inv", lang), ["" for _ in pers], "section"))
    for i in range(len(res[0]["investing"])):
        rows.append(_row(lab(res[0]["investing"][i]["label"]), [r["investing"][i]["amount"] for r in res], "item", 1))
    rows.append(_row(L("net_inv", lang), [r["net_investing"] for r in res], "total"))
    rows.append(_row(L("cf_fin", lang), ["" for _ in pers], "section"))
    for i in range(len(res[0]["financing"])):
        vals = [r["financing"][i]["amount"] for r in res]
        if any(vals) or i < 4:
            rows.append(_row(lab(res[0]["financing"][i]["label"]), vals, "item", 1))
    rows.append(_row(L("net_fin", lang), [r["net_financing"] for r in res], "total"))
    rows.append(_row(L("beginning_cash", lang), [r["beginning_cash"] for r in res], "item"))
    rows.append(_row(L("change_cash", lang), [r["change_in_cash"] for r in res], "item"))
    rows.append(_row(L("ending_cash", lang), [r["ending_cash"] for r in res], "grand"))
    ok = all(r["reconciles"] for r in res)
    rows.append(_row(L("balanced_ok" if ok else "unbalanced", lang), ["" for _ in pers], "note"))
    return {"title": L("cash_flow", lang), "subtitle": _subtitle(db, lang, pers, bu_id), "columns": _cols(pers, lang),
            "rows": rows, "lang": lang, "ok": ok}


def trial_balance_table(db, date_from=None, date_to=None, bu_id=None, lang="en"):
    tb = acc.trial_balance(db, date_from, date_to, bu_id)
    cols = [{"label": L("account", lang), "align": "left"}]
    if date_from:
        cols += [{"label": f"{L('opening', lang)} {L('debit', lang)}", "align": "right", "fmt": "money"},
                 {"label": f"{L('opening', lang)} {L('credit', lang)}", "align": "right", "fmt": "money"},
                 {"label": f"{L('movement', lang)} {L('debit', lang)}", "align": "right", "fmt": "money"},
                 {"label": f"{L('movement', lang)} {L('credit', lang)}", "align": "right", "fmt": "money"}]
    cols += [{"label": f"{L('closing', lang)} {L('debit', lang)}" if date_from else L("debit", lang), "align": "right",
              "fmt": "money"},
             {"label": f"{L('closing', lang)} {L('credit', lang)}" if date_from else L("credit", lang), "align": "right",
              "fmt": "money"}]
    keys = (["opening_debit", "opening_credit", "period_debit", "period_credit"] if date_from else []) + ["debit", "credit"]
    nm = _names(db, lang)
    rows = [_row(f"{r['code']}  {nm.get(r['code'], r['name'])}", [r[k] or "" for k in keys], "item", 0, r["code"])
            for r in tb["rows"]]
    rows.append(_row(L("total", lang), [tb["totals"][k] for k in keys], "grand"))
    rows.append(_row(L("balanced_ok" if tb["balanced"] else "unbalanced", lang), ["" for _ in keys], "note"))
    pers = [("", date_from or "…", acc.clean_date(date_to))]
    return {"title": L("trial_balance", lang), "subtitle": _subtitle(db, lang, pers, bu_id, as_of=not date_from),
            "columns": cols, "rows": rows, "lang": lang, "ok": tb["balanced"]}


def general_ledger_table(db, account=None, date_from=None, date_to=None, bu_id=None, party=None, lang="en"):
    gl = acc.general_ledger(db, account, date_from, date_to, bu_id, party, limit=100000)
    nm = _names(db, lang)
    cols = [{"label": L("date", lang), "align": "left"}, {"label": L("voucher", lang), "align": "left"},
            {"label": L("description", lang), "align": "left"}, {"label": L("party", lang), "align": "left"},
            {"label": L("bu", lang), "align": "left"}]
    if not account:
        cols.append({"label": L("account", lang), "align": "left"})
    cols += [{"label": L("debit", lang), "align": "right", "fmt": "money"},
             {"label": L("credit", lang), "align": "right", "fmt": "money"}]
    if account:
        cols.append({"label": L("balance", lang), "align": "right", "fmt": "money"})
    rows = []
    if account and date_from:
        rows.append({"cells": ["", "", L("opening", lang), "", "", "", "", gl["opening"]], "style": "subtotal", "indent": 0})
    for r in gl["rows"]:
        cells = [r["date"], f"#{r['id']} {r['kind']}", r["memo"], r["party"], r["bu_code"] or ""]
        if not account:
            cells.append(f"{r['account']} {nm.get(r['account'], r['account_name'])}")
        cells += [r["debit"] or "", r["credit"] or ""]
        if account:
            cells.append(r["balance"])
        rows.append({"cells": cells, "style": "item", "indent": 0})
    td = round(sum(r["debit"] for r in gl["rows"]), 2)
    tc = round(sum(r["credit"] for r in gl["rows"]), 2)
    tot = [L("total", lang)] + [""] * (len(cols) - (4 if account else 3)) + [td, tc] + ([gl["closing"]] if account else [])
    rows.append({"cells": tot, "style": "grand", "indent": 0})
    title = L("general_ledger", lang) + (f" — {gl['account']['code']} {nm.get(gl['account']['code'], gl['account']['name'])}"
                                         if gl["account"] else "")
    pers = [("", date_from or "…", acc.clean_date(date_to))]
    return {"title": title, "subtitle": _subtitle(db, lang, pers, bu_id), "columns": cols, "rows": rows, "lang": lang}


def bu_table(db, date_from=None, date_to=None, status=None, lang="en"):
    data = acc.bu_summaries(db, date_from, date_to, status)
    heads = [("Code", "الكود", "left"), ("Name", "الاسم", "left"), ("Basis", "أساس الفوترة", "left"),
             ("Revenue", "الإيراد", "right"), ("Direct costs", "التكاليف المباشرة", "right"),
             ("Direct overhead", "مصروفات مباشرة", "right"), ("Allocated overhead", "مصروفات محملة", "right"),
             ("Net profit", "صافي الربح", "right"), ("Margin", "الهامش", "right"),
             ("Receivable", "المدينون", "right"), ("To collect", "المتبقي للتحصيل", "right")]
    cols = [{"label": h[1] if lang == "ar" else h[0], "align": h[2], "fmt": "money" if h[2] == "right" else None}
            for h in heads]
    cols[8]["fmt"] = "pct"
    rows = [{"cells": [b["code"], b["name"], b["basis"], b["revenue"], b["cogs"], b["direct_overhead"],
                       b["allocated_overhead"], b["net_profit"], b["margin"], b["receivable"], b["to_collect"]],
             "style": "item", "indent": 0} for b in data]
    tot = lambda k: round(sum(b[k] for b in data), 2)  # noqa: E731
    rows.append({"cells": [L("total", lang), "", "", tot("revenue"), tot("cogs"), tot("direct_overhead"),
                           tot("allocated_overhead"), tot("net_profit"), "", tot("receivable"), tot("to_collect")],
                 "style": "grand", "indent": 0})
    if any(b["expected_total_cost"] for b in data):      # projects with an expected cost: what is still to come
        cols += [{"label": "التكلفة المتبقية" if lang == "ar" else "Cost still to come", "align": "right", "fmt": "money"},
                 {"label": "الإنجاز (بالتكلفة)" if lang == "ar" else "% complete (by cost)", "align": "right", "fmt": "pct"}]
        for r, b in zip(rows, data):
            r["cells"] += [b["cost_still_to_come"], b["pct_complete"] if b["pct_complete"] is not None else ""]
        rows[-1]["cells"] += [tot("cost_still_to_come"), ""]
    pers = [("", date_from or "…", acc.clean_date(date_to))]
    return {"title": L("bu_pnl", lang), "subtitle": _subtitle(db, lang, pers, None), "columns": cols, "rows": rows,
            "lang": lang}


def project_cash(db, date_from=None, date_to=None):
    """Raw cash per project: what came into any cash / bank (or the funding partner's) account for it and what went out,
    what is still owed to suppliers and staff for it, and where it stands once they are paid. Money with no project
    (overhead, loans to partners, distributions) is one company line."""
    from .funders import ROLE_SQL
    money = {r["code"] for r in db.q(f"SELECT code FROM accounts WHERE subtype='cash' OR {ROLE_SQL}")}
    owe = {r["code"] for r in db.q(f"SELECT code FROM accounts WHERE type='Liability' AND subtype IN ('payable','accrued') AND NOT {ROLE_SQL}")}
    w, p = acc._where(date_from=date_from, date_to=date_to)
    agg = {}
    for r in db.q(f"""SELECT e.bu_id, l.account, SUM(l.debit) dr, SUM(l.credit) cr FROM journal_lines l
                      JOIN journal_entries e ON e.id=l.entry_id WHERE {w} GROUP BY e.bu_id, l.account""", p):
        g = agg.setdefault(r["bu_id"], {"in": 0, "out": 0, "owed": 0})
        if r["account"] in money:
            g["in"] += r["dr"]
            g["out"] += r["cr"]
        elif r["account"] in owe:
            g["owed"] += r["cr"] - r["dr"]
    rows = []
    for b in db.q("SELECT id, code, name FROM business_units ORDER BY code"):
        g = agg.get(b["id"])
        if g and (g["in"] or g["out"] or g["owed"]):
            rows.append({"id": b["id"], "code": b["code"], "name": b["name"], **g})
    if agg.get(None):
        rows.append({"id": None, "code": "", "name": "", **agg[None]})
    for r in rows:
        r["net"] = r["in"] - r["out"]
        r["after"] = r["net"] - r["owed"]
        r["spent"] = round(r["out"] / r["in"], 4) if r["in"] else None
        for k in ("in", "out", "owed", "net", "after"):
            r[k] = acc.to_major(r[k])
    return rows


def project_cash_table(db, date_from=None, date_to=None, lang="en"):
    ar = lang == "ar"
    data = project_cash(db, date_from, date_to)
    heads = [("Code", "الكود", "left"), ("Project", "المشروع", "left"), ("Cash in", "النقدية الواردة", "right"),
             ("Cash out", "النقدية المنصرفة", "right"), ("Net cash", "صافي النقدية", "right"), ("Cash out vs in", "المنصرف من الوارد", "right"),
             ("Still to pay suppliers", "المتبقي للموردين", "right"), ("After paying suppliers", "بعد سداد الموردين", "right"),
             ("Where it stands", "الموقف", "left")]
    cols = [{"label": h[1] if ar else h[0], "align": h[2], "fmt": "money" if h[2] == "right" else None} for h in heads]
    cols[5]["fmt"] = "pct"

    def stands(r):
        if r["after"] < -0.005:
            return (f"يحتاج {-r['after']:,.0f} (من أموال الشريك)" if ar else f"Short by {-r['after']:,.0f} (partner money)")
        return (f"متقدم بمقدار {r['after']:,.0f}" if ar else f"Ahead by {r['after']:,.0f}")
    rows = [{"cells": [r["code"] or ("الشركة (بدون مشروع)" if ar else "Company (no project)"),
                       r["name"] or ("مصروفات عامة وقروض شركاء وتوزيعات" if ar else "overhead, partner loans, distributions"),
                       r["in"], r["out"], r["net"], r["spent"] if r["spent"] is not None else "", r["owed"], r["after"], stands(r)],
             "style": "item", "indent": 0} for r in data]
    tot = {k: round(sum(r[k] for r in data), 2) for k in ("in", "out", "net", "owed", "after")}
    rows.append({"cells": [L("total", lang), "", tot["in"], tot["out"], tot["net"], "", tot["owed"], tot["after"], stands(tot)],
                 "style": "grand", "indent": 0})
    pers = [("", date_from or "…", acc.clean_date(date_to))]
    return {"title": L("project_cash", lang), "subtitle": _subtitle(db, lang, pers, None), "columns": cols, "rows": rows,
            "lang": lang}


def payables_table(db, date_to=None, lang="en"):
    data = acc.party_balances(db, subtypes=("payable", "accrued"), date_to=date_to)
    cols = [{"label": L("party", lang), "align": "left"}, {"label": L("account", lang), "align": "left"},
            {"label": "Billed" if lang != "ar" else "المستحق", "align": "right", "fmt": "money"},
            {"label": "Paid" if lang != "ar" else "المدفوع", "align": "right", "fmt": "money"},
            {"label": L("balance", lang), "align": "right", "fmt": "money"},
            {"label": "Last activity" if lang != "ar" else "آخر حركة", "align": "left"}]
    rows = [{"cells": [r["party"] or "—", r["account"], r["owed"], r["paid"], r["balance"], r["last_date"]],
             "style": "item", "indent": 0} for r in data]
    rows.append({"cells": [L("total", lang), "", "", "", round(sum(r["balance"] for r in data), 2), ""],
                 "style": "grand", "indent": 0})
    pers = [("", "…", acc.clean_date(date_to))]
    return {"title": L("payables", lang), "subtitle": _subtitle(db, lang, pers, None, as_of=True), "columns": cols,
            "rows": rows, "lang": lang}


def stock_table(db, warehouse_id=None, q=None, lang="en"):
    data = acc.stock_report(db, warehouse_id, q)
    heads = [("Code", "الكود", "left", None), ("Item", "الصنف", "left", None), ("Warehouse", "المخزن", "left", None),
             ("Unit", "الوحدة", "left", None), ("On hand", "الرصيد", "right", "num"), ("Avg cost", "متوسط التكلفة", "right", "money"),
             ("Last purchase price", "آخر سعر شراء", "right", "money"), ("Last purchased on", "تاريخ آخر شراء", "left", None),
             ("Value", "القيمة", "right", "money")]
    cols = [{"label": h[1] if lang == "ar" else h[0], "align": h[2], "fmt": h[3]} for h in heads]
    rows = [{"cells": [r["sku"] or "", r["name"], r["warehouse"], r["unit"], r["qty"], r["avg_cost"], r["last_rate"],
                       r["last_date"] or "", r["value"]], "style": "item", "indent": 0} for r in data]
    rows.append({"cells": [L("total", lang), "", "", "", "", "", "", "", round(sum(r["value"] for r in data), 2)], "style": "grand",
                 "indent": 0})
    if not any(r["sku"] for r in data):          # no item has a code: no empty Code column
        cols = cols[1:]
        for r in rows:
            r["cells"] = r["cells"][1:]
        rows[-1]["cells"][0] = L("total", lang)
    return {"title": L("stock", lang), "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


def fixed_assets_table(db, date_to=None, q=None, state=None, lang="en"):
    from . import assets
    data = assets.register(db, q, None, state or "active", acc.clean_date(date_to) if date_to else None)
    heads = [("Code", "الكود", "left", None), ("Asset", "الأصل", "left", None), ("Account", "الحساب", "left", None),
             ("Acquired", "تاريخ الشراء", "left", None), ("Life (months)", "العمر (شهور)", "right", "num"),
             ("Cost", "التكلفة", "right", "money"), ("Accumulated depreciation", "مجمع الإهلاك", "right", "money"),
             ("Book value", "القيمة الدفترية", "right", "money"), ("Monthly", "الإهلاك الشهري", "right", "money")]
    cols = [{"label": h[1] if lang == "ar" else h[0], "align": h[2], "fmt": h[3]} for h in heads]
    rows = [{"cells": [r["code"], r["name"] + (" (disposed)" if r["disposed"] and lang != "ar" else " (مستبعد)" if r["disposed"] else ""),
                       r["account"], r["acquired"], r["life_months"], r["cost"], r["accumulated"], r["book_value"],
                       r["monthly"]], "style": "item", "indent": 0} for r in data["rows"]]
    t = data["totals"]
    rows.append({"cells": [L("total", lang), "", "", "", "", t["cost"], t["accumulated"], t["book_value"], ""],
                 "style": "grand", "indent": 0})
    title = "سجل الأصول الثابتة" if lang == "ar" else "Fixed asset register"
    return {"title": title, "subtitle": f"{L('as_of', lang)} {acc.clean_date(date_to) if date_to else acc.today()} · "
                                        f"{db.setting('currency')}", "columns": cols, "rows": rows, "lang": lang}


def receivables_table(db, date_from=None, date_to=None, lang="en"):
    data = [b for b in acc.bu_summaries(db, date_from, date_to) if b["receivable"] or b["advances"] or b["to_collect"]]
    heads = [("Code", "الكود", "left"), ("Name", "الاسم", "left"), ("Invoiced receivable", "مديونية الفواتير", "right"),
             ("Client advances", "دفعات مقدمة", "right"), ("Expected billing", "الفوترة المتوقعة", "right"),
             ("Collected", "المحصل", "right"), ("To collect", "المتبقي للتحصيل", "right")]
    cols = [{"label": h[1] if lang == "ar" else h[0], "align": h[2], "fmt": "money" if h[2] == "right" else None}
            for h in heads]
    rows = [{"cells": [b["code"], b["name"], b["receivable"], b["advances"], b["expected_billing"], b["collected"],
                       b["to_collect"]], "style": "item", "indent": 0} for b in data]
    t = lambda k: round(sum(b[k] for b in data), 2)  # noqa: E731
    rows.append({"cells": [L("total", lang), "", t("receivable"), t("advances"), t("expected_billing"), t("collected"),
                           t("to_collect")], "style": "grand", "indent": 0})
    pers = [("", date_from or "…", acc.clean_date(date_to))]
    return {"title": L("receivables", lang), "subtitle": _subtitle(db, lang, pers, None), "columns": cols, "rows": rows,
            "lang": lang}


def vouchers_table(db, entries, title, lang="en", subtitle=""):
    """A filtered list of vouchers (any module screen) as an exportable table."""
    heads = [("No.", "رقم"), ("Date", "التاريخ"), ("Type", "النوع"), ("Description", "البيان"), ("Party", "الطرف"),
             ("Business unit", "وحدة الأعمال"), ("Debit lines", "الحسابات المدينة"), ("Credit lines", "الحسابات الدائنة"),
             ("Amount", "المبلغ"), ("Status", "الحالة")]
    cols = [{"label": h[1] if lang == "ar" else h[0], "align": "right" if h[0] == "Amount" else "left",
             "fmt": "money" if h[0] == "Amount" else None} for h in heads]
    rows = []
    for e in entries:
        dr = "; ".join(f"{ln['account']} {ln['name']} {ln['debit']:,.2f}" for ln in e["lines"] if ln["debit"])
        cr = "; ".join(f"{ln['account']} {ln['name']} {ln['credit']:,.2f}" for ln in e["lines"] if ln["credit"])
        status = "void" if e["status"] == "void" else ("reversed" if e["reversed_by"] else "posted")
        rows.append({"cells": [e["id"], e["date"], e["kind"], e["memo"], e["party"], e["bu_code"] or "", dr, cr,
                               e["amount"], status], "style": "item", "indent": 0})
    rows.append({"cells": [L("total", lang), "", "", "", "", "", "", "",
                           round(sum(e["amount"] for e in entries if e["status"] == "posted"), 2), ""],
                 "style": "grand", "indent": 0})
    return {"title": title, "subtitle": subtitle, "columns": cols, "rows": rows, "lang": lang}


def negotiations_table(db, date_from=None, date_to=None, bu_id=None, lang="en"):
    """Negotiated discounts per business unit and deal: kept as company earnings or passed on as a lower cost."""
    n = acc.negotiations(db, date_from, date_to, bu_id)
    ar = lang == "ar"
    heads = [("Date", "التاريخ", "left"), ("No.", "رقم", "left"), ("Supplier", "المورد", "left"),
             ("Deal reference", "مرجع الاتفاق", "left"), ("Description", "البيان", "left"),
             ("Quoted", "السعر المعروض", "right"), ("Negotiated discount", "خصم التفاوض", "right"),
             ("Company earnings", "أرباح للشركة", "right"), ("Passed to client (lower cost)", "لصالح العميل (تكلفة أقل)", "right")]
    cols = [{"label": h[1] if ar else h[0], "align": h[2], "fmt": "money" if h[2] == "right" else None} for h in heads]
    rows = []
    for g in n["by_bu"]:
        rows.append({"cells": [f"{g['bu_code']} — {g['bu_name']}" if g["bu_code"] else ("مشترك" if ar else "Shared"),
                               "", "", "", "", "", "", "", ""], "style": "section", "indent": 0})
        for r in (x for x in n["rows"] if x["bu_id"] == g["bu_id"]):
            rows.append({"cells": [r["date"], r["id"], r["party"], r["deal_ref"], r["memo"], r["quoted"] or "",
                                   r["savings"], r["kept"] or "", r["passed"] or ""], "style": "item", "indent": 1})
        rows.append({"cells": ["المجموع" if ar else "Subtotal", "", "", "", "", g["quoted"], g["savings"], g["kept"],
                               g["passed"]], "style": "subtotal", "indent": 0})
    t = n["total"]
    rows.append({"cells": [L("total", lang), "", "", "", "", t["quoted"], t["savings"], t["kept"], t["passed"]],
                 "style": "grand", "indent": 0})
    pers = [("", date_from or "…", acc.clean_date(date_to))]
    return {"title": "مكاسب التفاوض مع الموردين" if ar else "Negotiation earnings", "subtitle": _subtitle(db, lang, pers, None),
            "columns": cols, "rows": rows, "lang": lang}


def materials_table(db, p, lang="en"):
    ar = lang == "ar"
    d = acc.materials_used(db, p.get("bu_id"), p.get("q"), p.get("date_from"), p.get("date_to"), item=p.get("item"), source=p.get("source"))
    src = {"issued": ("Issued from store", "صرف من المخزن"), "bought": ("Bought for the project", "شراء للمشروع"),
           "direct": ("Direct cost", "تكلفة مباشرة")}
    heads = [("Date", "التاريخ", "left"), ("No.", "رقم", "left"), ("Project", "المشروع", "left"), ("Material", "الخامة", "left"),
             ("Qty", "الكمية", "right", "num"), ("Unit", "الوحدة", "left"), ("Unit cost", "سعر الوحدة", "right"),
             ("Amount", "القيمة", "right"), ("How", "الطريقة", "left"), ("From", "من", "left")]
    rows = [{"cells": [r["date"], f"#{r['entry_id']}", r["bu_code"], r["description"], r["qty"], r["unit"], r["rate"], r["amount"],
                       src[r["source"]][1 if ar else 0], r["from"]], "style": "item", "entry": r["entry_id"]} for r in d["rows"]]
    one = d.get("one_material")
    rows.append({"cells": [L("total", lang), "", "", "", one["qty"] if one else "", one["unit"] if one else "",
                           round(d["total"] / one["qty"], 2) if one and one["qty"] else "", d["total"], "", ""], "style": "grand"})
    return {"title": "الخامات المستخدمة في المشروعات" if ar else "Materials used on projects",
            "subtitle": _subtitle(db, lang, [("", p.get("date_from") or "…", acc.clean_date(p.get("date_to")))], p.get("bu_id")),
            "columns": _mkcols(heads, lang), "rows": rows, "lang": lang,
            "summary": {k: d[k] for k in ("count", "total", "materials", "one_material")}}


def _mkcols(heads, lang):
    return [{"label": h[1] if lang == "ar" else h[0], "align": h[2], "fmt": h[3] if len(h) > 3 else ("money" if h[2] == "right" else None)}
            for h in heads]


def loans_table(db, direction=None, state=None, lang="en"):
    from . import loans
    ar = lang == "ar"
    reg = loans.register(db, direction=direction, state=state)
    cols = _mkcols([("Code", "الكود", "left"), ("Lender / borrower", "الجهة", "left"), ("Start", "البداية", "left"),
                  ("Rate %", "الفائدة %", "right", "num"), ("Amount", "المبلغ", "right"), ("Outstanding", "المتبقي", "right"),
                  ("Accrued interest", "فوائد مستحقة", "right"), ("Overdue", "متأخرات", "right"), ("Next due", "القسط القادم", "left"),
                  ("Maturity", "الاستحقاق النهائي", "left"), ("Status", "الحالة", "left")], lang)
    rows = []
    for d, title in (("borrowed", ("LOANS THE COMPANY OWES", "قروض على الشركة")), ("lent", ("LOANS GIVEN BY THE COMPANY", "قروض ممنوحة من الشركة"))):
        rs = [r for r in reg["rows"] if r["direction"] == d]
        if not rs:
            continue
        rows.append({"cells": [title[1] if ar else title[0]] + [""] * 10, "style": "section", "indent": 0})
        for r in rs:
            rows.append({"cells": [r["code"], r["party"], r["start_date"], r["rate_pct"] if r["interest_type"] in ("declining", "flat") else "",
                                   r["principal"] or r["opening_balance"], r["outstanding"], r["accrued"], r["overdue"],
                                   f"{r['next_due']} ({r['next_amount']:,.2f})" if r["next_due"] else "", r["maturity"],
                                   ("مغلق" if ar else "Closed") if r["status"] == "closed" else ("نشط" if ar else "Active")],
                         "style": "item", "indent": 1})
        t = reg["totals"][d]
        rows.append({"cells": [L("total", lang), "", "", "", "", t["outstanding"], t["accrued"], t["overdue"], "", "", ""],
                     "style": "subtotal", "indent": 0})
    return {"title": "سجل القروض" if ar else "Loans register", "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


def loan_schedule_table(db, code, lang="en"):
    from . import loans
    ar = lang == "ar"
    d = loans.detail(db, code)
    ln = d["loan"]
    cols = _mkcols([("No.", "رقم", "left"), ("Due", "تاريخ الاستحقاق", "left"), ("Principal", "أصل القرض", "right"),
                  ("Interest", "الفائدة", "right"), ("Installment", "القسط", "right"), ("Balance after", "الرصيد بعده", "right"),
                  ("Still due", "المتبقي منه", "right"), ("Status", "الحالة", "left")], lang)
    st = {"paid": ("Paid", "مسدد"), "partial": ("Part paid", "مسدد جزئياً"), "open": ("Open", "قادم"), "overdue": ("Overdue", "متأخر")}
    rows = [{"cells": [r["n"], r["due"], r["principal"], r["interest"], r["payment"], r["balance"], r["left"] or "",
                       st[r["state"]][1 if ar else 0]], "style": "item", "indent": 0} for r in d["schedule"]]
    rows.append({"cells": [L("total", lang), "", round(sum(r["principal"] for r in d["schedule"]), 2),
                           round(sum(r["interest"] for r in d["schedule"]), 2), round(sum(r["payment"] for r in d["schedule"]), 2),
                           "", round(sum(r["left"] for r in d["schedule"]), 2), ""], "style": "grand", "indent": 0})
    who = ("المقرض" if ln["direction"] == "borrowed" else "المقترض") if ar else ("Lender" if ln["direction"] == "borrowed" else "Borrower")
    return {"title": (f"جدول سداد القرض {ln['code']}" if ar else f"Loan schedule {ln['code']}"),
            "subtitle": f"{who}: {ln['party']} · {ln['start_date']} · {db.setting('currency')}", "columns": cols, "rows": rows, "lang": lang}


def custody_table(db, state=None, lang="en"):
    from . import custody
    ar = lang == "ar"
    reg = custody.register(db, state=state)
    cols = _mkcols([("Code", "الكود", "left"), ("Held by", "صاحب العهدة", "left"), ("Type", "النوع", "left"),
                  ("Given on", "تاريخ الصرف", "left"), ("Settle by", "تسوية قبل", "left"), ("Given", "المصروف له", "right"),
                  ("Spent (reports)", "المصروفات (تسويات)", "right"), ("Returned", "المرتجع", "right"),
                  ("Deducted", "المخصوم", "right"), ("Balance", "الرصيد", "right"), ("Status", "الحالة", "left")], lang)
    kinds = {"temporary": ("Temporary", "مؤقتة"), "permanent": ("Permanent float", "مستديمة")}
    rows = [{"cells": [r["code"], r["holder"], kinds[r["kind"]][1 if ar else 0], r["issued_on"], r["due_date"], r["issued"], r["spent"],
                       r["returned"], r["deducted"], r["balance"],
                       (("مغلقة" if ar else "Closed") if r["status"] == "closed" else
                        ("متأخرة" if ar else "Overdue") if r["overdue"] else ("مفتوحة" if ar else "Open"))],
             "style": "item", "indent": 0} for r in reg["rows"]]
    rows.append({"cells": [L("total", lang)] + [""] * 8 + [reg["totals"]["balance"], ""], "style": "grand", "indent": 0})
    return {"title": "العهد النقدية" if ar else "Cash custody", "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


def custody_statement_table(db, code, lang="en"):
    from . import custody
    ar = lang == "ar"
    d = custody.statement(db, code)
    c = d["custody"]
    cols = _mkcols([("Date", "التاريخ", "left"), ("Type", "النوع", "left"), ("Description", "البيان", "left"),
                  ("Given", "صرف", "right"), ("Spent / returned", "مصروف / مرتجع", "right"), ("Balance", "الرصيد", "right")], lang)
    rows = []
    for r in d["rows"]:
        rows.append({"cells": [r["date"], r["kind"], r["memo"], r["in"] or "", r["out"] or "", r["balance"]], "style": "item", "indent": 0})
        for x in r.get("lines") or []:
            rows.append({"cells": ["", x.get("receipt", ""), f"{x.get('description', '')}" + (f" — {x['supplier']}" if x.get("supplier") else ""),
                                   "", x.get("amount", 0), ""], "style": "note", "indent": 2})
    rows.append({"cells": [L("total", lang), "", "", c["issued"], round(c["spent"] + c["returned"] + c["deducted"], 2), c["balance"]],
                 "style": "grand", "indent": 0})
    return {"title": (f"كشف حساب العهدة {c['code']} — {c['holder']}" if ar else f"Custody statement {c['code']} — {c['holder']}"),
            "subtitle": f"{c['purpose']} · {db.setting('currency')}", "columns": cols, "rows": rows, "lang": lang}


def rent_fleet_table(db, lang="en"):
    from . import rentals
    ar = lang == "ar"
    d = rentals.list_items(db)
    cols = _mkcols([("Code", "الكود", "left"), ("Item", "الصنف", "left"), ("Unit", "الوحدة", "left"), ("In the yard", "بالمخزن", "right", "num"),
                    ("On rent", "مؤجر", "right", "num"), ("Owned", "المملوك", "right", "num"), ("Utilisation %", "نسبة التشغيل", "right", "pct"),
                    ("Avg cost (FIFO)", "متوسط التكلفة", "right"), ("Cost", "التكلفة", "right"), ("Accumulated depreciation", "مجمع الإهلاك", "right"),
                    ("Book value", "القيمة الدفترية", "right"), ("Rent / month", "الإيجار الشهري", "right"),
                    ("Rent % of cost / month", "الإيجار % من التكلفة", "right", "pct"), ("Months to pay back", "شهور استرداد التكلفة", "right", "num"),
                    ("Loss price", "سعر الفقد", "right")], lang)
    rows = [{"cells": [i["code"], (i["name_ar"] if ar and i["name_ar"] else i["name"]), i["unit"], i["in_yard"], i["on_rent"], i["owned"],
                       i["utilisation"], i["avg_cost"], i["value"], i["accum"], i["book_value"], i["monthly_rate"], i["rent_pct"], i["payback_months"], i["loss_price"]],
             "style": "item", "indent": 0} for i in d["items"]]
    rows.append({"cells": [L("total", lang), "", "", "", "", "", "", "", d["totals"]["value"], d["totals"]["accum"], d["totals"]["book_value"], "", "", "", ""],
                 "style": "grand", "indent": 0})
    return {"title": "معدات التأجير" if ar else "Rental equipment", "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


def rentals_table(db, status=None, lang="en"):
    from . import rentals
    ar = lang == "ar"
    d = rentals.list_rentals(db, status=status)
    st = {"quote": ("Price offer", "عرض سعر"), "ordered": ("PO received", "أمر شراء"), "active": ("On site", "بالموقع"),
          "returned": ("Returned", "تم الإرجاع"), "closed": ("Closed", "مغلق"), "cancelled": ("Cancelled", "ملغي")}
    cols = _mkcols([("Code", "الكود", "left"), ("Client", "العميل", "left"), ("Site", "الموقع", "left"), ("Status", "الحالة", "left"),
                    ("PO", "أمر الشراء", "left"), ("Start", "البداية", "left"), ("Rent / month on site", "الإيجار الشهري بالموقع", "right"),
                    ("Pieces out", "القطع بالموقع", "right", "num"), ("Invoiced to", "تمت الفوترة حتى", "left")], lang)
    rows = [{"cells": [r["code"], r["client"], r["site"], st.get(r["status"], (r["status"],) * 2)[1 if ar else 0], r["po_number"], r["start_date"],
                       r["monthly_on_site"], r["pieces_out"], r["billed_to"]], "style": "item", "indent": 0} for r in d["rows"]]
    rows.append({"cells": [L("total", lang), "", "", "", "", "", d["totals"]["monthly"], "", ""], "style": "grand", "indent": 0})
    return {"title": "عقود التأجير" if ar else "Rental contracts", "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


SALE_ST = {"draft": ("Draft", "مسودة"), "confirmed": ("Confirmed", "مؤكد"), "delivered": ("Delivered", "تم التسليم"),
           "cancelled": ("Cancelled", "ملغي")}
INST_ST = {"paid": ("Paid", "مسدد"), "cheque": ("By cheque", "بشيك"), "partial": ("Part paid", "مسدد جزئياً"),
           "overdue": ("Overdue", "متأخر"), "due": ("Due", "مستحق"), "cancelled": ("Cancelled", "ملغي")}


def sales_table(db, p, lang="en"):
    from . import sales
    ar = lang == "ar"
    d = sales.register(db, p.get("q"), p.get("status"), p.get("customer"), p.get("date_from"), p.get("date_to"))
    cols = _mkcols([("Sale", "البيع", "left"), ("Date", "التاريخ", "left"), ("Customer", "العميل", "left"), ("Status", "الحالة", "left"),
                    ("Price", "الثمن", "right"), ("Interest", "فوائد التقسيط", "right"), ("Total", "الإجمالي", "right"),
                    ("Paid", "المسدد", "right"), ("In cheques", "بشيكات", "right"), ("Balance", "المتبقي", "right"),
                    ("Overdue", "متأخرات", "right"), ("Next due", "القسط القادم", "left")], lang)
    rows = [{"cells": [r["code"], r["date"], r["customer"], SALE_ST.get(r["status"], (r["status"],) * 2)[1 if ar else 0], r["price"], r["markup"],
                       r["total"], r["paid"], r["in_cheques"], r["balance"], r["overdue"],
                       f"{r['next_due']} ({r['next_amount']:,.2f})" if r["next_due"] and r["status"] != "cancelled" else ""],
             "style": "item", "indent": 0} for r in d["rows"]]
    t = d["totals"]
    rows.append({"cells": [L("total", lang), "", "", "", "", "", t["total"], t["paid"], t["in_cheques"], t["balance"], t["overdue"], ""],
                 "style": "grand", "indent": 0})
    return {"title": "المبيعات" if ar else "Sales", "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


def installments_table(db, p, lang="en"):
    from . import sales
    ar = lang == "ar"
    d = sales.installments(db, p.get("as_of") or p.get("date_to_asof"), p.get("status") or "open", p.get("date_from"), p.get("date_to"),
                           p.get("customer"))
    cols = _mkcols([("Due", "الاستحقاق", "left"), ("Sale", "البيع", "left"), ("Customer", "العميل", "left"), ("Phone", "الهاتف", "left"),
                    ("Payment", "الدفعة", "left"), ("Amount", "المبلغ", "right"), ("Paid", "المسدد", "right"),
                    ("By cheque", "بشيك", "right"), ("Left", "المتبقي", "right"), ("Days late", "أيام التأخير", "right", "num"),
                    ("Status", "الحالة", "left")], lang)
    rows = [{"cells": [r["due"], r["code"], r["customer"], r["phone"] or r["whatsapp"], r["label"], r["amount"], r["paid"], r["by_cheque"],
                       r["left"], r["days_late"] or "", INST_ST.get(r["status"], (r["status"],) * 2)[1 if ar else 0]],
             "style": "item", "indent": 0} for r in d["rows"]]
    t = d["totals"]
    rows.append({"cells": [L("total", lang), "", "", "", "", t["amount"], t["paid"], t["by_cheque"], t["left"], "", ""], "style": "grand", "indent": 0})
    return {"title": "الأقساط والتحصيل" if ar else "Installments & collections", "subtitle": f"{L('as_of', lang)} {d['as_of']} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


def products_table(db, p, lang="en"):
    from . import sales
    ar = lang == "ar"
    d = sales.list_products(db, p.get("q"), p.get("kind"), p.get("category"), p.get("status"))
    cols = _mkcols([("Code", "الكود", "left"), ("Name", "الاسم", "left"), ("Kind", "النوع", "left"), ("Category", "التصنيف", "left"),
                    ("Price", "السعر", "right"), ("Status", "الحالة", "left"), ("In stock", "بالمخزن", "right", "num"),
                    ("Details", "المواصفات", "left")], lang)
    rows = [{"cells": [r["code"], (r["name_ar"] if ar and r["name_ar"] else r["name"]), r["kind"], r["category"], r["price"], r["status"],
                       r["in_stock"] if r["in_stock"] is not None else "", " · ".join(f"{a}: {b}" for a, b in r["specs"][:6])],
             "style": "item", "indent": 0} for r in d["rows"]]
    return {"title": "المنتجات والوحدات" if ar else "Products & units", "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


def bankrec_table(db, sid, lang="en"):
    """The classic reconciliation statement: bank balance ↔ book balance, with the items in between."""
    from . import bankrec
    ar = lang == "ar"
    d = bankrec.detail(db, sid)
    st, sm = d["statement"], d["summary"]
    cols = _mkcols([("Item", "البند", "left"), ("Date", "التاريخ", "left"), ("Voucher / reference", "القيد / المرجع", "left"),
                    ("Amount", "المبلغ", "right")], lang)
    R = lambda en, a_, v="", style="item", date="", ref="": {"cells": [a_ if ar else en, date, ref, v], "style": style, "indent": 1 if style == "item" else 0}  # noqa: E731
    rows = [R("Balance per bank statement", "الرصيد حسب كشف البنك", st["closing"], "subtotal", st["date_to"])]
    rows.append(R("Add: deposits in transit (in the books, not yet in the bank)", "يضاف: إيداعات بالطريق (بالدفاتر ولم تظهر بالبنك)", sm["deposits_in_transit"], "section"))
    rows += [R(i["memo"] or i["kind"], i["memo"] or i["kind"], i["amount"], date=i["date"], ref=f"#{i['id']} {i['party']}".strip()) for i in sm["outstanding_in"]]
    rows.append(R("Less: outstanding payments / cheques (in the books, not yet in the bank)", "يخصم: مدفوعات / شيكات معلقة (بالدفاتر ولم تُصرف)", -sm["outstanding_payments"], "section"))
    rows += [R(i["memo"] or i["kind"], i["memo"] or i["kind"], i["amount"], date=i["date"], ref=f"#{i['id']} {i['party']}".strip()) for i in sm["outstanding_out"]]
    rows.append(R("Adjusted bank balance", "الرصيد المعدل حسب البنك", sm["adjusted_bank"], "grand"))
    rows.append(R("Balance per books", "الرصيد حسب الدفاتر", sm["book"], "subtotal", st["date_to"]))
    rows.append(R("Bank items not yet in the books (charges, interest…)", "حركات بالبنك لم تُسجل بالدفاتر (عمولات، فوائد…)", sm["unrecorded"], "section"))
    rows += [R(ln["description"], ln["description"], ln["amount"], date=ln["date"], ref=ln["ref"]) for ln in d["lines"] if ln["status"] == "open"]
    if sm["late_booked"]:
        rows.append(R("In the bank by the end date, booked after it", "بالبنك قبل نهاية الفترة وسُجلت بعدها", sm["late_booked"], "section"))
    rows.append(R("Adjusted book balance", "الرصيد المعدل حسب الدفاتر", sm["adjusted_book"], "grand"))
    rows.append(R("Difference", "الفرق", sm["difference"], "grand"))
    acct = d["account"]
    title = ("مذكرة تسوية البنك" if ar else "Bank reconciliation") + f" — {acct['code']} {(acct.get('name_ar') if ar and acct.get('name_ar') else acct['name'])}"
    status = ("مسوّاة" if ar else "Reconciled") if st["status"] == "reconciled" else ("مفتوحة" if ar else "Open")
    return {"title": title, "subtitle": f"{st['date_from']} → {st['date_to']} · {status} · {db.setting('currency')}", "columns": cols, "rows": rows, "lang": lang}


def party_statement_table(db, party, date_from=None, date_to=None, lang="en"):
    from . import statements
    ar = lang == "ar"
    d = statements.statement(db, party, date_from, date_to)
    cols = _mkcols([("Date", "التاريخ", "left"), ("Voucher", "المستند", "left"), ("Ref", "المرجع", "left"),
                    ("Description", "البيان", "left"), ("Debit", "مدين", "right"),
                    ("Credit", "دائن", "right"), ("Balance", "الرصيد", "right")], lang)
    rows = [{"cells": ["", "", "", "رصيد أول المدة" if ar else "Opening balance", "", "", d["opening"]], "style": "subtotal", "indent": 0}]
    rows += [{"cells": [r["date"], f"#{r['id']} {r['kind']}", r["ref"], r["memo"], r["debit"] or "", r["credit"] or "", r["balance"]],
              "style": "item", "indent": 0} for r in d["rows"]]
    rows.append({"cells": [L("total", lang), "", "", "", d["debits"], d["credits"], d["closing"]], "style": "grand", "indent": 0})
    c = d["closing"]
    say = (("له علينا" if c < 0 else "عليه لنا") if ar else ("we owe them" if c < 0 else "they owe us")) if c else ("لا رصيد" if ar else "settled")
    rows.append({"cells": ["", "", "", ("الرصيد الختامي: " if ar else "Closing balance: ") + say, "", "", abs(c)], "style": "subtotal", "indent": 0})
    if c:
        rows.append({"cells": ["", "", "", "أعمار الرصيد" if ar else "AGE OF THE OPEN BALANCE", "", "", ""], "style": "section", "indent": 0})
        for b, v in zip(d["buckets_ar"] if ar else [x + " days" for x in d["buckets"]], d["aging"]):
            if v:
                rows.append({"cells": ["", "", "", b, "", "", abs(v)], "style": "item", "indent": 1})
    _later_rows(rows, d, ar, 7, 3)
    per = f"{d['date_from'] or '…'} → {d['date_to']}"
    return {"title": (f"كشف حساب — {d['party']}" if ar else f"Statement of account — {d['party']}"),
            "subtitle": f"{per} · {db.setting('currency')} · " + ("+ مدين = مستحق لنا، − = مستحق له" if ar else
                                                               "+ = they owe us, − = we owe them"),
            "columns": cols, "rows": rows, "lang": lang}


def _later_rows(rows, d, ar, ncols, at):
    """Entries dated after the as-of date are not in the balance above: list them, since a wrong year or month in a
    date otherwise shows up as a balance that looks reversed."""
    if not d.get("later"):
        return
    def row(text, style, indent):
        cells = [""] * ncols
        cells[at] = text
        rows.append({"cells": cells, "style": style, "indent": indent})
    row((f"⚠ قيود بتاريخ بعد {d['date_to']} — غير محسوبة أعلاه (راجع التاريخ إن كان خطأ)" if ar else
         f"⚠ DATED AFTER {d['date_to']} — NOT COUNTED ABOVE (check the date if this is a typo)"), "section", 0)
    for g in d["later"]:
        n, when = g["entries"], (g["first"] if g["first"] == g["last"] else f"{g['first']} → {g['last']}")
        side = ("عليه لنا" if g["net"] > 0 else "له علينا") if ar else ("they owe us" if g["net"] > 0 else "we owe them")
        amt = f"{abs(g['net']):,.2f} {side}, " if g["net"] else ""
        row(f"{g['party']} — {amt}{n} {'قيد' if ar else ('entry' if n == 1 else 'entries')}, {when}", "note", 1)


def aging_table(db, date_to=None, side="all", lang="en"):
    from . import statements
    ar = lang == "ar"
    d = statements.aging(db, date_to, side or "all")
    labels = d["buckets_ar"] if ar else [b + " days" for b in d["buckets"]]
    cols = _mkcols([("Contact", "الجهة", "left"), ("Last activity", "آخر حركة", "left")] +
                   [(b, b, "right") for b in labels] + [("Balance", "الرصيد", "right")], lang)
    rows = []
    for key, title in (("receivable", ("THEY OWE US", "مستحق لنا")), ("payable", ("WE OWE THEM", "مستحق علينا"))):
        rs = [r for r in d["rows"] if (r["balance"] > 0) == (key == "receivable")]
        if not rs:
            continue
        rows.append({"cells": [title[1] if ar else title[0]] + [""] * (len(labels) + 2), "style": "section", "indent": 0})
        for r in rs:
            rows.append({"cells": [r["party"], r["last"]] + [abs(v) or "" for v in r["aging"]] + [abs(r["balance"])],
                         "style": "item", "indent": 1})
        t = d[key]
        rows.append({"cells": [L("total", lang), ""] + [abs(v) for v in t["aging"]] + [abs(t["balance"])], "style": "subtotal", "indent": 0})
    _later_rows(rows, d, ar, len(labels) + 3, 0)
    return {"title": "أعمار الديون (العملاء والموردين)" if ar else "Aging — receivables & payables",
            "subtitle": f"{L('as_of', lang)} {d['date_to']} · {db.setting('currency')}", "columns": cols, "rows": rows, "lang": lang}


MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTHS_AR = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]


def budget_table(db, year=None, bu=None, through=None, lang="en"):
    from . import budgets
    ar = lang == "ar"
    d = budgets.report(db, year, bu, through)
    cols = _mkcols([("Account", "الحساب", "left"), ("Budget to date", "الموازنة حتى تاريخه", "right"), ("Actual to date", "الفعلي حتى تاريخه", "right"),
                    ("Variance (+ good)", "الانحراف (+ لصالحنا)", "right"), ("Variance %", "الانحراف %", "right", "num"),
                    ("Used %", "نسبة الاستهلاك %", "right", "num"), ("Full-year budget", "موازنة السنة", "right")], lang)
    rows = []
    for sec in d["sections"]:
        if not sec["rows"]:
            continue
        rows.append({"cells": [sec["label_ar"] if ar else sec["label"]] + [""] * 6, "style": "section", "indent": 0})
        for r in sec["rows"]:
            rows.append({"cells": [f"{r['account']} {r['name_ar'] if ar and r['name_ar'] else r['name']}", r["budget"], r["actual"],
                                   r["variance"], "" if r["pct"] is None else r["pct"], "" if r["used_pct"] is None else r["used_pct"],
                                   r["year_budget"]], "style": "item", "indent": 1})
        t = sec["total"]
        rows.append({"cells": [("إجمالي " + sec["label_ar"]) if ar else ("Total " + sec["label"].lower()), t["budget"], t["actual"],
                               t["variance"], "", "", t["year_budget"]], "style": "subtotal", "indent": 0})
    pr = d["profit"]
    rows.append({"cells": ["صافي الربح" if ar else "Net profit", pr["budget"], pr["actual"], pr["variance"], "", "", pr["year_budget"]],
                 "style": "grand", "indent": 0})
    bu_name = ""
    if d["bu_id"]:
        u = db.one("SELECT code, name FROM business_units WHERE id=?", (d["bu_id"],))
        bu_name = f" · {u['code']} {u['name']}"
    elif d["bu_id"] == 0:
        bu_name = " · " + ("المستوى العام للشركة" if ar else "company level (no business unit)")
    m = (MONTHS_AR if ar else MONTHS)[d["through"] - 1]
    return {"title": (f"الموازنة مقابل الفعلي {d['year']}" if ar else f"Budget vs actual {d['year']}"),
            "subtitle": (f"يناير – {m}" if ar else f"Jan – {m}") + f"{bu_name} · {db.setting('currency')}" +
                        ("" if d["has_budget"] else (" · لا توجد موازنة لهذه السنة بعد" if ar else " · no budget entered for this year yet")),
            "columns": cols, "rows": rows, "lang": lang}


PO_STATUS = {"draft": ("Draft — awaiting approval", "مسودة — بانتظار الاعتماد"), "approved": ("Approved", "معتمد"),
             "partial": ("Partly received", "استلام جزئي"), "received": ("Received", "تم الاستلام"),
             "closed": ("Closed", "مغلق"), "cancelled": ("Cancelled", "ملغى")}


def purchase_order_table(db, pid, lang="en"):
    from . import purchasing
    ar = lang == "ar"
    po = purchasing.get(db, pid)
    cols = _mkcols([("No.", "م", "left"), ("Description", "البيان", "left"), ("Unit", "الوحدة", "left"), ("Qty", "الكمية", "right", "num"),
                    ("Unit price", "سعر الوحدة", "right"), ("Amount", "القيمة", "right"), ("Received", "المستلم", "right", "num")], lang)
    rows = [{"cells": [i, ln["description"], ln["unit"], ln["qty"], ln["rate"], ln["amount"], ln["received"] or ""], "style": "item", "indent": 0}
            for i, ln in enumerate(po["lines"], start=1)]
    rows.append({"cells": ["", "الإجمالي قبل الضريبة" if ar else "Total before tax", "", "", "", po["total"], ""], "style": "subtotal", "indent": 0})
    if po["vat"]:
        rows.append({"cells": ["", f"{'ضريبة القيمة المضافة' if ar else 'VAT'} ({po['vat_code']})", "", "", "", po["vat"], ""], "style": "item", "indent": 0})
    rows.append({"cells": ["", "الإجمالي" if ar else "Order total", "", "", "", po["grand_total"], ""], "style": "grand", "indent": 0})
    info = [("Supplier", "المورد", po["supplier"]), ("Date", "التاريخ", po["date"]), ("Delivery by", "التسليم قبل", po["expected"]),
            ("Project", "المشروع", f"{po['bu_code']} {po['bu_name']}" if po["bu_code"] else ""), ("Deliver to", "التسليم إلى", po["warehouse"] or ""),
            ("Status", "الحالة", PO_STATUS[po["status"]][1 if ar else 0]), ("Approved by", "اعتمده", po["approved_by"]),
            ("Terms", "الشروط", po["terms"])]
    for en, a, v in info:
        if v:
            rows.append({"cells": ["", f"{a if ar else en}: {v}", "", "", "", "", ""], "style": "note", "indent": 0})
    return {"title": (f"أمر شراء {po['code']}" if ar else f"Purchase order {po['code']}"),
            "subtitle": f"{po['supplier']} · {po['date']} · {db.setting('currency')}", "columns": cols, "rows": rows, "lang": lang}


def purchase_orders_table(db, p, lang="en"):
    from . import purchasing
    ar = lang == "ar"
    d = purchasing.listing(db, p.get("q"), p.get("status"), p.get("supplier"), p.get("bu_id"))
    cols = _mkcols([("Order", "الأمر", "left"), ("Date", "التاريخ", "left"), ("Supplier", "المورد", "left"), ("Project", "المشروع", "left"),
                    ("Delivery by", "التسليم", "left"), ("Total", "الإجمالي", "right"), ("Received", "المستلم", "right"),
                    ("Still on order", "المتبقي", "right"), ("Status", "الحالة", "left")], lang)
    rows = [{"cells": [r["code"], r["date"], r["supplier"], r["bu_code"] or "", r["expected"], r["grand_total"], r["received_value"],
                       r["open_value"] or "", PO_STATUS[r["status"]][1 if ar else 0]], "style": "item", "indent": 0} for r in d["rows"]]
    rows.append({"cells": [L("total", lang), "", "", "", "", round(sum(r["grand_total"] for r in d["rows"]), 2),
                           round(sum(r["received_value"] for r in d["rows"]), 2), d["open_value"], ""], "style": "grand", "indent": 0})
    return {"title": "أوامر الشراء" if ar else "Purchase orders", "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')}",
            "columns": cols, "rows": rows, "lang": lang}


SUB_KINDS = {"maintenance": ("Maintenance", "صيانة"), "warranty": ("Warranty", "ضمان"), "insurance": ("Insurance", "تأمين"),
             "software": ("Software / app", "برنامج / تطبيق"), "support": ("Support", "دعم فني"), "addon": ("Add-on", "إضافة"),
             "other": ("Other", "أخرى")}


def subscriptions_table(db, p, lang="en"):
    from . import services
    ar = lang == "ar"
    d = services.listing(db, p.get("q"), p.get("status"), p.get("kind"), p.get("customer"))
    st = {"active": ("Active", "سارية"), "paused": ("Paused", "موقوفة مؤقتاً"), "cancelled": ("Cancelled", "ملغاة"), "ended": ("Ended", "منتهية")}
    cols = _mkcols([("Code", "الكود", "left"), ("Customer", "العميل", "left"), ("Service", "الخدمة", "left"), ("Covers", "يغطي", "left"),
                    ("Price", "السعر", "right"), ("Every (months)", "كل (شهور)", "right", "num"), ("Monthly value", "القيمة الشهرية", "right"),
                    ("Next invoice", "الفاتورة القادمة", "left"), ("Ends", "تنتهي", "left"), ("Status", "الحالة", "left")], lang)
    rows = [{"cells": [r["code"] + (f" ↳ {r['parent_code']}" if r["parent_code"] else ""), r["customer"], r["title"], r["covers"], r["price"],
                       r["period_months"], r["monthly"] or "", r["next_bill"] if r["status"] == "active" else "", r["end"],
                       st[r["status"]][1 if ar else 0]], "style": "item", "indent": 1 if r["parent_code"] else 0} for r in d["rows"]]
    rows.append({"cells": [("الإيراد الشهري المتكرر" if ar else "Monthly recurring revenue"), "", "", "", "", "", d["mrr"], "", "", ""],
                 "style": "grand", "indent": 0})
    return {"title": "الاشتراكات وعقود الصيانة" if ar else "Subscriptions & maintenance contracts",
            "subtitle": f"{L('as_of', lang)} {acc.today()} · {db.setting('currency')} · " +
                        (f"إيراد سنوي متكرر {d['arr']:,.2f}" if ar else f"annual recurring {d['arr']:,.2f}"),
            "columns": cols, "rows": rows, "lang": lang}


def tax_table(db, date_from=None, date_to=None, lang="en"):
    t = acc.tax_summary(db, date_from, date_to)
    ar = lang == "ar"
    cols = [{"label": "البند" if ar else "Item", "align": "left"}, {"label": "الرقم الضريبي" if ar else "Tax ID",
                                                                  "align": "left"},
            {"label": "الأساس الخاضع" if ar else "Taxable base", "align": "right", "fmt": "money"},
            {"label": "الضريبة" if ar else "Tax", "align": "right", "fmt": "money"}]
    rows = [{"cells": ["ضريبة القيمة المضافة" if ar else "VALUE ADDED TAX", "", "", ""], "style": "section", "indent": 0}]
    for b in t["vat_by_code"]:
        side = ("مبيعات" if ar else "Sales") if b["side"] == "sales" else ("مشتريات" if ar else "Purchases")
        rows.append({"cells": [f"{side} — {b['code']} ({b['documents']})", "", b["net"], b["vat"]], "style": "item",
                     "indent": 1})
    rows += [{"cells": ["ضريبة المخرجات" if ar else "Output VAT (on sales)", "", "", t["output_vat"]], "style": "subtotal",
              "indent": 0},
             {"cells": ["ضريبة المدخلات" if ar else "Input VAT (on purchases)", "", "", -t["input_vat"]],
              "style": "subtotal", "indent": 0},
             {"cells": ["صافي الضريبة المستحقة" if ar else "Net VAT payable / (refundable)", "", "", t["net_vat"]],
              "style": "grand", "indent": 0},
             {"cells": ["ضريبة الخصم والتحصيل — المخصومة من الموردين" if ar else
                        "WITHHOLDING TAX DEDUCTED FROM SUPPLIERS (to remit)", "", "", ""], "style": "section", "indent": 0}]
    for x in t["wht_withheld_from_suppliers"]:
        rows.append({"cells": [x["party"] or "—", x["tax_id"], "", x["amount"]], "style": "item", "indent": 1})
    rows.append({"cells": ["الرصيد المستحق للمصلحة" if ar else "Balance owed to the Tax Authority", "", "",
                           t["wht_payable_balance"]], "style": "total", "indent": 0})
    rows.append({"cells": ["ضريبة الخصم والتحصيل — المخصومة بواسطة العملاء" if ar else
                           "WITHHOLDING TAX DEDUCTED BY CLIENTS (tax credit)", "", "", ""], "style": "section", "indent": 0})
    for x in t["wht_deducted_by_clients"]:
        rows.append({"cells": [x["party"] or "—", x["tax_id"], "", x["amount"]], "style": "item", "indent": 1})
    rows.append({"cells": ["رصيد دائن لدى المصلحة" if ar else "Tax credit balance", "", "", t["wht_receivable_balance"]],
                 "style": "total", "indent": 0})
    pers = [("", date_from or "…", acc.clean_date(date_to))]
    return {"title": "الإقرار الضريبي — ضريبة القيمة المضافة والخصم والتحصيل" if ar else "Tax report — VAT & withholding",
            "subtitle": _subtitle(db, lang, pers, None), "columns": cols, "rows": rows, "lang": lang}


def build(db, name, p):
    """Dispatch a report by name with query params dict p."""
    lang = p.get("lang") or "en"
    bu = p.get("bu_id") or None
    if name == "payroll_register":
        from . import payroll
        return payroll.register_table(db, p.get("run"), lang)
    if name == "partner_pack":
        from . import partner_pack
        return partner_pack.table(db, p.get("date_from") or None, p.get("date_to") or None, lang)
    if name == "funding_schedule":
        from . import fundsched
        return fundsched.table(db, p.get("date_to") or None, lang)
    if name == "profit_categories":
        from . import categories
        return categories.table(db, p.get("date_from") or None, p.get("date_to") or None, bu, lang)
    if name == "income_statement":
        return income_statement_table(db, p.get("date_from"), p.get("date_to"), p.get("compare"), bu,
                                      p.get("detail", "1") != "0", lang, p.get("allocate", "1") != "0")
    if name == "balance_sheet":
        return balance_sheet_table(db, p.get("date_to"), p.get("date_from"), p.get("compare"), bu, lang)
    if name == "cash_flow":
        return cash_flow_table(db, p.get("date_from"), p.get("date_to"), p.get("compare"), bu, lang)
    if name == "trial_balance":
        return trial_balance_table(db, p.get("date_from"), p.get("date_to"), bu, lang)
    if name == "general_ledger":
        return general_ledger_table(db, p.get("account"), p.get("date_from"), p.get("date_to"), bu, p.get("party"), lang)
    if name == "bu_pnl":
        return bu_table(db, p.get("date_from"), p.get("date_to"), p.get("status"), lang)
    if name == "payables":
        return payables_table(db, p.get("date_to"), lang)
    if name == "project_cash":
        return project_cash_table(db, p.get("date_from"), p.get("date_to"), lang)
    if name == "stock":
        return stock_table(db, p.get("warehouse_id"), p.get("q"), lang)
    if name == "tax":
        return tax_table(db, p.get("date_from"), p.get("date_to"), lang)
    if name == "receivables":
        return receivables_table(db, p.get("date_from"), p.get("date_to"), lang)
    if name == "rent_fleet":
        return rent_fleet_table(db, lang)
    if name == "rentals":
        return rentals_table(db, p.get("status"), lang)
    if name == "sales":
        return sales_table(db, p, lang)
    if name == "bank_reconciliation":
        if not str(p.get("id") or "").isdigit():
            raise acc.ValidationError("Choose a bank statement")
        return bankrec_table(db, p.get("id"), lang)
    if name == "installments":
        return installments_table(db, p, lang)
    if name == "products":
        return products_table(db, p, lang)
    if name == "loans":
        return loans_table(db, p.get("direction"), p.get("state"), lang)
    if name == "loan_schedule":
        return loan_schedule_table(db, p.get("code") or p.get("id"), lang)
    if name == "custody":
        return custody_table(db, p.get("state"), lang)
    if name == "custody_statement":
        return custody_statement_table(db, p.get("code") or p.get("id"), lang)
    if name == "materials_used":
        return materials_table(db, p, lang)
    if name == "negotiations":
        return negotiations_table(db, p.get("date_from"), p.get("date_to"), bu, lang)
    if name == "party_statement":
        return party_statement_table(db, p.get("party"), p.get("date_from"), p.get("date_to"), lang)
    if name == "aging":
        return aging_table(db, p.get("date_to"), p.get("side"), lang)
    if name == "budget":
        return budget_table(db, p.get("year"), p.get("bu_id", ""), p.get("through"), lang)
    if name == "subscriptions":
        return subscriptions_table(db, p, lang)
    if name == "purchase_order":
        return purchase_order_table(db, p.get("id") or p.get("code"), lang)
    if name == "purchase_orders":
        return purchase_orders_table(db, p, lang)
    if name == "fixed_assets":
        return fixed_assets_table(db, p.get("date_to"), p.get("q"), p.get("state"), lang)
    raise acc.ValidationError(f"Unknown report {name!r}")


REPORT_PERMS = {"payroll_register": "payroll.runs", "partner_pack": ("reports.partner_pack", "partners.equity"), "funding_schedule": ("partners.equity", "reports.partner_pack"), "income_statement": "reports.income_statement", "profit_categories": "reports.income_statement", "balance_sheet": "reports.balance_sheet",
                "cash_flow": "reports.cash_flow", "trial_balance": "reports.trial_balance",
                "general_ledger": "accounting.ledger", "bu_pnl": "business_units.profitability", "project_cash": "business_units.profitability",
                "payables": "purchasing.payables", "stock": "stock.balances", "receivables": "sales.receivables",
                "tax": "tax.report", "fixed_assets": "assets.register",
                "negotiations": "purchasing.negotiations", "materials_used": ("projects.overview", "stock.moves", "purchasing.costs"), "loans": "accounting.loans", "loan_schedule": "accounting.loans",
                "custody": "purchasing.custody", "custody_statement": "purchasing.custody", "rent_fleet": "rentals.fleet",
                "rentals": "rentals.contracts", "sales": "shop.orders", "installments": "shop.installments",
                "products": "shop.products", "bank_reconciliation": "accounting.bankrec",
                "party_statement": ("reports.statement", "sales.parties", "purchasing.contacts", "sales.receivables", "purchasing.payables"),
                "aging": ("reports.aging", "sales.receivables", "purchasing.payables"), "budget": ("reports.budget", "accounting.budgets"),
                "purchase_order": "purchasing.orders", "purchase_orders": "purchasing.orders",
                "subscriptions": "shop.subscriptions"}


# ==================================================================== writers
def _fmt(v, fmt, lang="en"):
    if v in ("", None):
        return ""
    if fmt == "pct" or (isinstance(v, float) and fmt == "money" and False):
        return f"{v * 100:.1f}%"
    if isinstance(v, (int, float)):
        if fmt == "num":
            return f"{v:,.2f}".rstrip("0").rstrip(".")
        s = f"{abs(v):,.2f}"
        return f"({s})" if v < 0 else s
    return str(v)


def _cell_fmt(table, row, i):
    if row["style"] == "note" and i > 0 and isinstance(row["cells"][i], float) and abs(row["cells"][i]) <= 10:
        return "pct"
    return table["columns"][i].get("fmt")


def to_html(table, company="", auto_print=False, font="system", size=12, logo=None, wm=None):
    rtl = table.get("lang") == "ar"
    head = "".join(f'<th class="{c.get("align", "left")}">{html.escape(str(c["label"]))}</th>'
                   for c in table["columns"])
    body = []
    for r in table["rows"]:
        tds = []
        for i, v in enumerate(r["cells"]):
            f = _cell_fmt(table, r, i)
            cls = table["columns"][i].get("align", "left")
            if cls == "right" and isinstance(v, str) and len(v) > 24:
                cls = "right txt"   # a sentence in a number column: wrap it, keep its own direction
            neg = " neg" if isinstance(v, (int, float)) and v < 0 else ""
            pad = f' style="padding-inline-start:{8 + 16 * r.get("indent", 0)}px"' if i == 0 else ""
            rest = r["cells"][i + 1:]
            span = len(rest) + 1 if "txt" in cls and not any(x not in (None, "") for x in rest) else 1
            tds.append(f'<td class="{cls}{neg}"{pad}{f" colspan={span}" if span > 1 else ""}>{html.escape(_fmt(v, f))}</td>')
            if span > 1:
                break
        body.append(f'<tr class="{r["style"]}">{"".join(tds)}</tr>')
    fam = "Arial, 'Segoe UI', Tahoma, sans-serif" if font in ("system", "", None) else f"'{font}', Arial, sans-serif"
    return f"""<!doctype html><html lang="{'ar' if rtl else 'en'}" dir="{'rtl' if rtl else 'ltr'}"><head><meta charset="utf-8">
<title>{html.escape(table['title'])}</title><style>
@page {{ size: A4 {"landscape" if len(table['columns']) > 6 else "portrait"}; margin: 14mm 12mm; }}
body{{font-family:{fam};font-size:{size}px;color:#111;margin:24px}}
h1{{font-size:{size + 8}px;margin:0}} .co{{font-weight:700;color:#1f3a5f;letter-spacing:.5px}}
.sub{{color:#555;margin:4px 0 14px}} table{{border-collapse:collapse;width:100%}}
th{{background:#1f3a5f;color:#fff;padding:7px 8px;font-size:{size - 1}px}} td{{padding:5px 8px;border-bottom:1px solid #e5e7eb}}
.right{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap;direction:ltr;unicode-bidi:isolate}} .left{{text-align:start}}
[dir=rtl] .right{{text-align:left}} .right.txt{{white-space:normal;direction:inherit;text-align:start}}
tr.section td{{background:#dfe6ee;font-weight:700;text-transform:uppercase}} tr.header td{{font-weight:700;color:#1f3a5f}}
tr.subtotal td{{font-weight:600;border-top:1px solid #9ca3af}} tr.total td{{font-weight:700;background:#f3f4f6;border-top:1px solid #111}}
tr.grand td{{font-weight:800;background:#e8eef6;border-top:2px solid #111;border-bottom:3px double #111}}
tr.note td{{color:#6b7280;font-style:italic;font-size:{size - 1}px}} .neg{{color:#b00020}}
.foot{{margin-top:16px;color:#888;font-size:{size - 2}px}} @media print{{body{{margin:0}} .noprint{{display:none}} .wm{{position:fixed}}}}
.hd{{display:flex;justify-content:space-between;align-items:flex-start;gap:16px}} .hd img{{max-height:80px;max-width:220px}}
.wm{{position:fixed;top:50%;left:50%;transform:translate(-50%,-50%);width:55%;max-width:520px;z-index:-1;pointer-events:none}}
{"table{table-layout:fixed} td{word-wrap:break-word;white-space:normal} tr.header td{white-space:normal} td:first-child,th:first-child{width:" + table["label_width"] + "}" if table.get("label_width") else ""}
</style></head><body>{f'<img class="wm" src="data:image/png;base64,{wm["data"]}" alt="">' if wm else ''}
<div class="hd"><div><div class="co">{html.escape(company)}</div><h1>{html.escape(table['title'])}</h1>
<div class="sub">{html.escape(table.get('subtitle', ''))}</div></div>{f'<img src="data:{logo["mime"]};base64,{logo["data"]}" alt="">' if logo else ''}</div>
<button class="noprint" onclick="window.print()" style="margin-bottom:10px">Print / Save as PDF</button>
<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>
<div class="foot">Generated {_dt.datetime.now():%Y-%m-%d %H:%M}</div>
{'<script>window.onload=()=>setTimeout(()=>window.print(),300)</script>' if auto_print else ''}</body></html>"""


def to_xlsx(table, company=""):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        raise acc.ValidationError("Excel export needs openpyxl:  python -m pip install openpyxl") from None
    wb = Workbook()
    ws = wb.active
    ws.title = table["title"][:31].replace("/", "-").replace("(", "").replace(")", "")[:31]
    rtl = table.get("lang") == "ar"
    ws.sheet_view.rightToLeft = rtl
    ws.sheet_view.showGridLines = False
    n = len(table["columns"])
    navy, band, light = "1F3A5F", "DFE6EE", "F3F4F6"
    ws.cell(1, 1, company).font = Font(name="Arial", bold=True, size=11, color=navy)
    ws.cell(2, 1, table["title"]).font = Font(name="Arial", bold=True, size=14)
    ws.cell(3, 1, table.get("subtitle", "")).font = Font(name="Arial", size=9, color="555555")
    thin = Side(style="thin", color="111111")
    for i, c in enumerate(table["columns"], start=1):
        cell = ws.cell(5, i, c["label"])
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="right" if c.get("align") == "right" else "left", wrap_text=True)
    r0 = 6
    for ri, row in enumerate(table["rows"]):
        st = row["style"]
        for ci, v in enumerate(row["cells"], start=1):
            f = _cell_fmt(table, row, ci - 1)
            cell = ws.cell(r0 + ri, ci, v if v != "" else None)
            bold = st in ("section", "header", "subtotal", "total", "grand")
            cell.font = Font(name="Arial", bold=bold, italic=st == "note", color="6B7280" if st == "note" else "111111")
            if isinstance(v, (int, float)):
                cell.number_format = "0.0%" if f == "pct" else ("#,##0.00;[Red](#,##0.00);-" if f == "money" else
                                                                  ("#,##0" if float(v).is_integer() else "#,##0.###"))   # no stray "." on whole numbers
                cell.alignment = Alignment(horizontal="right")
            if ci == 1 and row.get("indent"):
                cell.alignment = Alignment(indent=row["indent"] * 2)
            if st == "section":
                cell.fill = PatternFill("solid", fgColor=band)
            elif st in ("total", "grand"):
                cell.fill = PatternFill("solid", fgColor=light if st == "total" else band)
                cell.border = Border(top=thin, bottom=Side(style="double", color="111111") if st == "grand" else None)
            elif st == "subtotal":
                cell.border = Border(top=Side(style="thin", color="9CA3AF"))
    ws.column_dimensions["A"].width = 48
    from openpyxl.utils import get_column_letter
    for i in range(2, n + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.page_setup.orientation = "landscape" if n > 6 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    buf = io.BytesIO()
    xlsafe.no_formulas(wb)
    wb.save(buf)
    return buf.getvalue()


def _x(s):
    return html.escape(str(s), quote=False)


def to_docx(table, company="", logo=None, wm=None):
    """Minimal but valid Word document (stdlib only); optional logo (top) and watermark (page header)."""
    rtl = table.get("lang") == "ar"
    bidi = "<w:bidi/>" if rtl else ""
    rtl_run = "<w:rtl/>" if rtl else ""
    n = len(table["columns"])
    landscape = n > 6
    page_w, page_h = (16838, 11906) if landscape else (11906, 16838)
    usable = page_w - 1440
    first = int(usable * (0.34 if n > 3 else 0.5))
    other = int((usable - first) / max(n - 1, 1))
    widths = [first] + [other] * (n - 1)

    def para(t, bold=False, size=20, color="111111", align=None, italic=False):
        jc = f'<w:jc w:val="{align}"/>' if align else ""
        return (f'<w:p><w:pPr>{bidi}{jc}<w:spacing w:after="60"/></w:pPr><w:r><w:rPr>'
                f'<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>{"<w:b/><w:bCs/>" if bold else ""}'
                f'{"<w:i/>" if italic else ""}<w:color w:val="{color}"/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/>'
                f'{rtl_run}</w:rPr><w:t xml:space="preserve">{_x(t)}</w:t></w:r></w:p>')

    def cell(t, w, bold=False, fill=None, align="left", color="111111", italic=False, indent=0, top=None, bottom=None):
        shd = f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>' if fill else ""
        borders = ""
        if top or bottom:
            borders = "<w:tcBorders>" + (f'<w:top w:val="{top}" w:sz="6" w:color="111111"/>' if top else "") + \
                      (f'<w:bottom w:val="{bottom}" w:sz="6" w:color="111111"/>' if bottom else "") + "</w:tcBorders>"
        jc = {"right": "end" if not rtl else "start", "left": "start" if not rtl else "end"}.get(align, "start")
        jc = "right" if align == "right" else "left"
        if rtl:
            jc = "left" if align == "right" else "right"
        ind = f'<w:ind w:start="{indent * 240}"/>' if indent else ""
        return (f'<w:tc><w:tcPr><w:tcW w:w="{w}" w:type="dxa"/>{borders}{shd}</w:tcPr><w:p><w:pPr>{bidi}'
                f'<w:jc w:val="{jc}"/>{ind}<w:spacing w:after="0"/></w:pPr><w:r><w:rPr>'
                f'<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>{"<w:b/><w:bCs/>" if bold else ""}'
                f'{"<w:i/>" if italic else ""}<w:color w:val="{color}"/><w:sz w:val="17"/><w:szCs w:val="17"/>{rtl_run}'
                f'</w:rPr><w:t xml:space="preserve">{_x(t)}</w:t></w:r></w:p></w:tc>')

    rows = ['<w:tr><w:trPr><w:tblHeader/></w:trPr>' +
            "".join(cell(c["label"], widths[i], True, "1F3A5F", c.get("align", "left"), "FFFFFF")
                    for i, c in enumerate(table["columns"])) + "</w:tr>"]
    for r in table["rows"]:
        st = r["style"]
        fill = {"section": "DFE6EE", "total": "F3F4F6", "grand": "DFE6EE"}.get(st)
        bold = st in ("section", "header", "subtotal", "total", "grand")
        rows.append("<w:tr>" + "".join(
            cell(_fmt(v, _cell_fmt(table, r, i)), widths[i], bold, fill, table["columns"][i].get("align", "left"),
                 "6B7280" if st == "note" else ("B00020" if isinstance(v, (int, float)) and v < 0 else "111111"),
                 st == "note", r.get("indent", 0) if i == 0 else 0,
                 "single" if st in ("subtotal", "total", "grand") else None, "double" if st == "grand" else None)
            for i, v in enumerate(r["cells"])) + "</w:tr>")
    grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    tbl = (f'<w:tbl><w:tblPr>{"<w:bidiVisual/>" if rtl else ""}<w:tblW w:w="{usable}" w:type="dxa"/>'
           f'<w:tblBorders><w:insideH w:val="single" w:sz="2" w:color="E5E7EB"/></w:tblBorders>'
           f'<w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid>{grid}</w:tblGrid>{"".join(rows)}</w:tbl>')
    logo_xml = ""
    if logo:
        h = 60
        w = round(h * (logo["w"] or 1) / (logo["h"] or 1))
        logo_xml = (f'<w:p><w:pPr>{bidi}<w:jc w:val="{"left" if rtl else "right"}"/><w:spacing w:after="0"/></w:pPr><w:r><w:drawing>'
                    f'<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{w * 9525}" cy="{h * 9525}"/><wp:docPr id="1" '
                    'name="Logo"/><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
                    '<pic:nvPicPr><pic:cNvPr id="1" name="Logo"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill><a:blip r:embed="rIdLogo"/>'
                    '<a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/>'
                    f'<a:ext cx="{w * 9525}" cy="{h * 9525}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
                    '</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>')
    body = (logo_xml + para(company, True, 22, "1F3A5F") + para(table["title"], True, 32) +
            para(table.get("subtitle", ""), size=18, color="555555") + tbl +
            para(f"Generated {_dt.datetime.now():%Y-%m-%d %H:%M}", size=16, color="888888"))
    orient = ' w:orient="landscape"' if landscape else ""
    from .documents import REL, SHAPETYPE75, b64bytes, rels
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/'
          'officeDocument/2006/relationships" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/'
          '2006/picture" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
          'xmlns:w10="urn:schemas-microsoft-com:office:word"')
    doc_rels, parts, overrides = [], [], ""
    hdr_ref = ""
    if logo:
        ext = "png" if logo["mime"] == "image/png" else "jpeg"
        parts.append((f"word/media/logo.{ext}", b64bytes(logo["data"])))
        doc_rels.append(f'<Relationship Id="rIdLogo" Type="{REL}image" Target="media/logo.{ext}"/>')
    if wm:
        parts.append(("word/media/watermark.png", b64bytes(wm["data"])))
        hpt = 300
        parts.append(("word/_rels/header1.xml.rels", rels([f'<Relationship Id="rIdWm" Type="{REL}image" Target="media/watermark.png"/>'])))
        parts.append(("word/header1.xml", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:hdr {ns}><w:p><w:r><w:pict>'
                      f'{SHAPETYPE75}<v:shape id="WordPictureWatermark1" type="#_x0000_t75" style="position:absolute;margin-left:0;'
                      f'margin-top:0;width:{hpt}pt;height:{round(hpt * (wm["h"] or 1) / (wm["w"] or 1))}pt;z-index:-251656192;'
                      'mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;'
                      'mso-position-vertical-relative:margin" o:allowincell="f"><v:imagedata r:id="rIdWm" o:title="watermark"/>'
                      '</v:shape></w:pict></w:r></w:p></w:hdr>'))
        doc_rels.append(f'<Relationship Id="rIdHdr" Type="{REL}header" Target="header1.xml"/>')
        overrides = ('<Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                     'wordprocessingml.header+xml"/>')
        hdr_ref = '<w:headerReference w:type="default" r:id="rIdHdr"/>'
    doc = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>' + body +
           f'<w:sectPr>{hdr_ref}<w:pgSz w:w="{page_w}" w:h="{page_h}"{orient}/><w:pgMar w:top="850" w:right="720" '
           'w:bottom="850" w:left="720" w:header="400" w:footer="400" w:gutter="0"/></w:sectPr></w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-'
                   'package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" '
                   'ContentType="image/png"/><Default Extension="jpeg" ContentType="image/jpeg"/><Override '
                   'PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                   'wordprocessingml.document.main+xml"/>' + overrides + '</Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.'
                   'openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.'
                   'openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                   '</Relationships>')
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels", rels(doc_rels))
        for name, data in parts:
            z.writestr(name, data)
    return buf.getvalue()


def to_csv(table):
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in table["columns"]])
    for r in table["rows"]:
        w.writerow(["" if v is None else v for v in r["cells"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")
