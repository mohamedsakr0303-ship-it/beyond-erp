"""Egyptian Tax Authority (ETA) e-invoicing.

Two ways to use it:
  1. Generate for import — ETA v1.0 document JSON (single or batch) and an Excel batch
     (Header / Lines / Taxes) for the ETA portal uploader or an integrator.
  2. Direct link — OAuth2 client-credentials login, signing through a local signing
     service (your USB token / HSM), submission, status refresh and cancellation.

Signing: ETA requires a CAdES-BES signature made with the taxpayer's e-seal/token for
documentTypeVersion "1.0". Tokens are hardware, so this app calls a *signer service* you run
next to the token (see README → E-invoicing). Contract:
    POST <signer_url>  {"serialized": "<canonical string>", "sha256": "<hex>", "document": {...}}
    ->  {"signature": "<base64 CAdES-BES>"}
Version "0.9" documents need no signature and are accepted by the pre-production system only.
"""
import datetime as _dt
import hashlib
import io
import json
import re
import urllib.error
import urllib.parse
import urllib.request

from . import accounting as acc

ENDPOINTS = {
    "preprod": ("https://id.preprod.eta.gov.eg", "https://api.preprod.invoicing.eta.gov.eg"),
    "prod": ("https://id.eta.gov.eg", "https://api.invoicing.eta.gov.eg"),
}
# app units -> ETA unit codes (UN/ECE Rec 20 as used by ETA)
UNIT_CODES = {"pcs": "EA", "عدد": "EA", "ea": "EA", "task": "EA", "lump sum": "EA", "مقطوعية": "EA",
              "m2": "MTK", "م٢": "MTK", "m²": "MTK", "sqm": "MTK", "m.t": "MTR", "ml": "MTR", "م.ط": "MTR",
              "m": "MTR", "kg": "KGM", "ton": "TNE", "l": "LTR", "day": "DAY", "hour": "HUR", "month": "MON"}


def _r(v, n=5):
    return round(float(v or 0), n)


def settings(db):
    s = {r["key"]: r["value"] for r in db.q("SELECT key, value FROM settings WHERE key LIKE 'eta_%' OR key LIKE 'company_%'")}
    return s


def public_settings(db):
    s = settings(db)
    s = {k: v for k, v in s.items() if k.startswith("eta_")}
    from .db import ETA_SECRETS
    for k in ETA_SECRETS:
        s[k] = "••••••••" if s.get(k) else ""
    return s


def _party(db, name):
    return db.one("SELECT * FROM parties WHERE name=?", (name,)) or {}


def build_document(db, extract_id, version=None):
    """Return (document, problems). problems is a list of human-readable blocking issues."""
    x = acc.get_extract(db, extract_id)
    if not x:
        raise acc.ValidationError("Invoice not found")
    s = settings(db)
    version = version or s.get("eta_doc_version") or "1.0"
    problems = []
    rin = (s.get("company_tax_id") or "").strip()
    if not re.fullmatch(r"\d{9}", rin):
        problems.append("Company tax registration number (9 digits) is missing — Settings → Company")
    if not s.get("eta_activity_code"):
        problems.append("Taxpayer activity code is missing — Settings → E-invoicing")
    for k in ("eta_governate", "eta_region", "eta_street", "eta_building"):
        if not s.get(k):
            problems.append(f"Issuer address field '{k[4:]}' is missing — Settings → E-invoicing")
    cust_name = x["customer"] or x.get("bu_customer") or x["bu_name"]
    p = _party(db, cust_name)
    rtype = p.get("receiver_type") or "B"
    receiver_id = (p.get("tax_id") if rtype == "B" else p.get("national_id")) or ""
    if not p:
        problems.append(f"Customer '{cust_name}' is not in Customers & Suppliers — add it with tax ID and address")
    elif rtype == "B" and not re.fullmatch(r"\d{9}", receiver_id or ""):
        problems.append(f"Customer '{cust_name}' needs a 9-digit tax registration number")
    elif rtype == "P" and x["gross"] >= 50000 and not receiver_id:
        problems.append("Invoices of 50,000 EGP or more to individuals need the customer's national ID")
    if x["status"] != "posted":
        problems.append("Post the invoice before issuing it as an e-invoice")
    vat = db.one("SELECT * FROM tax_codes WHERE code=?", (x["vat_code"],)) if x["vat_code"] else None
    wht = db.one("SELECT * FROM tax_codes WHERE code=?", (x["wht_code"],)) if x["wht_code"] else None
    codes = {}
    for r in db.q("SELECT item_no, description, eta_code FROM rates WHERE eta_code<>''"):
        codes[r["item_no"]] = codes[r["description"]] = r["eta_code"]
    lines, tax_totals = [], {}
    for ln in x["lines"]:
        net = _r(ln["qty"] * ln["price"])
        items = []
        if vat:
            amt = _r(net * vat["rate"])
            items.append({"taxType": vat["eta_type"] or "T1", "amount": amt, "subType": vat["eta_subtype"] or "V009",
                          "rate": _r(vat["rate"] * 100, 2)})
        if wht:
            amt = _r(net * wht["rate"])
            items.append({"taxType": wht["eta_type"] or "T4", "amount": amt, "subType": wht["eta_subtype"] or "W001",
                          "rate": _r(wht["rate"] * 100, 2)})
        for it in items:
            tax_totals[it["taxType"]] = _r(tax_totals.get(it["taxType"], 0) + it["amount"])
        total = _r(net + sum(it["amount"] for it in items if it["taxType"] != "T4")
                   - sum(it["amount"] for it in items if it["taxType"] == "T4"))
        code = codes.get(ln["item_no"]) or codes.get(ln["description"]) or s.get("eta_default_item_code") or ""
        if not code:
            problems.append(f"Line '{ln['description'][:40]}' has no ETA item code (price list → ETA code, "
                            "or a default in Settings → E-invoicing)")
        lines.append({
            "description": ln["description"][:500], "itemType": "GS1" if code.isdigit() else "EGS",
            "itemCode": code, "unitType": UNIT_CODES.get((ln["unit"] or "").strip().lower(),
                                                        s.get("eta_default_unit") or "EA"),
            "quantity": _r(ln["qty"]), "internalCode": str(ln["item_no"] or ln["seq"] + 1),
            "salesTotal": net, "total": total, "valueDifference": 0.0, "totalTaxableFees": 0.0, "netTotal": net,
            "itemsDiscount": 0.0, "unitValue": {"currencySold": "EGP", "amountEGP": _r(ln["price"])},
            "discount": {"rate": 0.0, "amount": 0.0}, "taxableItems": items})
    net_amount = _r(sum(line["netTotal"] for line in lines))
    doc = {
        "issuer": {"address": {"branchID": s.get("eta_branch_id") or "0", "country": "EG",
                               "governate": s.get("eta_governate", ""), "regionCity": s.get("eta_region", ""),
                               "street": s.get("eta_street", ""), "buildingNumber": s.get("eta_building", ""),
                               "postalCode": s.get("eta_postal_code", ""), "floor": "", "room": "", "landmark": "",
                               "additionalInformation": ""},
                   "type": "B", "id": rin, "name": s.get("company_name", "")},
        "receiver": {"address": {"country": p.get("country") or "EG", "governate": p.get("governate", ""),
                                 "regionCity": p.get("region", ""), "street": p.get("street", ""),
                                 "buildingNumber": p.get("building", ""), "postalCode": p.get("postal_code", ""),
                                 "floor": "", "room": "", "landmark": "", "additionalInformation": ""},
                     "type": rtype, "id": receiver_id, "name": cust_name},
        "documentType": "I", "documentTypeVersion": version,
        "dateTimeIssued": f"{x['date']}T{_dt.datetime.utcnow():%H:%M:%S}Z",
        "taxpayerActivityCode": s.get("eta_activity_code", ""), "internalID": x["number"],
        "purchaseOrderReference": "", "purchaseOrderDescription": "", "salesOrderReference": x.get("boq_title") or "",
        "salesOrderDescription": "", "proformaInvoiceNumber": "",
        "invoiceLines": lines, "totalDiscountAmount": 0.0, "totalSalesAmount": net_amount, "netAmount": net_amount,
        "taxTotals": [{"taxType": k, "amount": v} for k, v in tax_totals.items()],
        "totalAmount": _r(sum(line["total"] for line in lines)), "extraDiscountAmount": 0.0,
        "totalItemsDiscountAmount": 0.0,
    }
    if x["date"] > acc.today():
        problems.append("The issue date cannot be in the future")
    return doc, sorted(set(problems))


def serialize(obj, name=None):
    """ETA canonical serialization (the string that is hashed and signed)."""
    if isinstance(obj, dict):
        out = []
        for k, v in obj.items():
            if isinstance(v, list):
                out.append(f'"{k.upper()}"' + "".join(f'"{k.upper()}"' + serialize(e) for e in v))
            else:
                out.append(f'"{k.upper()}"' + serialize(v))
        return "".join(out)
    if isinstance(obj, list):
        return "".join(serialize(e) for e in obj)
    return '"' + (json.dumps(obj, ensure_ascii=False) if not isinstance(obj, str) else obj.replace('"', '\\"')) + '"'


def sign(db, doc):
    s = settings(db)
    if doc.get("documentTypeVersion") == "0.9":
        return doc
    url = s.get("eta_signer_url")
    if not url:
        raise acc.ValidationError("Set the signing service URL (Settings → E-invoicing) or use version 0.9 "
                                  "on pre-production for testing")
    ser = serialize({k: v for k, v in doc.items() if k != "signatures"})
    body = json.dumps({"serialized": ser, "sha256": hashlib.sha256(ser.encode("utf-8")).hexdigest(),
                       "document": doc}).encode("utf-8")
    res = _http("POST", url, body, {"Content-Type": "application/json"})
    sig = res.get("signature") if isinstance(res, dict) else None
    if not sig:
        raise acc.ValidationError("The signing service did not return a signature")
    return {**doc, "signatures": [{"signatureType": "I", "value": sig}]}


def _http(method, url, body=None, headers=None, timeout=40):
    import sys
    if sys.platform == "emscripten":
        raise acc.ValidationError("The direct ETA link needs the desktop or self-hosted version — browsers block "
                                  "these connections. Use 'Generate for import' (JSON / Excel) here instead.")
    req = urllib.request.Request(url, data=body, method=method, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            raw = r.read()
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            detail = json.loads(raw.decode("utf-8"))
        except (ValueError, UnicodeDecodeError):
            detail = raw.decode("utf-8", "replace")[:400]
        raise acc.ValidationError(f"{url.split('/')[2]} answered {e.code}: {detail}") from None
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        raise acc.ValidationError(f"Could not reach {url.split('/')[2]}: {e}") from None
    try:
        return json.loads(raw.decode("utf-8")) if raw else {}
    except ValueError:
        return {"raw": raw.decode("utf-8", "replace")}


def _token(db):
    s = settings(db)
    if not s.get("eta_client_id") or not s.get("eta_client_secret"):
        raise acc.ValidationError("Enter the ETA client ID and secret (Settings → E-invoicing)")
    id_base, _ = ENDPOINTS.get(s.get("eta_env") or "preprod", ENDPOINTS["preprod"])
    body = urllib.parse.urlencode({"grant_type": "client_credentials", "client_id": s["eta_client_id"],
                                   "client_secret": s["eta_client_secret"], "scope": "InvoicingAPI"}).encode()
    res = _http("POST", f"{id_base}/connect/token", body, {"Content-Type": "application/x-www-form-urlencoded"})
    if not res.get("access_token"):
        raise acc.ValidationError("ETA login failed — check client ID / secret and environment")
    return res["access_token"]


def _api_base(db):
    return ENDPOINTS.get(settings(db).get("eta_env") or "preprod", ENDPOINTS["preprod"])[1]


def test_connection(db):
    _token(db)
    return {"ok": True, "environment": settings(db).get("eta_env")}


def submit(db, extract_ids, user=""):
    docs, results = [], []
    for xid in extract_ids:
        doc, problems = build_document(db, xid)
        if problems:
            raise acc.ValidationError(f"Invoice #{xid}: " + "; ".join(problems))
        docs.append((xid, sign(db, doc)))
    token = _token(db)
    res = _http("POST", f"{_api_base(db)}/api/v1/documentsubmissions",
                json.dumps({"documents": [d for _, d in docs]}, ensure_ascii=False).encode("utf-8"),
                {"Content-Type": "application/json", "Authorization": f"Bearer {token}"})
    by_internal = {d["internalID"]: xid for xid, d in docs}
    for a in res.get("acceptedDocuments", []):
        xid = by_internal.get(a.get("internalId"))
        if xid:
            db.x("UPDATE extracts SET eta_uuid=?, eta_status='Submitted', eta_submission=?, eta_error='' WHERE id=?",
                 (a.get("uuid", ""), res.get("submissionId", ""), xid))
            results.append({"id": xid, "status": "Submitted", "uuid": a.get("uuid")})
    for r in res.get("rejectedDocuments", []):
        xid = by_internal.get(r.get("internalId"))
        err = json.dumps(r.get("error", {}), ensure_ascii=False)[:1500]
        if xid:
            db.x("UPDATE extracts SET eta_status='Rejected', eta_error=? WHERE id=?", (err, xid))
            results.append({"id": xid, "status": "Rejected", "error": err})
    db.log(user, "eta.submit", json.dumps(results, ensure_ascii=False)[:500])
    return {"submission_id": res.get("submissionId"), "results": results}


def refresh_status(db, extract_id):
    x = db.one("SELECT eta_uuid FROM extracts WHERE id=?", (extract_id,))
    if not x or not x["eta_uuid"]:
        raise acc.ValidationError("This invoice has not been submitted")
    res = _http("GET", f"{_api_base(db)}/api/v1/documents/{x['eta_uuid']}/details",
                headers={"Authorization": f"Bearer {_token(db)}"})
    status = res.get("status") or res.get("documentStatus") or "Unknown"
    errs = json.dumps(res.get("validationResults", {}), ensure_ascii=False)[:1500] if status == "Invalid" else ""
    db.x("UPDATE extracts SET eta_status=?, eta_error=? WHERE id=?", (status, errs, extract_id))
    return {"status": status, "error": errs}


def cancel(db, extract_id, reason):
    x = db.one("SELECT eta_uuid FROM extracts WHERE id=?", (extract_id,))
    if not x or not x["eta_uuid"]:
        raise acc.ValidationError("This invoice has not been submitted")
    _http("PUT", f"{_api_base(db)}/api/v1/documents/state/{x['eta_uuid']}/state",
          json.dumps({"status": "cancelled", "reason": acc.text(reason, "reason", required=True)}).encode(),
          {"Content-Type": "application/json", "Authorization": f"Bearer {_token(db)}"})
    db.x("UPDATE extracts SET eta_status='Cancel requested' WHERE id=?", (extract_id,))
    return {"status": "Cancel requested"}


# ---------------------------------------------------------------- batch export
def batch(db, extract_ids, signed=False):
    docs, problems = [], {}
    for xid in extract_ids:
        doc, pr = build_document(db, xid)
        if pr:
            problems[doc["internalID"]] = pr
        docs.append(sign(db, doc) if signed and not pr else doc)
    return docs, problems


def batch_json(db, extract_ids, signed=False):
    docs, problems = batch(db, extract_ids, signed)
    return json.dumps({"documents": docs, "_problems": problems} if problems else {"documents": docs},
                      ensure_ascii=False, indent=2).encode("utf-8")


def batch_xlsx(db, extract_ids):
    """Excel batch: Header, Lines and Taxes sheets, one row per document / line / tax item."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise acc.ValidationError("Excel export needs openpyxl:  python -m pip install openpyxl") from None
    docs, problems = batch(db, extract_ids)
    wb = Workbook()
    head = wb.active
    head.title = "Header"
    hcols = ["internalID", "documentType", "documentTypeVersion", "dateTimeIssued", "taxpayerActivityCode",
             "issuer.id", "issuer.name", "issuer.branchID", "receiver.type", "receiver.id", "receiver.name",
             "receiver.country", "receiver.governate", "receiver.regionCity", "receiver.street",
             "receiver.buildingNumber", "totalSalesAmount", "totalDiscountAmount", "netAmount", "T1 (VAT)",
             "T4 (WHT)", "totalAmount", "problems"]
    lines = wb.create_sheet("Lines")
    lcols = ["internalID", "line", "description", "itemType", "itemCode", "unitType", "quantity",
             "unitValue.amountEGP", "salesTotal", "discount.amount", "netTotal", "total", "internalCode"]
    taxes = wb.create_sheet("Taxes")
    tcols = ["internalID", "line", "taxType", "subType", "rate", "amount"]
    for ws, cols in ((head, hcols), (lines, lcols), (taxes, tcols)):
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.freeze_panes = "A2"
    for d in docs:
        tt = {t["taxType"]: t["amount"] for t in d["taxTotals"]}
        r, i = d["receiver"], d["issuer"]
        head.append([d["internalID"], d["documentType"], d["documentTypeVersion"], d["dateTimeIssued"],
                     d["taxpayerActivityCode"], i["id"], i["name"], i["address"]["branchID"], r["type"], r["id"],
                     r["name"], r["address"]["country"], r["address"]["governate"], r["address"]["regionCity"],
                     r["address"]["street"], r["address"]["buildingNumber"], d["totalSalesAmount"],
                     d["totalDiscountAmount"], d["netAmount"], tt.get("T1", 0), tt.get("T4", 0), d["totalAmount"],
                     "; ".join(problems.get(d["internalID"], []))])
        for n, ln in enumerate(d["invoiceLines"], start=1):
            lines.append([d["internalID"], n, ln["description"], ln["itemType"], ln["itemCode"], ln["unitType"],
                          ln["quantity"], ln["unitValue"]["amountEGP"], ln["salesTotal"], ln["discount"]["amount"],
                          ln["netTotal"], ln["total"], ln["internalCode"]])
            for t in ln["taxableItems"]:
                taxes.append([d["internalID"], n, t["taxType"], t["subType"], t["rate"], t["amount"]])
    for ws in (head, lines, taxes):
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 2, 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
