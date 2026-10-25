"""Funding facility (the funding partner's account is the company's bank and what he fronts is owed to him), the
facility limit with a temporary line, weekly statement, cash in & out by project, and the V2.0 Cost Allocation sheet."""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from erp import accounting as a  # noqa: E402
from erp import interchange, partner_pack, reports, settlement  # noqa: E402
from erp.db import Database  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAVE_XL = True
except ImportError:
    HAVE_XL = False


def books(mode="facility"):
    db = Database(":memory:")
    for code in ("P-1", "P-2"):
        db.x("INSERT INTO business_units(code,name,basis,status,start_date,overhead_pct) VALUES (?,?,?,?,?,0.5)",
             (code, code, "Fixed Price", "In Progress", "2026-01-01"))
    db.x("INSERT INTO partners(name,default_share) VALUES ('Funder',0.5),('Other',0.5)")
    db.set_setting("funding_partner", "Funder")
    db.set_setting("funding_mode", mode)
    settlement.apply_funding_mode(db, mode)
    fund = db.one("SELECT code FROM accounts WHERE role='funding_partner'")["code"]
    # every payment leaves his account, every client payment lands in it
    a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Works", "qty": 1, "rate": 1000,
                               "paid": 600, "paid_from": fund, "date": "2026-02-02"})       # Monday
    a.record_project_cost(db, {"bu_code": "P-2", "vendor": "Sup", "description": "Works", "qty": 1, "rate": 500,
                               "paid": 500, "paid_from": fund, "date": "2026-02-10"})
    a.record_journal(db, {"memo": "Client payment", "date": "2026-02-11", "bu_code": "P-1",
                          "lines": [{"account": fund, "debit": 300}, {"account": "4010", "credit": 300}]})
    return db, fund


class FundingFacility(unittest.TestCase):
    def test_facility_is_a_liability_and_capital_mode_is_equity(self):
        db, fund = books("facility")
        acct = db.one("SELECT type, subtype FROM accounts WHERE code=?", (fund,))
        self.assertEqual((acct["type"], acct["subtype"]), ("Liability", "short_term_debt"))
        bs = a.balance_sheet(db)
        line = next(x for x in bs["current_liabilities"] if x["code"] == fund)
        self.assertEqual(line["amount"], 800)                 # fronted: 1,100 paid out − 300 in
        self.assertTrue(bs["balanced"])
        settlement.apply_funding_mode(db, "capital")
        bs = a.balance_sheet(db)
        self.assertFalse(any(x["code"] == fund for x in bs["current_liabilities"]))
        self.assertTrue(any(x["code"] == fund for x in bs["equity"]))
        self.assertTrue(bs["balanced"])

    def test_company_money_in_his_account_is_an_asset(self):
        db, fund = books("facility")
        a.record_journal(db, {"memo": "Big client payment", "date": "2026-02-12", "bu_code": "P-2",
                              "lines": [{"account": fund, "debit": 5000}, {"account": "4010", "credit": 5000}]})
        bs = a.balance_sheet(db)
        held = next(x for x in bs["current_assets"] if x["code"] == fund)
        self.assertEqual(held["amount"], 4200)
        self.assertFalse(any(x["code"] == fund for x in bs["current_liabilities"]))
        self.assertTrue(bs["balanced"])

    def test_limit_temporary_line_status_and_weeks(self):
        db, fund = books("facility")
        db.set_setting("funding_limit", "1000")
        db.set_setting("funding_extra", "600")
        db.set_setting("funding_extra_shrinks", "yes")
        f = settlement.facility(db, "2026-02-20")
        self.assertEqual((f["limit"], f["extra_now"]), (1000, 0))          # no partner loan outstanding: no extra line
        self.assertEqual(f["fronted"], 800)
        self.assertEqual(f["owed"], 400)                                    # 1,000 cost − 600 paid on the first deal
        self.assertEqual(f["after_suppliers"], 1200)
        self.assertEqual(f["state"], "over")
        self.assertEqual(f["peak"], 1100)
        self.assertEqual([(w["week"], w["out"], w["in"], w["fronted"]) for w in f["weeks"]],
                         [("2026-02-02", 600, 0, 600), ("2026-02-09", 500, 300, 800), ("2026-02-16", 0, 0, 800)])
        # a loan to a partner opens the temporary line, which shrinks to what is still owed
        from erp import loans
        loans.save(db, {"direction": "lent", "party": "Other", "account": "1310",
                        "start_date": "2026-02-15", "principal": 400, "opening_balance": 400, "schedule_type": "none"})
        db.set_setting("funding_limit", "1300")
        f = settlement.facility(db, "2026-02-20")
        self.assertEqual((f["partner_loans"], f["extra_now"], f["limit"]), (400, 400, 1700))
        self.assertEqual((f["fronted"], f["after_suppliers"]), (1200, 1600))    # the loan left his account too
        self.assertEqual(f["state"], "near")                                # 1,600 of 1,700 = 94% ≥ 90% (critical)
        flags = " | ".join(x["flag"] for x in a.review_flags(db, None, "2026-02-20"))
        self.assertIn("funding is near its", flags)

    def test_distribution_test_uses_room_in_the_facility(self):
        db, fund = books("facility")
        db.set_setting("funding_limit", "10000")
        d = partner_pack.compute(db, None, "2026-02-20")
        self.assertEqual(d["dist"]["facility_room"], 10000 - 800)
        self.assertNotIn("repay the funding partner first", d["dist"]["verdict"][0])

    def test_cash_in_and_out_by_project(self):
        db, fund = books("facility")
        rows = {r["code"]: r for r in reports.project_cash(db)}
        self.assertEqual((rows["P-1"]["in"], rows["P-1"]["out"], rows["P-1"]["owed"], rows["P-1"]["after"]), (300, 600, 400, -700))
        self.assertEqual((rows["P-2"]["out"], rows["P-2"]["owed"]), (500, 0))
        t = reports.project_cash_table(db)
        self.assertEqual(t["rows"][-1]["cells"][4], -800)

    def test_settings_round_trip_keeps_the_mode(self):
        db, fund = books("facility")
        db.set_setting("funding_extra", "600")
        dst = Database(":memory:")
        interchange.import_data(dst, interchange.read_json(interchange.to_json(db).encode()), mode="replace", user="t")
        self.assertEqual(dst.one("SELECT type FROM accounts WHERE code=?", (fund,))["type"], "Liability")
        self.assertEqual(dst.setting("funding_extra"), "600")
        self.assertEqual(a.balance_sheet(dst)["total_liabilities"], a.balance_sheet(db)["total_liabilities"])


class CostPlusOverheadAndPool(unittest.TestCase):
    def test_project_overhead_in_cost_plus_base_is_optional(self):
        db = Database(":memory:")
        db.x("INSERT INTO business_units(code,name,basis,fee_pct,status) VALUES ('P-1','P','Cost-Plus %',0.1,'In Progress')")
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "S", "description": "W", "qty": 1, "rate": 1000, "date": "2026-02-01"})
        a.record_expense(db, {"bu_code": "P-1", "description": "Site office", "amount": 200, "account": "6010", "date": "2026-02-01"})
        exp = lambda: a.bu_summaries(db)[0]["expected_billing"]  # noqa: E731
        self.assertEqual(exp(), 1100)
        db.set_setting("costplus_overhead", "yes")
        self.assertEqual(exp(), 1320)

    def test_shared_direct_cost_with_no_project_is_pooled(self):
        db = Database(":memory:")
        for c in ("P-1", "P-2"):
            db.x("INSERT INTO business_units(code,name,status,overhead_pct) VALUES (?,?,'In Progress',0.5)", (c, c))
        a.record_journal(db, {"memo": "Shared materials", "date": "2026-02-01",
                              "lines": [{"account": "5010", "debit": 100}, {"account": "2010", "credit": 100}]})
        from erp import allocation
        r = allocation.allocate(db)
        self.assertEqual(r["pool"], 10000)
        self.assertEqual(r["company"], 0)


class LegacyV20(unittest.TestCase):
    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_weights_default_rule_and_facility(self):
        import openpyxl
        from erp import legacy
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Project Register"
        ws.append(["Project ID", "Client Name", "Project Type", "Location", "Total Area (Sqm)", "Start Date", "Status",
                   "Total Contract Value", "Billing Basis / Fee Type", "Supervision Fee %", "Cost-Plus fee % (type it for Cost-Plus projects)"])
        ws.append(["PRJ-1", "A", "Villa", "X", 1, "2026-01-01", "In Progress", 0, "Cost-Plus %", 0.42, 0.2])
        ws.append(["PRJ-2", "B", "Villa", "X", 1, "2026-01-01", "In Progress", 0, "Flexible / Settled", 0.1, None])
        ca = wb.create_sheet("Cost Allocation")
        ca.append(["COST ALLOCATION"])
        ca.append(["Allocation rule", "Billable to client?", "Pooled amount", "PRJ-1", "PRJ-2", "Total weight", "What this rule does"])
        ca.append(["Main", "No", 0, 3, 1, 4, "split"])
        ca.append(["Supervision", "Yes", 0, 1, 1, 2, "split"])
        ca.append([None, None, 0, 0, 0, 0, None])
        ca.append(["TOTAL POOLED", None, 0])
        ca.append([])
        ca.append(["Default rule: shared lines with no Alloc Rule", None, "Supervision"])
        ca.append(["Count project-tagged overhead (6xxx lines with a project) in the Cost-Plus base?", None, "Yes"])
        ca.append([])
        ca.append(["RESULT — amount each project receives"])
        ca.append(["Rule", "Billable?", "Pooled", "PRJ-1", "PRJ-2"])
        ca.append(["Main", "No", 0, 0, 0])
        coa = wb.create_sheet("Chart of Accounts")
        coa.append(["Code", "Account Name", "Type", "Normal"])
        coa.append(["2400", "Funding partner facility — money he has fronted", "Liability", "Credit"])
        bt = wb.create_sheet("Balances Tracker")
        bt.append(["SECTION 5 - THE FUNDING PARTNER'S ACCOUNT: MONEY IN AND OUT  (Funder)"])
        for label, val in (("Base line he always covers", 250000), ("Temporary extra line, agreed for the partner loan", 180000),
                           ("Does the extra line shrink as the loan is repaid? (Yes / No)", "Yes")):
            bt.append([None] * 5 + [label, None, None, val])
        bt.append([None] * 5 + ["Warn me at (amber) · critical at (red) →", None, 0.75, 0.9])
        gj = wb.create_sheet("General Journal")
        gj.append(["GENERAL JOURNAL"])
        gj.append(["JE #", "Date", "Source / Type", "Project", "Stage", "Account", "Description", "Vendor / Party",
                   "Deal Ref", "Qty", "Unit", "Rate", "Debit", "Credit", "Note", "Alloc Rule"])
        for r in (["JE-1", "2026-01-05", "Overhead", None, None, "6010", "Rent", None, None, None, None, None, 400, None, None, None],
                  ["JE-1", "2026-01-05", "Overhead", None, None, "1020", "Rent", None, None, None, None, None, None, 400, None, None],
                  ["JE-2", "2026-01-06", "Overhead", None, None, "6010", "Fuel", None, None, None, None, None, 100, None, None, "Typo rule"],
                  ["JE-2", "2026-01-06", "Overhead", None, None, "1020", "Fuel", None, None, None, None, None, None, 100, None, None],
                  ["JE-AUTO-FUND", "2026-01-07", "Funding facility (auto)", None, None, "1020", "Funding", "Funder", None, None, None, None, 500, None, None, None],
                  ["JE-AUTO-FUND", "2026-01-07", "Funding facility (auto)", None, None, "2400", "Funding", "Funder", None, None, None, None, None, 500, None, None]):
            gj.append(r)
        path = str(Path(tempfile.mkdtemp()) / "v20.xlsx")
        wb.save(path)
        data = legacy.read_legacy_workbook(path)
        rules = {r["name"]: r for r in data["Allocation Rules"]}
        self.assertEqual(set(rules), {"Main", "Supervision"})                # the results table is not read as rules
        self.assertEqual(rules["Main"]["split"], "PRJ-1: 75%; PRJ-2: 25%")
        self.assertEqual((rules["Supervision"]["is_default"], rules["Main"]["is_default"]), ("Yes", "No"))
        settings = {r["key"]: r["value"] for r in data["Settings"]}
        self.assertEqual((settings["funding_mode"], settings["funding_limit"], settings["funding_extra"], settings["funding_extra_shrinks"],
                          settings["funding_warn"], settings["funding_critical"], settings["costplus_overhead"]),
                         ("facility", "250000", "180000", "yes", "0.75", "0.9", "yes"))
        self.assertEqual(next(b for b in data["Business Units"] if b["code"] == "PRJ-1")["fee_pct"], 0.2)
        db = Database(":memory:")
        interchange.import_data(db, data, mode="replace", user="t")
        fund = db.one("SELECT code, type FROM accounts WHERE role='funding_partner'")
        self.assertEqual(fund["type"], "Liability")
        self.assertIsNone(db.one("SELECT 1 FROM journal_entries WHERE ref='JE-AUTO-FUND'"))
        self.assertEqual(a.balance_sheet(db)["current_liabilities"][0]["amount"], 500)   # fronted 500, owed to him
        from erp import allocation
        by = {db.one("SELECT code FROM business_units WHERE id=?", (k,))["code"]: v for k, v in allocation.allocate(db)["by_bu"].items()}
        # rent (no rule) and fuel (unknown rule) both follow the default 'Supervision' rule: half each, billable
        self.assertEqual(by["PRJ-1"], {"billable": 25000, "non_billable": 0})


if __name__ == "__main__":
    unittest.main()
