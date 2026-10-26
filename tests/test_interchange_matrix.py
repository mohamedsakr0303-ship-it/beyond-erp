"""Import / export matrix: every importable sheet's templates, hostile files, encodings and formula-looking text."""
import io
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from erp import accounting as a  # noqa: E402
from erp import bankrec, demo, importkit, industries, interchange, reports, workbook  # noqa: E402
from erp.db import Database  # noqa: E402

PARTIES = "Parties"


def header():
    return importkit.template_csv(PARTIES).decode("utf-8-sig").splitlines()[0]


def read_parties(db, raw, fn):
    rows, _ = importkit.read(db, PARTIES, raw, fn)
    return rows


class Templates(unittest.TestCase):
    def test_every_sheet_has_templates_and_an_empty_one_is_refused_clearly(self):
        db = Database(":memory:")
        names = importkit.importable()
        self.assertGreater(len(names), 40)
        for n in names:
            with self.subTest(sheet=n):
                for raw, fn in ((importkit.template_xlsx(db, n, "en"), "t.xlsx"), (importkit.template_xlsx(db, n, "ar"), "t.xlsx"),
                                (importkit.template_csv(n), "t.csv")):
                    with self.assertRaisesRegex(a.ValidationError, "No rows to import"):
                        importkit.read(db, n, raw, fn)

    def test_the_csv_templates_example_row_is_never_imported(self):
        db = Database(":memory:")
        raw = importkit.template_csv(PARTIES) + "Acme,customer\n".encode()
        self.assertEqual([r["name"] for r in read_parties(db, raw, "t.csv")], ["Acme"])
        raw = importkit.template_csv("Receipt Vouchers") + "2026-10-06,,RV-9,Paid,,500,Cash,,2100,,\n".encode()
        rows, _ = importkit.read(db, "Receipt Vouchers", raw, "t.csv")
        self.assertEqual([r["ref"] for r in rows], ["RV-9"])

    def test_demo_layouts_read_back_in_full_in_both_languages(self):
        for ind in (None, "scaffolding", "contracting"):
            for n, rows in demo.sheets(ind).items():
                if n not in importkit.importable() or not rows:
                    continue
                for lang in ("en", "ar"):
                    with self.subTest(industry=ind, sheet=n, lang=lang):
                        db = Database(":memory:")
                        if ind:
                            industries.apply(db, ind, True)
                        given = demo.translate_sheets({n: rows}, lang)[n]
                        got, _ = importkit.read(db, n, importkit.template_xlsx(db, n, lang, given), "t.xlsx")
                        self.assertEqual(len(got), len(given))


class HostileFiles(unittest.TestCase):
    def setUp(self):
        self.db = Database(":memory:")

    def test_broken_files_are_refused_with_a_clear_message(self):
        good = importkit.template_xlsx(self.db, PARTIES, "en")
        for raw in (good[:500], b"PK\x03\x04" + b"x" * 300):
            with self.assertRaisesRegex(a.ValidationError, "not a readable Excel workbook"):
                read_parties(self.db, raw, "x.xlsx")
        with self.assertRaisesRegex(a.ValidationError, "neither an Excel workbook nor a text"):
            read_parties(self.db, b"\x00\x01garbage" * 50, "x.xlsx")
        with self.assertRaisesRegex(a.ValidationError, r"old Excel file \(\.xls\)"):
            read_parties(self.db, b"\xd0\xcf\x11\xe0" + b"\x00" * 600, "x.xls")
        with self.assertRaisesRegex(a.ValidationError, "Match a column"):
            read_parties(self.db, b"foo,bar\n1,2\n", "x.csv")
        # the whole-books import and the bank statement reader say the same
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            f.write(good[:500])
        with self.assertRaisesRegex(a.ValidationError, "not a readable Excel workbook"):
            interchange.read_xlsx(f.name)
        Path(f.name).unlink()
        with self.assertRaisesRegex(a.ValidationError, "not a readable Excel workbook"):
            bankrec._table(b"PK\x03\x04" + b"x" * 300, "s.xlsx")

    def test_a_csv_renamed_xlsx_is_still_read(self):
        rows = read_parties(self.db, (header() + "\nAcme,customer\n").encode(), "x.xlsx")
        self.assertEqual([r["name"] for r in rows], ["Acme"])

    def test_arabic_text_in_every_encoding_excel_saves(self):
        name = "شركة الأمل للتجارة"
        for enc in ("utf-8", "utf-8-sig", "utf-16", "cp1256"):
            with self.subTest(enc=enc):
                raw = (header() + f"\n{name},customer\n").encode(enc)
                self.assertEqual([r["name"] for r in read_parties(self.db, raw, "x.csv")], [name])
                self.assertEqual([r["name"] for r in interchange.read_csv(PARTIES, raw)[PARTIES]], [name])
        tab = (header().replace(",", "\t") + f"\n{name}\tcustomer\n").encode("utf-16")       # Excel's "Unicode Text"
        self.assertEqual([r["name"] for r in read_parties(self.db, tab, "x.txt")], [name])

    def test_semicolon_csv(self):
        raw = (header().replace(",", ";") + "\nAcme;customer\n").encode()
        self.assertEqual([r["name"] for r in read_parties(self.db, raw, "x.csv")], ["Acme"])

    def test_the_same_party_twice_is_one_record(self):
        demo.load(self.db, "t")
        rows = read_parties(self.db, (header() + "\nAcme,customer\nAcme,customer\n").encode(), "x.csv")
        interchange.import_data(self.db, {PARTIES: rows}, user="t")
        self.assertEqual(self.db.q("SELECT COUNT(*) n FROM parties WHERE name='Acme'")[0]["n"], 1)


class FormulaText(unittest.TestCase):
    """Text typed as a formula is exported as text and imports back unchanged."""

    def setUp(self):
        self.db = Database(":memory:")
        self.risky = ["=1+1", '=HYPERLINK("http://example.com","x")', "@SUM(A1)", "+20 100 000", "-not a number"]
        rows = read_parties(self.db, (header() + "".join(f'\n"{v.replace(chr(34), chr(34) * 2)}",customer' for v in self.risky)).encode(), "x.csv")
        interchange.import_data(self.db, {PARTIES: rows}, user="t")

    def names(self, rows):
        return sorted(r["name"] for r in rows if r["name"] in self.risky)

    def test_workbook_cells_are_text(self):
        import openpyxl
        raw = workbook.build_workbook(self.db, template=False)
        ws = openpyxl.load_workbook(io.BytesIO(raw))["Parties"]
        cells = [c for row in ws.iter_rows() for c in row if c.value in self.risky]
        self.assertEqual(len(cells), len(self.risky))
        self.assertTrue(all(c.data_type == "s" for c in cells), [(c.coordinate, c.data_type) for c in cells])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            f.write(raw)
        self.assertEqual(self.names(interchange.read_xlsx(f.name)[PARTIES]), sorted(self.risky))
        Path(f.name).unlink()

    def test_csv_cells_are_quoted_and_come_back_unchanged(self):
        raw = interchange.to_csv(self.db, PARTIES)
        lines = raw.decode("utf-8-sig").splitlines()
        self.assertFalse([ln for ln in lines if ln[:1] in "=+-@"])
        self.assertEqual(self.names(interchange.read_csv(PARTIES, raw)[PARTIES]), sorted(self.risky))
        self.assertEqual(self.names(read_parties(self.db, raw, "p.csv")), sorted(self.risky))

    def test_report_export_has_no_formulas(self):
        import openpyxl
        table = {"title": "T", "columns": [{"label": "Name", "align": "left"}, {"label": "Note", "align": "left"}],
                 "rows": [reports._row(v, [v]) for v in self.risky]}
        raw = reports.to_xlsx(table)
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        self.assertFalse([c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"])


if __name__ == "__main__":
    unittest.main()
