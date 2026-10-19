"""BeyondERP test-suite (stdlib unittest — run:  python -m unittest discover tests)."""
import http.cookiejar
import json
import sys
import tempfile
import threading
import datetime
import types
import unittest
import urllib.error
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from erp import accounting as a  # noqa: E402
from erp import assets, auth, entries, equity, interchange, reports  # noqa: E402
from erp.db import Database  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAVE_XL = True
except ImportError:
    HAVE_XL = False


def seeded():
    db = Database(":memory:")
    db.x("INSERT INTO business_units(code,name,basis,fee_pct,overhead_pct) VALUES ('PRJ-1','Villa','Cost-Plus %',0.25,0.6)")
    db.x("INSERT INTO business_units(code,name,basis,overhead_pct) VALUES ('PRJ-2','Flat','Fixed Price',0.4)")
    db.x("INSERT INTO partners(name,default_share) VALUES ('A',0.5),('B',0.5)")
    db.x("INSERT INTO warehouses(name) VALUES ('Site Store')")
    a.record_partner_tx(db, {"partner": "A", "type": "Capital", "amount": 100000, "date": "2026-01-02"})
    a.record_project_cost(db, {"bu_code": "PRJ-1", "description": "Cement", "vendor": "Sup", "qty": 10, "rate": 500,
                               "paid": 3000, "savings": 200, "date": "2026-02-01"})
    a.record_expense(db, {"description": "Rent", "amount": 5000, "account": "6010", "date": "2026-02-01"})
    a.record_payroll(db, {"bu_code": "PRJ-1", "worker": "Ali", "qty": 5, "rate": 300, "paid": 1000, "date": "2026-02-05"})
    a.record_collection(db, {"bu_code": "PRJ-1", "amount": 20000, "date": "2026-02-10"})
    a.stock_move(db, {"type": "Receipt", "warehouse": "Main Store", "item": "Tiles", "qty": 100, "rate": 50,
                      "paid": 5000, "date": "2026-02-11"})
    a.stock_move(db, {"type": "Transfer", "warehouse": "Main Store", "to_warehouse": "Site Store", "item": "Tiles",
                      "qty": 30, "date": "2026-02-12"})
    a.stock_move(db, {"type": "Issue", "warehouse": "Site Store", "item": "Tiles", "qty": 10, "bu_code": "PRJ-1",
                      "date": "2026-02-13"})
    x = a.save_extract(db, {"bu_code": "PRJ-1", "number": "INV-1", "date": "2026-03-01",
                            "lines": [{"description": "Tiling", "qty": 10, "price": 1000}]})
    a.post_extract(db, x)
    a.record_payment(db, {"party": "Sup", "amount": 1000, "date": "2026-03-02"})
    a.record_journal(db, {"memo": "Depreciation", "date": "2026-03-31",
                          "lines": [{"account": "7010", "debit": 500}, {"account": "1490", "credit": 500}]})
    a.record_journal(db, {"memo": "Tax", "date": "2026-03-31",
                          "lines": [{"account": "7050", "debit": 700}, {"account": "2220", "credit": 700}]})
    a.record_partner_tx(db, {"partner": "B", "type": "Distribution", "amount": 2000, "date": "2026-03-31"})
    return db


class Engine(unittest.TestCase):
    def setUp(self):
        self.db = seeded()

    def test_unbalanced_entry_rejected(self):
        with self.assertRaises(a.ValidationError):
            a.post_entry(self.db, date="2026-01-01", kind="Journal Voucher",
                         lines=[{"account": "1020", "debit": 10}, {"account": "4010", "credit": 9}])

    def test_income_statement_waterfall(self):
        pl = a.income_statement(self.db)
        self.assertEqual(pl["net_sales"], 30000.0)
        self.assertEqual(pl["cost_of_sales"], 4800 + 1500 + 500)
        self.assertEqual(pl["ebitda"], round(pl["gross_profit"] - pl["operating_expenses"], 2))
        self.assertEqual(pl["ebit"], round(pl["ebitda"] - 500, 2))
        self.assertEqual(pl["net_profit"], round(pl["ebt"] - 700, 2))

    def test_statements_tie_out(self):
        for check in a.audit(self.db):
            self.assertTrue(check["ok"], check)
        cf = a.cash_flow(self.db, "2026-02-01", "2026-03-31")
        self.assertTrue(cf["reconciles"])
        self.assertTrue(a.trial_balance(self.db, "2026-02-01", "2026-03-31")["balanced"])

    def test_overhead_allocation(self):
        bus = {b["code"]: b for b in a.bu_summaries(self.db)}
        self.assertEqual(bus["PRJ-1"]["allocated_overhead"], 3000.0)  # 60% of 5,000 shared rent
        self.assertEqual(bus["PRJ-2"]["allocated_overhead"], 2000.0)

    def test_account_renumbering_follows_role(self):
        a.save_account(self.db, {"code": "1025", "name": "Bank CIB", "type": "Asset", "subtype": "cash"}, old_code="1020")
        self.assertEqual(self.db.acct("bank"), "1025")
        a.record_collection(self.db, {"bu_code": "PRJ-2", "amount": 10})
        self.assertTrue(all(c["ok"] for c in a.audit(self.db)))

    def test_system_account_cannot_be_deleted(self):
        with self.assertRaises(a.ValidationError):
            a.delete_account(self.db, "1020")

    def test_void_and_reverse(self):
        e = a.record_expense(self.db, {"description": "Oops", "amount": 99})
        a.void_entry(self.db, e, "typo")
        rid = a.reverse_entry(self.db, a.record_expense(self.db, {"description": "Oops2", "amount": 50}))
        self.assertTrue(rid)
        self.assertTrue(all(c["ok"] for c in a.audit(self.db)))

    def test_stock_cannot_go_negative(self):
        with self.assertRaises(a.ValidationError):
            a.stock_move(self.db, {"type": "Issue", "warehouse": "Site Store", "item": "Tiles", "qty": 1000,
                                   "bu_code": "PRJ-1"})

    def test_payment_limited_to_balance(self):
        with self.assertRaises(a.ValidationError):
            a.record_payment(self.db, {"party": "Sup", "amount": 999999})

    def test_voiding_stock_entry_restores_levels(self):
        before = {(r["name"], r["warehouse"]): r["qty"] for r in a.stock_report(self.db)}
        e = a.stock_move(self.db, {"type": "Issue", "warehouse": "Main Store", "item": "Tiles", "qty": 5,
                                   "bu_code": "PRJ-2"})
        a.void_entry(self.db, e)
        after = {(r["name"], r["warehouse"]): r["qty"] for r in a.stock_report(self.db)}
        self.assertEqual(before, after)


class Interchange(unittest.TestCase):
    def test_live_fee_is_paid_against_actual_costs(self):
        """Live fee: what the client paid so far over what the project cost so far — advances count, a negotiated
        supplier discount and a refundable deposit do not; nothing is left "to collect"."""
        db = Database(":memory:")
        db.x("INSERT INTO business_units(code,name,basis,fee_pct) VALUES ('LF-1','Villa',?,0.2)", (a.LIVE_FEE,))
        a.record_project_cost(db, {"bu_code": "LF-1", "description": "Steel", "vendor": "Sup", "qty": 1, "rate": 10000,
                                   "paid": 9000, "savings": 1000, "date": "2026-03-01"})
        a.record_collection(db, {"bu_code": "LF-1", "amount": 9000, "date": "2026-03-05"})
        b = next(x for x in a.bu_summaries(db) if x["code"] == "LF-1")
        self.assertEqual((b["client_paid"], b["actual_costs"], b["to_collect"]), (9000.0, b["cogs"] + b["direct_overhead"], 0.0))
        self.assertAlmostEqual(b["live_fee"], round((9000 - b["actual_costs"]) / b["actual_costs"], 4))
        self.assertEqual(b["expected_billing"], 9000.0)
        # books where the funding partner's account replaces the bank: the client's money lands there, and still counts
        fdb = Database(":memory:")
        fdb.x("INSERT INTO business_units(code,name,basis) VALUES ('LF-3','Flat',?)", (a.LIVE_FEE,))
        fdb.x("INSERT INTO partners(name,default_share) VALUES ('Funder',1)")
        fdb.set_setting("funding_partner", "Funder")
        a.record_project_cost(fdb, {"bu_code": "LF-3", "description": "Tiles", "vendor": "Sup", "qty": 1, "rate": 5000, "paid": 5000, "date": "2026-03-01"})
        a.record_collection(fdb, {"bu_code": "LF-3", "amount": 6000, "date": "2026-03-05"})
        landed = fdb.q("SELECT a.role, a.subtype FROM journal_lines l JOIN accounts a ON a.code=l.account JOIN journal_entries e ON e.id=l.entry_id "
                       "WHERE e.kind='Receipt Voucher' AND l.debit>0")
        self.assertTrue(any(r["role"] == "funding_partner" for r in landed), landed)
        f = a.bu_summaries(fdb)[0]
        self.assertEqual((f["client_paid"], f["live_fee"]), (6000.0, 0.2))
        # VAT the client paid is the state's, not the fee: 11,400 received incl. 14% VAT counts as 10,000
        vdb = Database(":memory:")
        vdb.x("INSERT INTO business_units(code,name,basis) VALUES ('LF-4','Shop',?)", (a.LIVE_FEE,))
        a.record_project_cost(vdb, {"bu_code": "LF-4", "description": "Wood", "vendor": "Sup", "qty": 1, "rate": 8000, "paid": 8000, "date": "2026-03-01"})
        a.record_collection(vdb, {"bu_code": "LF-4", "amount": 11400, "vat_code": "VAT14", "date": "2026-03-05"})
        v = a.bu_summaries(vdb)[0]
        self.assertEqual((v["client_paid_gross"], v["client_paid_vat"], v["client_paid"]), (11400.0, 1400.0, 10000.0))
        empty = Database(":memory:")
        empty.x("INSERT INTO business_units(code,name,basis) VALUES ('LF-2','New',?)", (a.LIVE_FEE,))
        self.assertIsNone(a.bu_summaries(empty)[0]["live_fee"])        # no costs yet: no fee to show

    def test_json_round_trip_is_exact(self):
        db = seeded()
        a.save_account(db, {"code": "1025", "name": "Bank CIB", "type": "Asset", "subtype": "cash"}, old_code="1020")
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        tb1 = {r["code"]: (r["debit"], r["credit"]) for r in a.trial_balance(db)["rows"]}
        tb2 = {r["code"]: (r["debit"], r["credit"]) for r in a.trial_balance(db2)["rows"]}
        self.assertEqual(tb1, tb2)
        self.assertEqual(db2.acct("bank"), "1025")

    def test_import_is_all_or_nothing(self):
        db = Database(":memory:")
        data = {"Business Units": [{"code": "X", "name": "X"}],
                "Operating Expenses": [{"date": "2026-01-01", "description": "ok", "amount": 5},
                                       {"date": "2026-01-01", "description": "bad", "amount": "abc"}]}
        with self.assertRaises(interchange.ImportErrors) as cm:
            interchange.import_data(db, data)
        self.assertEqual(cm.exception.errors[0]["row"], 6)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM business_units")["n"], 0)

    def test_dry_run_changes_nothing(self):
        db = Database(":memory:")
        interchange.import_data(db, {"Business Units": [{"code": "X", "name": "X"}]}, dry_run=True)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM business_units")["n"], 0)

    def test_csv_template_has_bilingual_headers(self):
        csv = interchange.to_csv(Database(":memory:"), "Direct Costs", template=True).decode("utf-8-sig")
        self.assertIn("Date / التاريخ", csv.splitlines()[0])

    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_workbook_round_trip(self):
        from erp import workbook
        db = seeded()
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "wb.xlsx"
            p.write_bytes(workbook.build_workbook(db))
            data = interchange.read_xlsx(p)
        db2 = Database(":memory:")
        interchange.import_data(db2, data, mode="replace")
        self.assertEqual(a.income_statement(db)["net_profit"], a.income_statement(db2)["net_profit"])
        self.assertEqual(a.balance_sheet(db)["total_assets"], a.balance_sheet(db2)["total_assets"])


class Reports(unittest.TestCase):
    def test_every_report_every_format(self):
        db = seeded()
        for name in ("income_statement", "balance_sheet", "cash_flow", "trial_balance", "general_ledger", "bu_pnl",
                     "payables", "stock", "receivables", "partner_pack"):
            for lang in ("en", "ar"):
                tb = reports.build(db, name, {"date_from": "2026-01-01", "date_to": "2026-12-31", "compare": "quarter",
                                              "lang": lang})
                self.assertTrue(tb["rows"], name)
                self.assertTrue(reports.to_html(tb).startswith("<!doctype html>"))
                self.assertEqual(reports.to_docx(tb)[:2], b"PK")
                if HAVE_XL:
                    self.assertEqual(reports.to_xlsx(tb)[:2], b"PK")

    def test_partner_pack(self):
        from erp import partner_pack as pp
        self.assertEqual(pp.billed_profit(0, 0, -500), -500)          # no revenue: net as is
        self.assertEqual(pp.billed_profit(1000, 200, -50), -50)        # a loss is never deferred
        self.assertAlmostEqual(pp.billed_profit(1000, 250, 400), 300)  # 25% still in work in progress
        db = seeded()
        d = pp.compute(db, "2026-01-01", "2026-12-31")
        net = a.income_statement(db, "2026-01-01", "2026-12-31")["net_profit"]
        self.assertAlmostEqual(d["billed_total"] + d["wip_profit"], net, places=2)
        self.assertAlmostEqual(sum(p["net_profit"] for p in d["projects"]) + d["company_left"], net, places=2)
        bs = a.balance_sheet(db, "2026-12-31")
        self.assertAlmostEqual(d["bs"]["ta"], bs["total_assets"], places=2)
        self.assertTrue(d["bs"]["balanced"])
        self.assertTrue(d["audit_ok"])
        self.assertEqual(reports.REPORT_PERMS["partner_pack"], ("reports.partner_pack", "partners.equity"))
        tb = reports.build(db, "partner_pack", {"date_from": "2026-01-01", "date_to": "2026-12-31", "lang": "ar"})
        self.assertIn("<table", reports.to_html(tb, auto_print=True))
        self.assertEqual(reports.to_xlsx(tb)[:2], b"PK")

    def test_comparative_columns(self):
        tb = reports.income_statement_table(seeded(), "2026-01-01", "2026-06-30", "month")
        self.assertEqual(len(tb["columns"]), 1 + 6)


class Auth(unittest.TestCase):
    def test_password_hash(self):
        h = auth.hash_password("correct horse")
        self.assertTrue(auth.check_password("correct horse", h))
        self.assertFalse(auth.check_password("wrong", h))

    def test_last_admin_protected(self):
        db = Database(":memory:")
        uid = auth.create_user(db, {"username": "boss", "password": "12345678"}, first=True)
        with self.assertRaises(a.ValidationError):
            auth.update_user(db, uid, {"is_admin": False}, acting_user_id=999)


class Api(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from erp.server import ERPServer
        cls.tmp = tempfile.TemporaryDirectory()
        cls.db = Database(Path(cls.tmp.name) / "t.sqlite3")
        cls.srv = ERPServer(("127.0.0.1", 0), cls.db, Path(cls.tmp.name) / "backups")
        cls.base = f"http://127.0.0.1:{cls.srv.server_address[1]}"
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()
        cls.db.close()
        cls.tmp.cleanup()

    def client(self):
        return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

    def call(self, op, method, path, body=None, headers=True):
        req = urllib.request.Request(self.base + path, method=method,
                                     data=json.dumps(body).encode() if body is not None else None,
                                     headers={"Content-Type": "application/json",
                                              **({"X-Requested-With": "BeyondERP"} if headers else {})})
        try:
            with op.open(req) as r:
                return r.status, json.loads(r.read() or b"{}")
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"{}")

    def test_supplier_discount_api(self):
        op = self.client()
        if self.call(op, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0] != 200:
            self.assertEqual(self.call(op, "POST", "/api/login", {"username": "admin", "password": "secret123"})[0], 200)
        bu = self.call(op, "POST", "/api/bus", {"code": "SD1", "name": "Disc"})[1]["id"]
        deal = {"bu_id": bu, "description": "Ceramic", "vendor": "Hani", "qty": 1, "rate": 16000, "savings": 1000,
                "paid": 14000, "deal_ref": "HANI-9", "deal_status": "Settled", "account": "5020"}
        st, r = self.call(op, "POST", "/api/docs/direct_cost", deal)
        self.assertEqual(st, 400)
        self.assertIn("Negotiation discount", r["error"])
        self.assertEqual(self.call(op, "POST", "/api/docs/direct_cost", {**deal, "deal_status": "Open"})[0], 200)
        st, d = self.call(op, "GET", "/api/supplier-discount?party=Hani")
        self.assertEqual((st, d["credit_account"], d["owed"], d["deal_ref"]), (200, "5020", 1000, "HANI-9"))
        self.assertEqual(self.call(op, "POST", "/api/docs/supplier_discount", {"party": "Hani", "amount": 1000})[0], 200)
        self.assertEqual(self.call(op, "GET", "/api/supplier-discount?party=Hani")[1]["owed"], 0)

    def test_stock_counts_delete_rules(self):
        """Open and cancelled counts can be deleted (stock.counts edit, deleting not switched off — also in bulk);
        a posted count cannot."""
        sup = self.client()
        if self.call(sup, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0] != 200:
            self.assertEqual(self.call(sup, "POST", "/api/login", {"username": "admin", "password": "secret123"})[0], 200)
        wh = self.call(sup, "POST", "/api/warehouses", {"name": "Count Store"})[1]["id"]
        st, r = self.call(sup, "POST", "/api/docs/stock_move", {"type": "Receipt", "warehouse_id": wh, "item": "Bolt", "qty": 5, "rate": 2, "vendor": "Supplier"})
        self.assertEqual(st, 200, r)
        start = lambda: self.call(sup, "POST", "/api/stock-counts", {"warehouse_id": wh})[1]["id"]  # noqa: E731
        c1 = start()
        self.assertEqual(self.call(sup, "POST", f"/api/stock-counts/{c1}/cancel", {})[0], 200)
        for user, perms in (("cview", {"stock": "add"}), ("cnodel", {"stock": "edit", "stock.counts!delete": "no"})):
            self.assertEqual(self.call(sup, "POST", "/api/users", {"username": user, "password": user + "pass1", "perms": perms})[0], 200)
            u = self.client()
            self.assertEqual(self.call(u, "POST", "/api/login", {"username": user, "password": user + "pass1"})[0], 200)
            self.assertEqual(self.call(u, "DELETE", f"/api/stock-counts/{c1}")[0], 403)
            st, r = self.call(u, "POST", "/api/bulk", {"target": "stock_counts", "action": "delete", "ids": [str(c1)]})
            self.assertEqual((st, r.get("done", 0)), (200, 0) if st == 200 else (st, 0))
        self.assertEqual(self.call(sup, "DELETE", f"/api/stock-counts/{c1}")[0], 200)              # cancelled
        c2 = start()
        self.assertEqual(self.call(sup, "POST", "/api/bulk", {"target": "stock_counts", "action": "delete", "ids": [str(c2)]})[1]["done"], 1)  # open
        c3 = start()
        self.assertEqual(self.call(sup, "PUT", f"/api/stock-counts/{c3}", {"lines": [{"item": "Bolt", "counted": 4}]})[0], 200)
        self.assertEqual(self.call(sup, "POST", f"/api/stock-counts/{c3}/post", {})[0], 200)
        st, r = self.call(sup, "DELETE", f"/api/stock-counts/{c3}")
        self.assertEqual(st, 400)
        self.assertIn("posted", r["error"])
        self.assertEqual([c["id"] for c in self.call(sup, "GET", "/api/stock-counts")[1]], [c3])

    def test_people_and_arranging_rules(self):
        """Administrators add people; only the super administrator edits or removes them. Arranging the menu and the
        forms needs its own permission."""
        sup = self.client()
        if self.call(sup, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0] != 200:
            self.assertEqual(self.call(sup, "POST", "/api/login", {"username": "admin", "password": "secret123"})[0], 200)
        self.assertEqual(self.call(sup, "POST", "/api/users", {"username": "adm2", "password": "adm2pass1", "is_admin": True})[0], 200)
        adm = self.client()
        self.assertEqual(self.call(adm, "POST", "/api/login", {"username": "adm2", "password": "adm2pass1"})[0], 200)
        st, u = self.call(adm, "POST", "/api/users", {"username": "staff2", "password": "staff2pass", "perms": {"sales": "view"}})
        self.assertEqual(st, 200)                                                       # an administrator may add a user
        self.assertEqual(self.call(adm, "PUT", f"/api/users/{u['id']}", {"full_name": "Changed"})[0], 403)   # but not edit them
        self.assertEqual(self.call(adm, "DELETE", f"/api/users/{u['id']}")[0], 403)                          # nor remove them
        self.assertEqual(self.call(sup, "PUT", f"/api/users/{u['id']}", {"full_name": "Changed", "perms": {"sales": "view"}})[0], 200)
        staff = self.client()
        self.call(staff, "POST", "/api/login", {"username": "staff2", "password": "staff2pass"})
        layout = {"menu_layout": {"hide": ["m_budgets"]}}
        self.assertEqual(self.call(staff, "PUT", "/api/settings", layout)[0], 403)
        self.assertEqual(self.call(adm, "PUT", "/api/settings", layout)[0], 200)
        self.assertEqual(self.call(sup, "PUT", f"/api/users/{u['id']}", {"perms": {"sales": "view", "settings.layout": "edit"}})[0], 200)
        self.assertEqual(self.call(staff, "PUT", "/api/settings", {"menu_layout": {}, "form_custom": {}})[0], 200)
        self.assertEqual(self.call(sup, "DELETE", f"/api/users/{u['id']}")[0], 200)

    def test_maker_checker(self):
        """A person whose entries need approval only creates drafts; an approver in the same group posts them (maybe
        corrected) or sends them back; nobody approves their own drafts."""
        sup = self.client()
        if self.call(sup, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0] != 200:
            self.assertEqual(self.call(sup, "POST", "/api/login", {"username": "admin", "password": "secret123"})[0], 200)
        jv = {"date": "2026-03-01", "memo": "Office rent", "party": "Landlord",
              "lines": [{"account": "5010", "debit": 700}, {"account": "1010", "credit": 700}]}
        add = lambda name, **kw: self.call(sup, "POST", "/api/users", {"username": name, "password": name + "pass1", **kw})
        self.assertEqual(add("maker", perms={"accounting": "add", "purchasing": "add"}, needs_approval=True, group_key="finance")[0], 200)
        self.assertEqual(add("checker", perms={"accounting": "edit", "purchasing": "edit", "accounting.approvals": "none"}, group_key="finance")[0], 200)
        self.assertEqual(add("outsider", perms={"accounting": "edit"}, group_key="management")[0], 200)
        maker, checker, outsider = self.client(), self.client(), self.client()
        for c, n in ((maker, "maker"), (checker, "checker"), (outsider, "outsider")):
            self.assertEqual(self.call(c, "POST", "/api/login", {"username": n, "password": n + "pass1"})[0], 200)
        before = self.call(sup, "GET", "/api/journal?q=Landlord")[1]["total"]
        st, r = self.call(maker, "POST", "/api/journal", jv)
        self.assertEqual((st, r["pending"], r["id"]), (200, True, None))
        aid = r["approval"]
        self.assertEqual(self.call(sup, "GET", "/api/journal?q=Landlord")[1]["total"], before)     # not in the books
        bad = dict(jv, lines=[{"account": "5010", "debit": 700}, {"account": "1010", "credit": 600}])
        self.assertEqual(self.call(maker, "POST", "/api/journal", bad)[0], 400)                   # checked when sent
        # who sees it: the checker (same group, approval permission) — not the outsider, never the maker themself
        self.assertEqual(self.call(maker, "POST", f"/api/approvals/{aid}/approve", {})[0], 403)
        self.assertEqual(self.call(checker, "POST", f"/api/approvals/{aid}/approve", {})[0], 403)  # no approval permission yet
        self.assertEqual(self.call(sup, "PUT", "/api/users/" + str(self.call(sup, "GET", "/api/users")[1][-2]["id"]),
                                   {"perms": {"accounting": "edit", "purchasing": "edit", "accounting.approvals": "edit"}})[0], 200)
        al = self.call(checker, "GET", "/api/alerts")[1]
        self.assertEqual(al["approvals"], [aid])
        self.assertEqual(self.call(outsider, "GET", "/api/alerts")[1]["approvals"], [])
        self.assertEqual(self.call(outsider, "GET", "/api/approvals")[1]["rows"], [])
        # sent back (needs a reason), fixed by the maker, then approved with a correction
        self.assertEqual(self.call(checker, "POST", f"/api/approvals/{aid}/reject", {})[0], 400)
        self.assertEqual(self.call(checker, "POST", f"/api/approvals/{aid}/reject", {"note": "Wrong month"})[0], 200)
        self.assertEqual(self.call(maker, "GET", "/api/alerts")[1]["rows"][0]["kind"], "approvals_back")
        self.assertEqual(self.call(maker, "PUT", f"/api/approvals/{aid}", {"payload": dict(jv, date="2026-02-28")})[0], 200)
        self.assertEqual(self.call(maker, "GET", f"/api/approvals/{aid}")[1]["status"], "pending")
        # the receipt goes with the draft: its maker adds it, the approver sees it, someone else's group does not
        st, att = self.call(maker, "POST", "/api/attachments", {"owner": "approval", "id": aid, "title": "receipt.pdf", "data": "JVBERi0xLjQ="})
        self.assertEqual(st, 200)
        self.assertEqual(len(self.call(checker, "GET", f"/api/attachments?owner=approval&id={aid}")[1]["rows"]), 1)
        self.assertEqual(self.call(outsider, "GET", f"/api/attachments?owner=approval&id={aid}")[0], 403)
        self.assertEqual(self.call(checker, "POST", "/api/attachments", {"owner": "approval", "id": aid, "title": "x", "data": "eA=="})[0], 403)
        fixed = dict(jv, date="2026-02-28", memo="Office rent Feb")
        st, r = self.call(checker, "POST", f"/api/approvals/{aid}/approve", {"payload": fixed})
        self.assertEqual(st, 200)
        e = self.call(sup, "GET", f"/api/journal/{r['id']}")[1]
        self.assertEqual((e["date"], e["memo"], e["status"], e["created_by"]), ("2026-02-28", "Office rent Feb", "posted", "maker"))
        self.assertEqual((e["meta"]["approved_by"], e["meta"]["entered_by"]), ("checker", "maker"))
        self.assertEqual([x["title"] for x in self.call(sup, "GET", f"/api/attachments?owner=entry&id={r['id']}")[1]["rows"]], ["receipt.pdf"])
        self.assertEqual(self.call(checker, "POST", f"/api/approvals/{aid}/approve", {})[0], 400)  # only once
        # the journal's filters: entered by, approved by, amounts
        ids = lambda q: [x["id"] for x in self.call(sup, "GET", "/api/journal?" + q)[1]["rows"]]
        self.assertIn(r["id"], ids("approved_by=checker"))
        self.assertIn(r["id"], ids("created_by=maker&min_amount=700&max_amount=700"))
        self.assertNotIn(r["id"], ids("min_amount=701"))
        # a correction to a posted entry also waits; vouchers (direct costs, expenses…) too
        st, r2 = self.call(maker, "PUT", f"/api/journal/{r['id']}", dict(fixed, memo="Rent changed"))
        self.assertTrue(r2["pending"])
        self.assertEqual(self.call(sup, "GET", f"/api/journal/{r['id']}")[1]["memo"], "Office rent Feb")
        st, r3 = self.call(maker, "POST", "/api/docs/expense", {"date": "2026-03-02", "description": "Paper", "amount": 50})
        self.assertEqual((st, r3["pending"]), (200, True))
        self.assertEqual(self.call(maker, "POST", f"/api/approvals/{r3['approval']}/withdraw", {})[0], 200)
        self.assertEqual(self.call(sup, "POST", f"/api/approvals/{r2['approval']}/approve", {})[0], 200)   # administrators: any group
        self.assertEqual(self.call(sup, "GET", f"/api/journal/{r['id']}")[1]["memo"], "Rent changed")
        # nobody lifts their own approval requirement
        me = [u for u in self.call(sup, "GET", "/api/users")[1] if u["username"] == "maker"][0]
        self.call(maker, "PUT", f"/api/users/{me['id']}", {"needs_approval": False, "full_name": "M"})
        self.assertTrue([u for u in self.call(sup, "GET", "/api/users")[1] if u["username"] == "maker"][0]["needs_approval"])

    def test_flow(self):
        admin = self.client()
        self.assertEqual(self.call(admin, "GET", "/api/journal")[0], 401)
        self.assertEqual(self.call(admin, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0], 200)
        self.assertEqual(self.call(admin, "POST", "/api/setup", {"username": "x", "password": "secret123"})[0], 400)
        self.assertEqual(self.call(admin, "POST", "/api/bus", {"code": "B1", "name": "Unit"}, headers=False)[0], 403)
        st, r = self.call(admin, "POST", "/api/bus", {"code": "B1", "name": "Unit", "overhead_pct": 0.5})
        self.assertEqual(st, 200)
        st, _ = self.call(admin, "POST", "/api/docs/receipt", {"bu_id": r["id"], "amount": 1000})
        self.assertEqual(st, 200)
        st, tb = self.call(admin, "GET", "/api/reports/trial_balance")
        self.assertEqual(st, 200)
        self.call(admin, "POST", "/api/users", {"username": "clerk", "password": "clerk1234",
                                                "perms": {"sales": "edit"}})
        clerk = self.client()
        self.assertEqual(self.call(clerk, "POST", "/api/login", {"username": "clerk", "password": "clerk1234"})[0], 200)
        self.assertEqual(self.call(clerk, "POST", "/api/docs/receipt", {"bu_id": r["id"], "amount": 5})[0], 200)
        self.assertEqual(self.call(clerk, "POST", "/api/docs/expense", {"description": "x", "amount": 5})[0], 403)
        self.assertEqual(self.call(clerk, "GET", "/api/users")[0], 403)
        self.assertEqual(self.call(clerk, "GET", "/api/reports/vouchers?kinds=Receipt%20Voucher")[0], 200)
        self.assertEqual(self.call(clerk, "GET", "/api/reports/vouchers")[0], 403)
        # screen-level permissions: no Purchasing section, but full access to Direct Costs only
        self.call(admin, "POST", "/api/users", {"username": "buyer", "password": "buyer1234",
                                                "perms": {"purchasing": "none", "purchasing.costs": "edit"}})
        buyer = self.client()
        self.call(buyer, "POST", "/api/login", {"username": "buyer", "password": "buyer1234"})
        self.assertEqual(self.call(buyer, "POST", "/api/docs/direct_cost",
                                   {"bu_id": r["id"], "description": "c", "qty": 1, "rate": 5, "paid": 5})[0], 200)
        self.assertEqual(self.call(buyer, "POST", "/api/docs/expense", {"description": "x", "amount": 5})[0], 403)
        self.assertEqual(self.call(buyer, "GET", "/api/docs/payment")[0], 403)
        # manual shared-overhead allocation, capped at 100%
        self.assertEqual(self.call(admin, "PUT", "/api/bus-allocation", {"allocations": {str(r["id"]): 0.7}})[0], 200)
        self.assertEqual(self.call(admin, "PUT", "/api/bus-allocation", {"allocations": {str(r["id"]): 1.2}})[0], 400)
        # deleting: users (never yourself / the last admin), and only records nothing depends on
        me = self.call(admin, "GET", "/api/bootstrap")[1]["user"]
        self.assertIn("your own account", self.call(admin, "DELETE", f"/api/users/{me['id']}")[1]["error"])
        st, tmp = self.call(admin, "POST", "/api/users", {"username": "temp", "password": "temp12345"})
        self.assertEqual(self.call(admin, "DELETE", f"/api/users/{tmp['id']}")[0], 200)
        self.assertIsNone(self.db.one("SELECT 1 FROM users WHERE username='temp'"))
        self.assertEqual(self.call(admin, "POST", "/api/tax-codes", {"code": "TMP", "name": "t", "kind": "vat", "rate": 0.1})[0], 200)
        self.assertEqual(self.call(admin, "DELETE", "/api/tax-codes/TMP")[0], 200)
        self.call(admin, "POST", "/api/parties-master", {"name": "Cust", "type": "customer", "default_vat": "VAT14"})
        self.assertIn("Clear Active", self.call(admin, "DELETE", "/api/tax-codes/VAT14")[1]["error"])
        pid = self.db.one("SELECT id FROM parties WHERE name='Cust'")["id"]
        self.call(admin, "POST", "/api/extracts", {"bu_id": r["id"], "number": "INV-D", "customer": "Cust",
                                                    "lines": [{"description": "x", "qty": 1, "price": 5}]})
        self.assertIn("invoice INV-D", self.call(admin, "DELETE", f"/api/parties-master/{pid}")[1]["error"])
        st, mv = self.call(admin, "POST", "/api/docs/stock_move", {"type": "Receipt", "item": "Paint", "qty": 2, "rate": 10, "warehouse": "Main Store", "date": "2026-05-01", "paid": 20})
        self.assertEqual(st, 200, mv)
        iid = self.db.one("SELECT id FROM stock_items WHERE name='Paint'")["id"]
        self.assertIn("movements", self.call(admin, "DELETE", f"/api/items/{iid}")[1]["error"])
        self.db.x("INSERT INTO stock_items(name) VALUES ('Typo')")
        self.assertEqual(self.call(admin, "DELETE", "/api/items/" + str(self.db.one("SELECT id FROM stock_items WHERE name='Typo'")["id"]))[0], 200)
        # "view + add": may create new records on that screen, but not change or delete them
        self.call(admin, "POST", "/api/users", {"username": "store", "password": "store1234",
                                                "perms": {"stock.moves": "add", "stock.warehouses": "add"}})
        store = self.client()
        self.call(store, "POST", "/api/login", {"username": "store", "password": "store1234"})
        st, w = self.call(store, "POST", "/api/warehouses", {"name": "Yard"})
        self.assertEqual(st, 200)
        self.assertEqual(self.call(store, "PUT", f"/api/warehouses/{w['id']}", {"name": "Yard 2"})[0], 403)
        self.assertEqual(self.call(store, "DELETE", f"/api/warehouses/{w['id']}")[0], 403)
        self.assertEqual(self.call(store, "GET", "/api/stock/moves")[0], 200)
        self.assertEqual(self.call(store, "GET", "/api/accounts")[0], 403)
        # offers, documents and branding
        st, o = self.call(admin, "GET", "/api/offers")
        self.assertEqual(st, 200)
        self.assertTrue(o["types"])
        st, new = self.call(admin, "POST", "/api/offers", {**o["types"][0], "type_id": o["types"][0]["id"], "client": "C"})
        self.assertEqual(st, 200)
        req = urllib.request.Request(self.base + f"/api/offers/{new['id']}/file?format=docx")
        with admin.open(req) as resp:
            self.assertEqual(resp.read()[:2], b"PK")
        st, b = self.call(admin, "GET", "/api/branding")
        self.assertEqual(st, 200)
        self.assertTrue(b["images"]["logo"])
        self.assertEqual(self.call(store, "PUT", "/api/branding", {"dark": "#000000"})[0], 403)
        self.assertEqual(self.call(admin, "PUT", "/api/branding", {"dark": "red"})[0], 400)
        # erase: administrator only, typed confirmation, keeps users
        self.assertEqual(self.call(store, "POST", "/api/erase", {"scope": "all", "confirm": "ERASE"})[0], 403)
        self.assertEqual(self.call(admin, "POST", "/api/erase", {"scope": "all", "confirm": "no"})[0], 400)
        st, er = self.call(admin, "POST", "/api/erase", {"scope": "transactions", "confirm": "ERASE"})
        self.assertEqual(st, 200)
        self.assertEqual(self.db.one("SELECT COUNT(*) n FROM journal_entries")["n"], 0)
        self.assertEqual(self.db.one("SELECT COUNT(*) n FROM business_units")["n"], 1)
        self.assertEqual(self.call(admin, "POST", "/api/erase", {"scope": "all", "confirm": "erase"})[0], 200)
        self.assertEqual(self.db.one("SELECT COUNT(*) n FROM business_units")["n"], 0)
        self.assertTrue(self.db.one("SELECT COUNT(*) n FROM accounts")["n"] > 50)
        self.assertEqual(self.call(store, "GET", "/api/stock/moves")[0], 200)  # users survive


class Workspace(unittest.TestCase):
    """Two desktop servers sharing one workspace through a stand-in Supabase."""

    @classmethod
    def setUpClass(cls):
        from erp.server import ERPServer
        sys.path.insert(0, str(Path(__file__).parent))
        import mock_supabase
        cls.sb = mock_supabase.start()
        cls.sb_url = f"http://127.0.0.1:{cls.sb.server_address[1]}"
        cls.tmp = tempfile.TemporaryDirectory()
        cls.nodes = []
        for name in ("a", "b"):
            db = Database(Path(cls.tmp.name) / f"{name}.sqlite3")
            srv = ERPServer(("127.0.0.1", 0), db, Path(cls.tmp.name) / f"bk-{name}")
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            cls.nodes.append((srv, db, f"http://127.0.0.1:{srv.server_address[1]}"))

    @classmethod
    def tearDownClass(cls):
        for srv, db, _ in cls.nodes:
            srv.cloud.stop()
            srv.shutdown()
            srv.server_close()
            db.close()
        cls.sb.shutdown()
        cls.tmp.cleanup()

    call = Api.call
    client = Api.client

    def node_call(self, i, op, method, path, body=None):
        self.base = self.nodes[i][2]
        return self.call(op, method, path, body)

    def dump(self, db):
        out = {t: db.q(f"SELECT * FROM {t} ORDER BY 1") for t in
               ("business_units", "journal_entries", "journal_lines", "users", "offers", "partners", "accounts",
                "settings", "audit_log", "ownership_changes", "fixed_assets", "partner_shares", "loans", "custodies",
                "rent_items", "rent_lots", "rentals", "rental_lines", "rental_handovers", "rent_moves", "rental_charges",
                "products", "product_media", "sales_orders", "sale_lines", "sale_installments", "sale_cheques", "extracts",
                "bank_statements", "bank_lines", "bank_rules", "recurring", "recurring_runs")}
        for u in out["users"]:
            u.pop("pw_hash")          # passwords stay on the device where they were set
        out["audit_log"] = [r for r in out["audit_log"] if r["action"] not in ("login", "setup")]  # device-local
        for r in out["audit_log"]:
            r.pop("id")
        return out

    def test_two_devices_converge(self):
        a, b = self.client(), self.client()
        # device A: local admin, connects and creates the workspace from its books
        self.node_call(0, a, "POST", "/api/setup", {"username": "admin", "password": "secret123"})
        self.assertEqual(self.node_call(0, a, "POST", "/api/bus", {"code": "PRE", "name": "Before sharing"})[0], 200)
        self.assertEqual(self.node_call(0, a, "PUT", "/api/cloud/config", {"url": self.sb_url, "key": "anon"})[0], 200)
        st, r = self.node_call(0, a, "POST", "/api/cloud/signin", {"email": "Owner@x.com", "password": "pw123456", "create": True})
        self.assertEqual(st, 200, r)
        st, r = self.node_call(0, a, "POST", "/api/cloud/claim", {})
        self.assertEqual(st, 200, r)
        self.assertTrue(r["is_admin"])
        # changes now go through the log
        st, bu = self.node_call(0, a, "POST", "/api/bus", {"code": "B1", "name": "Villa"})
        self.assertEqual(st, 200)
        self.assertEqual(len(self.sb.state.all_ops()), 1)
        self.assertEqual(self.node_call(0, a, "POST", "/api/users", {"username": "clerk", "password": "clerk1234",
                                                                     "perms": {"sales": "edit"}})[0], 200)
        self.assertNotIn("clerk1234", json.dumps(self.sb.state.all_ops()))  # passwords never leave the device
        # sign-in problems get plain answers; with "Confirm email" on, sign-up asks the person to open the email first
        st, r = self.node_call(0, a, "POST", "/api/cloud/signin", {"email": "owner@x.com", "password": "wrong-pass"})
        self.assertEqual(st, 400)
        self.assertIn("Wrong email or password", r["error"])
        self.sb.state.confirm_required = True
        st, r = self.node_call(0, a, "POST", "/api/cloud/signin", {"email": "new@x.com", "password": "pw123456", "create": True})
        self.assertIn("confirmation email", r["error"])
        st, r = self.node_call(0, a, "POST", "/api/cloud/signin", {"email": "new@x.com", "password": "pw123456"})
        self.assertIn("confirm your email", r["error"])
        st, r = self.node_call(0, a, "POST", "/api/cloud/signin", {"email": "owner@x.com", "password": "pw123456", "create": True})
        self.assertIn("already exists", r["error"])
        self.sb.state.confirm_required = False
        self.assertEqual(self.node_call(0, a, "POST", "/api/cloud/signin", {"email": "owner@x.com", "password": "pw123456"})[0], 200)
        # a user saved with an email is linked to the workspace; they can then sign in with their username
        st, r = self.node_call(0, a, "POST", "/api/users", {"username": "ali", "email": "Ali@X.com", "password": "ali12345"})
        self.assertEqual(st, 200, r)
        self.assertEqual(self.node_call(0, a, "POST", "/api/users", {"username": "ali2", "email": "ali@x.com"})[1]["error"],
                         "Another user already has this email address")
        self.assertEqual(self.node_call(0, a, "PUT", "/api/cloud/members/link", {"email": "ali@x.com", "username": "ali"})[0], 200)
        self.assertEqual(self.sb.state.username_of("ali@x.com"), "ali")
        self.sb.state.accounts["ali@x.com"] = "alipass1"
        probe = self.nodes[0][0].cloud.client()
        self.assertEqual(probe.sign_in_username("ALI", "alipass1")["user"]["email"], "ali@x.com")
        # this device takes the workspace password too — no separate local password to remember
        d = self.client()
        st, r = self.node_call(0, d, "POST", "/api/login", {"username": "ali", "password": "alipass1"})
        self.assertEqual((st, r["user"]["username"]), (200, "ali"))
        self.assertEqual(self.node_call(0, d, "POST", "/api/login", {"username": "ali", "password": "nope"})[0], 400)
        with self.assertRaises(RuntimeError):
            probe.sign_in_username("ali", "wrong")
        with self.assertRaises(RuntimeError):
            probe.sign_in_username("nobody", "alipass1")
        # the administrator creates a sign-in account from Users — no email, no confirmation — and it works at once
        self.assertEqual(self.node_call(0, a, "POST", "/api/users", {"username": "sara", "_no_password": True,
                                                                     "perms": {"stock": "view"}})[0], 200)
        st, r = self.node_call(0, a, "PUT", "/api/cloud/account", {"username": "sara", "password": "sarapw1"})
        self.assertEqual(st, 200, r)
        self.assertEqual(r["email"], "sara@users.beyond-erp.invalid")
        sess = probe.sign_in_username("sara", "sarapw1")
        self.assertEqual(sess["user"]["email"], "sara@users.beyond-erp.invalid")
        from erp.cloud import username_for
        self.assertEqual(username_for(sess["user"]["email"]), "sara")
        st, r = self.node_call(0, a, "PUT", "/api/cloud/account", {"username": "sara", "active": False})
        with self.assertRaises(RuntimeError):   # untick Active: sign-in blocked
            probe.sign_in_username("sara", "sarapw1")
        # A (admin) adds B's account as a gateway member; B joins and gets A's books
        self.assertEqual(self.node_call(0, a, "PUT", "/api/cloud/members",
                                        {"email": "desk@x.com", "gateway": True})[0], 200)
        self.node_call(1, b, "POST", "/api/setup", {"username": "admin", "password": "secret123"})
        self.node_call(1, b, "PUT", "/api/cloud/config", {"url": self.sb_url, "key": "anon"})
        st, r = self.node_call(1, b, "POST", "/api/cloud/signin", {"email": "desk@x.com", "password": "pw123456", "create": True})
        self.assertEqual(st, 200, r)
        st, r = self.node_call(1, b, "POST", "/api/cloud/join", {})
        self.assertEqual(st, 200, r)
        b = self.client()  # sessions are cleared by joining: sign in again (password hashes travel in the copy)
        self.assertEqual(self.node_call(1, b, "POST", "/api/login", {"username": "admin", "password": "secret123"})[0], 200)
        # both devices write "at the same time"; one gets a conflict, pulls and retries transparently
        results = []

        def w(i, op, code):
            results.append(self.nodes[i] and Api.call(type("C", (), {"base": self.nodes[i][2]})(), op, "POST", "/api/bus",
                                                      {"code": code, "name": code})[0])
        ts = [threading.Thread(target=w, args=(i % 2, (a, b)[i % 2], f"C{i}")) for i in range(8)]
        [t.start() for t in ts]
        [t.join() for t in ts]
        self.assertEqual(results, [200] * 8)
        st, r = self.node_call(1, b, "POST", "/api/docs/receipt", {"bu_id": bu["id"], "amount": 1500})
        self.assertEqual(st, 200)
        # a user without permission is refused before anything reaches the log
        c = self.client()
        self.node_call(0, c, "POST", "/api/login", {"username": "clerk", "password": "clerk1234"})
        n = len(self.sb.state.all_ops())
        self.assertEqual(self.node_call(0, c, "POST", "/api/docs/expense", {"description": "x", "amount": 5})[0], 403)
        self.assertEqual(len(self.sb.state.all_ops()), n)
        self.assertEqual(self.node_call(0, c, "POST", "/api/docs/receipt", {"bu_id": bu["id"], "amount": 7})[0], 200)
        # gateway relays B's local user; the log records that user, and both copies end up identical
        self.assertEqual(self.sb.state.all_ops()[-2]["app_user"], "admin")
        # partners' equity and fixed assets replay the same on every device (previews are never logged)
        for name in ("P1", "P2"):
            st, r = self.node_call(0, a, "POST", "/api/partners", {"name": name, "default_share": 0.5})
            self.assertEqual(st, 200, r)
            self.node_call(0, a, "POST", "/api/docs/partner_tx", {"partner": name, "type": "Capital", "amount": 50000})
        n = len(self.sb.state.all_ops())
        self.assertEqual(self.node_call(0, a, "POST", "/api/equity-preview/transfer", {"seller": "P1", "buyer": "P3", "share": 10})[0], 200)
        self.assertEqual(len(self.sb.state.all_ops()), n)
        st, r = self.node_call(1, b, "POST", "/api/equity/transfer", {"seller": "P1", "buyer": "P3", "share": 10})
        self.assertEqual(st, 200, r)
        st, r = self.node_call(0, a, "POST", "/api/assets", {"name": "Laptop", "account": "1430", "cost": 36000,
                                                             "life_months": 36, "acquired": "2026-01-10", "pay": "cash"})
        self.assertEqual(st, 200, r)
        st, r = self.node_call(1, b, "POST", "/api/assets/depreciation", {"through": "2026-03"})
        self.assertEqual(st, 200, r)
        # loans and cash custody too: made on one device, paid / settled on the other
        st, r = self.node_call(0, a, "POST", "/api/loans", {"direction": "borrowed", "party": "CIB", "principal": 60000, "rate": 12,
                                                            "interest_type": "declining", "term_months": 6, "start_date": "2026-01-01"})
        self.assertEqual(st, 200, r)
        self.nodes[1][0].cloud.pull()   # previews run on the local copy
        self.assertEqual(self.node_call(1, b, "POST", "/api/loans-preview/pay", {"id": r["id"], "principal": 100})[0], 200)
        st, r2 = self.node_call(1, b, "POST", f"/api/loans/{r['id']}/pay", {"date": "2026-02-01", "principal": 9752.9, "interest": 600})
        self.assertEqual(st, 200, r2)
        st, r = self.node_call(1, b, "POST", "/api/custody", {"holder": "Samir", "amount": 3000, "issued_on": "2026-02-01"})
        self.assertEqual(st, 200, r)
        st, r2 = self.node_call(0, a, "POST", f"/api/custody/{r['id']}/settle",
                                {"date": "2026-02-05", "lines": [{"description": "Cement", "account": "5010", "amount": 1140, "vat_code": "VAT14"}]})
        self.assertEqual(st, 200, r2)
        # equipment rental: stock bought on one device, rented out and delivered from the other
        self.assertEqual(self.node_call(0, a, "POST", "/api/rent-starter", {})[0], 200)
        self.nodes[1][0].cloud.pull()
        fid = self.nodes[1][1].one("SELECT id FROM rent_items WHERE code='SC-FRM2'")["id"]
        self.assertEqual(self.node_call(1, b, "POST", f"/api/rent-items/{fid}/lots", {"qty": 50, "unit_cost": 900, "source": "opening"})[0], 200)
        st, r = self.node_call(0, a, "POST", "/api/rentals", {"client": "C", "lines": [{"kind": "item", "item_id": fid, "qty": 10, "rate": 60}]})
        self.assertEqual(st, 200, r)
        rid = r["id"]
        self.node_call(0, a, "POST", f"/api/rentals/{rid}/po", {"po_number": "P1"})
        self.node_call(0, a, "POST", f"/api/rentals/{rid}/sign", {"signed_us": "x", "signed_client": "y"})
        self.nodes[1][0].cloud.pull()
        st, t = self.node_call(1, b, "GET", f"/api/rentals/{rid}/handover-template?kind=delivery")
        st, r = self.node_call(1, b, "POST", f"/api/rentals/{rid}/handovers", {**t, "signed": 1, "decision": "safe",
                                                                               "checklist": [dict(c, ok="yes") for c in t["checklist"]]})
        self.assertEqual(st, 200, r)
        st, r = self.node_call(1, b, "POST", f"/api/rental-handovers/{r['id']}/confirm", {})
        self.assertEqual(st, 200, r)
        # a unit sold on installments on one device; cheques cleared and money collected on the other
        st, r = self.node_call(0, a, "POST", "/api/products", {"name": "Flat 4", "kind": "unit", "price": 900000})
        self.assertEqual(st, 200, r)
        self.node_call(0, a, "POST", f"/api/products/{r['id']}/media", {"kind": "photo", "data": "data:image/png;base64,iVBORw0KGgo="})
        bu = self.nodes[0][1].one("SELECT code FROM business_units LIMIT 1")["code"]
        st, r = self.node_call(0, a, "POST", "/api/sales", {"customer": "Hany", "bu_code": bu, "date": "2026-03-01", "lines": [{"product_id": r["id"]}],
                                                           "plan": {"type": "installments", "down_pct": 20, "n": 4, "rate": 12}})
        self.assertEqual(st, 200, r)
        sid = r["id"]
        self.assertEqual(self.node_call(0, a, "POST", f"/api/sales/{sid}/confirm", {"collect": 180000, "method": "Cash"})[0], 200)
        self.nodes[1][0].cloud.pull()
        st, d = self.node_call(1, b, "GET", f"/api/sales/{sid}")
        st, r = self.node_call(1, b, "POST", f"/api/sales/{sid}/collect", {"date": "2026-03-02", "method": "Cheque",
                                                                       "cheques": [{"number": "9", "due": "2026-04-01", "amount": d["schedule"][1]["amount"]}]})
        self.assertEqual(st, 200, r)
        cid = self.nodes[1][1].one("SELECT id FROM sale_cheques")["id"]
        self.assertEqual(self.node_call(1, b, "POST", f"/api/sale-cheques/{cid}/clear", {"date": "2026-04-02"})[0], 200)
        # recurring entries posted on one device are never posted again by the other; a statement matched across both
        st, r = self.node_call(0, a, "POST", "/api/recurring", {"name": "Rent", "doc_type": "expense", "start": "2026-01-31", "month_end": 1,
                                                              "mode": "auto", "payload": {"description": "Rent {month}", "amount": 100, "account": "6010"}})
        self.assertEqual(st, 200, r)
        self.assertEqual(self.node_call(0, a, "POST", "/api/recurring-run", {"through": "2026-03-31"})[1]["posted"], 3)
        self.nodes[1][0].cloud.pull()
        self.assertEqual(self.node_call(1, b, "POST", "/api/recurring-run", {"through": "2026-03-31"})[1]["posted"], 0)
        bank = self.nodes[1][1].acct("bank")
        st, r = self.node_call(1, b, "POST", "/api/bank-statements", {"account": bank, "date_from": "2026-01-01", "date_to": "2026-03-31",
                                                                    "lines": [{"date": "2026-01-31", "description": "Rent", "amount": -100}]})
        self.assertEqual(st, 200, r)
        self.assertEqual(self.node_call(1, b, "POST", f"/api/bank-statements/{r['id']}/automatch", {})[1]["matched"], 1)
        for srv, _, _ in self.nodes:
            srv.cloud.pull()
        self.assertEqual(self.dump(self.nodes[0][1]), self.dump(self.nodes[1][1]))
        self.assertEqual(self.nodes[0][1].one("SELECT COUNT(*) n FROM business_units")["n"], 11)
        # offline: changes are refused with a clear message instead of forking the books
        url = self.sb_url
        self.nodes[1][0].cloud.configure("http://127.0.0.1:9", "anon")
        st, r = self.node_call(1, b, "POST", "/api/bus", {"code": "OFF", "name": "x"})
        self.assertEqual(st, 503)
        self.assertIn("workspace", r["error"])
        self.nodes[1][0].cloud.configure(url, "anon")


class ExportRoundTrip(unittest.TestCase):
    """Everything exported (JSON backup or the Excel workbook) comes back the same: every module, settings,
    appearance, logos and icon, partner terms, chat and profiles."""
    SKIP = {"id", "created_at", "updated_at", "created_by", "entry_id", "ts", "history", "edited_at"}
    LOCAL = {"sessions", "audit_log", "cloud_local", "ai_usage", "sqlite_sequence", "users", "ai_threads", "chat_members",
             "chat_reads", "user_profiles"}

    def rows(self, db, t):
        import collections
        cols = [c[1] for c in db.conn.execute(f"PRAGMA table_info({t})").fetchall()]
        keep = [c for c in cols if c not in self.SKIP and not c.endswith("_id")]
        return collections.Counter(json.dumps(list(r), default=str) for r in db.conn.execute(f"SELECT {','.join(keep) or '1'} FROM {t}"))

    def people(self, db, order):
        for n in order:
            db.x("INSERT INTO users(username,pw_hash) VALUES (?,?)", (n, "x"))

    def test_json_and_workbook_round_trip(self):
        from erp import demo, chat, workbook
        src = Database(":memory:")
        demo.load(src, "admin")
        src.set_setting("appearance", json.dumps({"app_name": "RT", "brand_align": "center", "app_icon": "seal"}))
        src.conn.execute("INSERT OR REPLACE INTO assets(key,mime,w,h,data) VALUES ('app_icon','image/png',1,1,'AAAA')")
        src.conn.execute("UPDATE partners SET salary=300000, drawings_interest=0.1, profit_share=0.5")
        self.people(src, ["admin", "sara"])
        admin, sara = (dict(src.one("SELECT * FROM users WHERE username=?", (n,))) for n in ("admin", "sara"))
        chat.save_profile(src, sara["id"], {"display_name": "Sara K"})
        chat.send(src, admin, 1, {"body": "Hello", "files": [{"name": "a.txt", "mime": "text/plain", "data": "aGk="}]})
        chat.send(src, sara, chat.direct(src, admin, sara["id"]), {"body": "dm"})
        path = str(Path(tempfile.mkdtemp()) / "wb.xlsx")
        with open(path, "wb") as f:
            f.write(workbook.build_workbook(src))
        for label, data in (("json", interchange.read_json(interchange.to_json(src).encode())), ("xlsx", interchange.read_xlsx(path))):
            dst = Database(":memory:")
            self.people(dst, ["sara", "admin"])          # same people, other user numbers
            interchange.import_data(dst, data, mode="replace", user="admin")
            tables = [r[0] for r in src.conn.execute("SELECT name FROM sqlite_master WHERE type='table'") if r[0] not in self.LOCAL]
            differ = {t for t in tables if self.rows(src, t) != self.rows(dst, t)}
            # journal_entries: internal invoice-link numbers are renumbered; settings: defaults added; chat_rooms: the
            # direct-message key holds user numbers
            self.assertLessEqual(differ, {"journal_entries", "settings", "chat_rooms"}, label)
            self.assertEqual(a.income_statement(src)["net_profit"], a.income_statement(dst)["net_profit"])
            self.assertEqual(self.rows(src, "journal_lines"), self.rows(dst, "journal_lines"))
            self.assertEqual(dst.one("SELECT data FROM assets WHERE key='app_icon'")["data"], "AAAA")
            self.assertEqual([m["body"] for m in dst.q("SELECT body FROM chat_messages ORDER BY id")], ["Hello", "dm"])
            self.assertEqual(dst.one("SELECT COUNT(*) n FROM chat_files")["n"], 1)


class ProfitAppropriation(unittest.TestCase):
    """Sharing the year's profit: legal reserve → salaries + interest on capital → interest on drawings → the rest by
    the profit-sharing ratio, onto the partners' current accounts (capital untouched)."""
    def setUp(self):
        db = self.db = Database(":memory:")
        db.x("INSERT INTO partners(name,default_share,capital_interest,salary,profit_share,drawings_interest) "
             "VALUES ('A',0.6,0.10,200000,0.5,0),('B',0.4,0,0,0.5,0.12)")
        a.record_partner_tx(db, {"partner": "A", "type": "Capital", "amount": 60000, "date": "2026-01-01"})
        a.record_partner_tx(db, {"partner": "B", "type": "Capital", "amount": 40000, "date": "2026-01-01"})
        a.post_entry(db, date="2026-06-30", kind="Journal Voucher", lines=[{"account": "1020", "debit": 100000},
                                                                           {"account": "4010", "credit": 100000}])
        a.post_entry(db, date="2026-06-30", kind="Journal Voucher", lines=[{"account": "6010", "debit": 20000},
                                                                           {"account": "1020", "credit": 20000}])
        a.record_partner_tx(db, {"partner": "B", "type": "Distribution", "amount": 10000, "date": "2026-07-01"})
        db.set_setting("legal_reserve_pct", "0.05")

    def test_year_end_sharing(self):
        db = self.db
        d = {"date_from": "2026-01-01", "date_to": "2026-12-31"}
        st = equity.preview(db, "appropriate", d)["statement"]
        self.assertEqual((st["profit"], st["reserve"], st["salaries"], st["interest_on_capital"]), (80000, 4000, 24000, 6000))
        self.assertEqual(st["interest_on_drawings"], 604.93)          # 10,000 × 12% × 184 days / 365
        self.assertEqual(st["residual"], 46604.93)                     # 80,000 − 4,000 − 30,000 + 604.93
        by = {r["partner"]: r for r in st["partners"]}
        self.assertEqual(round(by["A"]["share_of_rest"] + by["B"]["share_of_rest"], 2), 46604.93)
        self.assertEqual(by["A"]["total"], round(24000 + 6000 + by["A"]["share_of_rest"], 2))
        before = {n: equity.balances(db, n)["capital"] for n in "AB"}
        equity.post(db, "appropriate", d)
        self.assertEqual({n: equity.balances(db, n)["capital"] for n in "AB"}, before)   # capital stays as agreed
        self.assertEqual(equity.balances(db, "A")["current"], by["A"]["current_after"])
        self.assertEqual(equity.balances(db, "B")["current"], round(by["B"]["total"] - 10000, 2))
        self.assertEqual(equity.balances(db, "B")["drawings"], 0)                          # drawings closed in
        bs = a.balance_sheet(db, "2026-12-31")
        self.assertEqual(bs["balanced"] if "balanced" in bs else True, True)
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        with self.assertRaises(a.ValidationError):                                         # the year is shared once
            equity.preview(db, "appropriate", d)
        nxt = equity.preview(db, "appropriate", {"date_to": "2027-03-31", "mode": "full"})["statement"]   # next period starts after it
        self.assertEqual(nxt["period_from"], "2027-01-01")

    def test_loss_and_short_profit(self):
        db = self.db
        a.post_entry(db, date="2027-01-10", kind="Journal Voucher", lines=[{"account": "6010", "debit": 5000},
                                                                           {"account": "1020", "credit": 5000}])
        q = {"date_from": "2027-01-01", "date_to": "2027-03-31", "close_drawings": False}
        st = equity.preview(db, "appropriate", q)["statement"]
        self.assertEqual((st["reserve"], st["salaries"], st["interest_on_capital"]), (0, 0, 0))   # nothing out of a loss
        self.assertEqual([r["share_of_rest"] for r in st["partners"]], [-2500, -2500])
        full = equity.preview(db, "appropriate", {**q, "mode": "full"})["statement"]
        self.assertEqual(full["salaries"], 6000)                      # 3 months × 2,000
        self.assertEqual(round(sum(r["total"] for r in full["partners"]), 2), -5000)   # the shortfall is shared as a loss
        equity.post(db, "appropriate", {**q, "mode": "full"})
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        eid = equity.appropriations(db)[0]["id"]
        a.void_entry(db, eid, "redo")                                 # undo = void the entry
        self.assertEqual(equity.appropriations(db), [])

    def test_ratio_must_total_100(self):
        self.db.x("UPDATE partners SET profit_share=0.3 WHERE name='B'")
        with self.assertRaises(a.ValidationError):
            equity.preview(self.db, "appropriate", {"date_to": "2026-12-31"})


class Equity(unittest.TestCase):
    def setUp(self):
        self.db = Database(":memory:")
        self.db.x("INSERT INTO partners(name,default_share) VALUES ('A',0.6),('B',0.4)")
        self.db.x("INSERT INTO business_units(code,name) VALUES ('U1','Unit')")
        self.db.x("INSERT INTO partner_shares(partner_id,bu_id,share) VALUES (1,1,0.5),(2,1,0.5)")
        a.record_partner_tx(self.db, {"partner": "A", "type": "Capital", "amount": 60000, "date": "2026-01-01"})
        a.record_partner_tx(self.db, {"partner": "B", "type": "Capital", "amount": 40000, "date": "2026-01-01"})

    def shares(self):
        return {p["name"]: round(p["default_share"], 6) for p in self.db.q("SELECT * FROM partners")}

    def check_books(self):
        self.assertEqual([c["check"] for c in a.audit(self.db) if not c["ok"]], [])

    def test_transfer_admit_exit_and_undo(self):
        db = self.db
        self.assertEqual(equity.preview(db, "transfer", {"seller": "A", "buyer": "C", "share": 1})["ownership"][-1]["after"], 0.01)
        r = equity.post(db, "transfer", {"seller": "A", "buyer": "C", "share": 30, "price": 99999, "date": "2026-02-01"})
        self.assertEqual(self.shares(), {"A": 0.3, "B": 0.4, "C": 0.3})
        self.assertEqual(equity.balances(db, "C")["capital"], 30000)   # half of A's capital moved with half of A's share
        self.assertEqual(equity.balances(db, "A")["capital"], 30000)
        bu = {r_["partner_id"]: r_["share"] for r_ in db.q("SELECT * FROM partner_shares WHERE bu_id=1")}
        self.assertEqual(bu, {1: 0.25, 2: 0.5, 3: 0.25})            # the unit's own shares move in proportion
        equity.post(db, "admit", {"name": "D", "share": 20, "amount": 50000, "basis": "bonus", "date": "2026-03-01"})
        self.assertAlmostEqual(sum(self.shares().values()), 1)
        self.assertEqual(equity.balances(db, "D")["capital"], 30000)   # 20% × (100,000 + 50,000)
        self.assertEqual(equity.total_capital(db), 150000)
        x = equity.post(db, "exit", {"partner": "B", "settlement": 45000, "paid": 5000, "date": "2026-04-01"})
        self.assertEqual(self.shares()["B"], 0)
        self.assertEqual(db.one("SELECT active FROM partners WHERE name='B'")["active"], 0)
        self.assertEqual(equity.balances(db, "B")["due_to"], 40000)
        self.assertAlmostEqual(sum(self.shares().values()), 1)
        self.check_books()
        with self.assertRaises(a.ValidationError):   # only the latest change can be undone
            equity.undo_change(db, r["change"])
        a.void_entry(db, x["id"], "wrong")              # voiding the exit restores B
        self.assertGreater(self.shares()["B"], 0)
        self.assertEqual(db.one("SELECT active FROM partners WHERE name='B'")["active"], 1)
        self.check_books()

    def test_loans_interest_distribution_and_capital_interest(self):
        db = self.db
        equity.post(db, "loan", {"partner": "A", "amount": 36500, "rate": 10, "date": "2026-01-01"})
        p = equity.preview(db, "interest", {"partner": "A", "date": "2026-01-11"})
        self.assertEqual(p["lines"][0]["debit"], 100.0)                 # 10 days × 36,500 × 10% / 365
        equity.post(db, "interest", {"partner": "A", "date": "2026-01-11"})
        with self.assertRaises(a.ValidationError):
            equity.post(db, "interest", {"partner": "A", "date": "2026-01-11"})   # already booked to that day
        self.assertEqual(equity.balances(db, "A")["interest"], 100)
        equity.post(db, "repay", {"partner": "A", "what": "interest", "amount": 100, "date": "2026-01-12"})
        equity.post(db, "contribute", {"partner": "A", "amount": 36500, "source": "loan", "ownership": "by_capital",
                                       "date": "2026-01-13"})
        self.assertEqual(equity.balances(db, "A")["loan"], 0)
        self.assertEqual(self.shares(), {"A": round(96500 / 136500, 6), "B": round(40000 / 136500, 6)})
        equity.post(db, "distribute", {"amount": 1000, "date": "2026-01-31"})
        self.assertEqual(equity.balances(db, "B")["distributions_payable"], round(1000 * 40000 / 136500, 2))
        db.x("UPDATE partners SET capital_interest=0.1 WHERE name='B'")
        rep = a.partner_report(db, "2026-01-01", "2026-12-31")
        self.assertGreater(next(r for r in rep["partners"] if r["name"] == "B")["capital_interest"], 3900)
        self.check_books()

    def test_same_account_both_sides_refused(self):
        with self.assertRaises(a.ValidationError):
            a.record_journal(self.db, {"date": "2026-01-02", "lines": [{"account": "1020", "debit": 5}, {"account": "1020", "credit": 5}]})
        # the same account is fine when the lines belong to different partners (a capital transfer)
        a.record_journal(self.db, {"date": "2026-01-02", "lines": [{"account": "3010", "debit": 5, "party": "A"},
                                                                    {"account": "3010", "credit": 5, "party": "B"}]})

    def test_old_books_get_the_new_accounts(self):
        db = Database(":memory:")
        db.x("DELETE FROM accounts WHERE code IN ('3020','3110')")
        db.x("INSERT INTO accounts(code,name,type,subtype,normal) VALUES ('3020','Mine','Equity','capital','Credit')")
        db.x("DELETE FROM settings WHERE key='coa_v2'")
        db._migrate()
        self.assertEqual(db.one("SELECT name FROM accounts WHERE code='3020'")["name"], "Mine")
        self.assertEqual(db.acct("partner_current"), "3021")
        self.assertEqual(db.acct("legal_reserve"), "3110")

    def test_json_round_trip_keeps_line_parties_and_history(self):
        equity.post(self.db, "transfer", {"seller": "A", "buyer": "C", "share": 30, "date": "2026-02-01"})
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(self.db).encode()), mode="replace")
        for n in ("A", "B", "C"):
            self.assertEqual(equity.balances(db2, n), equity.balances(self.db, n))
        self.assertEqual(len(equity.history(db2)), 1)


class FixedAssets(unittest.TestCase):
    def test_purchase_depreciation_disposal(self):
        db = Database(":memory:")
        a.record_journal(db, {"date": "2026-01-01", "lines": [{"account": "1020", "debit": 900000}, {"account": "3010", "credit": 900000}]})
        aid = assets.save(db, {"name": "Truck", "account": "1420", "acquired": "2026-01-15", "cost": 600000, "salvage": 60000,
                               "life_months": 60, "pay": "payable", "vendor": "Dealer", "paid": 100000, "vat_code": "VAT14"})
        self.assertEqual(a.party_balance(db, db.acct("ap"), "Dealer"), 584000)   # 600,000 + 14% VAT − 100,000 paid
        r = assets.run_depreciation(db, "2026-04")
        self.assertEqual(r["total"], 36000)                              # 4 months × 9,000
        with self.assertRaises(a.ValidationError):
            assets.run_depreciation(db, "2026-04")
        row = assets.register(db)["rows"][0]
        self.assertEqual((row["accumulated"], row["book_value"], row["booked_through"]), (36000, 564000, "2026-04"))
        a.void_entry(db, r["entries"][0], "redo")                       # voided depreciation is booked again next run
        self.assertEqual(assets.run_depreciation(db, "2026-04")["total"], 36000)
        assets.dispose(db, aid, {"date": "2026-06-20", "proceeds": 500000})   # May + June first, then the sale
        row = assets.register(db)["rows"][0]
        self.assertTrue(row["disposed"])
        self.assertEqual(row["accumulated"], 54000)
        self.assertEqual(next(x for x in a.account_totals(db) if x["code"] == "7045")["net"], 4600000)   # loss 46,000.00
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        # land is never depreciated; an existing asset can be registered with its past depreciation
        assets.save(db, {"name": "Plot", "account": "1440", "acquired": "2025-01-01", "cost": 100000, "life_months": 0})
        assets.save(db, {"name": "Desk", "account": "1400", "acquired": "2024-01-01", "cost": 12000, "life_months": 60,
                         "opening_depr": 4800})
        p = assets.plan_depreciation(db, "2026-06")
        self.assertEqual([(x["name"], x["from"]) for x in p["rows"]], [("Desk", "2026-01")])   # 24 months already taken
        # the register and its depreciation survive a JSON round trip
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        self.assertEqual([(r["code"], r["accumulated"], r["disposed"]) for r in assets.register(db2)["rows"]],
                         [(r["code"], r["accumulated"], r["disposed"]) for r in assets.register(db)["rows"]])
        self.assertEqual(reports.build(db, "fixed_assets", {})["rows"][-1]["cells"][5], 112000)


class EntryTypes(unittest.TestCase):
    def test_types_menu_and_posting(self):
        db = Database(":memory:")
        types = entries.load(db)
        self.assertIn("legal_reserve", [t["id"] for t in types])
        cash = entries.resolve_side(db, "subtype:cash")
        self.assertIn("1020", cash["choose"])
        eid = entries.post(db, "cash_transfer", {"amount": 500, "debit_account": "1010", "credit_account": "1020"})
        self.assertEqual(a.get_entry(db, eid)["amount"], 500)
        with self.assertRaises(a.ValidationError):
            entries.post(db, "cash_transfer", {"amount": 500, "debit_account": "1020", "credit_account": "1020"})
        entries.save(db, [{"name": "Pay rent", "kind": "Rent", "debit": "6010", "credit": "role:bank", "party": "required"}])
        with self.assertRaises(a.ValidationError):
            entries.post(db, "pay_rent", {"amount": 100})           # party required
        eid = entries.post(db, "pay_rent", {"amount": 100, "party": "Landlord"})
        self.assertEqual(a.get_entry(db, eid)["kind"], "Rent")
        self.assertNotIn("legal_reserve", [t["id"] for t in entries.load(db)])   # removed built-ins stay removed
        entries.save_menu(db, [{"key": "entry:pay_rent", "show": True}, {"key": "doc:journal", "show": True}, {"key": "bad", "show": True}])
        m = entries.menu(db)
        self.assertEqual([x["key"] for x in m if x["show"]], ["entry:pay_rent", "doc:journal"])


class DepreciationPolicy(unittest.TestCase):
    """Company depreciation rules: method, first month, defaults per asset account; changes apply from now on."""

    def test_policy(self):
        from erp import assets as A
        db = seeded()
        fa = db.acct("fixed_asset")
        A.save_policy(db, {"convention": "next_month", "by_account": {fa: {"life_months": 36, "residual_pct": 10}}})
        aid = A.save(db, {"name": "Laptop", "account": fa, "acquired": "2026-01-15", "cost": 36000})
        x = A.get(db, aid)
        self.assertEqual((x["life_months"], x["salvage"], x["depr_start"]), (36, 360000, "2026-02"))
        self.assertEqual(A._row(db, x)["monthly"], 900)                         # (36000 − 3600) ÷ 36
        # declining: front-loaded, still ends exactly at the residual value
        did = A.save(db, {"name": "Crane", "account": fa, "acquired": "2026-01-01", "cost": 120000, "salvage": 0,
                          "life_months": 60, "depr_start": "2026-01", "depr_method": "declining", "depr_factor": 2})
        sch = A.schedule(db, A.get(db, did))
        self.assertGreater(sch["rows"][0]["depreciation"], sch["rows"][-1]["depreciation"])
        self.assertEqual(sch["rows"][-1]["book_value"], 0)
        self.assertEqual(sch["rows"][-1]["year"], "2030")                        # 60 months from 2026-01 → 2030-12
        self.assertEqual(round(sum(r["depreciation"] for r in sch["rows"]), 2), 120000)
        A.run_depreciation(db, "2026-06", [did])
        st = A._row(db, A.get(db, did))
        self.assertEqual(st["booked_through"], "2026-06")
        # straight line: shortening the life re-spreads only what is left
        A.run_depreciation(db, "2026-06", [aid])                               # 2026-02..06 = 5 × 900
        A.save(db, {**A._row(db, A.get(db, aid)), "life_months": 18, "depr_method": "", "depr_factor": ""}, aid)
        r = A._row(db, A.get(db, aid))
        self.assertEqual((r["accumulated"], r["months_left"]), (4500, 13))
        self.assertEqual(r["monthly"], round((32400 - 4500) / 13, 2))
        last = A.schedule(db, A.get(db, aid))["rows"][-1]
        self.assertEqual(last["book_value"], 3600)
        # company default method flips the assets that follow the policy, not the ones with their own
        A.save_policy(db, {"method": "declining", "factor": 1.5, "convention": "next_month",
                           "by_account": {fa: {"life_months": 36, "residual_pct": 10}}})
        self.assertEqual(A._row(db, A.get(db, aid))["method"], "declining")
        self.assertEqual(A._row(db, A.get(db, did))["factor"], 2)
        with self.assertRaises(a.ValidationError):
            A.save_policy(db, {"method": "sum_of_digits"})
        self.assertTrue(A.preview_schedule(db, {"account": fa, "cost": 1000, "acquired": "2026-03-01"})["rows"])
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])


class ProductSales(unittest.TestCase):
    """Catalog with media, sales on payment plans (installments with interest), cheques, delivery, refund and cancel."""

    def setUp(self):
        from erp import sales
        self.S, self.db = sales, seeded()
        self.unit = sales.save_product(self.db, {"name": "Apartment A-301", "kind": "unit", "price": 2400000, "bu_code": "PRJ-1",
                                                 "specs": [["Area", "150 m²"]], "cost": 1500000, "inventory_account": "1210"})
        self.item = sales.save_product(self.db, {"name": "Tiles box", "kind": "item", "price": 120, "stock_item": "Tiles", "unit": "box"})

    def test_schedule(self):
        rows, markup = self.S.build_schedule(240000000, {"type": "installments", "down_pct": 10, "n": 24, "rate": 10,
                                                          "extras": [{"label": "On delivery", "pct": 10, "due": "2027-06-01"}]}, "2026-03-01")
        self.assertEqual(markup, 38400000)                     # (2.4M − 240k − 240k) × 10% × 2 years
        self.assertEqual(sum(r["amount"] for r in rows), 240000000 + markup)
        self.assertEqual(rows[0]["kind"], "down")
        self.assertEqual([r for r in rows if r["kind"] == "installment"][-1]["due"], "2028-03-01")
        rows, _ = self.S.build_schedule(100000, {"type": "installments", "n": 3, "round_to": 10}, "2026-01-31")
        self.assertEqual([r["amount"] for r in rows], [33000, 33000, 34000])
        self.assertEqual([r["due"] for r in rows], ["2026-02-28", "2026-03-28", "2026-04-28"])  # month-end clamps
        with self.assertRaises(a.ValidationError):
            self.S.build_schedule(1000, {"type": "custom", "rows": [{"due": "2026-01-01", "amount": 5}]}, "2026-01-01")

    def test_installment_sale_on_delivery(self):
        S, db = self.S, self.db
        sid = S.save(db, {"customer": "Mona", "bu_code": "PRJ-1", "date": "2026-03-01", "revenue_on": "delivery", "lines": [{"product_id": self.unit}],
                          "plan": {"type": "installments", "down_pct": 10, "n": 24, "rate": 10, "first_due": "2026-04-01"}})
        self.assertEqual(S.get_product(db, self.unit)["status"], "reserved")
        with self.assertRaises(a.ValidationError):                          # the unit cannot be sold twice
            S.save(db, {"customer": "X", "bu_code": "PRJ-1", "lines": [{"product_id": self.unit}]})
        S.confirm(db, sid, {"collect": 240000, "method": "Bank Transfer"})
        d = S.detail(db, sid, "2026-03-15")
        self.assertEqual((d["status"], d["paid"], d["invoiced"]), ("confirmed", 240000, False))
        self.assertEqual(a.party_balance(db, db.acct("client_advances"), "Mona"), 240000)
        amt = d["schedule"][1]["amount"]
        S.collect(db, sid, {"method": "Cheque", "date": "2026-03-02", "cheques": [{"number": "1", "due": "2026-04-01", "amount": amt},
                                                                                {"number": "2", "due": "2026-05-01", "amount": amt}]})
        c1, c2 = S.detail(db, sid)["cheques"]
        self.assertEqual(S.detail(db, sid, "2026-04-15")["schedule"][1]["status"], "cheque")
        S.cheque(db, c1["id"], "clear", {"date": "2026-04-02"})
        S.cheque(db, c2["id"], "bounce", {"date": "2026-05-03"})
        S.deliver(db, sid, {"date": "2026-07-01"})
        d = S.detail(db, sid, "2026-07-01")
        self.assertEqual((d["status"], d["paid"], d["invoiced"], d["overdue_count"]), ("delivered", 348000, True, 2))
        self.assertAlmostEqual(d["interest_earned"], round(432000 * 348000 / 2832000, 2), places=2)
        self.assertEqual(a.party_balance(db, db.acct("client_advances"), "Mona"), 0)
        self.assertEqual(S.get_product(db, self.unit)["status"], "delivered")
        tb = {r["code"]: r for r in a.trial_balance(db)["rows"]}
        self.assertEqual(tb["5005"]["debit"], 1500000)              # cost of the unit
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        self.assertEqual(S.installments(db, "2026-07-01", "overdue")["totals"]["left"], 216000)
        from erp import interchange
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        d2 = S.detail(db2, sid, "2026-07-01")                              # payments, cheques and interest survive a restore
        self.assertEqual((d2["paid"], d2["interest_earned"], d2["overdue_count"], len(d2["entries"])),
                         (d["paid"], d["interest_earned"], 2, len(d["entries"])))

    def test_item_sale_refund_cancel_and_backup(self):
        from erp import interchange, sales_docs
        S, db = self.S, self.db
        sid = S.save(db, {"customer": "Shop client", "bu_code": "PRJ-1", "date": "2026-03-05", "vat_code": "VAT14",
                          "lines": [{"product_id": self.item, "qty": 10, "discount": 50}]})
        self.assertEqual(S.detail(db, sid)["total"], 1311)                  # 10 × 115 + 14% VAT
        S.confirm(db, sid, {"collect": 1311, "method": "Cash"})
        stock = lambda: sum(r["qty"] for r in a.stock_report(db))  # noqa: E731
        self.assertEqual(stock(), 80)
        d = S.detail(db, sid)
        self.assertEqual((d["balance"], d["invoice"]["status"]), (0, "posted"))
        for k in ("quote", "contract", "statement"):
            self.assertIn("SO-0001", sales_docs.sale_html(db, sid, k, "ar"))
        mid = S.add_media(db, self.unit, {"kind": "photo", "data": "data:image/png;base64,iVBORw0KGgo="})
        S.add_media(db, self.unit, {"kind": "model", "url": "https://example.com/flat.glb", "title": "3D"})
        self.assertEqual(S.product_row(db, S.get_product(db, self.unit))["cover"], mid)
        self.assertIn("<img", sales_docs.product_html(db, self.unit))
        with self.assertRaises(a.ValidationError):
            S.cancel(db, sid, {"date": "2026-03-06"})                          # money must be refunded first
        S.refund(db, sid, {"amount": 1000, "keep": 311, "method": "Cash", "date": "2026-03-06"})
        S.cancel(db, sid, {"date": "2026-03-06"})
        self.assertEqual((S.get(db, sid)["status"], stock()), ("cancelled", 90))
        tb = {r["code"]: r for r in a.trial_balance(db)["rows"]}
        self.assertEqual(tb["4090"]["credit"], 311)                  # cancellation charge kept
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        s2 = S.save(db, {"customer": "B", "bu_code": "PRJ-1", "lines": [{"product_id": self.unit}], "date": "2026-04-01"})
        S.confirm(db, s2, {})
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        self.assertEqual(S.detail(db2, "SO-0002")["balance"], 2400000)
        self.assertEqual(S.detail(db2, "SO-0001")["paid"], S.detail(db, "SO-0001")["paid"])
        self.assertEqual(S.get_product(db2, "P-0001")["status"], "sold")
        self.assertEqual(len(S.media_list(db2, S.get_product(db2, "P-0001")["id"])), 2)


class LegacyJournalWorkbook(unittest.TestCase):
    """V1.3+ workbooks: the General Journal is the source of truth; archive sheets are ignored; codes are remapped."""

    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_v13_journal(self):
        import openpyxl
        import tempfile
        from erp import legacy, interchange
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Project Register"
        ws.append(["Project ID", "Client Name", "Project Type", "Location", "Total Area (Sqm)", "Start Date", "Status",
                   "Total Contract Value", "Billing Basis / Fee Type", "Supervision Fee %"])
        ws.append(["PRJ-2026-001", "Villa A", "Villa", "X", 100, "2026-01-01", "In Progress", 0, "Cost-Plus %", 0.25])
        gj = wb.create_sheet("General Journal")
        gj.append(["BEYOND | GENERAL JOURNAL"])
        gj.append(["JE #", "Date", "Source / Type", "Project", "Stage", "Account", "Description", "Vendor / Party",
                   "Deal Ref", "Qty", "Unit", "Rate", "Debit", "Credit"])
        for r in (["JE-1", "2026-01-02", "Partner funding (auto)", None, None, "1020", "Funding", "P1", None, None, None, None, 1000, None],
                  ["JE-1", "2026-01-02", "Partner funding (auto)", None, None, "3150", "Funding", "P1", None, None, None, None, None, 1000],
                  ["JE-2", "2026-01-05", "Project cost", "PRJ-2026-001", "All", "5020", "Works", "Sub", "SUB-1", 1, "Task", 600, 600, None],
                  ["JE-2", "2026-01-05", "Project cost", "PRJ-2026-001", "All", "1020", "Works", "Sub", "SUB-1", None, None, None, None, 400],
                  ["JE-2", "2026-01-05", "Project cost", "PRJ-2026-001", "All", "2010", "Works", "Sub", "SUB-1", None, None, None, None, 200],
                  ["JE-3", "2026-01-06", "Custody", None, None, "1320", "Custody", "Ali", None, None, None, None, 100, 0],
                  ["JE-3", "2026-01-06", "Custody", None, None, "1020", "Custody", "Ali", None, None, None, None, 0, 100],
                  ["JE-4", "2026-01-07", "Accrual (auto)", "PRJ-2026-001", "All", "1150", "Accrued", None, None, None, None, None, 0, None],
                  ["JE-4", "2026-01-07", "Accrual (auto)", "PRJ-2026-001", "All", "4010", "Accrued", None, None, None, None, None, None, 0]):
            gj.append(r)
        arch = wb.create_sheet("Project Costing (COGS)")   # archive: must not be imported again
        arch.append(["Date", "Project ID", "Cost Category", "Item Description", "Vendor/Supplier", "Quantity", "Unit Cost", "Total Cost"])
        arch.append(["2026-01-05", "PRJ-2026-001", "Subcontractor", "Works", "Sub", 1, 600, 600])
        reg = wb.create_sheet("Invoice Register")
        reg.append(["Invoice No", "Date", "Project", "Client", "Status"])
        reg.append(["INV-1", "2026-01-10", "PRJ-2026-001", "Villa A", "Draft"])
        il = wb.create_sheet("Invoice Lines")
        il.append(["Invoice No", "Item No", "Category", "Sub-category", "Description", "Unit", "Contract Qty", "Qty this invoice", "Rate"])
        il.append(["INV-1", "1.1", "Works", None, "Walls", "m2", 10, 4, 50])
        il.append(["INV-1", "V-01", "Works", "Variation", "Extra", "LS", None, 1, 300])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            data = legacy.read_legacy_workbook(f.name)
        db = Database(":memory:")
        interchange.import_data(db, data, mode="replace")
        tb = {r["code"]: round(r["debit"] - r["credit"], 2) for r in a.trial_balance(db)["rows"]}
        self.assertEqual((tb["1020"], tb["3010"], tb["5020"], tb["2010"], tb["1035"]), (500, -1000, 600, -200, 100))
        self.assertNotIn("1320", tb)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries WHERE kind<>'Work in Progress'")["n"], 3)  # no snapshots
        tbw = {r["code"]: round(r["debit"] - r["credit"], 2) for r in a.trial_balance(db)["rows"]}
        self.assertEqual(tbw.get(db.acct("unbilled_revenue")), 750)                     # cost-plus: 600 × 1.25, kept live
        self.assertEqual(a.get_entry(db, db.one("SELECT id FROM journal_entries WHERE ref='JE-2'")["id"])["meta"]["deal_ref"], "SUB-1")
        self.assertEqual([(x["number"], x["total"]) for x in a.list_extracts(db)], [("INV-1", 500)])
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])

    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_v14_invoice_sheet_new_accounts_and_measured_basis(self):
        """V1.4: no Invoice Lines sheet (the register points at 'Invoice 1'), new accounts in the chart, Measured basis."""
        import openpyxl
        import tempfile
        from erp import legacy, interchange
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Project Register"
        ws.append(["Project ID", "Client Name", "Project Type", "Location", "Total Area (Sqm)", "Start Date", "Status",
                   "Total Contract Value", "Billing Basis / Fee Type", "Supervision Fee %"])
        ws.append(["PRJ-2026-002", "Villa B", "Villa", "Site A", 950, "2026-09-17", "In Progress", 0, "Measured (BOQ rates)", 0.5757])
        ws.append(["PRJ-2026-003", "Flat C", "Flat", "Site B", 100, "2026-09-20", "In Progress", 0, "Cost-Plus %", 0.25])
        ca = wb.create_sheet("Cost Allocation")
        ca.append(["BEYOND DESIGN | COST ALLOCATION RULES"])
        ca.append(["Allocation rule", "Billable to client?", "Pooled amount (from journal)", "PRJ-2026-002", "PRJ-2026-003"])
        ca.append(["Even split", "No", 0, 0.75, 0.25])
        coa = wb.create_sheet("Chart of Accounts")
        coa.append(["Code", "Account Name", "Type", "Normal", "Income Statement line"])
        coa.append(["5040", "Site Preparation", "COGS", "Debit", "Cost of sales"])
        coa.append(["5045", "Site Survey", "COGS", "Debit", "Cost of sales"])
        gj = wb.create_sheet("General Journal")
        gj.append(["JE #", "Date", "Source / Type", "Project", "Stage", "Account", "Description", "Vendor / Party",
                   "Deal Ref", "Qty", "Unit", "Rate", "Debit", "Credit", "Note", "Alloc Rule", "Account name", "Entry check", "dist #"])
        for r in (["JE-1", "2026-09-23", "Payroll", "PRJ-2026-002", "All", "5045", "Survey", "Eng", None, 1, "LS", 2900, 2900, None],
                  ["JE-1", "2026-09-23", "Payroll", "PRJ-2026-002", "All", "1020", "Survey", "Eng", None, None, None, None, None, 2900],
                  ["JE-2", "2026-09-29", "Client advance", "PRJ-2026-002", "All", "1020", "Advance", "Client", None, None, None, None, 50000, None],
                  ["JE-2", "2026-09-29", "Client advance", "PRJ-2026-002", "All", "2100", "Advance", "Client", None, None, None, None, None, 50000]):
            gj.append(r)
        reg = wb.create_sheet("Invoice Register")
        reg.append(["Invoice No", "Date", "Project", "Client", "Location", "Type", "Status", "Gross this invoice", "Posted to journal?"])
        reg.append(["Invoice 01", "2026-09-23", "PRJ-2026-002", "Villa B", "Site A", "Interim", "Draft", 3500, "No — draft"])
        inv = wb.create_sheet("Invoice 1")
        inv.append(["BEYOND DESIGN", None, "23/9/2026", "التاريخ"])
        inv.append(["Villa B", "المشروع"])
        inv.append(["Site A", "الموقع"])
        inv.append(["Invoice  01", "الاصدار"])
        inv.append(["1- أعمال الهدم والبناء"])
        inv.append(["م", "البيان", "الوحدة", "الكميه بالعقد", "المنفذ فعلي", "السعر", "الاجمالي بالعقد", "الإجمالي الفعلي"])
        inv.append([1.1, "Walls", "م٢", 10, 4, 500, 5000, 2000])
        inv.append([1.2, "Not started", "م٢", 5, None, 100, 500, None])
        inv.append([None, "أعمال مستجدة: extra slab", "مقطوعية", 1, 1, 1500, 0, 1500])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            data = legacy.read_legacy_workbook(f.name)
        self.assertEqual(data["Business Units"][0]["basis"], "Fixed Price")
        self.assertEqual([b["overhead_pct"] for b in data["Business Units"]], [0.75, 0.25])   # the sheet's first rule
        self.assertEqual([(x["item_no"], x["qty"], x["price"]) for x in data["Invoice Lines"]],
                         [("1.1", 4, 500), ("V-01", 1, 1500)])
        db = Database(":memory:")
        interchange.import_data(db, data, mode="replace")
        tb = {r["code"]: round(r["debit"] - r["credit"], 2) for r in a.trial_balance(db)["rows"]}
        self.assertEqual((tb["5045"], tb["2100"]), (2900, -50000))
        self.assertEqual(db.one("SELECT kind FROM journal_entries WHERE ref='JE-2'")["kind"], "Receipt Voucher")
        self.assertEqual([(x["number"], x["total"]) for x in a.list_extracts(db)], [("Invoice 01", 3500)])

    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_clashing_account_number_and_partner_loan(self):
        """1330 is 'Partner loans' in the workbook but withholding tax in the app: it keeps its meaning under a new number,
        the partner owes us in aging, and the Balances Tracker's loan lands in the Loans register without posting twice."""
        import openpyxl
        import tempfile
        from erp import legacy, interchange, loans, statements
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Project Register"
        ws.append(["Project ID", "Client Name"])
        ws.append(["PRJ-2026-001", "Villa A"])
        coa = wb.create_sheet("Chart of Accounts")
        coa.append(["Code", "Account Name", "Type", "Normal"])
        coa.append(["1330", "Partner loans — owed by partners (سلف الشركاء)", "Asset", "Debit"])
        coa.append(["1020", "Bank Account", "Asset", "Debit"])
        gj = wb.create_sheet("General Journal")
        gj.append(["JE #", "Date", "Source / Type", "Project", "Stage", "Account", "Description", "Vendor / Party",
                   "Deal Ref", "Qty", "Unit", "Rate", "Debit", "Credit"])
        gj.append(["JE-1", "2026-01-01", "Partner funding (auto)", None, None, "1020", "Funding", "P1", None, None, None, None, 500000, None])
        gj.append(["JE-1", "2026-01-01", "Partner funding (auto)", None, None, "3150", "Funding", "P1", None, None, None, None, None, 500000])
        gj.append(["JE-2", "2026-03-01", "Partner loan", None, None, "1330", "Partner loan", "Partner B", None, None, None, None, 180000, None])
        gj.append(["JE-2", "2026-03-01", "Partner loan", None, None, "1020", "Partner loan", "Partner B", None, None, None, None, None, 180000])
        bt = wb.create_sheet("Balances Tracker")
        bt.append(["SECTION 6 - PARTNER LOANS (سلف الشركاء): MONEY LENT TO PARTNERS, TO BE PAID BACK"])
        bt.append(["Partner", "Lent", "Paid back in cash", "Settled from profit share", "Still owed", "Plan start date", "Plan length (months)"])
        bt.append(["Partner A", 0, 0, 0, 0])
        bt.append(["Partner B", 180000, 0, 0, 180000, "2026-03-01", 8])
        bt.append(["TOTAL STILL OWED BY PARTNERS", 180000, 0, 0, 180000])
        bt.append(["SECTION 5 - THE FUNDING PARTNER'S ACCOUNT: MONEY IN AND OUT  (P1)", None, None, None, None, "FUNDING LIMIT"])
        bt.append([None, None, None, None, None, "Agreed funding limit (includes everything he has put in)", None, None, 430000])
        gj.append(["JE-3", "2026-03-02", "Project cost", "PRJ-2026-001", "All", "5020", "Tiles", "Hani", "HANI-2", 1, "Task", 15000, 14000, None])
        gj.append(["JE-3", "2026-03-02", "Project cost", "PRJ-2026-001", "All", "1020", "Tiles", "Hani", "HANI-2", None, None, None, None, 14000])
        ng = wb.create_sheet("Negotiation Register")
        ng.append(["Date", "Deal Ref", "Project", "Vendor / subcontractor", "What was negotiated", "Quoted", "Settled at", "Saving", "Treatment"])
        ng.append(["2026-03-02", "HANI-2", "PRJ-2026-001", "Hani", "Tiles", 15000, 14000, 1000, "Keep — bill the client the quoted price"])
        do = wb.create_sheet("Design Offer")
        for k, v in (("Date", "2026-10-07"), ("Client", "Client name"), ("Project / unit", "Unit 1"), ("Scope line 1", "Layout"),
                     ("Total cost", 60000), ("Discount", 10000), ("Advance %", 0.6), ("After renders %", 0.2), ("After approval %", 0.2)):
            do.append([None, k, v])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            data = legacy.read_legacy_workbook(f.name)
        db = Database(":memory:")
        interchange.import_data(db, data, mode="replace")
        self.assertEqual((db.setting("funding_partner"), float(db.setting("funding_limit"))), ("P1", 430000))
        self.assertEqual([(r["deal_ref"], r["quoted"], r["savings"], r["mode"]) for r in a.negotiations(db)["rows"]],
                         [("HANI-2", 15000, 1000, "keep")])
        from erp import documents
        offers = documents.list_offers(db)
        self.assertEqual([(o["client"], documents.offer_totals(documents.pricing(o))["net"]) for o in offers], [("Client name", 50000)])
        tb = {r["code"]: round(r["debit"] - r["credit"], 2) for r in a.trial_balance(db)["rows"]}
        self.assertEqual((tb.get("1331"), tb.get("1330")), (180000, None))
        self.assertEqual(db.one("SELECT name FROM accounts WHERE code='1331'")["name"], "Partner loans — owed by partners (سلف الشركاء)")
        self.assertIn(("Partner B", 180000), [(r["party"], r["balance"]) for r in statements.aging(db, "2026-03-31")["rows"]])
        ln = loans.register(db)["rows"]
        self.assertEqual([(x["party"], x["outstanding"], x["term_months"], x["next_amount"]) for x in ln], [("Partner B", 180000, 8, 22500)])
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries")["n"], 2)   # the loan was not posted again; no
        # funding snapshot either: the sheet's bank is P1's own account, so every payment is on his funding account
        self.assertEqual((tb.get("3015"), tb.get("1020", 0)), (-194000, 0))
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])


class DemoCompany(unittest.TestCase):
    """The demo company loads in one go through the import sheets and touches every module; its workbook re-imports."""

    def test_demo_per_business_type(self):
        """Each business type gets only its own kind of data, under its own name, and its books pass the audit."""
        from erp import demo, industries
        seen = {}
        for k in demo.INDUSTRY_KEYS:
            db = Database(":memory:")
            industries.apply(db, k)
            demo.load(db, "admin", industry=k)
            self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [], k)
            self.assertEqual(db.setting("company_name"), demo.PROFILES[k]["name"])
            seen[k] = {t: db.one(f"SELECT COUNT(*) n FROM {t}")["n"] for t in ("rentals", "sales_orders", "boqs", "business_units")}
        self.assertTrue(seen["scaffolding"]["rentals"] and not seen["scaffolding"]["sales_orders"] and not seen["scaffolding"]["boqs"])
        self.assertTrue(seen["development"]["sales_orders"] and not seen["development"]["rentals"] and not seen["development"]["boqs"])
        self.assertTrue(seen["interior"]["boqs"] and not seen["interior"]["rentals"] and not seen["interior"]["sales_orders"])
        self.assertFalse(seen["investment"]["rentals"] or seen["investment"]["sales_orders"] or seen["investment"]["boqs"])

    def test_demo(self):
        from erp import demo, interchange, sales, bankrec, recurring, loans, custody, rentals
        db = Database(":memory:")
        demo.load(db, "admin")
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        kinds = {r["kind"] for r in db.q("SELECT DISTINCT kind FROM journal_entries")}
        for k in ("Direct Cost", "Operating Expense", "Payroll Accrual", "Receipt Voucher", "Payment Voucher", "Sales Invoice",
                  "Goods Receipt", "Material Issue", "Stock Transfer" if False else "Stock Adjustment", "Depreciation", "Loan Received",
                  "Loan Repayment", "Custody Issued", "Custody Settlement", "Rental Deposit", "Cheque Cleared", "Cheque Bounced",
                  "Installment Interest", "Bank Entry", "Partner Loan", "Share Transfer", "Supplier Discount", "Asset Purchase"):
            self.assertIn(k, kinds)
        self.assertTrue(loans.register(db)["rows"] and custody.register(db)["rows"] and rentals.list_rentals(db)["rows"])
        self.assertTrue(any(s["status"] == "reconciled" for s in bankrec.overview(db)["statements"]))
        self.assertTrue(recurring.listing(db) and sales.register(db)["rows"])
        with self.assertRaises(a.ValidationError):
            demo.load(db)                                   # never on top of real books
        if HAVE_XL:
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=".xlsx") as f:
                f.write(demo.workbook_bytes(lambda: Database(":memory:")))
                f.flush()
                data = interchange.read_xlsx(f.name)
            db2 = Database(":memory:")
            interchange.import_data(db2, data)
            self.assertEqual([c["check"] for c in a.audit(db2) if not c["ok"]], [])
            self.assertTrue(db2.one("SELECT 1 FROM bank_lines"))

        self.assertTrue(db.one("SELECT 1 FROM rent_items WHERE code='SC-FRM2' AND monthly_rate>0"))  # calculator codes priced

    def test_scaffold_calculator_formulas_match_the_app(self):
        """The Excel calculator sheet uses the app's own model, translated formula by formula."""
        import math
        import re
        from erp import rentals
        spec = rentals.FRAME_MODEL
        for vals in ({}, {"height": 24, "spare_pct": 10, "length": 41}, {"board_levels": 0, "tubes_per_tie": 2}):
            names = {r["key"]: vals.get(r["key"], r["value"]) for r in spec["inputs"] + spec["params"]}
            refs = {k: f"V_{k}" for k in names}
            for r in spec["derived"]:
                refs[r["key"]] = f"V_{r['key']}"
                names[r["key"]] = rentals.safe_eval(r["expr"], names)
            env = {f"V_{k}": v for k, v in names.items()}
            env.update(ROUNDUP=lambda x, n: math.ceil(x), IF=lambda c, a, b: a if c else b)
            for r in spec["derived"] + [{"expr": c["expr"]} for c in spec["components"]]:
                xl = re.sub(r"(?<![<>!=])=(?!=)", "==", rentals.to_excel(r["expr"], refs)).replace("<>", "!=").replace("^", "**")
                self.assertAlmostEqual(eval(xl, env), rentals.safe_eval(r["expr"], names), msg=r["expr"])  # noqa: S307
        if HAVE_XL:
            from erp import workbook
            import io as _io2
            import openpyxl as _ox
            wb = _ox.load_workbook(_io2.BytesIO(workbook.build_workbook(Database(":memory:"), template=True)))
            ws = wb["Scaffold Calculator"]
            self.assertIn("SC-FRM2", [c.value for c in ws["A"]])

class StatementsAndBudgets(unittest.TestCase):
    """A contact's statement with aging, and budgets against actual results."""

    def test_statement_and_aging(self):
        from erp import statements, reports
        db2 = Database(":memory:")
        from erp import demo
        demo.load(db2)
        st = statements.statement(db2, "Atlas Builders")
        self.assertAlmostEqual(st["closing"], st["opening"] + st["debits"] - st["credits"], places=2)
        self.assertAlmostEqual(sum(st["aging"]), st["closing"], places=2)
        self.assertEqual(st["closing"], sum(b["balance"] for b in st["balances"]))
        ag = statements.aging(db2)
        self.assertAlmostEqual(sum(ag["receivable"]["aging"]), ag["receivable"]["balance"], places=2)
        self.assertTrue(ag["payable"]["balance"] < 0 < ag["receivable"]["balance"])
        for name, p in (("party_statement", {"party": "Atlas Builders", "lang": "ar"}), ("aging", {"side": "payable"})):
            self.assertTrue(reports.build(db2, name, p)["rows"])
        with self.assertRaises(a.ValidationError):
            statements.statement(db2, "")

    def test_budgets(self):
        from erp import budgets, interchange
        db = Database(":memory:")
        from erp import demo
        demo.load(db)
        y = int(a.today()[:4])
        r = budgets.report(db, y, None, 12)
        rev = next(s for s in r["sections"] if s["type"] == "Revenue")
        self.assertEqual(rev["total"]["year_budget"], 7200000)
        exp = next(s for s in r["sections"] if s["type"] == "Expense")
        rent = next(x for x in exp["rows"] if x["account"] == "6030")
        self.assertEqual(rent["year_budget"], 36000)
        self.assertEqual(rent["variance"], round(36000 - rent["actual"], 2))          # a cost under budget is favourable
        budgets.save(db, y, 0, [{"account": "6020", "total": 1000}])
        g = {x["code"]: x for x in budgets.grid(db, y, 0)["rows"]}
        self.assertEqual(g["6020"]["total"], 1000)
        self.assertTrue(budgets.seed(db, y + 1, 0, y, 10))
        with self.assertRaises(a.ValidationError):
            budgets.save(db, y, 0, [{"account": "1010", "total": 5}])                 # balance-sheet accounts have no budget
        rows = interchange.export_data(db)["Budgets"]
        db2 = Database(":memory:")
        interchange.import_data(db2, {"Budgets": rows})
        self.assertEqual(budgets.export_rows(db2), budgets.export_rows(db))

class Attachments(unittest.TestCase):
    """Files kept with vouchers and records; they survive a JSON backup and restore."""

    def test_attach_and_restore(self):
        import base64
        import json as _json
        from erp import attachments, demo, interchange
        db = Database(":memory:")
        demo.load(db)
        eid = db.one("SELECT id FROM journal_entries WHERE kind='Operating Expense' ORDER BY id LIMIT 1")["id"]
        pdf = "data:application/pdf;base64," + base64.b64encode(b"%PDF-1.4 test").decode()
        a1 = attachments.add(db, "entry", eid, {"data": pdf, "title": "Supplier invoice.pdf"}, "admin")
        attachments.add(db, "loan", db.one("SELECT code FROM loans LIMIT 1")["code"], {"url": "https://example.com/contract", "title": "Contract"})
        attachments.add(db, "party", db.one("SELECT id FROM parties WHERE name='Atlas Builders'")["id"], {"data": pdf})
        self.assertEqual(a.get_entry(db, eid)["files"], 1)
        self.assertEqual(attachments.get(db, a1)["size"], 13)
        attachments.rename(db, a1, "Invoice 77.pdf")
        with self.assertRaises(a.ValidationError):
            attachments.add(db, "entry", eid, {"url": "ftp://x"})
        with self.assertRaises(a.ValidationError):
            attachments.add(db, "entry", 999999, {"data": pdf})
        with self.assertRaises(a.ValidationError):
            attachments.add(db, "entry", eid, {"data": "data:application/pdf;base64," + base64.b64encode(b"x" * (attachments.MAX_FILE + 1)).decode()})
        backup = _json.loads(interchange.to_json(db))
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(_json.dumps(backup).encode()))
        titles = sorted(r["title"] for r in db2.q("SELECT title FROM attachments"))
        self.assertEqual(titles, ["Contract", "File", "Invoice 77.pdf"])
        e2 = db2.one("SELECT owner_id FROM attachments WHERE title='Invoice 77.pdf'")["owner_id"]
        self.assertEqual(db2.one("SELECT memo FROM journal_entries WHERE id=?", (e2,))["memo"],
                         db.one("SELECT memo FROM journal_entries WHERE id=?", (eid,))["memo"])
        attachments.delete(db, a1)
        self.assertEqual(a.get_entry(db, eid)["files"], 0)

class PurchaseOrders(unittest.TestCase):
    """Order -> approve -> receive in parts (vouchers posted) -> close; restore from a JSON backup."""

    def test_po_flow(self):
        import json as _json
        from erp import demo, interchange, purchasing as P
        db = Database(":memory:")
        demo.load(db)
        self.assertEqual(db.one("SELECT status FROM purchase_orders WHERE code='PO-D100'")["status"], "approved")
        db.set_setting("po_self_approve_limit", "10000")
        pid = P.save(db, {"supplier": "Spark Electric", "bu_code": "PRJ-D01", "vat_code": "VAT14",
                          "lines": [{"kind": "cost", "description": "Cable 4 mm", "qty": 10, "unit": "roll", "rate": 1500},
                                    {"kind": "expense", "description": "Site lighting", "qty": 1, "rate": 800}]}, "clerk")
        with self.assertRaises(a.ValidationError):
            P.receive(db, pid, {"lines": [{"line_id": 1, "qty": 1}]})          # not approved yet
        with self.assertRaises(a.ValidationError):
            P.approve(db, pid, "clerk")                                        # above the self-approval limit
        P.approve(db, pid, "boss")
        ln = P.get(db, pid)["lines"]
        with self.assertRaises(a.ValidationError):
            P.receive(db, pid, {"lines": [{"line_id": ln[0]["id"], "qty": 11}]})
        before = a.party_balance(db, db.acct("ap"), "Spark Electric")
        ids = P.receive(db, pid, {"lines": [{"line_id": ln[0]["id"], "qty": 4}]}, "clerk")
        e = a.get_entry(db, ids[0])
        self.assertEqual((e["kind"], e["ref"], e["amount"]), ("Direct Cost", P.get(db, pid)["code"], 6840.0))
        self.assertAlmostEqual(a.party_balance(db, db.acct("ap"), "Spark Electric") - before, 6840.0)
        po = P.get(db, pid)
        self.assertEqual((po["status"], po["open_value"]), ("partial", 9800.0))
        a.void_entry(db, ids[0], "wrong", "boss")                              # a voided receipt is on order again
        self.assertEqual(P.get(db, pid)["open_value"], 15800.0)
        P.set_status(db, pid, "closed", "rest not needed")
        self.assertEqual(P.get(db, pid)["open_value"], 0)
        with self.assertRaises(a.ValidationError):
            P.set_status(db, pid, "cancelled")
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()))
        self.assertEqual(sorted((r["code"], r["status"]) for r in db2.q("SELECT code, status FROM purchase_orders")),
                         sorted((r["code"], r["status"]) for r in db.q("SELECT code, status FROM purchase_orders")))
        live = "SELECT COUNT(*) n FROM po_receipts r JOIN journal_entries e ON e.id=r.entry_id WHERE e.status='posted'"
        self.assertEqual(db2.one(live)["n"], db.one(live)["n"])
        self.assertEqual({r["code"]: r["open_value"] for r in P.listing(db2)["rows"]}, {r["code"]: r["open_value"] for r in P.listing(db)["rows"]})

class ServiceSubscriptions(unittest.TestCase):
    """Maintenance / warranty / software subscriptions: invoices per period, deferred revenue, visits, add-ons."""

    def test_subscription_cycle(self):
        from erp import demo, interchange, services as S
        db = Database(":memory:")
        demo.load(db)
        self.assertTrue(S.listing(db)["mrr"] > 0)
        p = S.save_plan(db, {"name": "Yearly maintenance", "price": 12000, "period_months": 12, "visits": 2, "extra_visit_price": 500})
        m = S.save_plan(db, {"name": "Support", "kind": "software", "price": 1000, "period_months": 1})
        a1 = S.save(db, {"customer": "Emily Clark", "plan_id": p, "covers": "Villa 3", "start": "2025-01-10", "term_months": 12, "auto_renew": 0})
        s2 = S.save(db, {"customer": "James Carter", "plan_id": m, "start": "2025-01-01"})
        S.save(db, {"customer": "James Carter", "plan_id": m, "parent_id": s2, "title": "Extra module", "price": 300, "period_months": 1, "start": "2025-01-01"})
        for k in range(3):
            S.save_visit(db, a1, {"date": f"2025-0{3 + k}-05", "description": "Check"})
        third = db.one("SELECT charge FROM sub_visits WHERE sub_id=? ORDER BY date DESC LIMIT 1", (a1,))["charge"]
        self.assertEqual(third, 50000)                                                  # beyond the 2 included visits
        r = S.run(db, "2025-03-31")
        inv = db.q("SELECT x.number FROM sub_invoices i JOIN extracts x ON x.id=i.extract_id WHERE i.sub_id=?", (s2,))
        self.assertEqual(len(inv), 3)                                                   # Jan, Feb, Mar (add-on on the same invoices)
        self.assertEqual(db.one("SELECT COUNT(DISTINCT extract_id) n FROM sub_invoices i JOIN subscriptions s ON s.id=i.sub_id "
                                "WHERE s.customer='James Carter' AND i.period_from<'2025-04-01'")["n"], 3)
        dfr = db.acct("service_deferred")
        bal = lambda acct: db.one("SELECT COALESCE(SUM(credit-debit),0) n FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id "  # noqa: E731
                                  "WHERE e.status='posted' AND l.account=? AND e.ref LIKE 'SUB-0%'", (acct,))["n"]
        self.assertEqual(bal(dfr), (12000 - 3 * 1000) * 100)                           # 3 of 12 months earned
        self.assertGreaterEqual(r["recognized"], 3)
        S.run(db, "2026-02-01")
        self.assertEqual(S.get(db, a1)["status"], "ended")
        self.assertEqual(db.one("SELECT COUNT(*) n FROM sub_visits WHERE sub_id=? AND charge>0 AND extract_id IS NULL", (a1,))["n"], 0)
        S.set_status(db, s2, "cancelled")
        self.assertEqual({x["status"] for x in db.q("SELECT status FROM subscriptions WHERE parent_id=?", (s2,))}, {"cancelled"})
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()))
        self.assertEqual(S.listing(db2)["mrr"], S.listing(db)["mrr"])
        self.assertEqual(S.listing(db2)["deferred"], S.listing(db)["deferred"])
        self.assertEqual(S.run(db2, a.today())["invoices"], 0)                          # nothing billed twice after a restore

class CompanyNamesAndDelete(unittest.TestCase):
    """One person cannot have two companies with one name; a super administrator deletes one by typing its name
    (mirrors supabase/migrations/*_company_names_delete.sql, checked on Postgres too)."""

    def test_unique_names_and_delete(self):
        import json as _json
        import urllib.error
        import urllib.request
        sys.path.insert(0, str(Path(__file__).parent))
        import mock_supabase
        sb = mock_supabase.start()
        base = f"http://127.0.0.1:{sb.server_address[1]}"
        op = urllib.request.build_opener(urllib.request.ProxyHandler({}))

        def rpc(fn, body, who):
            r = urllib.request.Request(f"{base}/rest/v1/rpc/{fn}", data=_json.dumps(body).encode(), method="POST",
                                       headers={"Content-Type": "application/json", "Authorization": f"Bearer tok:{who}"})
            try:
                with op.open(r) as x:
                    return x.status, _json.loads(x.read() or b"null")
            except urllib.error.HTTPError as e:
                return e.code, _json.loads(e.read() or b"null")
        first = rpc("create_company", {"p_name": "Acme"}, "ali@x.com")[1]["id"]
        self.assertEqual(rpc("create_company", {"p_name": "  acme "}, "ali@x.com")[0], 400)      # same name, same person
        self.assertEqual(rpc("create_company", {"p_name": "Acme"}, "omar@x.com")[0], 200)        # someone else may
        other = rpc("create_company", {"p_name": "Acme 2"}, "ali@x.com")[1]["id"]
        cos = rpc("my_companies", {}, "ali@x.com")[1]
        self.assertEqual([(c["name"], c["is_super"], c["changes"]) for c in cos], [("Acme", True, 0), ("Acme 2", True, 0)])
        self.assertTrue(all(c["created_at"] for c in cos))
        self.assertEqual(rpc("delete_company", {"p_company": other, "p_name": "Acme 2"}, "omar@x.com")[0], 403)  # not his
        self.assertEqual(rpc("delete_company", {"p_company": other, "p_name": "Acme"}, "ali@x.com")[0], 400)     # wrong name
        self.assertEqual(rpc("delete_company", {"p_company": other, "p_name": "acme 2"}, "ali@x.com")[0], 200)
        self.assertEqual([c["id"] for c in rpc("my_companies", {}, "ali@x.com")[1]], [first])
        sb.shutdown()


class HostedBilling(unittest.TestCase):
    """Trial -> expired (changes refused) -> paid; mirrors supabase/migrations/*_billing.sql (checked on Postgres too)."""

    def test_trial_expiry_and_payment(self):
        import json as _json
        import urllib.request
        sys.path.insert(0, str(Path(__file__).parent))
        import mock_supabase
        sb = mock_supabase.start()
        base = f"http://127.0.0.1:{sb.server_address[1]}"
        op = urllib.request.build_opener(urllib.request.ProxyHandler({}))

        def rpc(fn, body, who):
            r = urllib.request.Request(f"{base}/rest/v1/rpc/{fn}", data=_json.dumps(body).encode(), method="POST",
                                       headers={"Content-Type": "application/json", "Authorization": f"Bearer tok:{who}"})
            try:
                with op.open(r) as x:
                    return x.status, _json.loads(x.read() or b"null")
            except urllib.error.HTTPError as e:
                return e.code, _json.loads(e.read() or b"null")
        st = sb.state
        st.owners.add("boss@x.com")
        cid = rpc("create_company", {"p_name": "Acme"}, "ali@x.com")[1]["id"]
        appended = lambda seq: rpc("append_op", {"p_company": cid, "p_seq": seq, "p_version": "1", "p_app_user": "", "p_device": "t",  # noqa: E731
                                                  "p_method": "POST", "p_path": "/api/x", "p_body": "{}", "p_content_type": "application/json",
                                                  "p_clock": "2026-01-01T00:00:00"}, "ali@x.com")[0]
        self.assertEqual(rpc("billing_status", {"p_company": cid}, "ali@x.com")[1]["status"], "free")
        self.assertEqual(rpc("billing_set_config", {"p": {"enabled": True}}, "ali@x.com")[0], 403)   # customers cannot
        rpc("billing_set_config", {"p": {"enabled": True}}, "boss@x.com")
        s = rpc("billing_status", {"p_company": cid}, "ali@x.com")[1]
        self.assertEqual((s["status"], s["days_left"], s["users_allowed"]), ("trial", 15, 8))
        self.assertEqual(appended(1), 200)
        st.days_shift = 16
        self.assertEqual(rpc("billing_status", {"p_company": cid}, "ali@x.com")[1]["status"], "expired")
        self.assertEqual(appended(2), 400)                                                  # nothing saved once expired
        self.assertEqual(rpc("billing_quote", {"p_company": cid, "p_plan": "business", "p_months": 12, "p_extra_users": 1}, "ali@x.com")[1]["amount"],
                         (4500 + 350) * 10)
        rpc("billing_record_payment", {"p_company": cid, "p_plan": "starter", "p_months": 1, "p_extra_users": 0, "p_amount": None,
                                       "p_reference": "INSTAPAY-1", "p_note": ""}, "boss@x.com")
        s = rpc("billing_status", {"p_company": cid}, "ali@x.com")[1]
        self.assertEqual((s["status"], s["plan"], s["payments"][0]["amount"]), ("active", "starter", 1900))
        self.assertEqual(appended(2), 200)
        sb.shutdown()

class RentalServices(unittest.TestCase):
    """Installation / dismantling (months of rent or per m²) and transport (per trip): price vs our cost, scope."""

    def test_service_quote_and_contract_cost(self):
        from erp import demo, rentals as R
        db = Database(":memory:")
        demo.load(db)
        q = R.service_quote(db, 1152, 23000, 60000, "both")
        by = {x["code"]: x for x in q["lines"]}
        self.assertEqual((by["install"]["amount"], by["dismantle"]["amount"]), (36000, 24000))   # 0.6 + 0.4 = one month of rent
        self.assertEqual(by["install"]["total_cost"], 1152 * 6)                                    # crew paid per m²
        self.assertEqual((q["trips"], by["transport_out"]["amount"]), (3, 3 * 3500))               # 23 t on 10 t trucks
        self.assertEqual({x["code"] for x in R.service_quote(db, 1152, 23000, 60000, "install")["lines"]}, {"install", "transport_out", "transport_back"})
        self.assertEqual({x["charge_type"] for x in R.service_quote(db, 1152, 23000, 60000, "none")["lines"]}, {"transport"})
        small = R.service_quote(db, 40, 500, 2000, "both")
        self.assertTrue({x["code"]: x for x in small["lines"]}["install"]["min_applied"])            # minimum charge
        R.save_service_rates(db, {"truck_tons": 8, "services": [{"code": "inst", "charge_type": "install", "name": "Install",
                                                                 "basis": "m2", "price": 25, "cost_basis": "m2", "cost": 15}]})
        self.assertEqual(R.service_quote(db, 100, 0, 0)["lines"][0]["profit"], 1000)
        with self.assertRaises(a.ValidationError):
            R.service_quote(db, 100, 0, 0, "half")
        rid = db.one("SELECT id FROM rentals ORDER BY id LIMIT 1")["id"]
        d = R.detail(db, rid)
        lines = [{**{k: ln[k] for k in ("id", "kind", "item_id", "kit_id", "model_id", "description", "unit", "qty", "rate",
                                         "charge_type", "charge_when", "inputs")},
                  "components": [{"item_id": c["item_id"], "per_unit": c["per_unit"]} for c in ln["components"]]} for ln in d["lines"]]
        lines.append({"kind": "charge", "charge_type": "install", "charge_when": "delivery", "description": "Installation", "qty": 1,
                      "rate": 30000, "cost": 20736})
        R.save(db, {**{k: d["rental"][k] for k in ("client", "site", "start_date", "duration_months")}, "lines": lines}, rid=rid)
        self.assertEqual(next(ln for ln in R.detail(db, rid)["lines"] if ln["description"] == "Installation" and ln["rate"] == 30000)["cost"], 20736)

class SuperAdmin(unittest.TestCase):
    """Super administrators manage administrators and can erase / replace the books; there is always one."""

    def test_levels(self):
        from erp import auth
        db = Database(":memory:")
        boss = auth.create_user(db, {"username": "boss", "password": "pw123456"}, first=True)
        self.assertTrue(auth.public_user(db.one("SELECT * FROM users WHERE id=?", (boss,)))["is_super"])
        adm = auth.create_user(db, {"username": "adm", "password": "pw123456", "is_admin": True})
        clerk = auth.create_user(db, {"username": "clerk", "password": "pw123456"})
        self.assertFalse(auth.public_user(db.one("SELECT * FROM users WHERE id=?", (adm,)))["is_super"])
        with self.assertRaises(a.ValidationError):
            auth.update_user(db, boss, {"is_super": False}, adm)              # the only super administrator
        auth.update_user(db, adm, {"is_super": True}, boss)
        auth.update_user(db, boss, {"is_super": False}, adm)                  # fine now: another one exists
        self.assertTrue(db.one("SELECT is_admin FROM users WHERE id=?", (boss,))["is_admin"])
        with self.assertRaises(a.ValidationError):
            auth.delete_user(db, adm, boss)                                   # last super administrator
        self.assertTrue(clerk)

    def test_server_rules(self):
        import io as _io
        from erp import auth, server
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "pw123456"}, first=True)
        auth.create_user(db, {"username": "adm", "password": "pw123456", "is_admin": True})
        clerk = auth.create_user(db, {"username": "clerk", "password": "pw123456"})
        class App:
            pass
        app = App()
        app.db, app.backup = db, (lambda *a_, **k: None)
        tok_adm, _ = auth.login(db, "adm", "pw123456")
        call = lambda m, p, b=None: server.handle(app, m, p, {"x-session": tok_adm, "x-requested-with": "BeyondERP",  # noqa: E731
                                                              "content-type": "application/json"}, json.dumps(b or {}).encode())
        r = call("POST", "/api/erase", {"confirm": "ERASE", "scope": "transactions"})
        self.assertEqual(r.status, 403)
        self.assertIn("super administrator", json.loads(r.body)["error"])
        self.assertEqual(call("POST", "/api/users", {"username": "x", "password": "pw123456", "is_admin": True}).status, 403)
        # administrators add people; changing someone else's account is for the super administrator
        self.assertEqual(call("PUT", f"/api/users/{clerk}", {"full_name": "Clerk K", "is_admin": False, "active": True}).status, 403)
        self.assertEqual(call("POST", "/api/users", {"username": "clerk2", "password": "pw123456"}).status, 200)
        boss = db.one("SELECT id FROM users WHERE username='boss'")["id"]
        self.assertEqual(call("PUT", f"/api/users/{boss}", {"active": False}).status, 403)       # admins cannot touch the owner

class SignInLockout(unittest.TestCase):
    def test_lockout_after_wrong_passwords(self):
        from erp import auth
        db = Database(":memory:")
        auth.create_user(db, {"username": "lock", "password": "rightpass1"}, first=True)
        auth._fails.clear()
        for _ in range(auth.MAX_FAILS):
            with self.assertRaises(a.ValidationError):
                auth.login(db, "lock", "wrong")
        with self.assertRaises(a.ValidationError) as e:
            auth.login(db, "lock", "rightpass1")          # even the right password waits until the lock passes
        self.assertIn("Too many", str(e.exception))
        auth._fails.clear()
        self.assertTrue(auth.login(db, "lock", "rightpass1")[0])


class ArabicTexts(unittest.TestCase):
    """Texts the app writes (memos, messages, starting records) read in the person's language."""

    def test_messages_and_memos(self):
        import re as _re
        from erp import errors_ar, textl10n
        for en, ar in errors_ar.MESSAGES.items():           # every value in the English has its place in the Arabic
            self.assertEqual(sorted(set(_re.findall(r"\{(\d)\}", ar))), [str(i) for i in range(1, en.count("{}") + 1)], en)
        self.assertEqual(textl10n.error("amount is required", "ar"), "المبلغ مطلوب")
        self.assertEqual(textl10n.error("voucher 12: description is required", "ar"), "السند 12: البيان مطلوب")
        self.assertEqual(textl10n.error("Account 1101 is inactive", "ar"), "الحساب 1101 غير نشط")
        self.assertEqual(textl10n.error("Please sign in", "en"), "Please sign in")
        self.assertEqual(textl10n.error("Something new", "ar"), "Something new")   # unknown text is left as it is
        self.assertEqual(textl10n.text("Supplier Discount", "ar"), "خصم مورد")
        self.assertEqual(textl10n.text("Opening balance", "en"), "Opening balance")

    def test_server_replies(self):
        from erp import auth, server
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "pw123456"}, first=True)

        class App:
            pass
        app = App()
        app.db, app.backup = db, (lambda *a_, **k: None)
        tok, _ = auth.login(db, "boss", "pw123456")
        h = {"x-session": tok, "x-requested-with": "BeyondERP", "content-type": "application/json"}
        r = server.handle(app, "POST", "/api/bus", {**h, "x-lang": "ar"}, b'{"code": "P1"}')
        self.assertEqual(r.status, 400)
        self.assertTrue(any("؀" <= c <= "ۿ" for c in json.loads(r.body)["error"]))
        r = server.handle(app, "POST", "/api/bus", h, b'{"code": "P1"}')
        self.assertFalse(any("؀" <= c <= "ۿ" for c in json.loads(r.body)["error"]))
        self.assertEqual(json.loads(server.handle(app, "GET", "/api/nope?lang=ar", h).body)["error"], "غير موجود")

    def test_starting_records_follow_the_language(self):
        from erp import textl10n
        db = Database(":memory:")
        names = lambda: {w["name"] for w in db.q("SELECT name FROM warehouses")}   # noqa: E731
        self.assertIn("Main Store", names())
        textl10n.seed_names(db, "ar")
        self.assertIn("المخزن الرئيسي", names())
        textl10n.seed_names(db, "en")
        self.assertIn("Main Store", names())

    def test_offer_types_are_general(self):
        from erp.documents import DEFAULT_OFFER_TYPES
        self.assertEqual(len(DEFAULT_OFFER_TYPES), 2)
        self.assertNotIn("Design", " ".join(t["name"] for t in DEFAULT_OFFER_TYPES))


class UndoAfterSaving(unittest.TestCase):
    """'Undo' right after saving: only your own change, only for a while, never over later work or closed books."""

    def setUp(self):
        from erp import auth, server
        self.db = db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "pw123456"}, first=True)
        auth.create_user(db, {"username": "other", "password": "pw123456", "is_admin": 1})

        class App:
            pass
        self.app = App()
        self.app.db, self.app.backup = db, (lambda *a_, **k: None)
        self.tok = {u: auth.login(db, u, "pw123456")[0] for u in ("boss", "other")}
        self.server = server

    def call(self, m, p, b=None, who="boss"):
        r = self.server.handle(self.app, m, p, {"x-session": self.tok[who], "x-requested-with": "BeyondERP",
                                                "content-type": "application/json"}, json.dumps(b or {}).encode())
        return r.status, json.loads(r.body or b"{}")

    def names(self):
        return {r["name"] for r in self.db.q("SELECT name FROM parties")}

    def test_own_change_is_undone(self):
        st, _ = self.call("POST", "/api/parties-master?_undo=uabc123", {"name": "Acme", "type": "customer"})
        self.assertEqual(st, 200)
        self.assertTrue(self.call("GET", "/api/ai-undo/uabc123")[1]["can_undo"])
        self.assertEqual(self.call("POST", "/api/undo/uabc123", who="other")[0], 403)      # not theirs
        self.assertEqual(self.call("POST", "/api/undo/uabc123")[0], 200)
        self.assertNotIn("Acme", self.names())
        self.assertEqual(self.call("POST", "/api/undo/uabc123")[0], 400)                   # only once
        self.assertEqual(self.call("POST", "/api/undo/not-a-tag")[0], 400)

    def test_delete_is_put_back(self):
        _, r = self.call("POST", "/api/parties-master", {"name": "Keep", "type": "supplier"})
        self.assertEqual(self.call("DELETE", f"/api/parties-master/{r['id']}?_undo=udel001")[0], 200)
        self.assertNotIn("Keep", self.names())
        self.assertEqual(self.call("POST", "/api/undo/udel001")[0], 200)
        self.assertIn("Keep", self.names())

    def test_later_work_and_age_are_respected(self):
        _, r = self.call("POST", "/api/parties-master?_undo=ulater1", {"name": "Old", "type": "customer"})
        self.call("PUT", f"/api/parties-master/{r['id']}", {"name": "Old", "type": "customer", "notes": "changed later"})
        st, e = self.call("POST", "/api/undo/ulater1")
        self.assertEqual(st, 400)
        self.assertIn("were changed after", e["error"])
        self.call("POST", "/api/parties-master?_undo=uaged01", {"name": "Aged", "type": "customer"})
        self.db.x("UPDATE ai_undo SET at='2000-01-01T00:00:00' WHERE change='uaged01'")
        self.assertIn("minutes", self.call("POST", "/api/undo/uaged01")[1]["error"])
        self.assertIn("Aged", self.names())

    def test_closed_books_and_protected_records(self):
        from erp import aiundo
        rows = [{"tbl": "journal_entries", "old": None, "new": json.dumps({"date": "'2026-01-15'"})}]
        self.db.set_setting("lock_date", "2026-03-31")
        self.assertIn("closed books", aiundo.own_problem(self.db, rows))
        self.db.set_setting("lock_date", "")
        self.assertIsNone(aiundo.own_problem(self.db, rows))
        self.assertIsNotNone(aiundo.own_problem(self.db, [{"tbl": "users", "old": None, "new": "{}"}]))
        sent = [{"tbl": "extracts", "old": json.dumps({"eta_uuid": "'ABC'"}), "new": json.dumps({"eta_uuid": "'ABC'", "status": "'posted'"})}]
        self.assertIn("tax authority", aiundo.own_problem(self.db, sent))


class OfflineLicence(unittest.TestCase):
    """Signed licences for the desktop edition: trial, read-only without a licence, users / installation / updates."""

    def setUp(self):
        import os
        import secrets as _s
        from erp import licence
        self.seed = _s.token_bytes(32)
        os.environ["BEYOND_LICENCE_PUBLIC_KEY"] = licence.public_from_private(self.seed).hex()

    def tearDown(self):
        import os
        os.environ.pop("BEYOND_LICENCE_PUBLIC_KEY", None)

    def test_lifecycle(self):
        import datetime as _dt
        import secrets as _s
        from erp import auth, licence, server
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "pw123456"}, first=True)
        st = licence.status(db)
        self.assertEqual((st["state"], st["can_write"], st["days_left"]), ("trial", True, 15))

        class App:
            pass
        app = App()
        app.db, app.backup, app.licensing = db, (lambda *a_, **k: None), True
        tok, _ = auth.login(db, "boss", "pw123456")
        call = lambda m, p, b=None: server.handle(app, m, p, {"x-session": tok, "x-requested-with": "BeyondERP",  # noqa: E731
                                                              "content-type": "application/json"}, json.dumps(b or {}).encode())
        self.assertEqual(call("POST", "/api/bus", {"code": "P1", "name": "One"}).status, 200)
        db.set_setting("first_run", (_dt.date.today() - _dt.timedelta(days=16)).isoformat())
        self.assertEqual(licence.status(db)["state"], "unlicensed")
        self.assertEqual(call("POST", "/api/bus", {"code": "P2", "name": "Two"}).status, 402)   # read-only
        self.assertEqual(call("GET", "/api/bus").status, 200)                                     # viewing still works
        iid = licence.install_id(db)
        base = {"company": "Acme", "edition": "business", "users": 2, "companies": 1, "updates_until": "2099-01-01", "install": iid}
        self.assertEqual(call("POST", "/api/licence", {"key": "BEYOND1-garbage.x"}).status, 400)
        forged = licence.issue(_s.token_bytes(32), base)                                          # signed with another key
        self.assertEqual(call("POST", "/api/licence", {"key": forged}).status, 400)
        other = licence.issue(self.seed, {**base, "install": "AAAA-BBBB-CCCC-DDDD"})
        self.assertEqual(call("POST", "/api/licence", {"key": other}).status, 400)
        old = licence.issue(self.seed, {**base, "updates_until": "2020-01-01"})
        db.set_setting("licence_key", old)
        self.assertEqual(licence.status(db)["state"], "unlicensed")                               # updates ended before this version
        good = licence.issue(self.seed, base)
        r = call("POST", "/api/licence", {"key": good})
        self.assertEqual((r.status, json.loads(r.body)["state"]), (200, "licensed"))
        self.assertEqual(call("POST", "/api/bus", {"code": "P2", "name": "Two"}).status, 200)
        self.assertEqual(call("POST", "/api/users", {"username": "u2", "password": "pw123456"}).status, 200)
        self.assertEqual(call("POST", "/api/users", {"username": "u3", "password": "pw123456"}).status, 400)  # 2 users bought

    def test_free_without_vendor_key(self):
        import os
        from erp import licence
        os.environ.pop("BEYOND_LICENCE_PUBLIC_KEY", None)
        if licence.PUBLIC_KEY_FILE.exists():
            self.skipTest("a vendor key is installed in this checkout")
        self.assertEqual(licence.status(Database(":memory:"))["state"], "free")

class ImportKit(unittest.TestCase):
    """Per-screen templates, and importing a user's own layout with column matching."""

    def test_own_layout_and_templates(self):
        from erp import importkit, interchange
        db = seeded()
        own = ("مصروفات الشهر\n\nالتاريخ,البيان,المبلغ,المورد,حساب,ملاحظة\n01/10/2026,Internet,٩٠٠,WE,6020,x\n"
               "05/10/2026,Electricity,\"1,250.50\",Elec Co,6040,\n06/10/2026,Bad,abc,,6040,\n").encode()
        an = importkit.analyze(db, "Operating Expenses", own, "mine.csv")
        self.assertEqual((an["header_row"], an["rows"], an["missing"]), (2, 3, []))
        self.assertEqual(an["mapping"], {"0": "date", "1": "description", "2": "amount", "3": "vendor", "4": "account"})
        rows, nums = importkit.read(db, "Operating Expenses", own, "mine.csv")
        self.assertEqual((rows[0]["amount"], rows[1]["amount"], nums), ("900", "1250.50", [4, 5, 6]))
        with self.assertRaises(interchange.ImportErrors) as e:
            interchange.import_data(db, {"Operating Expenses": rows}, dry_run=True)
        self.assertEqual(len(e.exception.errors), 1)
        interchange.import_data(db, {"Operating Expenses": rows[:2]})
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries WHERE memo LIKE 'Electricity%'")["n"], 1)
        with self.assertRaises(a.ValidationError):
            importkit.read(db, "Operating Expenses", own, "mine.csv", {"0": "date"})     # required columns not matched
        for name in importkit.importable():                                          # every layout has a template
            self.assertTrue(importkit.template_csv(name))
        if HAVE_XL:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(importkit.template_xlsx(db, "Products", "ar")))
            self.assertIn("How to fill", wb.sheetnames + ["How to fill"])
            ws = wb.worksheets[0]
            col = {c["key"]: i for i, c in enumerate(importkit.describe("Products")["columns"], start=1)}
            ws.cell(5, col["name"], "Villa 9")
            ws.cell(5, col["price"], 5000000)
            ws.cell(5, col["specs"], "Area: 300 m²; Floor: 2")
            buf = io.BytesIO()
            wb.save(buf)
            self.assertTrue(importkit.analyze(db, "Products", buf.getvalue(), "t.xlsx")["exact"])
            r, _ = importkit.read(db, "Products", buf.getvalue(), "t.xlsx")
            interchange.import_data(db, {"Products": r})
            from erp import sales
            p = sales.product_row(db, sales.get_product(db, "P-0001"))
            self.assertEqual((p["name"], p["price"], p["specs"]), ("Villa 9", 5000000, [["Area", "300 m²"], ["Floor", "2"]]))


class BankReconciliation(unittest.TestCase):
    """Import a statement, match it to the vouchers, book what the bank did on its own, reconcile and lock."""

    def test_parse_formats(self):
        from erp import bankrec as B
        ar = "كشف حساب\nالتاريخ,البيان,مدين,دائن,الرصيد\n٠٥/٠٣/٢٠٢٦,إيداع نقدي,,\"1,500.00\",\"11,500.00\"\n06/03/2026,رسوم,25,,11475\n"
        p = B.parse(ar)
        self.assertEqual([(r["date"], r["amount"]) for r in p["rows"]], [("2026-03-05", 1500), ("2026-03-06", -25)])
        self.assertEqual((p["opening"], p["closing"]), (10000, 11475))
        us = "Date\tDescription\tAmount\n03/25/2026\tPOS\t(12.50)\n03/01/2026\tSalary\t5000\n"
        p = B.parse(us)                                   # month/day detected; newest-first files are flipped
        self.assertEqual([(r["date"], r["amount"]) for r in p["rows"]], [("2026-03-01", 5000), ("2026-03-25", -12.5)])
        self.assertEqual(B.parse_amount("1.234,50"), 1234.5)
        self.assertEqual(B.parse_amount("250.00 DR"), -250)
        if HAVE_XL:
            import base64
            import io
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Value Date", "Narrative", "Withdrawals", "Deposits", "Balance"])
            ws.append([datetime.date(2026, 4, 1), "Fee", 10, None, 990])
            buf = io.BytesIO()
            wb.save(buf)
            p = B.parse("data:application/octet-stream;base64," + base64.b64encode(buf.getvalue()).decode(), "s.xlsx")
            self.assertEqual((p["rows"][0]["date"], p["rows"][0]["amount"], p["opening"]), ("2026-04-01", -10, 1000))
        with self.assertRaises(a.ValidationError):
            B.parse("nothing,here\n1,2\n")

    def test_reconcile(self):
        from erp import bankrec as B, interchange
        db = seeded()
        bank = db.acct("bank")
        rows = [{"date": "2026-01-02", "description": "Transfer partner A", "amount": 100000},
                {"date": "2026-02-01", "description": "Office rent", "amount": -5000},
                {"date": "2026-02-03", "description": "CHQ Sup", "amount": -3000},
                {"date": "2026-02-06", "description": "Salary Ali", "amount": -1000},
                {"date": "2026-02-11", "description": "Deposit", "amount": 20000},
                {"date": "2026-02-28", "description": "BANK CHARGES FEB", "amount": -50},
                {"date": "2026-02-28", "description": "Credit interest", "amount": 12}]
        sid = B.save(db, {"account": bank, "date_from": "2026-01-01", "date_to": "2026-02-28", "opening": 0, "closing": 110962, "lines": rows})
        self.assertEqual(B.auto_match(db, sid), 5)
        sm = B.detail(db, sid)["summary"]
        self.assertEqual((sm["outstanding_payments"], sm["unrecorded"], sm["difference"]), (5000, -38, 0))  # tiles not cleared yet
        with self.assertRaises(a.ValidationError):
            B.reconcile(db, sid)                                           # two bank lines still open
        open_ = [ln for ln in B.detail(db, sid)["lines"] if ln["status"] == "open"]
        B.create_entry(db, sid, open_[0]["id"], {"account": db.acct("bank_charges"), "remember": 1, "pattern": "bank charges"})
        B.create_entry(db, sid, open_[1]["id"], {"account": db.acct("interest_income")})
        with self.assertRaises(a.ValidationError):                         # the amounts must agree
            B.match(db, sid, open_[0]["id"], [1])
        B.reconcile(db, sid, {"date": "2026-03-01"})
        with self.assertRaises(a.ValidationError):
            B.unmatch(db, sid, open_[0]["id"])
        sid2 = B.save(db, {"account": bank, "date_from": "2026-03-01", "date_to": "2026-03-31", "opening": 110962, "closing": 102912,
                           "lines": [{"date": "2026-02-28", "description": "Tiles", "amount": -5000},
                                     {"date": "2026-03-03", "description": "Payment Sup", "amount": -1000},
                                     {"date": "2026-03-31", "description": "Bank charges MAR", "amount": -50},
                                     {"date": "2026-03-31", "description": "Drawing B", "amount": -2000}]})
        self.assertEqual((B.auto_match(db, sid2), B.apply_rules(db, sid2)), (3, 1))
        self.assertEqual(B.detail(db, sid2)["summary"]["difference"], 0)
        B.reconcile(db, sid2, {"date": "2026-04-01"})
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        # paper statement: tick vouchers
        d3 = B.save(db, {"account": db.acct("cash"), "date_from": "2026-01-01", "date_to": "2026-03-31", "opening": 0, "closing": 0})
        self.assertEqual(B.detail(db, d3)["lines"], [])
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        self.assertEqual(B.get(db2, sid2)["status"], "reconciled")
        self.assertEqual(B.detail(db2, sid2)["summary"]["difference"], 0)
        self.assertEqual(B.detail(db2, sid)["summary"]["open_lines"], 0)


class Recurring(unittest.TestCase):
    """Templates post on schedule exactly once; review entries take a changed amount or are skipped."""

    def test_recurring(self):
        from erp import recurring as R, interchange
        db = seeded()
        t = R.save(db, {"name": "Rent", "doc_type": "expense", "start": "2026-01-31", "month_end": 1, "mode": "auto",
                        "payload": {"description": "Office rent {month}", "amount": 5000, "account": "6010"}})
        self.assertEqual(R.occurrences(R.get(db, t), "2026-05-15"), ["2026-01-31", "2026-02-28", "2026-03-31", "2026-04-30"])
        self.assertEqual(R.run(db, "2026-04-30")["posted"], 4)
        self.assertEqual(R.run(db, "2026-04-30")["posted"], 0)              # never twice
        self.assertEqual([h["vouchers"][0]["memo"] for h in R.history(db, t)][0], "Office rent 2026-04")
        e = R.history(db, t)[0]["vouchers"][0]["id"]
        self.assertEqual(a.get_entry(db, e)["meta"]["form"]["description"], "Office rent 2026-04")   # editable in its form
        el = R.save(db, {"name": "Electricity", "doc_type": "expense", "start": "2026-02-10", "mode": "review",
                         "payload": {"description": "Electricity {month}", "amount": 0, "account": "6040"}})
        self.assertEqual(R.run(db, "2026-03-31")["posted"], 0)              # review ones wait
        R.post(db, el, "2026-02-10", {"amount": 812.5})
        R.skip(db, el, "2026-03-10")
        with self.assertRaises(a.ValidationError):
            R.post(db, el, "2026-03-10")
        with self.assertRaises(a.ValidationError):
            R.post(db, el, "2026-03-11")                                     # not one of its dates
        q = R.save(db, {"name": "Accrual", "doc_type": "journal", "start": "2026-01-31", "unit": "month", "every": 3, "month_end": 1,
                        "max_count": 2, "mode": "auto", "payload": {"memo": "Accrual {month}", "lines": [{"account": "6090", "debit": 100}, {"account": "2030", "credit": 100}]}})
        bad = R.save(db, {"name": "Broken", "doc_type": "payment", "start": "2026-01-01", "mode": "auto", "payload": {"party": "Nobody", "amount": 999999}})
        r = R.run(db, "2026-12-31")
        self.assertEqual((r["posted"], len(r["errors"])), (8 + 2, 1))      # rent May–Dec + 2 accruals; the bad payment stops
        self.assertTrue(R.get(db, bad)["last_error"])
        self.assertEqual(R.run(db, "2026-12-31")["posted"], 0)
        self.assertEqual(len(R.history(db, q)), 2)
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        self.assertEqual(R.run(db2, "2026-12-31")["posted"], 0)
        self.assertEqual(len(R.history(db2, t)), 12)


class Contacts(unittest.TestCase):
    """Suppliers, subcontractors and drivers saved with their contact details; what the books say about each."""

    def test_directory_and_summary(self):
        db = seeded()
        pid = a.save_party(db, {"name": "Osama Driver", "category": "Driver", "phone": "0001234567",
                                "vehicle": "ABC 123", "bank_details": "CIB 1000"})
        p = db.one("SELECT * FROM parties WHERE id=?", (pid,))
        self.assertEqual((p["category"], p["type"], p["vehicle"]), ("driver", "supplier", "ABC 123"))
        cid = a.save_party(db, {"name": "Client Co", "category": "customer"})
        self.assertEqual(db.one("SELECT type FROM parties WHERE id=?", (cid,))["type"], "customer")
        nid = a.save_party(db, {"name": "Plain"})                     # old callers: no category → customer as before
        self.assertEqual(db.one("SELECT type FROM parties WHERE id=?", (nid,))["type"], "customer")
        a.save_party(db, {"name": "Rapid Plumbing", "category": "Subcontractor", "type": "both"})
        self.assertEqual(db.one("SELECT type, category FROM parties WHERE name='Rapid Plumbing'"), {"type": "both", "category": "subcontractor"})
        a.record_expense(db, {"description": "Trip", "amount": 800, "vendor": "Osama Driver", "paid": 300, "date": "2026-03-01"})
        a.record_payment(db, {"party": "Osama Driver", "amount": 200, "date": "2026-03-05"})
        s = a.contact_summary(db, "Osama Driver")
        self.assertEqual(s["spent"], 800)
        self.assertEqual(sum(b["we_owe"] for b in s["balances"]), 300)
        self.assertEqual((s["vouchers"], len(s["recent"])), (2, 2))
        a.save_party(db, {**p, "name": "Osama the Driver", "type": ""}, pid)  # a rename keeps the history
        self.assertEqual(a.contact_summary(db, "Osama the Driver")["vouchers"], 2)


class SupplierDiscount(unittest.TestCase):
    """A deal settled for less than what is still owed: the rest is closed as a discount, not left in payables."""

    def test_settle_with_discount(self):
        db = seeded()
        deal = {"bu_code": "PRJ-1", "description": "Ceramic", "vendor": "Hani", "qty": 1, "rate": 16000, "savings": 1000,
                "paid": 14000, "deal_ref": "HANI-1", "deal_status": "Settled", "account": "5020", "date": "2026-03-01"}
        with self.assertRaises(a.ValidationError):          # new entries: "Settled" with money still owed is refused
            a.record_project_cost(db, deal, strict=True)
        a.record_project_cost(db, deal)                     # old entries / replays still load
        ap = db.acct("ap")
        self.assertEqual(a.party_balance(db, ap, "Hani"), 1000)
        dflt = a.supplier_discount_defaults(db, "Hani")
        self.assertEqual((dflt["credit_account"], dflt["deal_ref"], dflt["owed"]), ("5020", "HANI-1", 1000))
        cost_before = a.income_statement(db)["net_profit"]
        with self.assertRaises(a.ValidationError):          # cannot discount more than is owed
            a.record_supplier_discount(db, {"party": "Hani", "amount": 1500})
        a.record_supplier_discount(db, {"party": "Hani", "date": "2026-03-02"})
        self.assertEqual(a.party_balance(db, ap, "Hani"), 0)
        self.assertEqual(round(a.income_statement(db)["net_profit"] - cost_before, 2), 1000)
        g = [x for x in a.deals(db) if x["ref"] == "HANI-1"][0]
        self.assertEqual((g["quoted"], g["savings"], g["cost"], g["status"]), (16000, 2000, 14000, "Settled"))


class Negotiations(unittest.TestCase):
    """Negotiated discounts: company earnings (cost stays at full price) or a lower cost, per project and per entry."""

    def test_keep_and_pass(self):
        db = seeded()
        earn = db.acct("negotiation_income")
        self.assertEqual(db.one("SELECT type FROM accounts WHERE code=?", (earn,))["type"], "Revenue")
        base = {"bu_code": "PRJ-1", "description": "Tiles", "vendor": "Hani", "qty": 1, "rate": 15000, "savings": 1000,
                "paid": 14000, "account": "5020", "deal_ref": "D1", "date": "2026-03-01"}
        e1 = a.get_entry(db, a.record_project_cost(db, {**base, "negotiation": "keep"}))
        self.assertEqual({(ln["account"], ln["debit"], ln["credit"]) for ln in e1["lines"]},
                         {("5020", 15000, 0), (earn, 0, 1000), ("1020", 0, 14000)})
        e2 = a.get_entry(db, a.record_project_cost(db, {**base, "deal_ref": "D2"}))          # no choice, no BU setting
        self.assertEqual(e2["meta"]["negotiation"], "pass")
        self.assertIn(("5020", 14000, 0), {(ln["account"], ln["debit"], ln["credit"]) for ln in e2["lines"]})
        db.x("UPDATE business_units SET negotiation='keep' WHERE code='PRJ-1'")               # project default
        e3 = a.get_entry(db, a.record_project_cost(db, {**base, "deal_ref": "D3"}))
        self.assertEqual(e3["meta"]["negotiation"], "keep")
        with self.assertRaises(a.ValidationError):
            a.record_project_cost(db, {**base, "negotiation": "maybe"})
        n = a.negotiations(db)
        self.assertEqual((n["total"]["savings"], n["total"]["kept"], n["total"]["passed"]), (3200, 2000, 1200))
        prj = [g for g in n["by_bu"] if g["bu_code"] == "PRJ-1"][0]
        self.assertEqual((prj["kept"], prj["setting"]), (2000, "keep"))
        b = [x for x in a.bu_summaries(db) if x["code"] == "PRJ-1"][0]
        self.assertEqual(b["savings_kept"], 2000)
        # a late discount on a 'keep' project goes to Negotiation Earnings by default
        a.record_project_cost(db, {**base, "savings": 0, "paid": 14000, "deal_ref": "D4", "negotiation": "keep"})
        d = a.supplier_discount_defaults(db, "Hani")
        self.assertEqual((d["negotiation"], d["credit_account"], d["owed"]), ("keep", earn, 1000))
        a.record_supplier_discount(db, {"party": "Hani", "negotiation": "pass"})
        self.assertEqual(a.party_balance(db, db.acct("ap"), "Hani"), 0)
        self.assertEqual(a.negotiations(db)["rows"][-1]["mode"], "pass")

    def test_edit_through_form_and_search(self):
        from erp.server import ERPServer
        tmp = tempfile.TemporaryDirectory()
        db = Database(Path(tmp.name) / "t.sqlite3")
        srv = ERPServer(("127.0.0.1", 0), db, Path(tmp.name) / "backups")
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            api = Api("call")
            api.base = f"http://127.0.0.1:{srv.server_address[1]}"
            op = api.client()
            self.assertEqual(api.call(op, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0], 200)
            bu = api.call(op, "POST", "/api/bus", {"code": "P1", "name": "Villa", "negotiation": "keep"})[1]["id"]
            deal = {"bu_id": bu, "description": "Ceramic", "vendor": "Hani Ceramic", "qty": 1, "rate": 16000,
                    "savings": 1000, "paid": 14000, "deal_ref": "HANI-1", "deal_status": "Open", "account": "5020"}
            eid = api.call(op, "POST", "/api/docs/direct_cost", deal)[1]["id"]
            e = api.call(op, "GET", f"/api/journal/{eid}")[1]
            self.assertEqual((e["doc"], e["form"]["rate"], e["meta"]["negotiation"]), ("direct_cost", 16000, "keep"))
            # correct it through the form: the real quote was 15,000; settled; the entry keeps its number
            st, r = api.call(op, "PUT", f"/api/docs/direct_cost/{eid}", {**deal, "rate": 15000, "deal_ref": "HANI-2",
                                                                          "deal_status": "Settled", "negotiation": "keep"})
            self.assertEqual((st, r["id"]), (200, eid))
            e = api.call(op, "GET", f"/api/journal/{eid}")[1]
            self.assertEqual((e["ref"], e["meta"]["deal_status"], len(e["history"])), ("HANI-2", "Settled", 1))
            self.assertEqual(api.call(op, "GET", "/api/supplier-discount?party=Hani%20Ceramic")[1]["owed"], 0)
            self.assertEqual(api.call(op, "GET", "/api/negotiations")[1]["total"]["kept"], 1000)
            # search: account name, account code, amount, business unit
            for word in ("Subcontract", "5020", "14,000", "P1", "Savings"):
                rows = api.call(op, "GET", "/api/journal?q=" + urllib.parse.quote(word))[1]["rows"]
                self.assertIn(eid, [x["id"] for x in rows], word)
            self.assertEqual(api.call(op, "GET", "/api/journal?q=nothing-like-this")[1]["rows"], [])
            # journal editor can fix the reference and deal too
            self.assertEqual(api.call(op, "PUT", f"/api/journal/{eid}", {"ref": "X-9", "deal_ref": "X-9"})[0], 200)
            self.assertEqual(api.call(op, "GET", f"/api/journal/{eid}")[1]["meta"]["deal_ref"], "X-9")
            self.assertEqual(api.call(op, "GET", "/api/reports/negotiations?format=json")[0], 200)
        finally:
            srv.shutdown()
            srv.server_close()
            db.close()
            tmp.cleanup()

    def test_old_books_get_the_account(self):
        db = Database(":memory:")
        db.x("UPDATE accounts SET role=NULL WHERE role='negotiation_income'")
        db.x("DELETE FROM accounts WHERE code='4085'")
        db._migrate()
        self.assertEqual(db.acct("negotiation_income"), "4085")


class Loans(unittest.TestCase):
    """Loans taken and given: schedules, interest, repayments, write-offs, balances from the books."""

    def setUp(self):
        from erp import loans
        self.L = loans
        self.db = Database(":memory:")
        a.record_journal(self.db, {"date": "2026-01-01", "lines": [{"account": "1020", "debit": 1000000}, {"account": "3010", "credit": 1000000}]})

    def test_schedules(self):
        L = self.L
        base = {"principal": 120000, "start_date": "2026-01-01", "term_months": 12}
        eq = L.preview_schedule({**base, "interest_type": "declining", "rate": 12, "schedule_type": "equal"})
        self.assertEqual(len(eq["rows"]), 12)
        self.assertEqual(eq["rows"][0]["interest"], 1200)                        # 1 % of 120,000 in month 1
        self.assertAlmostEqual(eq["totals"]["principal"], 120000, places=2)
        self.assertEqual(eq["rows"][-1]["balance"], 0)
        self.assertLess(abs(eq["rows"][0]["payment"] - eq["rows"][5]["payment"]), 0.02)   # equal installments
        flat = L.preview_schedule({**base, "interest_type": "flat", "rate": 10, "schedule_type": "equal", "frequency": 3})
        self.assertEqual([r["interest"] for r in flat["rows"]], [3000] * 4)     # 12,000 flat over 4 quarters
        bullet = L.preview_schedule({**base, "interest_type": "declining", "rate": 12, "schedule_type": "bullet"})
        self.assertEqual([r["principal"] for r in bullet["rows"][:-1]], [0] * 11)
        self.assertEqual(bullet["rows"][-1]["principal"], 120000)
        fixed = L.preview_schedule({**base, "interest_type": "fixed", "fixed_interest": 1000, "schedule_type": "equal_principal", "term_months": 3})
        self.assertAlmostEqual(fixed["totals"]["interest"], 1000, places=2)
        self.assertEqual(L.preview_schedule({**base, "schedule_type": "none"})["rows"], [])

    def test_borrowed_loan_lifecycle(self):
        L, db = self.L, self.db
        lid = L.save(db, {"direction": "borrowed", "party": "CIB", "principal": 120000, "rate": 12, "interest_type": "declining",
                          "term_months": 12, "start_date": "2026-01-01", "schedule_type": "equal"})
        self.assertEqual(a.party_balance(db, "2300", "CIB"), 120000)              # money received, loan owed
        d = L.detail(db, lid, today="2026-02-15")
        self.assertEqual(d["loan"]["overdue"], d["schedule"][0]["payment"])       # installment 1 due 1 Feb, unpaid
        s = d["suggest"]
        self.assertEqual((s["principal"], s["interest"]), (d["schedule"][0]["principal"], 1200))
        L.post(db, lid, "pay", {"date": "2026-02-01", "principal": s["principal"], "interest": s["interest"]})
        d = L.detail(db, lid, today="2026-02-15")
        self.assertEqual((d["schedule"][0]["state"], d["loan"]["overdue"]), ("paid", 0))
        self.assertEqual(d["loan"]["interest_paid"], 1200)
        # accrue at month end, then a payment clears the accrued interest first
        L.post(db, lid, "interest", {"date": "2026-02-28", "amount": 1000})
        self.assertEqual(L.detail(db, lid)["loan"]["accrued"], 1000)
        L.post(db, lid, "pay", {"date": "2026-03-01", "principal": 9556.47, "interest": 1105.38})
        self.assertEqual(L.detail(db, lid)["loan"]["accrued"], 0)
        self.assertAlmostEqual(next(x for x in a.account_totals(db) if x["code"] == "7030")["net"] / 100, 1200 + 1105.38, places=2)
        with self.assertRaises(a.ValidationError):                                # more principal than is owed
            L.post(db, lid, "pay", {"date": "2026-03-02", "principal": 1e6})
        with self.assertRaises(a.ValidationError):                                # cannot close with a balance
            L.set_status(db, lid, True)
        with self.assertRaises(a.ValidationError):                                # entries exist: void them first
            L.delete(db, lid)
        L.post(db, lid, "writeoff", {"date": "2026-04-01"})                       # the bank forgave the rest
        L.set_status(db, lid, True)
        self.assertEqual(L.get(db, lid)["status"], "closed")
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])

    def test_lent_loan_and_opening_balance(self):
        L, db = self.L, self.db
        lid = L.save(db, {"direction": "lent", "party": "Ahmed", "principal": 24000, "interest_type": "none", "term_months": 6,
                          "start_date": "2026-01-01", "schedule_type": "equal"})
        self.assertEqual(-a.party_balance(db, db.acct("loans_receivable"), "Ahmed"), 24000)
        L.post(db, lid, "pay", {"date": "2026-02-01", "principal": 4000})
        self.assertEqual(L.detail(db, lid, today="2026-02-02")["loan"]["outstanding"], 20000)
        # a loan already in the books: nothing is posted, earlier installments count as paid
        old = L.save(db, {"direction": "borrowed", "party": "NBE", "principal": 60000, "term_months": 12, "interest_type": "none",
                          "start_date": "2025-07-01", "schedule_type": "equal", "booking": "opening", "opening_balance": 30000,
                          "opening_date": "2026-01-01"})
        d = L.detail(db, old, today="2026-01-15")
        self.assertEqual((d["loan"]["outstanding"], d["loan"]["overdue"], d["entries"]), (30000, 0, []))
        self.assertEqual(sum(1 for r in d["schedule"] if r["state"] == "paid"), 6)
        # flexible loan with daily interest: 36,500 at 10 % → 10 a day
        flex = L.save(db, {"direction": "lent", "party": "Partner Co", "principal": 36500, "rate": 10, "interest_type": "declining",
                           "schedule_type": "none", "term_months": 0, "start_date": "2026-03-01"})
        self.assertEqual(L.suggest(db, L.get(db, flex), "2026-03-31")["interest"], 300)

    def test_api_and_backup(self):
        from erp import interchange
        L, db = self.L, self.db
        db.x("INSERT INTO business_units(code,name) VALUES ('P1','Villa')")
        lid = L.save(db, {"direction": "borrowed", "party": "QNB", "principal": 50000, "rate": 8, "interest_type": "flat",
                          "term_months": 10, "start_date": "2026-01-01", "bu_code": "P1"})
        L.post(db, lid, "pay", {"date": "2026-02-01", "principal": 5000, "interest": 333.33})
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        self.assertEqual(L.detail(db2, "LN-0001")["loan"]["outstanding"], 45000)
        self.assertEqual(L.detail(db2, "LN-0001")["loan"]["bu_code"], "P1")
        from erp import reports
        for name, p in (("loans", {}), ("loan_schedule", {"code": "LN-0001"})):
            tb = reports.build(db, name, p)
            self.assertTrue(reports.to_html(tb).startswith("<!doctype html>"))


class FundingPartnerAndExit(unittest.TestCase):
    """The funding partner's own money is out − in across the company's bank; a partner's statement shows his capital
    and drawings; 'if he leaves' adds his billed profit and takes off what he owes; loans can be paid personally."""

    def test_funding_statement_exit_and_personal_loan(self):
        from erp import settlement as se, statements as st, loans
        db = Database(":memory:")
        db.x("INSERT INTO business_units(code,name) VALUES ('P1','Villa')")
        db.x("INSERT INTO partners(name,default_share) VALUES ('Funder',0.5),('Other',0.5)")
        db.set_setting("funding_partner", "Funder")
        db.set_setting("reserve_pct", "0")
        bank = db.acct("bank")
        a.record_journal(db, {"date": "2026-01-02", "bu_code": "P1", "party": "Client", "lines": [{"account": bank, "debit": 1000}, {"account": "4010", "credit": 1000}]})
        a.record_journal(db, {"date": "2026-01-03", "bu_code": "P1", "lines": [{"account": "5010", "debit": 600}, {"account": bank, "credit": 600}]})
        a.record_journal(db, {"date": "2026-01-04", "kind": "Owner Drawing", "party": "Other", "lines": [{"account": db.acct("drawings"), "debit": 100}, {"account": bank, "credit": 100}]})
        a.record_journal(db, {"date": "2026-01-05", "party": "Other", "lines": [{"account": "1175", "debit": 800}, {"account": bank, "credit": 800}]})
        led = se.funding_ledger(db, date_to="2026-01-31")
        self.assertEqual((led["money_in"], led["money_out"], led["own_money"], led["unbooked"]), (1000, 1500, 500, 500))
        se.book_funding(db, "2026-01-31")
        self.assertEqual(se.funding_ledger(db, date_to="2026-01-31")["unbooked"], 0)
        self.assertEqual(st.statement(db, "Funder", date_to="2026-01-31")["closing"], -500)      # we owe him his capital
        ex = {p["partner"]: p for p in se.exit_view(db, "2026-01-31")["partners"]}
        self.assertEqual(ex["Funder"]["owed_now"], 500 + 200)                 # capital + half of the 400 profit
        self.assertEqual(ex["Other"]["owed_now"], 200 - 100 - 800)            # profit − drawings − the loan he owes
        db.set_setting("reserve_pct", "0.25")                                  # a quarter of billed profit held back
        f = {p["partner"]: p for p in se.exit_view(db, "2026-01-31")["partners"]}["Funder"]
        self.assertEqual((f["undrawn_profit"], f["reserve_share"], f["owed_now"]), (150, 50, 700))   # same total, split
        self.assertNotIn("owed_when_billed", f)                                # unbilled-work profit is a memo only
        db.set_setting("reserve_pct", "0")
        lid = loans.save(db, {"direction": "lent", "party": "Other", "principal": 300, "start_date": "2026-02-01", "term_months": 1,
                              "booking": "money", "paid_from": "partner:Funder"})
        e = loans._entries(db, loans.get(db, lid)["code"])[0]
        self.assertIn((db.acct("capital"), 0, 300, "Funder"), [(l["account"], l["debit"], l["credit"], l["party"]) for l in a.get_entry(db, e["id"])["lines"]])
        ex = {p["partner"]: p for p in se.exit_view(db, "2026-02-28")["partners"]}
        self.assertEqual((ex["Funder"]["owed_now"], ex["Other"]["owed_now"]), (1000, -1000))
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])


class LiveWorkInProgress(unittest.TestCase):
    """No snapshots: a cost-plus project's work in progress follows its costs, a measured one its draft invoices, and
    posting the invoice moves the work to receivables without counting it twice. Money through the funding partner's
    account is his, tagged automatically."""

    def test_cost_plus_drafts_and_funding_account(self):
        db = Database(":memory:")
        db.x("INSERT INTO business_units(code,name,basis,fee_pct,wip_mode) VALUES ('CP','O West','Cost-Plus %',0.25,'cost_plus'),"
             "('MS','Site A','Fixed Price',0,'drafts'),('NO','Plain','Cost-Plus %',0.25,'')")
        db.x("INSERT INTO partners(name,default_share) VALUES ('Funder',1)")
        db.set_setting("funding_partner", "Funder")
        fund = db.acct("funding_partner")
        self.assertEqual(a.money_account(db), fund)                                # payments default to his account
        unb = db.acct("unbilled_revenue")
        wip = lambda code: round(sum(l["debit"] - l["credit"] for l in db.q(  # noqa: E731
            "SELECT l.debit, l.credit FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id JOIN business_units b ON b.id=e.bu_id "
            "WHERE e.status='posted' AND l.account=? AND b.code=?", (unb, code))) / 100, 2)
        e1 = a.record_journal(db, {"date": "2026-09-23", "bu_code": "CP", "lines": [{"account": "5045" if db.one("SELECT 1 FROM accounts WHERE code='5045'") else "5010", "debit": 6200}, {"account": fund, "credit": 6200}]})
        self.assertEqual(wip("CP"), 7750)                                          # 6,200 × 1.25
        a.record_journal(db, {"date": "2026-09-24", "bu_code": "NO", "lines": [{"account": "5010", "debit": 1000}, {"account": fund, "credit": 1000}]})
        self.assertEqual(wip("NO"), 0)                                             # not switched on
        self.assertEqual(db.one("SELECT party FROM journal_lines WHERE entry_id=? AND account=?", (e1, fund))["party"], "Funder")
        a.void_entry(db, e1, "mistake")
        self.assertEqual(wip("CP"), 0)                                             # follows the costs both ways
        a.record_journal(db, {"date": "2026-09-25", "bu_code": "CP", "lines": [{"account": "5010", "debit": 6200}, {"account": fund, "credit": 6200}]})
        xid = a.save_extract(db, {"bu_id": 2, "number": "Invoice 01", "date": "2026-09-23", "customer": "C",
                                  "lines": [{"description": "Works", "qty": 1, "price": 453767.75}]})
        self.assertEqual(wip("MS"), 453767.75)                                     # the draft's executed work
        net0 = a.income_statement(db)["net_profit"]
        a.post_extract(db, xid)
        self.assertEqual((wip("MS"), a.income_statement(db)["net_profit"]), (0, net0))   # billed, not earned twice
        cp = a.save_extract(db, {"bu_id": 1, "number": "OW-1", "date": "2026-09-30", "customer": "O",
                                 "lines": [{"description": "Works", "qty": 1, "price": 7750}]})
        a.post_extract(db, cp)
        self.assertEqual(wip("CP"), 0)                                             # cost-plus invoiced in full
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries WHERE kind='Work in Progress' AND status='posted'")["n"], 0)
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])


class InvoiceAgainstUnbilledWork(unittest.TestCase):
    """Billing work already recognised as 'done, not yet invoiced' moves it to receivables without earning it twice;
    a posted invoice can be reopened (entry voided, back to draft) unless the tax authority has it."""

    def test_bill_from_unbilled_and_reopen(self):
        from erp import server
        db = Database(":memory:")
        db.x("INSERT INTO business_units(code,name) VALUES ('P1','Villa')")
        unb = db.acct("unbilled_revenue")
        a.record_journal(db, {"date": "2026-10-07", "bu_code": "P1", "lines": [{"account": unb, "debit": 1000}, {"account": "4010", "credit": 1000}]})
        xid = a.save_extract(db, {"bu_id": 1, "number": "INV-1", "date": "2026-09-23", "customer": "C",
                                  "lines": [{"description": "Works", "qty": 1, "price": 1200}]})
        net = lambda: a.income_statement(db)["net_profit"]  # noqa: E731
        eid = a.post_extract(db, xid)
        lines = {l["account"]: (l["debit"], l["credit"]) for l in a.get_entry(db, eid)["lines"]}
        self.assertEqual((lines[unb], lines["4010"], lines[db.acct("ar")]), ((0, 1000), (0, 200), (1200, 0)))
        self.assertEqual(net(), 1200)                       # 1,000 earned before + only 200 new
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        ctx = types.SimpleNamespace(db=db, body={}, user={"username": "t"}, can=lambda *k: True, log=lambda *k: None)
        server.extract_reopen(ctx, str(xid))
        self.assertEqual((db.one("SELECT status FROM extracts WHERE id=?", (xid,))["status"], net()), ("draft", 1000))
        a.post_extract(db, xid)
        db.x("UPDATE extracts SET eta_uuid='U1', eta_status='Valid' WHERE id=?", (xid,))
        with self.assertRaises(a.ValidationError):
            server.extract_reopen(ctx, str(xid))


class Custody(unittest.TestCase):
    """Cash custody: issue, expense reports (split by business unit, with VAT), return, reimburse, deduct, close."""

    def test_custody_lifecycle(self):
        from erp import custody, interchange
        db = Database(":memory:")
        a.record_journal(db, {"date": "2026-01-01", "lines": [{"account": "1010", "debit": 50000}, {"account": "3010", "credit": 50000}]})
        db.x("INSERT INTO business_units(code,name) VALUES ('P1','Villa'),('P2','Flat')")
        acc_c = db.acct("custody")
        cid = custody.save(db, {"holder": "Samir", "amount": 5000, "issued_on": "2026-02-01", "due_date": "2026-02-20", "bu_code": "P1"})
        self.assertEqual(-a.party_balance(db, acc_c, "Samir"), 5000)
        ids = custody.post(db, cid, "settle", {"date": "2026-02-10", "report_no": "R-1", "lines": [
            {"description": "Cement", "account": "5010", "amount": 2280, "vat_code": "VAT14", "supplier": "Hardware"},
            {"description": "Taxi", "account": "6075", "amount": 300, "bu_id": 2}]})
        self.assertEqual(len(ids), 2)                                                  # one entry per business unit
        e = a.get_entry(db, ids[0])
        self.assertIn(("5010", 2000.0), [(ln["account"], ln["debit"]) for ln in e["lines"]])   # net of 14 % VAT
        st = custody.status(db, custody.get(db, cid), today="2026-02-15")
        self.assertEqual((st["balance"], st["spent"], st["overdue"]), (2420, 2580, False))
        self.assertTrue(custody.status(db, custody.get(db, cid), today="2026-03-01")["overdue"])
        with self.assertRaises(a.ValidationError):
            custody.post(db, cid, "return", {"amount": 9999})
        custody.post(db, cid, "return", {"date": "2026-02-11", "amount": 2000})
        custody.post(db, cid, "deduct", {"date": "2026-02-12", "amount": 420})          # kept from salary
        self.assertEqual(-a.party_balance(db, db.acct("accrued_payroll"), "Samir"), 420)
        with self.assertRaises(a.ValidationError):
            custody.delete(db, cid)
        custody.set_status(db, cid, True)
        # a standing float: overspending makes the company owe the holder; the report can top the float back up
        f = custody.save(db, {"holder": "Office", "kind": "permanent", "float_amount": 1000, "issued_on": "2026-03-01"})
        custody.post(db, f, "settle", {"date": "2026-03-05", "lines": [{"description": "Stationery", "account": "6080", "amount": 1200}],
                                       "replenish": True})
        st = custody.status(db, custody.get(db, f))
        self.assertEqual((st["balance"], st["spent"]), (1000, 1200))
        with self.assertRaises(a.ValidationError):
            custody.set_status(db, f, True)
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        self.assertEqual(custody.status(db2, custody.get(db2, "CU-0002"))["balance"], 1000)
        self.assertEqual(len(custody.statement(db2, "CU-0001")["rows"]), 5)

    def test_overspent_custody_is_owed_to_the_holder(self):
        """8,000 given, 9,800 of receipts: the company owes the holder 1,800 (not the reverse); reimbursing settles it.
        Entries dated after the as-of date stay out of the balance but are listed under it."""
        from erp import custody, statements, reports
        db = Database(":memory:")
        a.record_journal(db, {"date": "2026-01-01", "lines": [{"account": "1010", "debit": 50000}, {"account": "3010", "credit": 50000}]})
        cid = custody.save(db, {"holder": "Mohamed", "amount": 8000, "issued_on": "2026-09-29"})
        custody.post(db, cid, "settle", {"date": "2026-09-29", "lines": [{"description": "Electrical", "account": "5010", "amount": 9800}]})
        row = lambda d: next(r for r in statements.aging(db, d)["rows"] if r["party"] == "Mohamed")  # noqa: E731
        self.assertEqual(row("2026-09-30")["balance"], -1800)                         # − = we owe them
        self.assertIn("we owe them", str(reports.party_statement_table(db, "Mohamed", date_to="2026-09-30")["rows"]))
        custody.post(db, cid, "issue", {"date": "2026-10-01", "amount": 1800})       # reimburse
        self.assertFalse([r for r in statements.aging(db, "2026-10-02")["rows"] if r["party"] == "Mohamed"])
        # the same custody booked with a mistyped later date: out of the balance as of today, but flagged
        later = statements.aging(db, "2026-09-28")["later"]
        self.assertEqual([(g["party"], g["entries"], g["first"]) for g in later], [("Mohamed", 3, "2026-09-29")])
        self.assertIn("NOT COUNTED", str(reports.aging_table(db, "2026-09-28")["rows"]))
        from erp import server
        notes = server._future_dated({"Journal Vouchers": [{"voucher": "JE-1", "date": "2999-01-01"}, {"voucher": "JE-1", "date": "2999-01-01"}]})
        self.assertTrue(notes and "JE-1 (2999-01-01)" in notes[0] and "1 dated after today" in notes[0])


class Rentals(unittest.TestCase):
    """Equipment rental: FIFO stock, frames, the scaffold calculator, PO → signed contract → signed handovers → invoices."""

    def setUp(self):
        from erp import rentals
        self.R = rentals
        self.db = db = Database(":memory:")
        a.record_journal(db, {"date": "2026-01-01", "lines": [{"account": "1020", "debit": 5000000}, {"account": "3010", "credit": 5000000}]})
        rentals.starter(db)
        self.ids = {i["code"]: i["id"] for i in rentals.list_items(db)["items"]}
        rentals.add_lot(db, self.ids["SC-FRM2"], {"qty": 150, "unit_cost": 900, "date": "2025-01-01", "source": "opening"})
        rentals.add_lot(db, self.ids["SC-FRM2"], {"qty": 100, "unit_cost": 980, "date": "2026-01-05", "supplier": "Steel Co", "paid": 0})
        for c, q, cost in (("SC-BRC3", 250, 340), ("SC-PLT60", 80, 150), ("SC-JACK", 80, 220), ("SC-CLF", 100, 110), ("SC-CLS", 30, 110),
                           ("SC-TUB6", 20, 750), ("SC-BRD36", 60, 700), ("EL-HOIST", 1, 450000)):
            rentals.add_lot(db, self.ids[c], {"qty": q, "unit_cost": cost, "date": "2026-01-05", "source": "opening"})

    def test_calculator_and_fifo(self):
        R, db = self.R, self.db
        ev = R.evaluate(db, R.list_models(db)[0]["spec"])
        q = {c["item"]: c["qty"] for c in ev["components"]}
        self.assertEqual((ev["area"], q["SC-FRM2"], q["SC-JACK"], q["SC-CLF"], q["SC-BRD36"]), (1152, 200, 50, 52, 24))  # the handover form
        taller = R.evaluate(db, R.list_models(db)[0]["spec"], {"height": 24})
        self.assertEqual({c["item"]: c["qty"] for c in taller["components"]}["SC-BRC3"], 360)   # 12 lifts, +25 % bracing above 20 m
        for bad in ("__import__('os')", "length.real", "[1][0]", "2 ** 99"):
            with self.assertRaises(a.ValidationError):
                R.safe_eval(bad, {"length": 1})
        self.assertEqual(R.safe_eval("ceil(10 / 3) if h > 5 else 0", {"h": 6}), 4)
        st = R.item_stock(db, self.ids["SC-FRM2"])
        self.assertEqual((st["in_yard"], st["value"]), (250, 233000))
        self.assertEqual(round(R.unit_cost_next(db, self.ids["SC-FRM2"], 160) / 100, 2), 905)     # 150 × 900 + 10 × 980
        self.assertEqual(a.party_balance(db, db.acct("ap"), "Steel Co"), 98000)                    # the purchase is booked
        R.scrap(db, {"item_id": self.ids["SC-FRM2"], "qty": 5, "reason": "bent"})                 # oldest lot first
        self.assertEqual(R.item_detail(db, self.ids["SC-FRM2"])["lots"][0]["gone"], 5)

    def _contract(self):
        R, db = self.R, self.db
        ev = R.evaluate(db, R.list_models(db)[0]["spec"])
        kit = R.list_items(db)["kits"][0]
        rid = R.save(db, {"client": "Sawag", "site": "Agora — building 3", "vat_code": "VAT14", "lines": [
            {"kind": "sqm", "qty": ev["area"], "rate": 9, "components": [{"item_id": c["item_id"], "per_unit": c["qty"] / ev["area"]} for c in ev["components"]]},
            {"kind": "kit", "kit_id": kit["id"], "qty": 20, "rate": 120},
            {"kind": "item", "item_id": self.ids["EL-HOIST"], "qty": 1, "rate": 25000},
            {"kind": "charge", "charge_type": "install", "qty": 1, "rate": 15000},
            {"kind": "charge", "charge_type": "dismantle", "qty": 1, "rate": 10000},
            {"kind": "charge", "charge_type": "transport", "qty": 2, "rate": 3000}]})
        return rid

    def _deliver(self, rid, date="2026-08-01"):
        t = self.R.handover_template(self.db, rid, "delivery")
        t.update(date=date, signed=1, decision="safe", our_name="Hossam", client_name="Karim",
                 checklist=[dict(c, ok="yes") for c in t["checklist"]])
        return self.R.save_handover(self.db, rid, t)

    def test_contract_flow(self):
        R, db = self.R, self.db
        rid = self._contract()
        hid = self._deliver(rid)
        with self.assertRaises(a.ValidationError):                       # no purchase order yet
            R.confirm_handover(db, hid)
        R.action(db, rid, "po", {"po_number": "PO-778", "po_date": "2026-08-01"})
        with self.assertRaises(a.ValidationError):                       # contract not signed yet
            R.confirm_handover(db, hid)
        R.action(db, rid, "sign", {"signed_us": "Hossam", "signed_client": "Karim", "signed_on": "2026-08-01"})
        R.confirm_handover(db, hid)
        d = R.detail(db, rid)
        self.assertEqual(d["rental"]["status"], "active")
        self.assertEqual(sorted(c["kind"] for c in d["charges"]), ["install", "transport"])      # billed on delivery
        frames = R.item_detail(db, self.ids["SC-FRM2"])["lots"]
        self.assertEqual([(l["with_clients"], l["in_yard"]) for l in frames], [(150, 0), (70, 30)])   # FIFO: oldest first
        al = R.allocation(db, rid, d["lines"][0]["id"])                    # m² rent shared by FIFO cost
        self.assertAlmostEqual(sum(x["monthly"] for x in al["rows"]), al["rent"], delta=0.05)
        self.assertEqual(al["basis"], "cost")
        yields = {round(x["yield"], 4) for x in al["rows"]}
        self.assertEqual(len(yields), 1)                                  # every piece earns the same % of its cost
        # August invoice: a full calendar month = 1 month on every line
        p = R.next_period(db, rid)
        self.assertEqual((p["from"], p["to"]), ("2026-08-01", "2026-08-31"))
        xid = R.create_invoice(db, rid, {"from": p["from"], "to": p["to"]})
        x = a.get_extract(db, xid)
        self.assertEqual([ln["qty"] for ln in x["lines"][:3]], [1152, 20, 1])
        self.assertEqual(x["total"], 1152 * 9 + 20 * 120 + 25000 + 15000 + 6000)
        with self.assertRaises(a.ValidationError):                       # the same period twice
            R.create_invoice(db, rid, {"from": "2026-08-15", "to": "2026-09-10"})
        a.post_extract(db, xid)
        e = a.get_entry(db, a.get_extract(db, xid)["entry_id"])
        self.assertEqual({ln["account"] for ln in e["lines"] if ln["credit"]}, {"4045", "4046", "2200"})
        # return the frame sets on 15 Sep: 2 braces lost, 1 frame damaged
        t = R.handover_template(db, rid, "return")
        kl = d["lines"][1]["id"]
        t["date"] = "2026-09-15"
        t["data"]["units"] = {k: (v if int(k) == kl else 0) for k, v in t["data"]["units"].items()}
        items = [x for x in t["data"]["items"] if x["line_id"] == kl]
        for x in items:
            if x["code"] == "SC-BRC3":
                x["qty"], x["lost"] = x["qty"] - 2, 2
            if x["code"] == "SC-FRM2":
                x["qty"], x["damaged"] = x["qty"] - 1, 1
        t["data"]["items"] = items
        t["signed"] = 1
        with self.assertRaises(a.ValidationError):                       # not signed by both parties
            R.confirm_handover(db, R.save_handover(db, rid, {**t, "client_name": ""}))
        R.confirm_handover(db, R.save_handover(db, rid, t))
        loss = next(x for x in a.account_totals(db) if x["code"] == "7046")["net"]
        self.assertEqual(loss, 90000 + 2 * 34000)                         # written off at FIFO cost (frame 900, braces 340)
        sep = R.invoice_lines(db, rid, "2026-09-01", "2026-09-30")
        self.assertEqual([ln["qty"] for ln in sep["lines"] if ln["category"] == "Rent"], [1152, 10, 1])   # sets out 15 of 30 days
        self.assertAlmostEqual(sum(ln["qty"] * ln["price"] for ln in sep["lines"] if ln["category"] == "Compensation"), 2 * 391 + 1127)
        # everything back → dismantling is charged, then the contract can close once invoiced
        t = R.handover_template(db, rid, "return")
        t.update(date="2026-10-10", signed=1)
        R.confirm_handover(db, R.save_handover(db, rid, t))
        d = R.detail(db, rid)
        self.assertEqual(d["rental"]["status"], "returned")
        self.assertIn("dismantle", [c["kind"] for c in d["charges"]])
        with self.assertRaises(a.ValidationError):
            R.action(db, rid, "close", {})
        R.create_invoice(db, rid, {"from": "2026-09-01", "to": "2026-10-10"})
        R.action(db, rid, "close", {})
        self.assertEqual(R.get(db, rid)["status"], "closed")
        self.assertEqual(R.item_stock(db, self.ids["SC-FRM2"])["on_rent"], 0)
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])

    def test_depreciation(self):
        R, db = self.R, self.db
        fid = self.ids["SC-FRM2"]
        # lots: 150 × 900 (Jan 2025) and 100 × 980 (Jan 2026); 10 years to 10 %
        pl = R.plan_depreciation(db, "2026-06")
        lots = {r["lot"]: (r["from"], r["months"], r["amount"]) for r in pl["rows"] if r["item_id"] == fid}
        self.assertEqual(lots["LOT-00001"], ("2025-01", 18, int(round(150 * 90000 * 0.9 / 120 * 18))))
        self.assertEqual(lots["LOT-00002"], ("2026-01", 6, int(round(100 * 98000 * 0.9 / 120 * 6))))
        eid = R.run_depreciation(db, "2026-06")
        with self.assertRaises(a.ValidationError):
            R.run_depreciation(db, "2026-06")                 # nothing left for those months
        st = R.item_stock(db, fid)
        self.assertEqual(st["accum"], (lots["LOT-00001"][2] + lots["LOT-00002"][2]) / 100)
        self.assertEqual(R.plan_depreciation(db, "2026-07")["rows"][0]["months"], 1)   # next month only
        # scrapping 15 pieces of the oldest lot clears 10 % of that lot's depreciation
        R.scrap(db, {"item_id": fid, "qty": 15, "date": "2026-07-01"})
        e = a.get_entry(db, db.q("SELECT id FROM journal_entries WHERE kind='Rental Equipment Loss'")[-1]["id"])
        share = int(round(lots["LOT-00001"][2] * 15 / 150)) / 100
        self.assertIn((db.acct("rental_accum"), share, 0.0), [(ln["account"], ln["debit"], ln["credit"]) for ln in e["lines"]])
        self.assertIn(("1415", 0.0, 13500.0), [(ln["account"], ln["debit"], ln["credit"]) for ln in e["lines"]])
        a.void_entry(db, eid, "redo")                         # a voided run is simply booked again
        self.assertEqual(R.plan_depreciation(db, "2026-06")["rows"][0]["from"], "2025-01")
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])

    def test_minimum_period_and_deposit(self):
        R, db = self.R, self.db
        rid = R.save(db, {"client": "Client X", "min_days": 30, "deposit": 5000, "deposit_first": True,
                          "lines": [{"kind": "item", "item_id": self.ids["SC-FRM2"], "qty": 50, "rate": 60}]})
        R.action(db, rid, "po", {"po_number": "P"})
        R.action(db, rid, "sign", {"signed_us": "a", "signed_client": "b"})
        hid = self._deliver(rid)
        with self.assertRaises(a.ValidationError):            # deposit first
            R.confirm_handover(db, hid)
        R.action(db, rid, "deposit", {"type": "receive", "amount": 5000, "date": "2026-07-30"})
        R.confirm_handover(db, hid)
        t = R.handover_template(db, rid, "return")
        t.update(date="2026-08-10", signed=1, our_name="a", client_name="b")
        t["data"]["units"] = {k: 20 for k in t["data"]["units"]}
        t["data"]["items"][0]["qty"] = 20
        R.confirm_handover(db, R.save_handover(db, rid, t))
        mins = [c for c in R.detail(db, rid)["charges"] if c["kind"] == "minimum"]
        self.assertEqual(len(mins), 1)
        self.assertAlmostEqual(mins[0]["total"], 20 * 20 / 30 * 60, delta=0.05)   # 20 frames kept 10 of 30 days
        xid = R.create_invoice(db, rid, {"from": "2026-08-01", "to": "2026-08-31"})
        a.post_extract(db, xid)
        owed = -a.party_balance(db, db.acct("ar"), "Client X")
        with self.assertRaises(a.ValidationError):            # not more than they owe
            R.action(db, rid, "deposit", {"type": "apply", "amount": owed + 1})
        R.action(db, rid, "deposit", {"type": "apply", "amount": owed})
        self.assertEqual(a.party_balance(db, db.acct("ar"), "Client X"), 0)
        t = R.handover_template(db, rid, "return")
        t.update(date="2026-09-01", signed=1, our_name="a", client_name="b")
        R.confirm_handover(db, R.save_handover(db, rid, t))
        R.create_invoice(db, rid, {"from": "2026-09-01", "to": "2026-09-01"})
        with self.assertRaises(a.ValidationError):            # deposit still held
            R.action(db, rid, "close", {})
        left = R.detail(db, rid)["rental"]["deposit_held"]
        R.action(db, rid, "deposit", {"type": "refund", "amount": left})
        R.action(db, rid, "close", {})
        self.assertEqual(R.get(db, rid)["status"], "closed")
        self.assertEqual([c["check"] for c in a.audit(db) if not c["ok"]], [])

    def test_undo_short_stock_and_backup(self):
        from erp import interchange, rental_docs
        R, db = self.R, self.db
        rid = self._contract()
        R.action(db, rid, "po", {"po_number": "PO-1"})
        R.action(db, rid, "sign", {"signed_us": "A", "signed_client": "B"})
        t = R.handover_template(db, rid, "delivery")
        t.update(signed=1, decision="safe", date="2026-09-01", checklist=[dict(c, ok="yes") for c in t["checklist"]])
        t["data"]["items"][0]["qty"] = 999
        with self.assertRaises(a.ValidationError):                       # more than the yard holds
            R.confirm_handover(db, R.save_handover(db, rid, t))
        t["data"]["items"][0]["qty"] = 200
        t["checklist"][0]["ok"] = "no"
        with self.assertRaises(a.ValidationError):                       # a failed safety check needs 'not safe'
            R.confirm_handover(db, R.save_handover(db, rid, t))
        t["checklist"][0]["ok"] = "yes"
        hid = R.save_handover(db, rid, t)
        R.confirm_handover(db, hid)
        R.undo_handover(db, hid)
        self.assertEqual(R.item_stock(db, self.ids["SC-FRM2"])["on_rent"], 0)
        self.assertEqual(R.detail(db, rid)["charges"], [])
        R.confirm_handover(db, hid)
        for html in (rental_docs.offer_html(db, rid, "ar"), rental_docs.contract_html(db, rid, "en"), rental_docs.handover_html(db, hid, "ar")):
            self.assertTrue(html.startswith("<!doctype html>"))
        self.assertIn("الأرضية صلبة", rental_docs.handover_html(db, hid, "ar"))
        R.create_invoice(db, rid, {"from": "2026-09-01", "to": "2026-09-30"})
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
        self.assertEqual(R.detail(db2, "RC-0001")["rental"]["status"], "active")
        self.assertEqual(R.item_stock(db2, R.get_item(db2, "SC-FRM2")["id"])["on_rent"], 220)
        self.assertEqual(len(R.detail(db2, "RC-0001")["invoices"]), 1)


class Payroll(unittest.TestCase):
    """Payroll v2: Egyptian insurance and tax, posting, Excel round trip, crew sheets and business-type templates."""

    def _db(self):
        from erp import payroll as P
        db = Database(":memory:")
        db.x("INSERT INTO business_units(code,name,basis,status) VALUES ('P1','Site','Fixed Price','In Progress')")
        return db, P

    def test_tax_bands_and_high_income_rule(self):
        db, P = self._db()
        s = P.settings(db)
        self.assertEqual(P.annual_tax(40000, s), 0)
        self.assertAlmostEqual(P.annual_tax(100000, s), 1500 + 2250 + 6000)
        self.assertAlmostEqual(P.annual_tax(650000, s), 0.10 * 55000 + 0.15 * 15000 + 0.20 * 130000 + 0.225 * 200000 + 0.25 * 250000)
        self.assertAlmostEqual(P.annual_tax(1500000, s), 0.25 * 1200000 + 0.275 * 300000)
        P.save_settings(db, {"high_income": []})            # the rule can be switched off
        self.assertAlmostEqual(P.annual_tax(650000, P.settings(db)), 0 + 1500 + 2250 + 26000 + 45000 + 62500)

    def test_payslip_and_posting(self):
        db, P = self._db()
        e = P.save_employee(db, {"name": "Eng", "basic": 15000, "al_housing": 2000, "cost_type": "Direct Labour (Cost of Revenue)", "bu_code": "P1"})
        w = P.save_employee(db, {"name": "Worker", "emp_type": "daily", "daily_rate": 400, "insured": "No", "taxed": "No"})
        c = P.compute(P.get_employee(db, e), {"absent_days": 3, "ot_day": 4, "components": {"bonus": 500, "penalty": 100, "advance": 300}}, P.settings(db))
        self.assertAlmostEqual(c["gross"], 15000 * 27 / 30 + 2000 * 27 / 30 + round(15000 / 30 / 8 * 4 * 1.35, 2) + 500, places=2)
        self.assertEqual(c["si_emp"], round(16700 * 0.11, 2))      # 17,000 insurable, capped at the 2026 ceiling
        self.assertAlmostEqual(c["net"], c["gross"] - 100 - 300 - c["si_emp"] - c["tax"] - c["martyrs"], places=2)
        rid = P.create_run(db, {"period": "2026-09"})
        P.save_lines(db, rid, [{"employee_id": e, "absent_days": 3, "ot_day": 4, "components": {"bonus": 500, "penalty": 100, "advance": 300}},
                               {"employee_id": w, "units": 20, "paid": 8000, "method": "Cash"}])
        P.post_run(db, rid)
        tb = {r["account"]: r["n"] for r in db.q("SELECT account, SUM(debit-credit) n FROM journal_lines GROUP BY account")}
        self.assertEqual(sum(tb.values()), 0)
        self.assertEqual(tb["1010"], -800000)                         # the daily worker paid in cash
        self.assertEqual(db.one("SELECT bu_id FROM journal_entries WHERE ref='PR:2026-09:EMP-001'")["bu_id"], 1)
        P.unpost_run(db, rid)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries WHERE status='posted'")["n"], 0)
        self.assertEqual(reports.build(db, "payroll_register", {"run": rid})["rows"][-1]["style"], "grand")

    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_workbook_payroll_round_trip(self):
        from erp import payroll as P, workbook
        db = Database(":memory:")
        data = interchange.export_data(db)
        data["Employees"] = [{"code": "E1", "name": "A", "basic": 9000, "al_transport": 600, "insured": "Yes", "taxed": "Yes"},
                             {"code": "E2", "name": "B", "emp_type": "piece", "piece_rate": 17, "insured": "No", "taxed": "No"}]
        data["Payroll Run"] = [{"date": "2026-10-31", "employee": "E1", "bonus": 1000, "paid": 2000, "method": "Cash"},
                               {"date": "2026-10-31", "employee": "E2", "units": 1000}]
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "pay.xlsx"
            p.write_bytes(workbook.build_workbook(db, template=True, data=data))
            rows = interchange.read_xlsx(p)
        self.assertEqual(rows["Payroll Run"][0]["paid"], 2000)       # helper columns must not shadow the inputs
        db2 = Database(":memory:")
        interchange.import_data(db2, rows, mode="replace")
        r = P.get_run(db2, "2026-10")
        self.assertEqual(r["status"], "posted")
        self.assertEqual(r["totals"]["gross"], 9000 + 600 + 1000 + 17000)
        # the export carries employees and the payroll rules back out (and re-importing them doubles nothing)
        again = interchange.export_data(db2)
        self.assertEqual(len(again["Employees"]), 2)
        interchange.import_data(db2, {"Employees": again["Employees"], "Payroll Settings": again["Payroll Settings"]})
        self.assertEqual(db2.one("SELECT COUNT(*) n FROM employees")["n"], 2)

    def test_crew_sheet_stays_out_of_the_books(self):
        db, P = self._db()
        cid = P.save_crew(db, {"date": "2026-10-05", "crew": "Crew", "lines": [{"item": "Frames", "unit": "m²", "count": 260, "dim1": 2.5, "dim2": 2, "rate": 17}]})
        self.assertEqual(P.list_crew(db)[0]["total"], 22100)
        self.assertIsNone(db.one("SELECT 1 FROM journal_entries"))
        P.pay_crew(db, cid, {"amount": 22000, "method": "Cash", "bu_id": 1})
        self.assertEqual(db.one("SELECT SUM(credit) n FROM journal_lines")["n"], 2200000)

    def test_business_type_templates(self):
        from erp import industries
        for k in industries.ORDER:
            db = Database(":memory:")
            r = industries.apply(db, k)
            self.assertEqual(db.setting("business_type"), k)
            for code in r["accounts_added"]:
                self.assertTrue(db.one("SELECT 1 FROM accounts WHERE code=?", (code,)))
        self.assertIn("1225", industries.apply(Database(":memory:"), "development")["accounts_added"])


if __name__ == "__main__":
    unittest.main()


class Companies(unittest.TestCase):
    """One workspace, several companies: separate books and members, invite codes, choosing a company."""

    def test_companies(self):
        import os
        from erp.cloud import Supabase
        from erp.server import ERPServer
        sys.path.insert(0, str(Path(__file__).parent))
        import mock_supabase
        sb = mock_supabase.start()
        url = f"http://127.0.0.1:{sb.server_address[1]}"
        tmp = tempfile.TemporaryDirectory()
        nodes = []
        try:
            for name in ("a", "b"):
                db = Database(Path(tmp.name) / f"{name}.sqlite3")
                srv = ERPServer(("127.0.0.1", 0), db, Path(tmp.name) / f"bk-{name}")
                threading.Thread(target=srv.serve_forever, daemon=True).start()
                nodes.append((srv, db, f"http://127.0.0.1:{srv.server_address[1]}"))
            api = lambda i, op, m, p, b=None: Api.call(type("C", (), {"base": nodes[i][2]})(), op, m, p, b)  # noqa: E731
            a = Api.client(self)
            api(0, a, "POST", "/api/setup", {"username": "admin", "password": "secret123"})
            api(0, a, "PUT", "/api/cloud/config", {"url": url, "key": "anon"})
            self.assertEqual(api(0, a, "POST", "/api/cloud/signin", {"email": "owner@x.com", "password": "pw123456", "create": True})[0], 200)
            st, info = api(0, a, "POST", "/api/cloud/claim", {"name": "Alpha"})
            self.assertEqual((st, info.get("company_name")), (200, "Alpha"), info)
            self.assertEqual(api(0, a, "POST", "/api/bus", {"code": "A1", "name": "Alpha unit"})[0], 200)
            st, beta = api(0, a, "POST", "/api/cloud/company", {"action": "create", "name": "Beta"})
            self.assertEqual(st, 200, beta)
            alpha, beta = sb.state.company_named("Alpha"), beta["id"]
            self.assertEqual(api(0, a, "GET", "/api/cloud/info")[1]["company"], alpha)   # this device stays on Alpha
            # a newcomer signs up and asks to join Alpha with its invite code: pending until an admin approves
            cli = Supabase(url, "anon")
            tok = cli.sign_up("new@y.com", "pw654321")["access_token"]
            code = api(0, a, "GET", "/api/cloud/info")[1]["invite_code"]
            self.assertTrue(cli.rpc("join_company", {"p_code": code}, tok)["pending"])
            with self.assertRaises(RuntimeError):
                cli.rpc("join_company", {"p_code": "WRONG123"}, tok)
            self.assertEqual(cli.select("ops", f"company_id=eq.{alpha}&seq=gt.0", tok), [])     # nothing yet
            pending = [m for m in api(0, a, "GET", "/api/cloud/members")[1] if m["pending"]]
            self.assertEqual([m["email"] for m in pending], ["new@y.com"])
            self.assertEqual(api(0, a, "PUT", "/api/cloud/members", {"email": "new@y.com", "active": True})[0], 200)
            self.assertEqual(len(cli.select("ops", f"company_id=eq.{alpha}&seq=gt.0", tok)), 1)
            self.assertEqual(cli.select("ops", f"company_id=eq.{beta}&seq=gt.0", tok), [])      # Beta stays closed
            with self.assertRaises(PermissionError):
                cli.rpc("append_op", {"p_company": beta, "p_seq": 1, "p_version": "1", "p_app_user": "", "p_device": "x",
                                      "p_method": "POST", "p_path": "/api/bus", "p_body": "{}", "p_content_type": "",
                                      "p_clock": "2026-01-01T00:00:00"}, tok)
            # a fresh offline install: the owner belongs to two companies, so it asks which one to open
            os.environ.update(BEYOND_SUPABASE_URL=url, BEYOND_SUPABASE_ANON_KEY="anon")
            b = Api.client(self)
            st, r = api(1, b, "POST", "/api/cloud/first-join", {"username": "owner@x.com", "password": "pw123456"})
            self.assertEqual(sorted(c["name"] for c in r["choose"]), ["Alpha", "Beta"])
            st, r = api(1, b, "POST", "/api/cloud/first-join", {"username": "owner@x.com", "password": "pw123456",
                                                                "company": beta})
            self.assertEqual(st, 200, r)
            self.assertEqual(nodes[1][0].cloud.company, beta)
            self.assertEqual(nodes[1][1].setting("company_name"), "Beta")              # new company: fresh books
            self.assertIsNone(nodes[1][1].one("SELECT 1 FROM business_units WHERE code='A1'"))
            self.assertEqual(api(1, b, "POST", "/api/bus", {"code": "B1", "name": "Beta unit"})[0], 200)
            self.assertEqual([o["path"] for o in sb.state.ops[beta]], ["/api/bus"])
            self.assertEqual(len(sb.state.ops[alpha]), 1)                                     # Alpha untouched
            nodes[0][0].cloud.pull()
            self.assertIsNone(nodes[0][1].one("SELECT 1 FROM business_units WHERE code='B1'"))
            # usernames are shared by every company: Beta's administrator cannot create or grab Alpha's "ali"
            self.assertEqual(api(0, a, "POST", "/api/cloud/signin", {"email": "owner@x.com", "password": "pw123456"})[0], 200)
            self.assertEqual(api(0, a, "PUT", "/api/cloud/account", {"username": "ali", "password": "alipw123"})[0], 200)  # Alpha's ali
            beta_admin = Supabase(url, "anon")
            btok = beta_admin.sign_in("owner@x.com", "pw123456")["access_token"]
            fn = lambda body: beta_admin._call("POST", "/functions/v1/manage-user", {"company": beta, **body}, btok)  # noqa: E731
            for body in ({"username": "ali"}, {"username": "ali", "password": "other123"},
                         {"username": "ali", "email": "someone@else.com", "password": "other123"}):
                with self.assertRaises(RuntimeError) as cm:
                    fn(body)
                self.assertIn("already taken", str(cm.exception))
                self.assertIn("ali.beta", str(cm.exception))                              # a free name is suggested
            self.assertIsNone(sb.state.members.get((beta, "ali@users.beyond-erp.invalid")))
            self.assertEqual(sb.state.accounts["ali@users.beyond-erp.invalid"], "alipw123")  # Alpha's ali untouched
            self.assertEqual(fn({"username": "ali.beta", "password": "betapw1"})["username"], "ali.beta")
            # adding the very same person on purpose (by their email, no password) is allowed; they keep their password
            self.assertTrue(fn({"username": "ali", "email": "ali@users.beyond-erp.invalid"})["shared"])
            # a username can be the person's email address: the account is that address, and they sign in with it
            self.assertEqual(fn({"username": "Boss@Beta.com", "password": "bosspw12"})["username"], "boss@beta.com")
            self.assertEqual(sb.state.accounts["boss@beta.com"], "bosspw12")
            self.assertTrue(Supabase(url, "anon").sign_in("boss@beta.com", "bosspw12")["access_token"])
            for bad in ({"username": "boss@beta.com", "email": "other@beta.com", "password": "x123456"},
                        {"username": "two words", "password": "x123456"}):
                with self.assertRaises(RuntimeError):
                    fn(bad)
            # switching company on a connected device is refused (disconnect first)
            self.assertEqual(api(0, a, "POST", "/api/cloud/company", {"action": "choose", "company": beta})[0], 400)
        finally:
            os.environ.pop("BEYOND_SUPABASE_URL", None)
            os.environ.pop("BEYOND_SUPABASE_ANON_KEY", None)
            for srv, db, _ in nodes:
                srv.cloud.stop()
                srv.shutdown()
                srv.server_close()
                db.close()
            sb.shutdown()
            tmp.cleanup()


class Documents(unittest.TestCase):
    def setUp(self):
        self.db = seeded()

    def test_boq_invoice_offer_files(self):
        import io as _io
        import zipfile as _zip
        from erp import documents
        bid = a.save_boq(self.db, {"title": "Q", "client": "Cl", "lines": [
            {"category": "Demolition", "item_no": "1.1", "description": "Break wall", "unit": "m2", "qty": 10, "price": 100},
            {"category": "Painting", "item_no": "2.1", "description": "Paint", "unit": "m2", "qty": 0, "price": 50}]})
        data, name = documents.boq_file(self.db, bid, "xlsx")
        z = _zip.ZipFile(_io.BytesIO(data))
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        self.assertIn("rightToLeft", sheet)
        self.assertIn("SUM(", sheet)
        self.assertIn("xl/media/logo.png", z.namelist())
        html, _ = documents.boq_file(self.db, bid, "pdf")
        self.assertIn("إجمالي الأعمال الداخلية", html.decode())
        x = a.save_extract(self.db, {"bu_code": "PRJ-1", "boq_id": bid, "number": "Inv 2", "date": "2026-04-01",
                                     "lines": [{"category": "Demolition", "item_no": "1.1", "description": "Break wall", "qty": 4, "price": 100}]})
        data, name = documents.invoice_file(self.db, x, "xlsx")
        self.assertTrue(name.endswith("Inv. Inv 2.xlsx"))
        self.assertIn("E", _zip.ZipFile(_io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode())
        t = documents.list_offer_types(self.db)[0]
        secs = [dict(x, items=[{"desc": "Works", "amount": 60000}], discount=10000) if x["kind"] == "pricing" else x for x in t["sections"]]
        oid = documents.save_offer(self.db, dict(t, sections=secs, type_id=t["id"], client="Mr. X", project="Flat"))
        o = documents.get_offer(self.db, oid)
        self.assertEqual(o["net"], 50000)
        docx = _zip.ZipFile(_io.BytesIO(documents.offer_docx(self.db, o)))
        body = docx.read("word/document.xml").decode()
        self.assertIn("Fifty Thousand Egyptian Pounds", body)
        self.assertIn("word/media/watermark.png", docx.namelist())   # offers get the watermark by default
        self.assertEqual(documents.words_ar(60000), "ستون ألف")

    def test_banks_on_offers(self):
        from erp import documents
        self.db.set_setting("bank_name", "Old Bank")                 # books from before the bank list
        self.db.set_setting("bank_account", "111")
        self.assertEqual([(b["id"], b["acct"]) for b in documents.banks(self.db)], [("main", "111")])
        self.db.set_setting("banks", json.dumps([{"id": "cib", "label": "CIB", "bank": "CIB", "acct": "222"},
                                                 {"id": "nbe", "label": "NBE", "bank": "NBE", "acct": "333"}]))
        self.db.set_setting("bank_default", "nbe")
        self.assertEqual(documents.company(self.db)["bank"]["acct"], "333")
        t = documents.list_offer_types(self.db)[0]

        def bank_text(pick, details=None):
            secs = [dict(x) for x in t["sections"]]
            for x in secs:
                if x["kind"] == "bank":
                    x.update(bank_id=pick, **({"bank_details": details} if details else {}))
            oid = documents.save_offer(self.db, dict(t, type_id=t["id"], client="Z", sections=secs))
            blocks, _ = documents.offer_blocks(self.db, documents.get_offer(self.db, oid))
            return json.dumps(blocks, ensure_ascii=False)
        self.assertIn("333", bank_text(""))                               # default
        self.assertIn("222", bank_text("cib"))                            # one of the saved banks
        out = bank_text("custom", {"bank": "QNB", "acct": "999", "evil": "x"})
        self.assertIn("999", out)
        self.assertNotIn("333", out)

    def test_full_restore_brings_back_people_approvals_and_history(self):
        """A full JSON restore onto a new install: the other people (without passwords), their chats, approval
        requests, AI conversations and the activity log all come back; API keys stay out."""
        from erp import approvals, chat
        src = seeded()
        auth.create_user(src, {"username": "boss", "password": "Passw0rd!long1"}, first=True)
        auth.create_user(src, {"username": "sara", "password": "Passw0rd!long1", "full_name": "Sara", "perms": {"journal": "add"}})
        boss, sara = (src.one("SELECT * FROM users WHERE username=?", (n,)) for n in ("boss", "sara"))
        dm = chat.direct(src, boss, sara["id"])
        chat.send(src, sara, dm["id"] if isinstance(dm, dict) else dm, {"body": "hello"})
        approvals.submit(src, sara, "journal", {"date": "2026-03-01", "description": "Fuel", "lines": []}, module="journal")
        src.x("INSERT INTO ai_threads(user_id,title,messages) VALUES (?,?,?)", (sara["id"], "VAT", "[]"))
        auth.create_api_key(src, boss["id"], "BI", True, "boss")
        src.log("boss", "test.action", "kept")
        dst = Database(":memory:")
        auth.create_user(dst, {"username": "boss", "password": "another-long-pw"}, first=True)
        interchange.import_data(dst, interchange.read_json(interchange.to_json(src).encode()), mode="replace", user="boss")
        sara2 = dst.one("SELECT * FROM users WHERE username='sara'")
        self.assertEqual((sara2["full_name"], sara2["pw_hash"]), ("Sara", "!"))
        with self.assertRaises(Exception):
            auth.login(dst, "sara", "")
        self.assertEqual(dst.one("SELECT body FROM chat_messages")["body"], "hello")
        self.assertEqual(dst.one("SELECT created_by, status FROM approvals"), {"created_by": "sara", "status": "pending"})
        self.assertEqual(dst.one("SELECT title FROM ai_threads WHERE user_id=?", (sara2["id"],))["title"], "VAT")
        self.assertTrue(dst.one("SELECT 1 FROM audit_log WHERE action='test.action'"))
        self.assertIsNone(dst.one("SELECT 1 FROM api_keys"))
        self.assertEqual(dst.one("SELECT pw_hash FROM users WHERE username='boss'")["pw_hash"][:1] != "!", True)   # the restorer keeps their sign-in

    def test_backups_keep_connected_tools_keys_and_restore_from_sqlite(self):
        """A super administrator's backup keeps API keys (the same key works after a restore), the ETA credentials and
        the AI key; without the flag they stay out. A database backup file restores through the normal import."""
        import tempfile as _tf
        from erp import cloud
        src = seeded()
        auth.create_user(src, {"username": "boss", "password": "Passw0rd!long1"}, first=True)
        boss = src.one("SELECT * FROM users WHERE username='boss'")
        key = auth.create_api_key(src, boss["id"], "Power BI", True, "boss")
        key = key[1]
        src.set_setting("eta_client_secret", "eta-secret-1")
        cloud.local_set(src, "ai_key", "sk-test-0123456789abcdefghij")
        plain = interchange.to_json(src)
        self.assertNotIn("eta-secret-1", plain)
        self.assertNotIn("sk-test-0123456789", plain)
        dst = Database(":memory:")
        auth.create_user(dst, {"username": "boss", "password": "another-long-pw"}, first=True)
        interchange.import_data(dst, interchange.read_json(interchange.to_json(src, secrets=True).encode()), mode="replace", user="boss")
        self.assertEqual(auth.user_for_api_key(dst, key)[0]["username"], "boss")
        self.assertEqual(dst.setting("eta_client_secret"), "eta-secret-1")
        self.assertEqual(cloud.local_get(dst, "ai_key"), "sk-test-0123456789abcdefghij")
        # the database file itself
        with _tf.TemporaryDirectory() as d:
            p = Path(d) / "b.sqlite3"
            src.backup_to(p)
            raw = p.read_bytes()
        dst2 = Database(":memory:")
        auth.create_user(dst2, {"username": "boss", "password": "another-long-pw"}, first=True)
        interchange.import_data(dst2, interchange.read_sqlite(raw), mode="replace", user="boss")
        tb = lambda db: {r["code"]: (r["debit"], r["credit"]) for r in a.trial_balance(db)["rows"]}   # noqa: E731
        self.assertEqual(tb(src), tb(dst2))
        self.assertEqual(auth.user_for_api_key(dst2, key)[0]["username"], "boss")
        with self.assertRaises(a.ValidationError):
            interchange.read_sqlite(b"not a database at all")

    def test_lock_date_and_json_extras(self):
        self.db.set_setting("lock_date", "2026-03-31")
        with self.assertRaises(a.ValidationError):
            a.record_expense(self.db, {"description": "late", "amount": 5, "date": "2026-03-15"})
        a.record_expense(self.db, {"description": "ok", "amount": 5, "date": "2026-04-01"})
        from erp import documents
        t = documents.list_offer_types(self.db)[0]
        documents.save_offer(self.db, dict(t, type_id=t["id"], client="Y"))
        dump = interchange.to_json(self.db).encode()
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(dump), mode="replace")
        self.assertEqual(db2.one("SELECT COUNT(*) n FROM offers")["n"], 1)
        self.assertNotIn("eta_client_secret", dump.decode())


class Taxes(unittest.TestCase):
    def setUp(self):
        self.db = seeded()
        a.save_party(self.db, {"name": "Client Co", "type": "customer", "tax_id": "123456789", "governate": "Metro",
                               "region": "Uptown", "street": "Main Street", "building": "5"})

    def test_purchase_vat_and_withholding(self):
        a.record_project_cost(self.db, {"bu_code": "PRJ-1", "description": "tiles", "vendor": "TileCo", "qty": 1,
                                        "rate": 10000, "vat_code": "VAT14", "wht_code": "WHT1-S", "paid": 5000})
        self.assertEqual(a.party_balance(self.db, "2010", "TileCo"), 10000 + 1400 - 100 - 5000)
        t = a.tax_summary(self.db)
        self.assertEqual(t["input_vat"], 1400.0)
        self.assertEqual(t["wht_payable_balance"], 100.0)
        self.assertTrue(all(c["ok"] for c in a.audit(self.db)))

    def test_invoice_vat_and_client_withholding(self):
        x = a.save_extract(self.db, {"bu_code": "PRJ-1", "number": "INV-T", "customer": "Client Co",
                                     "vat_code": "VAT14", "wht_code": "WHT1-C",
                                     "lines": [{"description": "Work", "qty": 2, "price": 50000}]})
        a.post_extract(self.db, x)
        e = a.get_entry(self.db, self.db.one("SELECT entry_id FROM extracts WHERE id=?", (x,))["entry_id"])
        amounts = {ln["account"]: (ln["debit"], ln["credit"]) for ln in e["lines"]}
        self.assertEqual(amounts["1100"], (113000.0, 0.0))
        self.assertEqual(amounts["1330"], (1000.0, 0.0))
        self.assertEqual(amounts["2200"], (0.0, 14000.0))

    def test_vat_inclusive_cash_receipt(self):
        a.record_collection(self.db, {"bu_code": "PRJ-2", "amount": 11400, "vat_code": "VAT14"})
        self.assertEqual(a.tax_summary(self.db)["output_vat"], 1400.0)

    def test_party_validation(self):
        with self.assertRaises(a.ValidationError):
            a.save_party(self.db, {"name": "Bad", "tax_id": "12"})


class EInvoice(unittest.TestCase):
    def test_document_and_serialization(self):
        from erp import einvoice
        db = seeded()
        a.save_party(db, {"name": "Client Co", "tax_id": "123456789", "governate": "Metro", "region": "Nasr",
                          "street": "X", "building": "1"})
        for k, v in {"company_tax_id": "987654321", "eta_activity_code": "4330", "eta_governate": "Metro",
                     "eta_region": "Central District", "eta_street": "Y", "eta_building": "2",
                     "eta_default_item_code": "EG-987654321-1"}.items():
            db.set_setting(k, v)
        x = a.save_extract(db, {"bu_code": "PRJ-1", "number": "INV-E", "customer": "Client Co", "date": "2026-03-01",
                                "vat_code": "VAT14", "wht_code": "WHT1-C",
                                "lines": [{"description": "Work", "qty": 1, "price": 1000, "unit": "م٢"}]})
        a.post_extract(db, x)
        doc, problems = einvoice.build_document(db, x)
        self.assertEqual(problems, [])
        self.assertEqual(doc["totalAmount"], 1130.0)                 # 1000 + 140 VAT − 10 WHT
        self.assertEqual(doc["invoiceLines"][0]["unitType"], "MTK")
        self.assertEqual({t["taxType"] for t in doc["taxTotals"]}, {"T1", "T4"})
        ser = einvoice.serialize({"issuer": {"id": "1"}, "lines": [{"a": 1}, {"a": 2}]})
        self.assertEqual(ser, '"ISSUER""ID""1""LINES""LINES""A""1""LINES""A""2"')
        if HAVE_XL:
            self.assertEqual(einvoice.batch_xlsx(db, [x])[:2], b"PK")

    def test_missing_setup_is_reported(self):
        from erp import einvoice
        db = seeded()
        x = db.one("SELECT id FROM extracts")["id"]
        _, problems = einvoice.build_document(db, x)
        self.assertTrue(any("tax registration" in p for p in problems))


class ExcelSafety(unittest.TestCase):
    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_no_formula_exceeds_excel_limit_and_labels_are_english(self):
        import io as _io
        import re
        import openpyxl
        from erp import workbook
        wb = openpyxl.load_workbook(_io.BytesIO(workbook.build_workbook(seeded())))
        longest = max(len(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                      if isinstance(c.value, str) and c.value.startswith("="))
        self.assertLess(longest, 8192)
        self.assertFalse(re.search("[؀-ۿ]", str(wb["Income Statement"]["A6"].value)))
        self.assertEqual(wb["Direct Costs"].cell(4, 1).value, "Date")


class StandaloneWorkbook(unittest.TestCase):
    """The Excel workbook keeps its own books: every sheet that moves money posts to its Trial Balance."""

    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_new_sheets_and_postings(self):
        import io as _io
        import openpyxl
        from erp import demo, workbook
        wb = openpyxl.load_workbook(_io.BytesIO(demo.workbook_bytes(lambda: Database(":memory:"))))
        for name in ("Receivables", "Budget vs Actual", "Scaffold Calculator"):
            self.assertIn(name, wb.sheetnames)
        posting = {s["name"] for s in interchange.SHEETS if s.get("legs")}
        for name in ("Fixed Assets", "Invoice Lines", "Sales Contracts", "Loans", "Custody Expenses", "Service Subscriptions"):
            self.assertIn(name, posting)
        # the report period is always there, even when the given rows leave it out
        keys = [wb["Settings"].cell(r, 1).value for r in range(5, 8)]
        self.assertIn("report_to", keys)
        sc = wb["Scaffold Calculator"]
        self.assertTrue(any(str(c.value or "").startswith("SERVICES TOTAL") for c in sc["A"]))

    def test_app_export_does_not_book_depreciation_twice(self):
        db = seeded()
        assets.save(db, {"code": "FA-T1", "name": "Truck", "account": "1420", "acquired": "2026-01-10", "cost": 120000,
                         "life_months": 60, "pay": "none"})
        rows = interchange.export_data(db)["Fixed Assets"]
        self.assertTrue(rows and all(r["book_depr"] == "No" for r in rows))

    @unittest.skipUnless(HAVE_XL and __import__("shutil").which("soffice"), "LibreOffice not installed")
    def test_demo_workbook_matches_the_app(self):
        """Recalculate the demo workbook in LibreOffice and compare its Trial Balance with the same rows imported."""
        import subprocess
        import openpyxl
        from erp import demo, recurring, services
        sheets = demo.all_sheets(lambda: Database(":memory:"))
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "demo.xlsx"
            src.write_bytes(demo.workbook_bytes(lambda: Database(":memory:")))
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(Path(d) / "out"), str(src)],
                           capture_output=True, timeout=600)
            wb = openpyxl.load_workbook(Path(d) / "out" / "demo.xlsx", data_only=True)
        st = {r[0]: r[1] for r in wb["Settings"].iter_rows(min_row=5, values_only=True) if r[0]}
        to = str(st["report_to"])[:10]
        db = Database(":memory:")
        interchange.import_data(db, sheets)
        recurring.run(db, to, "t")
        services.run(db, to, "t")
        app = {r["account"]: (r["dr"] - r["cr"]) / 100 for r in db.q(
            "SELECT l.account, SUM(l.debit) dr, SUM(l.credit) cr FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id "
            "WHERE e.date<=? AND e.status!='void' GROUP BY l.account", (to,))}
        xl = {str(r[0]): r[7] or 0 for r in wb["Trial Balance"].iter_rows(min_row=5, values_only=True) if r[0]}
        for k in set(app) | set(xl):
            self.assertAlmostEqual(app.get(k, 0), xl.get(k, 0), delta=0.05, msg=f"account {k}")
        self.assertEqual([c.value for c in wb["Model Audit"]["E"] if c.value][-1], "ALL OK")


class Legacy(unittest.TestCase):
    SRC = ROOT.parent / "BEYOND_Accounting_System_V0_6_1.xlsx"

    @unittest.skipUnless(HAVE_XL and SRC.exists(), "legacy workbook not available")
    def test_legacy_boq_and_invoice(self):
        from erp import legacy
        data = legacy.read_legacy_workbook(self.SRC)
        db = Database(":memory:")
        interchange.import_data(db, data, mode="replace")
        self.assertEqual(db.one("SELECT COUNT(*) n FROM boqs")["n"], 2)
        x = a.get_extract(db, db.one("SELECT id FROM extracts")["id"])
        self.assertEqual(x["total"], 81256.0)
        self.assertTrue(all(ln["category"] for ln in x["lines"]))


def _sqljs_runner(raw):
    """sql.js's JSON contract (see web-src/bridge.js), emulated with real SQLite."""
    import sqlite3

    def runner(sql, params_json):
        try:
            params = json.loads(params_json)
            if params is None:
                raw.executescript(sql)
                return json.dumps({"rows": [], "cols": []})
            cur = raw.execute(sql, params)
            rows = [list(r) for r in cur.fetchall()]
            out = {"rows": rows, "cols": [d[0] for d in (cur.description or [])], "changes": raw.total_changes}
            if sql.lstrip().upper().startswith(("INSERT", "REPLACE")):
                out["lastId"] = cur.lastrowid
            return json.dumps(out)
        except sqlite3.Error as e:
            return json.dumps({"error": str(e)})
    return runner


class BrowserAdapter(unittest.TestCase):
    def test_jsdb_adapter_runs_the_engine(self):
        """Emulate sql.js's JSON contract with real SQLite and run the engine through erp.jsdb."""
        import sqlite3
        from erp import jsdb
        db = Database("/tmp/x.sqlite3", connection=jsdb.connect(_sqljs_runner(sqlite3.connect(":memory:", isolation_level=None))))
        db.x("INSERT INTO business_units(code,name) VALUES ('B','Unit')")
        a.record_collection(db, {"bu_code": "B", "amount": 100})
        a.record_expense(db, {"description": "rent", "amount": 40})
        self.assertEqual(a.income_statement(db)["net_profit"], 60.0)
        self.assertTrue(all(c["ok"] for c in a.audit(db)))

    def test_scratch_books_in_the_browser(self):
        """The browser has no sqlite3: Database(":memory:") (business-type templates, the demo company) must use a
        scratch sql.js database instead of failing with AttributeError."""
        import sqlite3
        import types
        from unittest import mock
        from erp import db as dbmod, industries, jsdb
        made, closed = [], []

        def scratch():
            made.append(1)
            return types.SimpleNamespace(run=_sqljs_runner(sqlite3.connect(":memory:", isolation_level=None)),
                                         close=lambda: closed.append(1))
        fake_js = types.SimpleNamespace(globalThis=types.SimpleNamespace(beyondSqlScratch=scratch))
        with mock.patch.object(dbmod, "sqlite3", None), mock.patch.dict(sys.modules, {"js": fake_js}):
            db2 = Database(":memory:")
            self.assertIsInstance(db2.conn, jsdb.Connection)
            industries.apply(db2, "contracting")
            self.assertEqual(db2.setting("business_type"), "contracting")
            self.assertTrue(db2.one("SELECT 1 x FROM accounts WHERE code='5025'"))
            db2.close()
        self.assertEqual((len(made), len(closed)), (1, 1))


class RentalsLanguageAssets(unittest.TestCase):
    """Rental contracts imported with delivery and billing marker; one-language workbooks and templates; the Assets section."""

    def books(self):
        db = Database(":memory:")
        interchange.import_data(db, {
            "Business Units": [{"code": "RNT-1", "name": "Building 3", "client": "Swag"}],
            "Rental Equipment": [{"code": "FR", "name": "Frame", "category": "scaffolding", "unit": "piece", "monthly_rate": 94.5,
                                  "std_cost": 1000, "opening_qty": 200, "opening_unit_cost": 1000, "opening_date": "2026-07-01"}],
            "Rental Contracts": [
                {"ref": "RC-1", "client": "Swag", "site": "Building 3", "quote_date": "2026-07-25", "start_date": "2026-08-01",
                 "duration_months": 12, "line_kind": "item", "item": "FR", "qty": 200, "rate": 94.5, "po_number": "EXT-1",
                 "signed_client": "Swag", "signed_us": "Rawasi", "bu_code": "RNT-1", "delivered_on": "2026-08-01", "invoiced_to": "2026-10-31"},
                {"ref": "RC-1", "client": "Swag", "line_kind": "charge", "charge_type": "install", "qty": 1, "rate": 38250, "cost": 27702}],
            "Sales Invoices": [{"number": "EXT-1", "date": "2026-09-30", "bu_code": "RNT-1", "status": "Posted", "customer": "Swag"}],
            "Invoice Lines": [{"number": "EXT-1", "description": "Rent Aug-Oct", "qty": 600, "price": 94.5}],
        })
        return db

    def test_rental_import_delivers_and_marks_billing(self):
        from erp import rentals
        db = self.books()
        r = rentals.list_rentals(db)["rows"][0]
        self.assertEqual(r["status"], "active")
        on = rentals.ongoing(db, "2026-10-05")["rows"]
        self.assertEqual(len(on), 1)
        self.assertEqual(on[0]["monthly"], 18900.0)
        self.assertEqual(on[0]["pieces_out"], 200)
        self.assertEqual(on[0]["billed_to"], "2026-10-31")

    def test_assets_overview_and_cash_transfer(self):
        from erp import asset_views
        db = Database(":memory:")
        a.record_journal(db, {"date": "2026-01-01", "memo": "opening", "lines": [{"account": "1010", "debit": 1000}, {"account": "3010", "credit": 1000}]})
        asset_views.transfer(db, {"date": "2026-01-02", "from": "1010", "to": "1020", "amount": 300})
        cash = {x["code"]: x["balance"] for x in asset_views.cash_accounts(db)["accounts"]}
        self.assertEqual((cash["1010"], cash["1020"]), (700.0, 300.0))
        ov = asset_views.overview(db)
        self.assertEqual(ov["total"], a.balance_sheet(db)["total_assets"])
        self.assertEqual(next(c for c in ov["classes"] if c["key"] == "cash")["amount"], 1000.0)
        with self.assertRaises(a.ValidationError):
            asset_views.transfer(db, {"from": "1010", "to": "1010", "amount": 5})

    def test_workbooks_and_templates_are_one_language(self):
        import io as _io
        import re as _re
        import openpyxl
        from erp import importkit, workbook
        AR = _re.compile("[؀-ۿ]")
        db = Database(":memory:")
        data_sheets = {s["name"] for s in interchange.SHEETS}
        en = openpyxl.load_workbook(_io.BytesIO(workbook.build_workbook(db, template=True)))
        bad = [(ws.title, c.coordinate) for ws in en.worksheets for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and not c.value.startswith("=") and AR.search(c.value)
               and not (ws.title in data_sheets and (c.row > 4 or c.row == 3))]
        self.assertEqual(bad, [])
        ar = openpyxl.load_workbook(_io.BytesIO(workbook.build_workbook(db, template=True, lang="ar")))
        self.assertEqual(ar.sheetnames[:2], ["الغلاف", "دليل الاستخدام"])
        self.assertTrue(ar["سندات القبض"].sheet_view.rightToLeft)
        for lang, want in (("en", "Date"), ("ar", "التاريخ")):
            t = openpyxl.load_workbook(_io.BytesIO(importkit.template_xlsx(db, "Receipt Vouchers", lang)))
            self.assertEqual(t.worksheets[0].cell(interchange.HEADER_ROW, 1).value, want)

    def test_arabic_workbook_imports_back(self):
        from erp import workbook
        db = self.books()
        a.record_collection(db, {"bu_code": "RNT-1", "amount": 5000})
        raw = workbook.build_workbook(db, lang="ar")
        d2 = Database(":memory:")
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
            fh.write(raw)
        interchange.import_data(d2, interchange.read_xlsx(fh.name), mode="replace")
        __import__("os").unlink(fh.name)
        self.assertEqual(a.balance_sheet(d2)["total_assets"], a.balance_sheet(db)["total_assets"])


class Extracts(unittest.TestCase):
    """Running extracts (مستخلص جاري): the client's totals, the next extract of a rental contract, follow-up, adjusting."""

    def setUp(self):
        self.db = RentalsLanguageAssets().books()      # rental RC-1 on RNT-1, extract EXT-1 (600 × 94.50) issued up to 31/10
        db = self.db
        db.x("INSERT INTO business_units(code,name) VALUES ('B9','Building 9')")
        self.xid = a.save_extract(db, {"number": "XT-1", "date": "2026-09-30", "bu_code": "B9", "customer": "Swag", "prev_paid": 65000,
                                       "seq_no": 1, "project": "اجورا", "work": "عمارة 9",
                                       "lines": [{"description": "تركيب و الفك", "qty": 1, "price": 38250},
                                                 {"description": "ايجار", "qty": 200, "price": 94.5, "pct": "50%"}]})

    def test_percent_and_client_totals(self):
        from erp import extracts as xt
        x = xt.detail(self.db, self.xid)
        self.assertEqual(x["lines"][1]["amount"], 9450.0)          # 200 × 94.50 × 50 %
        self.assertEqual(x["total"], 47700.0)
        self.assertEqual((x["prev"], x["prev_auto"]), (65000.0, False))
        self.assertEqual(x["due"], round(47700 - 65000, 2))

    def test_next_extract_for_rental_and_follow_up(self):
        from erp import extracts as xt
        db = self.db
        first = db.one("SELECT id FROM extracts WHERE number='EXT-1'")["id"]
        a.record_collection(db, {"bu_code": "RNT-1", "amount": 30000, "credit_account": db.acct("ar"), "date": "2026-10-02"})
        row = next(r for r in xt.board(db)["rows"] if r["id"] == first)
        self.assertEqual((row["collected"], row["owed"], row["pay_state"]), (30000.0, 26700.0, "part"))
        xid = xt.next_for_rental(db, 1, {"date": "2026-11-01"})
        x = xt.detail(db, xid)
        self.assertEqual((x["number"], x["seq_no"]), ("EXT-2", 2))  # follows the site's last number
        rent = [ln for ln in x["lines"] if ln["description"].startswith("ايجار")]
        self.assertTrue(rent and all("مدة 1/11/2026 الي 1/12/2026" in ln["description"] for ln in rent))
        self.assertEqual(rent[0]["qty"], 200.0)                     # frames on site × one month
        # running: works to date include extract 1; paid before = what was collected before the date
        self.assertEqual(x["works_to_date"], round(56700 + x["total"], 2))
        self.assertEqual(x["prev"], 30000.0)
        self.assertEqual(x["due"], round(56700 + x["total"] - 30000, 2))
        self.assertEqual(xt.xlsx(db, xid)[0][:2], b"PK")
        self.assertIn("المستحق".encode(), xt.html(db, xid))

    def test_extract_fields_round_trip_through_the_workbook(self):
        import tempfile
        from erp import workbook
        raw = workbook.build_workbook(self.db)
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
            fh.write(raw)
        d2 = Database(":memory:")
        interchange.import_data(d2, interchange.read_xlsx(fh.name), mode="replace")
        __import__("os").unlink(fh.name)
        x = d2.one("SELECT * FROM extracts WHERE number='XT-1'")
        self.assertEqual((x["seq_no"], x["project"], x["prev_paid"]), (1, "اجورا", 6500000))
        self.assertEqual(d2.one("SELECT pct FROM extract_lines WHERE extract_id=? AND description='ايجار'", (x["id"],))["pct"], 0.5)


class DashboardsPerGroup(unittest.TestCase):
    """Each group has its own dashboard; only super administrators change them; a card shows only to people allowed to
    see it unless a super administrator ticked "show even without access"; the server sends nothing else."""

    def setUp(self):
        from erp.server import handle
        from erp import demo
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = Database(":memory:")
        demo.load(self.db, "boss")
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "admin2", "password": "secret123", "is_admin": True})
        auth.create_user(self.db, {"username": "sam", "password": "secret123", "full_name": "Sam Sales", "group_key": "sales",
                                   "perms": {"shop": "view", "sales": "view"}})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "admin2", "sam")}

    def call(self, who, method, url, body=None):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json"}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        return r.status, json.loads(r.body)

    def test_groups_permissions_and_grants(self):
        st, v = self.call("sam", "GET", "/api/dashboard/view")
        self.assertEqual((st, v["group"]), (200, "sales"))
        ids = [w["w"] for w in v["widgets"]]
        self.assertIn("sales_target", ids)
        self.assertIn("installments_upcoming", ids)
        self.assertNotIn("catalog", v)                                     # the catalogue is for super administrators
        # another group's dashboard is not theirs to see — they get their own
        self.assertEqual(self.call("sam", "GET", "/api/dashboard/view?group=management")[1]["group"], "sales")
        # only super administrators change dashboards (not even an administrator)
        self.assertEqual(self.call("admin2", "PUT", "/api/dashboard/groups/sales", {"layout": []})[0], 403)
        self.assertEqual(self.call("sam", "POST", "/api/dashboard/groups", {"name": "x"})[0], 403)
        # a money card Sam may not see stays out, figures and all…
        self.assertEqual(self.call("boss", "PUT", "/api/dashboard/groups/sales",
                                   {"layout": [{"w": "sales_target", "opts": {"target": 50000, "period": "month"}}, {"w": "cash"}]})[0], 200)
        v = self.call("sam", "GET", "/api/dashboard/view")[1]
        self.assertEqual([w["w"] for w in v["widgets"]], ["sales_target"])
        self.assertEqual(v["widgets"][0]["data"]["target"], 50000)
        # …until a super administrator lets the group see it
        self.call("boss", "PUT", "/api/dashboard/groups/sales", {"layout": [{"w": "sales_target"}, {"w": "cash", "grant": True}]})
        v = self.call("sam", "GET", "/api/dashboard/view")[1]
        self.assertEqual([w["w"] for w in v["widgets"]], ["sales_target", "cash"])
        self.assertIn("value", v["widgets"][1]["data"])
        # the super administrator sees any group's dashboard and the catalogue to customise it
        v = self.call("boss", "GET", "/api/dashboard/view?group=hr")[1]
        self.assertEqual(v["group"], "hr")
        self.assertTrue(v["catalog"] and v["groups"])
        # new group, copied from Sales; deleting it moves its people to General
        st, r = self.call("boss", "POST", "/api/dashboard/groups", {"name": "Site engineers", "copy_from": "sales"})
        self.assertEqual(st, 200)
        self.assertEqual(self.call("boss", "PUT", f"/api/users/{self.db.one('SELECT id FROM users WHERE username=?', ('sam',))['id']}",
                                   {"group_key": r["key"]})[0], 200)
        self.assertEqual(self.call("sam", "GET", "/api/dashboard/view")[1]["group"], r["key"])
        self.call("boss", "DELETE", "/api/dashboard/groups/" + r["key"])
        self.assertEqual(self.call("sam", "GET", "/api/dashboard/view")[1]["group"], "general")
        self.assertEqual(self.call("boss", "DELETE", "/api/dashboard/groups/general")[0], 400)

    def test_every_card_works_on_real_books(self):
        from erp import dashboards
        lay = [{"w": w} for w in dashboards.WIDGETS]
        self.call("boss", "PUT", "/api/dashboard/groups/management", {"layout": lay})
        v = self.call("boss", "GET", "/api/dashboard/view?group=management")[1]
        self.assertEqual(len(v["widgets"]), len(dashboards.WIDGETS))
        self.assertEqual([w["w"] for w in v["widgets"] if w.get("error")], [])


class ChatProfilesAssistant(unittest.TestCase):
    """Chat rooms and their members, edit-until-seen, files, mentions, export; profiles; the AI assistant's policy."""

    def setUp(self):
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = Database(":memory:")
        auth.create_user(self.db, {"username": "boss", "password": "secret123", "full_name": "Boss"}, first=True)
        auth.create_user(self.db, {"username": "acc", "password": "secret123", "full_name": "Accountant",
                                   "perms": {"accounting": "edit"}})
        auth.create_user(self.db, {"username": "eng", "password": "secret123", "full_name": "Engineer"})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "acc", "eng")}

    def call(self, who, method, url, body=None, **headers):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json", **headers}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        try:
            return r.status, json.loads(r.body)
        except ValueError:
            return r.status, r.body

    def test_new_company_setup_until_done(self):
        boot = lambda who: self.call(who, "GET", "/api/bootstrap")[1]   # noqa: E731
        self.assertTrue(boot("boss")["onboarding"])        # an administrator of empty books gets the start options
        self.assertFalse(boot("eng")["onboarding"])        # other people never do
        self.assertEqual(self.call("eng", "POST", "/api/onboarding", {})[0], 403)
        self.assertTrue(boot("boss")["onboarding"])        # still pending until a choice is finished (resume later)
        self.assertEqual(self.call("boss", "POST", "/api/onboarding", {"business_type": "interior"})[0], 200)
        self.assertFalse(boot("boss")["onboarding"])
        self.assertEqual(self.db.setting("business_type"), "interior")

    def test_rooms_members_and_privacy(self):
        self.assertEqual(self.call("acc", "POST", "/api/chat/rooms", {"name": "x"})[0], 403)   # super administrators only
        st, r = self.call("boss", "POST", "/api/chat/rooms", {"name": "Accountants", "members": [2]})
        self.assertEqual(st, 200)
        rid = r["id"]
        self.assertEqual([x["name"] for x in self.call("eng", "GET", "/api/chat/rooms")[1]["rooms"]], ["General"])
        self.assertEqual(self.call("eng", "GET", f"/api/chat/rooms/{rid}/messages")[0], 403)
        self.assertEqual(self.call("eng", "POST", f"/api/chat/rooms/{rid}/messages", {"body": "hi"})[0], 403)
        self.assertEqual(self.call("acc", "POST", f"/api/chat/rooms/{rid}/messages", {"body": "hi"})[0], 200)
        # removing a member takes the room away from them; the history stays
        self.call("boss", "PUT", f"/api/chat/rooms/{rid}", {"members": [1]})
        self.assertEqual(self.call("acc", "GET", f"/api/chat/rooms/{rid}/messages")[0], 403)
        self.assertEqual(len(self.call("boss", "GET", f"/api/chat/rooms/{rid}/messages")[1]["messages"]), 1)
        # archived rooms are read-only; General cannot be archived or deleted
        self.call("boss", "POST", f"/api/chat/rooms/{rid}/archive", {"archived": True})
        self.assertEqual(self.call("boss", "POST", f"/api/chat/rooms/{rid}/messages", {"body": "x"})[0], 403)
        self.assertEqual(self.call("boss", "POST", "/api/chat/rooms/1/archive", {})[0], 400)
        self.assertEqual(self.call("boss", "DELETE", "/api/chat/rooms/1")[0], 400)

    def test_edit_until_seen_mentions_files_export(self):
        st, m = self.call("boss", "POST", "/api/chat/rooms/1/messages",
                          {"body": "hi @acc and @all", "files": [{"name": "a.txt", "mime": "text/plain", "data": "aGVsbG8="}]})
        self.assertEqual(st, 200)
        self.assertEqual(self.call("boss", "PUT", f"/api/chat/messages/{m['id']}", {"body": "hello @acc"})[0], 200)
        self.assertEqual(self.call("acc", "PUT", f"/api/chat/messages/{m['id']}", {"body": "hack"})[0], 403)
        up = self.call("acc", "GET", "/api/chat/updates?since=0")[1]
        self.assertTrue(up["messages"][0]["mention"])
        self.assertEqual(up["unread"], 1)
        msgs = self.call("acc", "GET", "/api/chat/rooms/1/messages")[1]["messages"]
        self.assertEqual(msgs[0]["files"][0]["name"], "a.txt")
        self.assertEqual(self.call("acc", "GET", f"/api/chat/files/{msgs[0]['files'][0]['id']}")[1], b"hello")
        self.call("acc", "POST", "/api/chat/rooms/1/read", {"last": m["id"]})
        st, err = self.call("boss", "PUT", f"/api/chat/messages/{m['id']}", {"body": "too late"})
        self.assertEqual(st, 400)
        self.assertIn("has been read", err["error"])
        self.assertEqual(self.call("acc", "GET", "/api/chat/updates?since=0")[1]["unread"], 0)
        st, dm = self.call("eng", "POST", "/api/chat/dm", {"user_id": 2})
        self.assertEqual(self.call("eng", "POST", "/api/chat/dm", {"user_id": 2})[1]["id"], dm["id"])   # one room per pair
        self.assertEqual(self.call("boss", "GET", f"/api/chat/rooms/{dm['id']}/messages")[0], 403)      # private, even for the super administrator
        r = self.handle(self.app, "GET", "/api/chat/rooms/1/export", {"x-requested-with": "BeyondERP", "x-session": self.tok["boss"]}, b"")
        import io
        import zipfile
        names = zipfile.ZipFile(io.BytesIO(r.body)).namelist()
        self.assertIn("transcript.html", names)
        self.assertTrue(any(n.startswith("files/") and n.endswith("a.txt") for n in names))
        self.assertEqual(self.call("acc", "GET", "/api/chat/rooms/1/export")[0], 403)

    def test_profile_defaults_and_prefs(self):
        p = self.call("eng", "GET", "/api/me/profile")[1]
        self.assertEqual(p["prefs"]["nav_mode"], "click")            # menus open on click unless the person chooses hover
        self.assertEqual(p["prefs"]["notify"]["level"], "all")
        p = self.call("eng", "PUT", "/api/me/profile", {"job_title": "Site engineer", "prefs": {"nav_mode": "hover",
                      "notify": {"level": "rooms", "rooms": [1], "mode": "vibrate"}}})[1]
        self.assertEqual((p["job_title"], p["prefs"]["nav_mode"], p["prefs"]["notify"]["mode"]), ("Site engineer", "hover", "vibrate"))
        self.assertEqual(self.call("eng", "PUT", "/api/me/profile", {"avatar": "javascript:alert(1)"})[0], 400)
        boot = self.call("eng", "GET", "/api/bootstrap")[1]
        self.assertEqual(boot["profile"]["prefs"]["nav_mode"], "hover")
        self.assertIn("Site engineer", [x["job_title"] for x in self.call("acc", "GET", "/api/people")[1]])

    def test_assistant_policy_and_permissions(self):
        from erp import ai, cloud
        aih = {"x-ai-call": "1"}
        self.assertEqual(self.call("acc", "GET", "/api/accounts", **aih)[0], 403)              # switched off
        self.assertEqual(self.call("acc", "PUT", "/api/ai-policy", {"enabled": True})[0], 403)  # super administrators only
        self.call("boss", "PUT", "/api/ai-policy", {"enabled": True, "write": "off", "blocked": ["payroll"]})
        self.assertEqual(self.call("acc", "GET", "/api/accounts", **aih)[0], 200)
        self.assertEqual(self.call("boss", "GET", "/api/users", **aih)[0], 403)                 # never user management
        self.assertEqual(self.call("boss", "PUT", "/api/ai-policy", {"enabled": False}, **aih)[0], 403)
        self.assertEqual(self.call("boss", "GET", "/api/payroll/runs", **aih)[0], 403)     # blocked section
        self.assertEqual(self.call("boss", "POST", "/api/ai/preview", {"method": "POST", "path": "/api/docs/payroll", "body": {}})[0], 403)
        # read-only policy: no changes at all, and no change tool offered
        self.assertEqual(self.call("acc", "POST", "/api/bus", {"code": "X", "name": "X"}, **aih)[0], 403)
        self.assertNotIn("erp_write", [t["name"] for t in self.call("acc", "GET", "/api/ai/context")[1]["tools"]])
        # changes allowed: still only what the person may do themselves, and only after the check and review
        self.call("boss", "PUT", "/api/ai-policy", {"write": "confirm"})
        bu = {"code": "X", "name": "X"}
        self.assertEqual(self.call("eng", "POST", "/api/bus", bu, **aih)[0], 403)
        st, err = self.call("boss", "POST", "/api/bus", bu, **aih)
        self.assertEqual(st, 403)                                                              # not reviewed: refused
        self.assertIn("validated and reviewed", err["error"])
        rv = self.call("eng", "POST", "/api/ai/preview", {"method": "POST", "path": "/api/bus", "body": bu})[1]
        self.assertFalse(rv["ok"])                                                             # the engineer may not
        self.assertNotIn("review", rv)
        rv = self.call("boss", "POST", "/api/ai/preview", {"method": "POST", "path": "/api/bus", "body": bu})[1]
        self.assertTrue(rv["ok"], rv["checks"])
        self.assertIsNone(self.db.one("SELECT 1 FROM business_units WHERE code='X'"))         # the check saved nothing
        rh = {**aih, "x-ai-review": rv["review"]}
        self.assertEqual(self.call("boss", "POST", "/api/bus", {"code": "Y", "name": "Y"}, **rh)[0], 403)  # not what was reviewed
        self.assertEqual(self.call("boss", "POST", "/api/bus", bu, **rh)[0], 200)
        # a posting: the check shows the entry and its effect; books changed meanwhile → review again
        a.record_collection(self.db, {"bu_code": "X", "amount": 100})
        exp = {"description": "Office rent", "amount": 40, "date": a.today()}
        rv = self.call("boss", "POST", "/api/ai/preview", {"method": "POST", "path": "/api/docs/expense", "body": exp})[1]
        self.assertTrue(rv["ok"], rv["checks"])
        self.assertEqual({c["key"] for c in rv["checks"]}, {"accepted", "balanced", "tb", "audit", "sheet"})
        self.assertEqual(len([e for e in rv["entries"] if e["status"] == "new"]), 1)
        self.assertTrue(any(f["key"] == "net_profit" and f["change"] == -40 for f in rv["figures"]), rv["figures"])
        a.record_expense(self.db, {"description": "someone else", "amount": 5})
        st, err = self.call("boss", "POST", "/api/docs/expense", exp, **{**aih, "x-ai-review": rv["review"]})
        self.assertEqual(st, 403)
        self.assertIn("changed since", err["error"])
        # a change the ERP itself rejects is blocked at the check, with the reason
        rv = self.call("boss", "POST", "/api/ai/preview", {"method": "POST", "path": "/api/docs/expense", "body": {"amount": -5}})[1]
        self.assertFalse(rv["ok"])
        self.assertFalse(rv["checks"][0]["ok"])
        # the shared workspace path enforces the same rule
        boss = auth.public_user(self.db.one("SELECT * FROM users WHERE username='boss'"))
        r, op = cloud.prepare(self.app, "POST", "/api/bus", {"x-requested-with": "BeyondERP", "content-type": "application/json",
                                                             "x-ai-call": "1"}, json.dumps({"code": "Z", "name": "Z"}).encode(), boss)
        self.assertEqual((r.status, op), (403, None))
        # who may use it
        self.call("boss", "PUT", "/api/ai-policy", {"access": "list", "users": ["acc"]})
        self.assertEqual(self.call("eng", "GET", "/api/ai/context")[0], 403)
        self.assertEqual(self.call("acc", "GET", "/api/ai/context")[0], 200)
        # the request is built on the server: model, instructions and tools from the policy, never from the page
        self.call("boss", "PUT", "/api/ai-policy", {"instructions": "Always answer in Arabic.", "model": "claude-opus-5-5"})
        params = self.call("acc", "POST", "/api/ai/params", {"messages": [{"role": "user", "content": "hi"}], "model": "x"})[1]["params"]
        self.assertEqual(params["model"], "claude-opus-5-5")
        self.assertIn("Always answer in Arabic.", params["system"][0]["text"])
        self.assertEqual(params["thinking"], {"type": "adaptive"})
        self.assertEqual(self.call("acc", "POST", "/api/ai/params", {"messages": [{"role": "system", "content": "x"}]})[0], 400)
        # the key and the conversations stay on the device: never replicated to the workspace
        self.assertFalse(cloud.replicated("PUT", "/api/ai/key"))
        self.assertFalse(cloud.replicated("POST", "/api/ai/threads"))
        self.assertTrue(cloud.replicated("PUT", "/api/ai-policy"))
        self.assertIn("ai_threads", cloud.LOCAL_TABLES)
        self.assertEqual(self.call("acc", "PUT", "/api/ai/key", {"key": "sk-ant-x" * 5})[0], 403)
        self.call("boss", "PUT", "/api/ai/key", {"key": "sk-ant-test-0123456789abcdef"})
        self.assertNotIn("sk-ant", json.dumps(self.call("boss", "GET", "/api/ai-policy")[1]))
        self.assertNotIn("sk-ant", json.dumps(self.call("boss", "GET", "/api/bootstrap")[1]))
        self.assertTrue(ai.has_key(self.db))
        # conversations are private to their author
        st, th = self.call("acc", "POST", "/api/ai/threads", {"title": "t", "messages": [{"role": "user", "content": "hi"}]})
        self.assertEqual(self.call("boss", "GET", f"/api/ai/threads/{th['id']}")[0], 400)


class AssistantAnyProvider(unittest.TestCase):
    """Other AI providers through the OpenAI-style chat-completions API: settings, translation both ways, errors."""

    def test_openai_compatible_round_trip(self):
        import threading as th
        from http.server import BaseHTTPRequestHandler, HTTPServer
        from erp import ai
        seen = []

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["content-length"])))
                seen.append((self.path, self.headers.get("Authorization"), body))
                if len(seen) == 1:
                    msg = {"role": "assistant", "content": None, "tool_calls": [
                        {"id": "c1", "type": "function", "function": {"name": "erp_read", "arguments": "{\"path\": \"/api/accounts\"}"}}]}
                    fin = "tool_calls"
                else:
                    msg, fin = {"role": "assistant", "content": "Here you go."}, "stop"
                out = json.dumps({"model": body["model"], "choices": [{"message": msg, "finish_reason": fin}]}).encode()
                self.send_response(200)
                self.send_header("content-type", "application/json")
                self.send_header("content-length", str(len(out)))
                self.end_headers()
                self.wfile.write(out)
        srv = HTTPServer(("127.0.0.1", 0), H)
        th.Thread(target=srv.serve_forever, daemon=True).start()
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "secret123", "full_name": "Boss"}, first=True)
        u = auth.public_user(db.one("SELECT * FROM users WHERE id=1"))
        with self.assertRaises(a.ValidationError):
            ai.save_policy(db, u, {"enabled": True, "provider": "custom", "base_url": "ftp://x", "model": "m"})
        ai.save_policy(db, u, {"enabled": True, "provider": "ollama", "base_url": f"http://127.0.0.1:{srv.server_port}/v1",
                               "model": "llama3.1:8b", "instructions": "Be brief."})
        m = ai.complete(db, u, {"messages": [{"role": "user", "content": "hi"}]})
        self.assertEqual(m["stop_reason"], "tool_use")
        self.assertEqual(m["content"][0], {"type": "tool_use", "id": "c1", "name": "erp_read", "input": {"path": "/api/accounts"}})
        path, authz, body = seen[0]
        self.assertEqual((path, authz, body["model"]), ("/v1/chat/completions", None, "llama3.1:8b"))   # local model: no key
        self.assertEqual(body["messages"][0]["role"], "system")
        self.assertIn("Be brief.", body["messages"][0]["content"])
        self.assertEqual({t["function"]["name"] for t in body["tools"]}, {"erp_find", "erp_read", "erp_write"})
        # the tool result goes back as a "tool" message tied to the call
        msgs = [{"role": "user", "content": "hi"}, {"role": "assistant", "content": m["content"]},
                {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "c1", "content": "[]"}]}]
        ai.key_set(db, u, "sk-test-0123456789abcdefghij")
        m2 = ai.complete(db, u, {"messages": msgs})
        self.assertEqual((m2["stop_reason"], m2["content"][0]["text"]), ("end_turn", "Here you go."))
        b2 = seen[1][2]["messages"]
        self.assertEqual(b2[2]["tool_calls"][0]["id"], "c1")
        self.assertEqual((b2[3]["role"], b2[3]["tool_call_id"]), ("tool", "c1"))
        self.assertEqual(seen[1][1], "Bearer sk-test-0123456789abcdefghij")
        srv.shutdown()
        # a provider that cannot be reached gives a clear message
        ai.save_policy(db, u, {"base_url": "http://127.0.0.1:9/v1"})
        with self.assertRaises(ai.AIError):
            ai.complete(db, u, {"messages": [{"role": "user", "content": "hi"}]})

    def test_kimi_style_reasoning_diagnostics_and_connection_test(self):
        import threading as th
        from http.server import BaseHTTPRequestHandler, HTTPServer
        from erp import ai
        seen = []

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["content-length"])))
                seen.append((self.path, body))
                if self.path.startswith("/web/"):            # a website, not an API
                    out, ctype = b"<!doctype html><title>Kimi</title>", "text/html"
                else:
                    msg = {"role": "assistant", "content": "OK", "reasoning_content": "The user wants OK."}
                    if body.get("tools"):
                        msg = {"role": "assistant", "content": "", "reasoning_content": "Look up accounts.",
                               "tool_calls": [{"id": "k1", "type": "function", "function": {"name": "erp_read", "arguments": "{\"path\":\"/api/accounts\"}"}}]}
                    out, ctype = json.dumps({"model": "kimi-k3", "choices": [{"message": msg, "finish_reason": "tool_calls" if body.get("tools") else "stop"}]}).encode(), "application/json"
                self.send_response(200)
                self.send_header("content-type", ctype)
                self.send_header("content-length", str(len(out)))
                self.end_headers()
                self.wfile.write(out)
        srv = HTTPServer(("127.0.0.1", 0), H)
        th.Thread(target=srv.serve_forever, daemon=True).start()
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "secret123", "full_name": "Boss"}, first=True)
        u = auth.public_user(db.one("SELECT * FROM users WHERE id=1"))
        self.assertEqual(ai.policy(db)["name"], "Kazuha")                          # not tied to one AI provider
        db.set_setting("ai_policy", json.dumps({"name": "Claude"}))
        self.assertEqual(ai.policy(db)["name"], "Kazuha")                          # the old default is renamed
        ai.save_policy(db, u, {"provider": "kimi", "base_url": f"http://127.0.0.1:{srv.server_port}/v1", "model": "kimi-k3"})
        # the connection test works while the assistant is still off, without tools
        m = ai.complete(db, u, {"test": True})
        self.assertEqual(m["content"][-1], {"type": "text", "text": "OK"})
        self.assertNotIn("tools", seen[-1][1])
        ai.save_policy(db, u, {"enabled": True})
        m = ai.complete(db, u, {"messages": [{"role": "user", "content": "accounts?"}]})
        self.assertEqual(m["content"][0], {"type": "reasoning", "text": "Look up accounts."})
        msgs = [{"role": "user", "content": "accounts?"}, {"role": "assistant", "content": m["content"]},
                {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "k1", "content": "[]"}]}]
        ai.complete(db, u, {"messages": msgs})
        sent = seen[-1][1]["messages"][2]
        self.assertEqual((sent["reasoning_content"], sent["tool_calls"][0]["id"]), ("Look up accounts.", "k1"))
        # switching to Claude drops the other provider's reasoning blocks
        ai.save_policy(db, u, {"provider": "anthropic", "model": "claude-opus-5-5"})
        _, params = ai.request_params(db, u, {"messages": msgs})
        self.assertFalse(any(b.get("type") == "reasoning" for m2 in params["messages"] if isinstance(m2["content"], list) for b in m2["content"]))
        # an address that is a website, not an API: say so, with what came back
        ai.save_policy(db, u, {"provider": "custom", "base_url": f"http://127.0.0.1:{srv.server_port}/web", "model": "kimi-k3"})
        with self.assertRaises(ai.AIError) as cm:
            ai.complete(db, u, {"messages": [{"role": "user", "content": "hi"}]})
        self.assertIn("did not respond as an AI API", str(cm.exception))
        self.assertIn("<!doctype html>", str(cm.exception))
        srv.shutdown()


class SiteModules(unittest.TestCase):
    """Projects, contracts & certificates, purchase requests, stocktakes, equipment, attendance and e-receipts."""

    def setUp(self):
        self.db = seeded()

    def test_variation_orders_change_the_project_contract_value(self):
        from erp import contracts
        cid = contracts.save(self.db, {"kind": "client", "bu_code": "PRJ-2", "party": "Client X", "value": 500000, "finish": "2026-12-31"})
        vid = contracts.save_vo(self.db, cid, {"description": "Extra room", "amount": 50000, "days": 20, "date": "2026-03-01"})
        self.assertEqual(contracts.client_value(self.db, 2)["revised"], 500000)   # pending: no change yet
        contracts.decide_vo(self.db, vid, "approved", "admin")
        self.assertEqual(contracts.client_value(self.db, 2)["revised"], 550000)
        self.assertEqual(self.db.one("SELECT contract_value FROM business_units WHERE id=2")["contract_value"], 55000000)
        self.assertEqual(contracts.get(self.db, cid)["revised_finish"], "2027-01-20")

    def test_subcontract_certificates_are_cumulative_and_recover_the_advance(self):
        from erp import contracts
        sid = contracts.save(self.db, {"kind": "sub", "bu_code": "PRJ-2", "party": "Sub Co", "retention_pct": 10, "advance_recovery_pct": 20,
                                       "lines": [{"description": "Plaster", "qty": 1000, "rate": 100}, {"description": "Paint", "qty": 1000, "rate": 50}]})
        lines = contracts.get(self.db, sid)["lines"]
        contracts.pay(self.db, sid, {"kind": "advance", "amount": 20000, "date": "2026-03-01"})
        x1 = contracts.save_certificate(self.db, sid, {"date": "2026-03-15", "lines": [{"line_id": lines[0]["id"], "qty_to_date": 300}]})
        with self.assertRaises(a.ValidationError):
            contracts.post_certificate(self.db, x1)          # must be approved first
        contracts.approve_certificate(self.db, x1)
        e = a.get_entry(self.db, contracts.post_certificate(self.db, x1))
        by = {ln["account"]: ln for ln in e["lines"]}
        self.assertEqual((by["5020"]["debit"], by["2015"]["credit"], by["2110"]["credit"], by["1310"]["credit"]), (30000, 21000, 3000, 6000))
        x2 = contracts.save_certificate(self.db, sid, {"date": "2026-04-15", "lines": [{"line_id": lines[0]["id"], "qty_to_date": 500},
                                                                                        {"line_id": lines[1]["id"], "qty_to_date": 200}]})
        contracts.approve_certificate(self.db, x2)
        contracts.post_certificate(self.db, x2)
        st = contracts.get(self.db, sid)["statement"]
        self.assertEqual((st["certified"], st["retention_held"], st["advance_outstanding"], st["payable"]), (60000, 6000, 8000, 42000))
        contracts.pay(self.db, sid, {"kind": "payment", "amount": 10000, "date": "2026-04-20"})
        contracts.pay(self.db, sid, {"kind": "release", "amount": 5000, "date": "2026-04-21"})
        with self.assertRaises(a.ValidationError):           # its retention was partly released
            contracts.reopen_certificate(self.db, x2)
        self.assertTrue(all(c["ok"] for c in a.audit(self.db)))

    def test_purchase_request_quotes_award_to_orders(self):
        from erp import procurement, purchasing
        pid = procurement.save(self.db, {"bu_code": "PRJ-1", "lines": [{"description": "Cement", "qty": 100, "est_rate": 90},
                                                                       {"description": "Sand", "qty": 10, "est_rate": 300}]}, "site")
        ids = [ln["id"] for ln in procurement.get(self.db, pid)["lines"]]
        procurement.save_quote(self.db, pid, {"supplier": "Sup A", "rates": {ids[0]: 85, ids[1]: 320}})
        procurement.save_quote(self.db, pid, {"supplier": "Sup B", "rates": {ids[0]: 88, ids[1]: 290}})
        with self.assertRaises(a.ValidationError):
            procurement.award(self.db, pid, {"strategy": "lowest"})   # not approved
        procurement.set_status(self.db, pid, "submitted")
        procurement.set_status(self.db, pid, "approved", "boss")
        cmp_ = procurement.get(self.db, pid)["comparison"]
        self.assertEqual((cmp_["lowest_split_total"], cmp_["best_full_total"]), (11400, 11700))
        codes = procurement.award(self.db, pid, {"strategy": "lowest"})
        self.assertEqual(len(codes), 2)
        self.assertEqual(sorted(purchasing.get(self.db, c)["supplier"] for c in codes), ["Sup A", "Sup B"])
        self.assertEqual(procurement.get(self.db, pid)["status"], "ordered")

    def test_stock_search_by_code_and_last_purchase_price(self):
        from erp import reports
        wh = self.db.one("SELECT id FROM warehouses WHERE name='Site Store'")["id"]
        for rate, date in ((100, "2026-01-05"), (140, "2026-03-05")):
            a.stock_move(self.db, {"type": "Receipt", "warehouse_id": wh, "item": "Steel pipe", "qty": 10, "rate": rate, "vendor": "S", "date": date}, "t")
        self.db.x("UPDATE stock_items SET sku='PIPE-01', category='Metals' WHERE name='Steel pipe'")
        for q in ("pipe-0", "metal", "steel 01"):                          # code, category, words across name and code
            self.assertEqual([r["name"] for r in a.stock_report(self.db, None, q)], ["Steel pipe"], q)
        self.assertEqual(a.stock_report(self.db, None, "steel xyz"), [])
        r = a.stock_report(self.db, None, "PIPE-01")[0]
        self.assertEqual((r["avg_cost"], r["last_rate"], r["last_date"]), (120, 140, "2026-03-05"))
        t = reports.stock_table(self.db, q="PIPE-01")
        self.assertEqual([c["label"] for c in t["columns"]][5:8], ["Avg cost", "Last purchase price", "Last purchased on"])
        self.assertEqual(t["rows"][0]["cells"][:1] + t["rows"][0]["cells"][5:8], ["PIPE-01", 120, 140, "2026-03-05"])

    def test_duplicate_items_found_merged_and_prevented(self):
        from erp import dedupe
        wh = self.db.one("SELECT id FROM warehouses WHERE name='Site Store'")["id"]
        ins = lambda n, sku="": self.db.x("INSERT INTO stock_items(name,unit,sku) VALUES (?,?,?)", (n, "pcs", sku)).lastrowid  # noqa: E731
        a1, a2 = ins("حنفية غسالة"), ins("حنفيه  غساله")                     # spelling only
        p25, p32 = ins("مواسير 25مم / 20بار"), ins("مواسير 32مم / 20بار")     # different sizes: never offered
        r1, r2 = ins("سلك 2مم احمر"), ins("سلك 2مم ازرق")                    # different colours: never offered
        c1, c2 = ins("Copper wire", "CW-10"), ins("Cu wire roll", "cw 10")    # same code
        groups = {tuple(g["ids"]): g for g in dedupe.find(self.db)}
        self.assertEqual(groups[(a1, a2)]["level"], "certain")
        self.assertTrue(groups[(c1, c2)]["reason"].startswith("same code"))
        self.assertFalse(any(p25 in k and p32 in k for k in groups) or any(r1 in k and r2 in k for k in groups))
        # merge: movements and balances move over, balances recomputed, the old name becomes an alias
        for iid, q, rate in ((a1, 4, 100), (a2, 6, 200)):
            self.db.x("INSERT INTO stock_moves(item_id,warehouse_id,date,qty,rate,kind) VALUES (?,?,?,?,?,?)", (iid, wh, "2026-03-01", q, rate * 100, "Receipt"))
            a._recompute_level(self.db, iid, wh)
        dedupe.merge(self.db, a1, [a2], "t")
        lvl = self.db.one("SELECT qty, avg_cost FROM stock_levels WHERE item_id=? AND warehouse_id=?", (a1, wh))
        self.assertEqual((lvl["qty"], lvl["avg_cost"]), (10, 16000))
        self.assertIsNone(self.db.one("SELECT 1 FROM stock_items WHERE id=?", (a2,)))
        self.assertEqual(a._item(self.db, {"item": "حنفيه  غساله"})["id"], a1)             # the merged name lands on it
        self.assertEqual(a._item(self.db, {"item": "حنفيّة غسّالة"}, create=True)["id"], a1)  # so does a new spelling
        self.assertEqual(a._item(self.db, {"item": "مواسير 40مم / 20بار"}, create=True)["name"], "مواسير 40مم / 20بار")
        dedupe.ignore(self.db, [c1, c2])
        self.assertNotIn((c1, c2), {tuple(g["ids"]) for g in dedupe.find(self.db)})
        with self.assertRaises(a.ValidationError):
            dedupe.check_sku(self.db, "CW 10")                                              # a code belongs to one item

    def test_stocktake_posts_only_the_differences(self):
        from erp import stocktake
        cid = stocktake.start(self.db, {"warehouse": "Site Store"})
        stocktake.save_counts(self.db, cid, {"lines": [{"item": "Tiles", "counted": 18}]})
        self.assertEqual(stocktake.post(self.db, cid), 1)
        lvl = self.db.one("SELECT qty FROM stock_levels l JOIN warehouses w ON w.id=l.warehouse_id WHERE w.name='Site Store'")
        self.assertEqual(lvl["qty"], 18)
        with self.assertRaises(a.ValidationError):
            stocktake.post(self.db, cid)

    def test_equipment_hours_charge_projects_and_maintenance_falls_due(self):
        from erp import equipment
        e = equipment.save(self.db, {"name": "Mixer", "hourly_rate": 100, "meter_start": 50})
        equipment.save_plan(self.db, {"equipment_id": e, "title": "Oil", "every_hours": 20})
        equipment.log_usage(self.db, {"equipment_id": e, "bu_code": "PRJ-1", "hours": 19, "fuel_cost": 300, "date": "2026-03-05"})
        self.assertEqual(equipment.alerts(self.db)["maintenance_due"][0]["state"], "due")
        equipment.log_usage(self.db, {"equipment_id": e, "bu_code": "PRJ-1", "hours": 2, "date": "2026-03-06"})
        self.assertEqual(equipment.alerts(self.db)["overdue"], 1)
        equipment.log_maintenance(self.db, {"equipment_id": e, "plan_id": 1, "cost": 500, "date": "2026-03-07"})
        self.assertEqual(equipment.alerts(self.db)["maintenance_due"], [])
        rep = equipment.report(self.db)
        self.assertEqual((rep["by_project"][0]["charge"], rep["pool_balance"]), (2100, -1600))

    def test_attendance_fills_the_payroll_run(self):
        from erp import labour, payroll
        w = payroll.save_employee(self.db, {"name": "Ali", "emp_type": "daily", "daily_rate": 400, "bu_code": "PRJ-1"})
        m = payroll.save_employee(self.db, {"name": "Omar", "emp_type": "monthly", "basic": 9000})
        labour.save_day(self.db, {"date": "2026-03-02", "rows": [{"employee_id": w, "status": "present", "overtime": 2, "qty": 20, "unit": "m2"},
                                                                 {"employee_id": m, "status": "absent"}]})
        labour.save_day(self.db, {"date": "2026-03-03", "rows": [{"employee_id": w, "status": "half"}, {"employee_id": m, "status": "present"}]})
        rid = payroll.create_run(self.db, {"period": "2026-03"})
        labour.fill_payroll(self.db, rid)
        got = {ln["employee_id"]: (ln["units"], ln["absent_days"], ln["ot_day"]) for ln in payroll.get_run(self.db, rid)["lines"]}
        self.assertEqual(got[w], (1.5, 0, 2))
        self.assertEqual(got[m], (0, 1, 0))

    def test_e_receipts_chain_and_post(self):
        from erp import ereceipt
        r1 = ereceipt.save(self.db, {"date": "2026-03-05", "time": "10:00",
                                     "lines": [{"description": "Paint", "qty": 2, "price": 250, "code": "EG-1", "vat_code": "VAT14"}]})
        u1 = ereceipt.issue(self.db, r1)["uuid"]
        r2 = ereceipt.save(self.db, {"date": "2026-03-06", "time": "10:00", "kind": "return", "ref_number": ereceipt.get(self.db, r1)["number"],
                                     "lines": [{"description": "Paint", "qty": 1, "price": 250, "code": "EG-1", "vat_code": "VAT14"}]})
        ereceipt.issue(self.db, r2)
        self.assertEqual(ereceipt.get(self.db, r2)["previous_uuid"], u1)
        self.assertTrue(ereceipt.verify_chain(self.db)["ok"])
        with self.assertRaises(a.ValidationError):
            ereceipt.delete(self.db, r1)
        self.db.x("UPDATE ereceipts SET document=replace(document,'250','260') WHERE id=?", (r1,))
        self.assertFalse(ereceipt.verify_chain(self.db)["ok"])   # tampering breaks the chain

    def test_demo_company_and_backup_round_trip(self):
        from erp import demo
        db = Database(":memory:")
        demo.load(db, "admin")
        db2 = Database(":memory:")
        interchange.import_data(db2, interchange.read_json(interchange.to_json(db).encode()), mode="replace", user="admin")
        for t in ("contracts", "variation_orders", "sub_certificates", "purchase_requests", "pr_quotes", "stock_counts", "equipment",
                  "equipment_usage", "maintenance_logs", "attendance", "ereceipts", "project_phases", "project_tasks"):
            n = db.one(f"SELECT COUNT(*) n FROM {t}")["n"]
            self.assertTrue(n, t)
            self.assertEqual(n, db2.one(f"SELECT COUNT(*) n FROM {t}")["n"], t)
        self.assertEqual(db2.one("SELECT COUNT(*) n FROM sub_certificates WHERE status='posted' AND entry_id IS NULL")["n"], 0)
        tb = lambda d: {r["code"]: (r.get("debit"), r.get("credit")) for r in a.trial_balance(d)["rows"]}  # noqa: E731
        self.assertEqual(tb(db), tb(db2))

    def test_professional_account_titles_only_replace_defaults(self):
        db = Database(":memory:")
        self.assertEqual(db.one("SELECT name FROM accounts WHERE code='1100'")["name"], "Trade Receivables — Customers")
        db.x("UPDATE accounts SET name='Bank Account' WHERE code='1020'")
        db.x("UPDATE accounts SET name='Our clients' WHERE code='1100'")
        db._migrate()
        self.assertEqual(db.one("SELECT name FROM accounts WHERE code='1020'")["name"], "Cash at Bank")
        self.assertEqual(db.one("SELECT name FROM accounts WHERE code='1100'")["name"], "Our clients")


class BulkActions(unittest.TestCase):
    """Select many records and edit / void / delete them together — each checked as if done by hand."""
    setUpClass, tearDownClass = classmethod(Api.setUpClass.__func__), classmethod(Api.tearDownClass.__func__)
    call, client = Api.call, Api.client

    def test_bulk_entries_and_lists(self):
        op = self.client()
        if self.call(op, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0] != 200:
            self.assertEqual(self.call(op, "POST", "/api/login", {"username": "admin", "password": "secret123"})[0], 200)
        b1 = self.call(op, "POST", "/api/bus", {"code": "BK1", "name": "Bulk one"})[1]["id"]
        b2 = self.call(op, "POST", "/api/bus", {"code": "BK2", "name": "Bulk two"})[1]["id"]
        ids = [self.call(op, "POST", "/api/docs/expense", {"description": f"Fuel {i}", "amount": 100 + i, "account": "6090",
                                                           "date": "2026-03-01"})[1]["id"] for i in range(3)]
        # edit: only the fields given change; the entries keep their history
        st, r = self.call(op, "POST", "/api/bulk", {"target": "entries", "action": "edit", "ids": ids, "changes": {"bu_id": b2, "ref": "BULK-1"}})
        self.assertEqual((st, r["done"], r["failed"]), (200, 3, []))
        e = self.call(op, "GET", f"/api/journal/{ids[0]}")[1]
        self.assertEqual((e["bu_id"], e["ref"], e["memo"], len(e["history"])), (b2, "BULK-1", "Fuel 0", 1))
        # delete only works on void entries: reported per record, the rest still done
        st, r = self.call(op, "POST", "/api/bulk", {"target": "entries", "action": "delete", "ids": ids[:2]})
        self.assertEqual((r["done"], len(r["failed"])), (0, 2))
        self.assertIn("Void an entry before deleting it", r["failed"][0]["error"])
        self.assertEqual(self.call(op, "POST", "/api/bulk", {"target": "entries", "action": "void", "ids": ids})[0], 400)  # needs a reason
        r = self.call(op, "POST", "/api/bulk", {"target": "entries", "action": "void", "ids": ids[:2], "reason": "duplicates"})[1]
        self.assertEqual(r["done"], 2)
        r = self.call(op, "POST", "/api/bulk", {"target": "entries", "action": "delete", "ids": ids})[1]
        self.assertEqual((r["done"], [f["id"] for f in r["failed"]]), (2, [str(ids[2])]))
        self.assertEqual(self.call(op, "GET", f"/api/journal/{ids[0]}")[0], 404)
        # business units: delete refuses one with transactions, status edit validates the value
        r = self.call(op, "POST", "/api/bulk", {"target": "bus", "action": "delete", "ids": [b1, b2]})[1]
        self.assertEqual((r["done"], r["failed"][0]["id"]), (1, str(b2)))
        self.assertEqual(self.call(op, "POST", "/api/bulk", {"target": "bus", "action": "edit", "ids": [b2], "changes": {"status": "Bogus"}})[1]["failed"][0]["error"][:10], "Choose one")
        self.assertEqual(self.call(op, "POST", "/api/bulk", {"target": "bus", "action": "edit", "ids": [b2], "changes": {"status": "On Hold"}})[1]["done"], 1)
        self.assertEqual(self.db.one("SELECT status FROM business_units WHERE id=?", (b2,))["status"], "On Hold")
        # someone without the voucher type's edit right cannot change it in bulk either
        self.call(op, "POST", "/api/users", {"username": "viewer", "password": "viewer123", "perms": {"purchasing": "view", "accounting": "view"}})
        v = self.client()
        self.assertEqual(self.call(v, "POST", "/api/login", {"username": "viewer", "password": "viewer123"})[0], 200)
        r = self.call(v, "POST", "/api/bulk", {"target": "entries", "action": "void", "ids": [ids[2]], "reason": "x"})[1]
        self.assertEqual((r["done"], len(r["failed"])), (0, 1))
        self.assertEqual(self.db.one("SELECT status FROM journal_entries WHERE id=?", (ids[2],))["status"], "posted")
        self.assertEqual(self.call(op, "POST", "/api/bulk", {"target": "nope", "action": "edit", "ids": [1]})[0], 400)


class BrandDefaults(unittest.TestCase):
    """The app is BeyondERP with a logo-only sidebar for old books too; a company's own logo is kept."""

    def _book(self, ap):
        path = Path(tempfile.mkdtemp()) / "b.db"
        d = Database(str(path))
        d.x("UPDATE settings SET value=? WHERE key='appearance'", (json.dumps(ap),))
        d.x("DELETE FROM settings WHERE key='brand_v2'")
        d.close()
        d = Database(str(path))
        return json.loads(d.one("SELECT value FROM settings WHERE key='appearance'")["value"])

    def test_old_books_move_to_the_new_defaults(self):
        ap = self._book({"app_name": "Beyond Design", "brand_show": "logo_name", "logo": "", "accent": "#123456"})
        self.assertEqual((ap["app_name"], ap["brand_show"], ap["logo"], ap["accent"]), ("", "logo", "", "#123456"))
        ap = self._book({"app_name": "Acme Books", "brand_show": "logo_name", "logo": "custom"})
        self.assertEqual((ap["app_name"], ap["brand_show"], ap["logo"]), ("Acme Books", "logo_name", "custom"))

    def test_runs_once(self):
        path = Path(tempfile.mkdtemp()) / "b.db"
        d = Database(str(path))
        d.x("UPDATE settings SET value=? WHERE key='appearance'", (json.dumps({"app_name": "Beyond Design"}),))
        d.close()
        d = Database(str(path))   # flag already set on a new book: a name chosen later is never touched
        self.assertEqual(json.loads(d.one("SELECT value FROM settings WHERE key='appearance'")["value"])["app_name"], "Beyond Design")


class ImportOnlyNew(unittest.TestCase):
    """'New rows only': a sheet exported earlier, edited in the app since, plus new rows — only the new rows come in."""

    def test_only_new_rows(self):
        from erp import demo
        db = Database(":memory:")
        demo.load(db)
        tables = ("journal_entries", "journal_lines", "parties", "stock_moves", "extracts", "employees", "attendance", "budgets")
        count = lambda: {t: db.one(f"SELECT COUNT(*) n FROM {t}")["n"] for t in tables}  # noqa: E731
        data = interchange.export_data(db)
        before = count()
        # the whole book again: nothing new
        res = interchange.import_data(db, data, mode="new")
        self.assertEqual(count(), before)
        self.assertFalse(any(res["counts"].values()))
        self.assertTrue(res["skipped"])
        # a party edited in the app since the export keeps the app's version; a new voucher and party come in
        first = data["Parties"][0]["name"]
        db.x("UPDATE parties SET tax_id='999' WHERE name=?", (first,))
        sheet = {"Journal Vouchers": data["Journal Vouchers"] + [
                     {"voucher": "NEW-1", "date": "2026-03-03", "kind": "Journal Entry", "memo": "brand new", "account": "6010", "debit": 123.45},
                     {"voucher": "NEW-1", "date": "2026-03-03", "kind": "Journal Entry", "memo": "brand new", "account": "1010", "credit": 123.45}],
                 "Parties": data["Parties"] + [{"name": "Brand New Client", "type": "customer"}]}
        res = interchange.import_data(db, sheet, mode="new")
        after = count()
        self.assertEqual(after["journal_entries"], before["journal_entries"] + 1)
        self.assertEqual(after["parties"], before["parties"] + 1)
        self.assertEqual(db.one("SELECT tax_id FROM parties WHERE name=?", (first,))["tax_id"], "999")
        # and the same file a second time adds nothing
        interchange.import_data(db, sheet, mode="new")
        self.assertEqual(count(), after)


class ImportUpdatesDocuments(unittest.TestCase):
    """BOQs and invoices are whole documents: re-importing one edited outside updates it in place (both 'add' and
    'new rows only'); an unchanged one is skipped; nothing is duplicated."""

    def test_changed_boq_and_invoice(self):
        from erp import demo
        db = Database(":memory:")
        demo.load(db)
        data = interchange.export_data(db)
        self.assertTrue(data.get("BOQ Lines") and data.get("Sales Invoices") and data.get("Invoice Lines"))
        counts = lambda: (db.one("SELECT COUNT(*) n FROM boqs")["n"], db.one("SELECT COUNT(*) n FROM extracts")["n"])  # noqa: E731
        before = counts()
        for mode in ("append", "new"):          # unchanged: nothing new, nothing duplicated
            res = interchange.import_data(db, {k: data[k] for k in ("BOQ Lines", "Sales Invoices", "Invoice Lines")}, mode=mode)
            self.assertEqual(counts(), before, mode)
            self.assertFalse(res.get("updated"), mode)
            self.assertTrue(res["skipped"].get("BOQ Lines") and res["skipped"].get("Sales Invoices"), mode)
        # one BOQ line's quantity and one invoice line's price change outside the app
        boq_rows = [dict(r) for r in data["BOQ Lines"]]
        title = boq_rows[0]["title"]
        boq_rows[0]["qty"] = float(boq_rows[0]["qty"] or 0) + 7
        posted = [r for r in data["Sales Invoices"] if str(r.get("status") or "Posted").lower() != "draft"][0]
        inv_lines = [dict(r) for r in data["Invoice Lines"]]
        ln = [r for r in inv_lines if r["number"] == posted["number"]][0]
        ln["price"] = float(ln["price"] or 0) + 100
        x0 = db.one("SELECT * FROM extracts WHERE number=?", (posted["number"],))
        old_total = a.get_extract(db, x0["id"])["total"]
        res = interchange.import_data(db, {"BOQ Lines": boq_rows, "Invoice Lines": inv_lines}, mode="new")
        self.assertEqual(counts(), before)
        self.assertEqual(res["updated"], {"BOQ Lines": 1, "Invoice Lines": 1})
        b = db.one("SELECT id FROM boqs WHERE title=? ORDER BY id DESC", (title,))
        self.assertIn(boq_rows[0]["qty"], [r["qty"] for r in db.q("SELECT qty FROM boq_lines WHERE boq_id=?", (b["id"],))])
        x = a.get_extract(db, x0["id"])
        self.assertEqual(x["status"], "posted")
        self.assertAlmostEqual(x["total"] - old_total, 100 * float(ln.get("qty") or 1) * float(ln.get("pct") or 1), places=1)
        self.assertEqual(db.one("SELECT status FROM journal_entries WHERE id=?", (x0["entry_id"],))["status"], "void")   # old entry kept, voided
        self.assertNotEqual(x["entry_id"], x0["entry_id"])
        # and importing that same file again changes nothing
        res = interchange.import_data(db, {"BOQ Lines": boq_rows, "Invoice Lines": inv_lines}, mode="append")
        self.assertFalse(res.get("updated"))


class LinesFromExcel(unittest.TestCase):
    """A BOQ / invoice edited outside the app comes back line for line, with each line's look; a plain sheet too."""

    def test_round_trip_and_plain_sheet(self):
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from erp import demo, documents, linesheet
        db = Database(":memory:")
        demo.load(db)
        bid = db.one("SELECT id FROM boqs ORDER BY id")["id"]
        B = a.get_boq(db, bid)
        lines = [dict(x) for x in B["lines"]]
        lines[1]["style"] = {"b": 1, "bg": "#FFFF00", "color": "#c00000", "align": "center", "junk": 1}
        a.save_boq(db, dict(B, lines=lines), boq_id=bid)
        B = a.get_boq(db, bid)
        self.assertEqual(B["lines"][1]["style"], {"b": 1, "bg": "#ffff00", "color": "#c00000", "align": "center"})
        data, _ = documents.boq_file(db, bid, "xlsx")
        got = linesheet.parse(data)["lines"]
        self.assertEqual([(x["category"], x["description"], x["qty"], x["price"]) for x in got],
                         [(x["category"], x["description"], x["qty"], x["price"]) for x in B["lines"]])
        self.assertEqual(got[1]["style"], {"b": 1, "bg": "#ffff00", "color": "#c00000", "align": "center"})
        self.assertEqual(got[0]["style"], {})
        # someone's own sheet: Arabic headers, a section bar, a total row, amounts with thousands separators
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["مقايسة"])
        ws.append(["م", "البيان", "الوحدة", "الكمية", "السعر", "الإجمالي"])
        ws.append(["1- أعمال الهدم"])
        ws.append([1, "هدم حوائط", "م2", "1,200", 50, 60000])
        ws.append([2, "نقل مخلفات", "نقلة", 10, None, 3000])
        ws["B5"].font = Font(bold=True)
        ws["B5"].fill = PatternFill("solid", fgColor="FFE2EFDA")
        ws.append(["إجمالي أعمال الهدم", None, None, None, None, 63000])
        buf = io.BytesIO()
        wb.save(buf)
        got = linesheet.parse(buf.getvalue())["lines"]
        self.assertEqual([(x["item_no"], x["category"], x["qty"], x["price"]) for x in got],
                         [("1", "أعمال الهدم", 1200, 50), ("2", "أعمال الهدم", 10, 300)])
        self.assertEqual(got[1]["style"], {"b": 1, "bg": "#e2efda"})
        with self.assertRaises(a.ValidationError):
            linesheet.parse(b"not a workbook")


class DashboardDesign(unittest.TestCase):
    """Cards and pages keep only what the designer allows (no unsafe links, no stray keys); building blocks may repeat;
    designing is for super administrators and whoever they allow."""

    def test_clean_and_permission(self):
        from erp import dashboards as dsh
        db = Database(":memory:")
        dsh.seed(db)
        lay = [{"w": "text", "look": {"title": "Hi", "body": "**b**", "style": "solid", "color": "#ABCDEF", "evil": 1}},
               {"w": "text", "look": {"style": "nope"}},
               {"w": "links", "look": {"links": [{"label": "Ok", "href": "#/invoices"}, {"label": "Bad", "href": "javascript:alert(1)"},
                                                 {"label": "Web", "href": "https://example.com/x"}]}},
               {"w": "monthly_chart", "look": {"chart": "area", "height": "tall"}}]
        dsh.save_group(db, "finance", {"layout": lay, "look": {"heading": "Hello {name}", "accent": "#123456", "gap": "airy", "x": 1}})
        g = db.one("SELECT layout, look FROM user_groups WHERE key='finance'")
        L = json.loads(g["layout"])
        self.assertEqual(L[0]["look"], {"title": "Hi", "body": "**b**", "style": "solid", "color": "#abcdef"})
        self.assertNotIn("look", L[1])
        self.assertEqual([x["href"] for x in L[2]["look"]["links"]], ["#/invoices", "https://example.com/x"])
        self.assertEqual(L[3]["look"], {"chart": "area", "height": "tall"})
        self.assertEqual(json.loads(g["look"]), {"heading": "Hello {name}", "accent": "#123456", "gap": "airy"})
        self.assertTrue(dsh.designer({"is_super": True}, lambda *a: False))
        self.assertTrue(dsh.designer({}, lambda m, lv="view": m == "settings.dashboards" and lv == "edit"))
        self.assertFalse(dsh.designer({"is_admin": True}, lambda *a: False))
        self.assertFalse(dsh.designer({"is_admin": True}, lambda *a: True))     # an administrator: only when super
        key = dsh.save_group(db, dsh.new_key(db, "Site engineers"), {"name": "Site engineers", "copy_from": "finance"})
        self.assertEqual(json.loads(db.one("SELECT look FROM user_groups WHERE key=?", (key,))["look"])["accent"], "#123456")


class FilesInBackups(unittest.TestCase):
    """The Excel export 'with files' is a zip: the workbook plus a folder per month of the attached files; importing that
    zip into empty books puts every file back on its entry."""

    def test_zip_round_trip(self):
        import base64
        import io
        import zipfile
        from erp import attachments, demo
        db = Database(":memory:")
        demo.load(db)
        e = db.one("SELECT id, date FROM journal_entries WHERE status='posted' AND date LIKE '2026-03%' ORDER BY id") or \
            db.one("SELECT id, date FROM journal_entries WHERE status='posted' ORDER BY id")
        pdf = b"%PDF-1.4 a supplier invoice"
        attachments.add(db, "entry", e["id"], {"title": "bill.pdf", "mime": "application/pdf", "data": base64.b64encode(pdf).decode()}, "boss")
        attachments.add(db, "entry", e["id"], {"title": "bill.pdf", "mime": "application/pdf", "data": base64.b64encode(b"second").decode()}, "boss")
        attachments.add(db, "entry", e["id"], {"title": "Drive folder", "url": "https://example.com/f"}, "boss")
        raw = interchange.workbook_zip(db)
        self.assertTrue(interchange.is_files_zip(raw))
        z = zipfile.ZipFile(io.BytesIO(raw))
        names = z.namelist()
        y, m = int(e["date"][:4]), int(e["date"][5:7])
        folder = f"Attachments/{y}-{m:02d} {interchange.MONTHS_EN[m - 1]} {y}/"
        files = [n for n in names if n.startswith(folder)]
        self.assertEqual(len(files), 2)                       # two files with the same name both kept
        self.assertIn(pdf, [z.read(n) for n in files])
        self.assertIn("Attachments/links.csv", names)
        self.assertTrue(any(n.endswith(".xlsx") for n in names))
        # the workbook inside no longer carries the files itself (they are in the folders)
        db2 = Database(":memory:")
        data = interchange.read_files_zip(raw, interchange.read_xlsx)
        self.assertEqual(len(data["_extras"]["attachments"]), 3)
        interchange.import_data(db2, data, mode="replace")
        got = db2.q("SELECT title, size, url FROM attachments WHERE owner='entry' ORDER BY id")
        self.assertEqual(sorted((g["title"], g["size"]) for g in got), [("Drive folder", 0), ("bill.pdf", 6), ("bill.pdf", len(pdf))])
        interchange.import_data(db2, data, mode="new")       # the same zip again: no file twice
        self.assertEqual(db2.one("SELECT COUNT(*) n FROM attachments")["n"], 3)


class Crm(unittest.TestCase):
    """Deals on a board, a contact's timeline (notes + what the books hold), follow-ups in the bell, and an offer turned
    into a project and its BOQ in one step."""

    def test_crm(self):
        from erp import crm, demo, documents
        db = Database(":memory:")
        demo.load(db)
        did = crm.save_deal(db, {"title": "Villa", "party": "Brand New Client", "value": 500000, "stage": "proposal"}, "boss")
        self.assertTrue(db.one("SELECT 1 FROM parties WHERE name='Brand New Client'"))        # a new contact is made
        b = crm.board(db, {})
        col = {c["key"]: c for c in b["columns"]}
        self.assertEqual((col["proposal"]["count"], col["proposal"]["value"], b["weighted"]), (1, 500000, 250000))
        crm.move_deal(db, did, "lost", "boss", "price")
        self.assertEqual(db.one("SELECT lost_reason FROM crm_deals WHERE id=?", (did,))["lost_reason"], "price")
        crm.add_activity(db, {"party": "Brand New Client", "kind": "call", "subject": "Asked for a discount"}, "boss")
        fid = crm.save_followup(db, {"title": "Send revised offer", "due": "2020-01-01", "party": "Brand New Client"}, "sara")
        self.assertEqual([f["id"] for f in crm.due_for(db, "sara")], [fid])
        self.assertEqual(crm.due_for(db, "boss"), [])
        tl = crm.timeline(db, "Brand New Client", lambda *a: True)
        kinds = [e["kind"] for e in tl["events"]]
        self.assertIn("call", kinds)
        self.assertIn("followup", kinds)
        crm.done_followup(db, fid, "sara", True, "sent")
        self.assertEqual(crm.due_for(db, "sara"), [])
        # a customer with invoices: they show on the timeline
        cust = db.one("SELECT customer FROM extracts WHERE customer<>'' LIMIT 1")
        if cust:
            self.assertIn("invoice", [e["kind"] for e in crm.timeline(db, cust["customer"], lambda *a: True)["events"]])
        # offer -> project + BOQ (once), linked deal won
        oid = documents.save_offer(db, {"number": "Q-77", "client": "Brand New Client", "project": "Villa", "lang": "en",
                                        "sections": [{"kind": "pricing", "items": [{"desc": "Design", "amount": 100000}, {"desc": "Supervision", "amount": 50000}],
                                                      "discount": 10000, "vat": 0}]})
        d2 = crm.save_deal(db, {"title": "Villa 2", "party": "Brand New Client", "value": 140000, "offer_id": oid}, "boss")
        made = []
        r = crm.convert_offer(db, oid, "boss", lambda o: made.append(o) or db.x("INSERT INTO business_units(code,name) VALUES ('PRJ-T','Villa')").lastrowid)
        boq = a.get_boq(db, r["boq_id"])
        self.assertEqual(round(boq["total"], 2), 140000)
        self.assertEqual(db.one("SELECT stage FROM crm_deals WHERE id=?", (d2,))["stage"], "won")
        self.assertEqual(crm.convert_offer(db, oid, "boss", lambda o: 1 / 0)["boq_id"], r["boq_id"])   # a second time: the same BOQ
        self.assertEqual(len(made), 1)


class SpecialRights(unittest.TestCase):
    """Super administrators hold every special right and give them to others (administrators included); the in-app
    BeyondERP name, logo and icon are fixed for everyone but the platform owner."""

    def test_grants_and_locked_brand(self):
        from erp import auth, dashboards
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "Passw0rd!long1"}, first=True)
        auth.create_user(db, {"username": "adm", "password": "Passw0rd!long2", "is_admin": True})
        u = lambda n: auth.public_user(db.one("SELECT * FROM users WHERE username=?", (n,)))  # noqa: E731
        self.assertEqual(set(u("boss")["grants"]), set(auth.GRANTS))
        self.assertEqual(u("adm")["grants"], [])
        self.assertFalse(dashboards.designer(u("adm"), lambda *a: True))
        auth.update_user(db, u("adm")["id"], {"grants": ["design_dashboards", "nonsense"]}, u("boss")["id"])
        self.assertEqual(u("adm")["grants"], ["design_dashboards"])
        self.assertTrue(dashboards.designer(u("adm"), lambda *a: False))
        self.assertTrue(auth.granted(u("boss"), "build_calculators"))
        self.assertFalse(auth.granted(u("adm"), "build_calculators"))
        self.assertTrue(auth.platform_owner({"username": "MohamedSakr0303@gmail.com "}))
        self.assertFalse(auth.platform_owner({"username": "boss", "email": "someone@example.com"}))
        # earlier custom app names / logos are cleared once
        db2 = Database(":memory:")
        db2.x("UPDATE settings SET value=? WHERE key='appearance'", (json.dumps({"app_name": "My Co", "logo": "custom", "accent": "#123456"}),))
        db2.x("DELETE FROM settings WHERE key='brand_v3'")
        db2._brand_v3()
        ap = json.loads(db2.one("SELECT value FROM settings WHERE key='appearance'")["value"])
        self.assertEqual(ap, {"accent": "#123456"})


class BackupSecretsAccess(unittest.TestCase):
    """Keys of connected tools travel only in a super administrator's backups and imports; the database file is theirs only."""

    def setUp(self):
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = seeded()
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "admin2", "password": "secret123", "is_admin": True})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "admin2")}
        auth.create_api_key(self.db, self.db.one("SELECT id FROM users WHERE username='boss'")["id"], "BI", True, "boss")

    def call(self, who, method, url, raw=b"", ctype="application/json"):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": ctype}
        return self.handle(self.app, method, url, h, raw)

    def test_secrets_only_for_super_administrators(self):
        boss = json.loads(self.call("boss", "GET", "/api/export/json").body)
        self.assertEqual(len(boss["extras"]["secrets"]["api_keys"]), 1)
        self.assertNotIn("secrets", json.loads(self.call("boss", "GET", "/api/export/json?secrets=0").body)["extras"])
        self.assertNotIn("secrets", json.loads(self.call("admin2", "GET", "/api/export/json").body)["extras"])
        self.assertEqual(self.call("admin2", "GET", "/api/backup").status, 403)
        # an administrator's import cannot plant a key for someone else
        boss["extras"]["secrets"] = {"api_keys": [{"user": "boss", "name": "planted", "key_hash": "0" * 64, "prefix": "berp_x", "read_only": 0}]}
        r = self.call("admin2", "POST", "/api/import?format=json&mode=append", json.dumps(boss).encode(), "application/octet-stream")
        self.assertEqual(r.status, 200, r.body[:300])
        self.assertIsNone(self.db.one("SELECT 1 FROM api_keys WHERE name='planted'"))
        self.assertEqual(self.call("admin2", "POST", "/api/import?format=sqlite&mode=append", b"SQLite format 3\x00" + b"\0" * 100,
                                   "application/octet-stream").status, 403)


class BillsAndAlerts(unittest.TestCase):
    """Company subscriptions: paying books expenses by their allocation and moves the due date; reminders carry their
    own time. Alerts: a person's own "notify" rules, checked after their time of day, land in their bell."""

    def setUp(self):
        from erp import demo
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = Database(":memory:")
        demo.load(self.db, "boss")
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "eng", "password": "secret123", "perms": {"team": "add", "projects": "view"}})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "eng")}

    def call(self, who, method, url, body=None):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json"}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        return r.status, json.loads(r.body)

    def test_subscription_pay_split_and_company_cost(self):
        from erp import billsubs
        st, r = self.call("boss", "POST", "/api/bills", {"name": "Software subscription", "vendor": "Software supplier", "amount": 1000, "next_due": "2026-01-31",
                                                         "alloc": "split", "split": [{"bu_code": "PRJ-D01", "pct": 60}, {"bu_code": "PRJ-D02", "pct": 40}],
                                                         "remind_days": 2, "remind_time": "10:30"})
        self.assertEqual(st, 200, r)
        hint = r["_push"][0]
        self.assertEqual((hint["day"] >= "2026-01-29", hint["time"], hint["to"], hint["pref"]), (True, "10:30", ["boss"], "reminders"))
        st, p = self.call("boss", "POST", f"/api/bills/{r['id']}/pay", {"date": "2026-01-31"})
        self.assertEqual((st, p["next_due"]), (200, "2026-02-28"))
        shares = sorted(x["amount"] for x in billsubs.get(self.db, r["id"])["payments"])
        self.assertEqual(shares, [400.0, 600.0])
        self.assertEqual(billsubs.next_after("2026-02-28", "monthly", 31), "2026-03-31")   # back to the 31st
        pool = a.shared_overhead_total(self.db)
        st, c = self.call("boss", "POST", "/api/bills", {"name": "Online service", "amount": 10000, "next_due": "2026-02-10", "alloc": "company"})
        self.call("boss", "POST", f"/api/bills/{c['id']}/pay", {"date": "2026-02-10"})
        self.assertEqual(a.shared_overhead_total(self.db), pool)            # a company cost stays out of the projects
        self.assertEqual(self.call("eng", "GET", "/api/bills")[0], 403)      # not his screen

    def test_alerts_after_time_feed_and_permissions(self):
        st, r = self.call("boss", "POST", "/api/my-alerts", {"name": "Attendance", "trigger": "attendance_missing", "settings": {"at": "09:30"}})
        self.assertEqual(st, 200, r)
        day = a.today()
        st, early = self.call("boss", "POST", "/api/automations/run", {"local": day + "T08:00"})
        self.assertEqual(early["done"], [])
        st, late = self.call("boss", "POST", "/api/automations/run", {"local": day + "T09:45"})
        self.assertEqual(late["done"][0]["acted"], 1)
        self.assertEqual(late["_push"][0]["to"], ["boss"])
        st, bell = self.call("boss", "GET", "/api/alerts")
        fid = bell["feed"][0]
        self.call("boss", "POST", "/api/my-alerts/read", {"ids": [fid]})
        self.assertNotIn(fid, self.call("boss", "GET", "/api/alerts")[1]["feed"])
        # someone without HR access can't be alerted about attendance, and can only alert himself
        self.assertEqual(self.call("eng", "POST", "/api/my-alerts", {"name": "x", "trigger": "attendance_missing"})[0], 403)
        st, mine = self.call("eng", "GET", "/api/my-alerts")
        self.assertNotIn("attendance_missing", [t["key"] for t in mine["catalog"]["triggers"]])


class CalculatorBuilder(unittest.TestCase):
    """Rental and general calculators: charges (installation, transport…) and results are formulas; only super
    administrators and the people they allow build or change formulas; anyone may open a general calculator."""

    def setUp(self):
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = Database(":memory:")
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "rita", "password": "secret123", "perms": {"rentals": "edit"}})
        auth.create_user(self.db, {"username": "sam", "password": "secret123", "perms": {"sales": "view"}})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "rita", "sam")}

    def call(self, who, method, url, body=None):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json"}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        return r.status, json.loads(r.body)

    def test_build_and_run(self):
        from erp import rentals
        spec = {"inputs": [{"key": "length", "value": 10}, {"key": "height", "value": 5}], "params": [{"key": "rate", "value": 4}],
                "area": "length * height", "derived": [], "components": [],
                "charges": [{"label": "Installation", "kind": "install", "qty": "area", "unit": "m²", "cost": "2", "price": "3"},
                            {"label": "Transport", "kind": "transport", "qty": "2", "unit": "trip", "cost": "100", "price": "150"}],
                "results": [{"label": "Rent", "expr": "area * rate", "money": "yes"}, {"label": "Total", "expr": "area * rate + charges_price", "money": "no"}]}
        r = rentals.evaluate(self.db, spec, {"length": 20})
        self.assertEqual(r["area"], 100)
        self.assertEqual([c["price_total"] for c in r["charges"]], [300, 300])
        self.assertEqual((r["charges_cost"], r["charges_price"]), (400, 600))
        self.assertEqual([(x["value"], x["money"]) for x in r["results"]], [(400, True), (1000, False)])
        # building is a special right: rita may use and tune the calculator but not add one
        gen = {"name": "Margin", "category": "general", "spec": {"inputs": [{"key": "q", "value": 2}], "params": [], "derived": [], "components": [],
               "area": "0", "results": [{"label": "Twice", "expr": "q * 2", "money": False}]}}
        self.assertEqual(self.call("rita", "POST", "/api/rent-models", gen)[0], 403)
        st, made = self.call("boss", "POST", "/api/rent-models", gen)
        self.assertEqual(st, 200)
        mid = made["id"]
        self.assertTrue(self.call("boss", "GET", "/api/rent-models")[1]["can_build"])
        self.assertFalse(self.call("rita", "GET", "/api/rent-models")[1]["can_build"])
        # new defaults are fine, a changed formula is not
        tuned = json.loads(json.dumps(gen)); tuned["spec"]["inputs"][0]["value"] = 5
        self.assertEqual(self.call("rita", "PUT", f"/api/rent-models/{mid}", tuned)[0], 200)
        tuned["spec"]["results"][0]["expr"] = "q * 3"
        self.assertEqual(self.call("rita", "PUT", f"/api/rent-models/{mid}", tuned)[0], 403)
        self.assertEqual(self.call("rita", "DELETE", f"/api/rent-models/{mid}")[0], 403)
        # once allowed, rita builds too
        auth.update_user(self.db, self.db.one("SELECT id FROM users WHERE username='rita'")["id"], {"grants": ["build_calculators"]},
                         self.db.one("SELECT id FROM users WHERE username='boss'")["id"])
        self.assertEqual(self.call("rita", "PUT", f"/api/rent-models/{mid}", tuned)[0], 200)
        # someone outside rentals sees only general calculators
        st, v = self.call("sam", "GET", "/api/rent-models")
        self.assertEqual((st, [m["category"] for m in v["models"]]), (200, ["general"]))
        st, r = self.call("sam", "POST", "/api/rent-models-preview/eval", {"spec": v["models"][0]["spec"], "values": {"q": 4}, "general": True})
        self.assertEqual((st, r["results"][0]["value"]), (200, 12))


class DesktopWindowMarker(unittest.TestCase):
    """The desktop window tells the app it is the desktop app (its own list of device permissions), keeping the page."""

    def test_marker(self):
        from unittest import mock
        from erp import appwindow
        seen = []
        with mock.patch.object(appwindow, "candidates", return_value=[]), mock.patch("webbrowser.open", side_effect=seen.append):
            for url in ("https://x.dev/", "https://x.dev/?a=1", "https://x.dev/#/chat/2", "https://x.dev/?app=desktop"):
                appwindow.open_window(url)
        self.assertEqual(seen, ["https://x.dev/?app=desktop", "https://x.dev/?a=1&app=desktop",
                                "https://x.dev/?app=desktop#/chat/2", "https://x.dev/?app=desktop"])


class ChatReactionsAndNotify(unittest.TestCase):
    """WhatsApp-style reactions (one per person, same again takes it back) and each person's notification settings."""

    def setUp(self):
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = Database(":memory:")
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "sara", "password": "secret123"})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "sara")}

    def call(self, who, method, url, body=None):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json"}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        return r.status, json.loads(r.body)

    def test_reactions(self):
        rid = 1   # the general room
        mid = self.call("boss", "POST", f"/api/chat/rooms/{rid}/messages", {"body": "Site visit at 10"})[1]["id"]
        rs = self.call("sara", "GET", "/api/chat/updates?since=0&rsince=0")[1]["reactions_last"]
        st, r = self.call("sara", "POST", f"/api/chat/messages/{mid}/react", {"emoji": "👍"})
        self.assertEqual((st, r["emoji"]), (200, "👍"))
        # the author's devices are told (through the workspace's push function), with the reactions setting honoured
        self.assertEqual([(h["kind"], h["to"], h["pref"], h["href"]) for h in r["_push"]], [("notify", ["boss"], "reactions", f"#/chat/{rid}")])
        self.assertEqual(self.call("boss", "POST", f"/api/chat/messages/{mid}/react", {"emoji": "👍"})[0], 200)
        msg = self.call("boss", "GET", f"/api/chat/rooms/{rid}/messages")[1]["messages"][0]
        self.assertEqual(msg["reactions"], [{"emoji": "👍", "count": 2, "users": [2, 1], "mine": True}])
        # another emoji replaces mine; the same again takes it back
        self.call("sara", "POST", f"/api/chat/messages/{mid}/react", {"emoji": "❤️"})
        r = self.call("sara", "GET", f"/api/chat/rooms/{rid}/messages")[1]["messages"][0]["reactions"]
        self.assertEqual([(x["emoji"], x["count"], x["mine"]) for x in r], [("👍", 1, False), ("❤️", 1, True)])
        self.assertEqual(self.call("sara", "POST", f"/api/chat/messages/{mid}/react", {"emoji": "❤️"}), (200, {"emoji": None}))
        # the author hears about reactions from others, never their own
        self.call("sara", "POST", f"/api/chat/messages/{mid}/react", {"emoji": "🙏"})
        up = self.call("boss", "GET", f"/api/chat/updates?since=0&rsince={rs}")[1]
        self.assertEqual([(x["emoji"], x["user_id"]) for x in up["reactions"]], [("🙏", 2)])
        # text or markup is not an emoji; a removed message takes its reactions with it
        self.assertEqual(self.call("sara", "POST", f"/api/chat/messages/{mid}/react", {"emoji": "<b>x</b>"})[0], 400)
        self.call("boss", "DELETE", f"/api/chat/messages/{mid}")
        self.assertEqual(self.db.one("SELECT COUNT(*) n FROM chat_reactions")["n"], 0)
        self.assertEqual(self.call("sara", "POST", f"/api/chat/messages/{mid}/react", {"emoji": "👍"})[0], 400)

    def test_notify_kinds(self):
        """Each kind of notification can be switched off; muted rooms, quiet days, the reminder time and a pause are kept."""
        st, p = self.call("sara", "PUT", "/api/me/profile", {"prefs": {"notify": {
            "level": "none", "links": False, "approvals": False, "decisions": True, "fu_new": False, "fu_due": True, "desk": False, "bell": False,
            "push": False, "muted": [3, "4", "x"], "quiet_days": [5, 9, "6"], "fu_time": "07:30", "snooze_until": 1893456000000}}})
        n = p["prefs"]["notify"]
        self.assertEqual((n["level"], n["links"], n["approvals"], n["decisions"], n["fu_new"], n["fu_due"], n["desk"], n["bell"], n["push"]),
                         ("none", False, False, True, False, True, False, False, False))
        self.assertEqual((n["muted"], n["quiet_days"], n["fu_time"], n["snooze_until"]), ([3, 4], [5, 6], "07:30", 1893456000000))
        st, p = self.call("sara", "PUT", "/api/me/profile", {"prefs": {"notify": {"fu_time": "9am", "snooze_until": "nope", "level": "loud"}}})
        n = p["prefs"]["notify"]
        self.assertEqual((n["fu_time"], n["snooze_until"], n["level"]), ("07:30", 1893456000000, "none"))     # nonsense is ignored

    def test_push_hints(self):
        """Replies tell the app whom to notify when the app is closed: a follow-up's owner now and on its morning
        (cancelled when done or deleted), an approval's maker when it is decided; never the person acting."""
        st, r = self.call("boss", "POST", "/api/crm/followups", {"title": "Call Mona", "due": "2030-01-05", "assigned_to": "sara"})
        fid = r["id"]
        self.assertEqual([(h.get("key"), h.get("day"), h["to"], h["title"]) for h in r["_push"]],
                         [(f"fu:{fid}", "2030-01-05", ["sara"], "Follow-up due today"), (None, None, ["sara"], "New follow-up from boss")])
        st, r = self.call("boss", "PUT", f"/api/crm/followups/{fid}", {"title": "Call Mona", "due": "2030-01-06", "assigned_to": "sara"})
        self.assertEqual([(h["key"], h["day"]) for h in r["_push"]], [(f"fu:{fid}", "2030-01-06")])       # rescheduled, no second "new"
        st, r = self.call("boss", "PUT", f"/api/crm/followups/{fid}", {"done": True})
        self.assertEqual(r["_push"], [{"cancel": f"fu:{fid}"}])
        st, r = self.call("boss", "POST", "/api/crm/followups", {"title": "Mine", "due": "2030-01-05"})
        self.assertEqual([h.get("title") for h in r["_push"]], ["Follow-up due today"])                  # my own: only the reminder
        st, r = self.call("boss", "DELETE", f"/api/crm/followups/{r['id']}")
        self.assertEqual(r["_push"][0]["cancel"][:3], "fu:")
        st, r = self.call("boss", "POST", "/api/chat/messages/999999/react", {"emoji": "👍"})
        self.assertNotIn("_push", r if isinstance(r, dict) else {})

    def test_voice_and_camera_messages(self):
        """A voice note / camera clip is an ordinary file message; the room list and the alerts say what kind it is."""
        import base64
        audio = base64.b64encode(b"\x1aE\xdf\xa3" + b"\0" * 300).decode()
        st, r = self.call("boss", "POST", "/api/chat/rooms/1/messages",
                          {"body": "", "files": [{"name": "Voice message (0:07).webm", "mime": "audio/webm", "data": audio}]})
        self.assertEqual(st, 200)
        room = next(x for x in self.call("sara", "GET", "/api/chat/rooms")[1]["rooms"] if x["id"] == 1)
        self.assertEqual((room["last"]["files"], room["last"]["file_mime"], room["last"]["text"]), (1, "audio/webm", ""))
        up = self.call("sara", "GET", "/api/chat/updates?since=0&rsince=0")[1]
        self.assertEqual([(m["files"], m["file_mime"]) for m in up["messages"]], [(1, "audio/webm")])
        f = self.call("sara", "GET", "/api/chat/rooms/1/messages")[1]["messages"][0]["files"][0]
        self.assertEqual((f["name"], f["mime"]), ("Voice message (0:07).webm", "audio/webm"))
        r = self.handle(self.app, "GET", f"/api/chat/files/{f['id']}", {"x-requested-with": "BeyondERP", "x-session": self.tok["sara"]}, b"")
        self.assertEqual((r.status, r.body[:4]), (200, b"\x1aE\xdf\xa3"))
        # a clip with a caption and its still
        thumb = "data:image/jpeg;base64," + base64.b64encode(b"\xff\xd8jpeg").decode()
        self.call("boss", "POST", "/api/chat/rooms/1/messages", {"body": "crack in slab B", "files": [
            {"name": "Clip 2026-10-13.mp4", "mime": "video/mp4", "data": base64.b64encode(b"\0\0\0 ftypmp42").decode(), "thumb": thumb}]})
        m = self.call("sara", "GET", "/api/chat/rooms/1/messages")[1]["messages"][-1]
        self.assertEqual((m["body"], m["files"][0]["mime"], m["files"][0]["thumb"]), ("crack in slab B", "video/mp4", thumb))

    def test_notify_settings(self):
        import base64
        st, p = self.call("sara", "PUT", "/api/me/profile", {"prefs": {"notify": {"sound": "marimba", "volume": 140, "vibrate": False, "vibe": "pulse",
                                                                                 "quiet": True, "quiet_from": "23:30", "quiet_to": "7:00", "ask": "never", "sound2": "x"}}})
        n = p["prefs"]["notify"]
        self.assertEqual((n["sound"], n["volume"], n["vibrate"], n["vibe"], n["quiet"], n["quiet_from"], n["quiet_to"], n["ask"]),
                         ("marimba", 100, False, "pulse", True, "23:30", "07:00", "never"))   # a bad time keeps the old one
        self.assertNotIn("sound2", n)
        clip = "data:audio/mpeg;base64," + base64.b64encode(b"ID3" + b"\0" * 1000).decode()
        self.assertEqual(self.call("sara", "PUT", "/api/me/profile", {"prefs": {"notify_sound": clip}})[1]["prefs"]["notify_sound"], clip)
        self.assertEqual(self.call("sara", "PUT", "/api/me/profile", {"prefs": {"notify_sound": "data:text/html;base64,AAAA"}})[0], 400)
        big = "data:audio/mpeg;base64," + "A" * 500_000
        self.assertEqual(self.call("sara", "PUT", "/api/me/profile", {"prefs": {"notify_sound": big}})[0], 400)
        self.assertNotIn("notify_sound", self.call("sara", "PUT", "/api/me/profile", {"prefs": {"notify_sound": ""}})[1]["prefs"])
        self.assertEqual(self.call("boss", "GET", "/api/me/profile")[1]["prefs"]["notify"]["sound"], "chime")   # only hers changed


class FinerPermissionsAndDemo(unittest.TestCase):
    """Per-screen action switches (export / import / delete), tabs inside screens with their own level, and demo data
    only ever into empty books after a confirmed warning."""

    def setUp(self):
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False, "server": None})()
        self.app.db = self.db = Database(":memory:")
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "sam", "password": "secret123", "perms": {
            "shop": "edit", "shop.products": "edit", "shop.subscriptions": "edit", "shop.subscriptions.plans": "none",
            "shop.products!delete": "no", "reports": "view", "reports.income_statement": "view", "reports.income_statement!export": "no",
            "reports.forecast": "view", "reports.forecast.loans": "none"}})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "sam")}

    def call(self, who, method, url, body=None):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json"}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        return r.status, (json.loads(r.body) if r.ctype.startswith("application/json") else r.body)

    def test_actions_and_tabs(self):
        u = auth.public_user(self.db.one("SELECT * FROM users WHERE username='sam'"))
        self.assertEqual(u["perms"]["shop.products!delete"], "no")
        self.assertEqual(auth.level_of(u, "shop.subscriptions.plans"), "none")
        self.assertEqual(auth.level_of(u, "shop"), "edit")
        st, made = self.call("sam", "POST", "/api/products", {"code": "U-1", "name": "Unit 1", "kind": "unit", "price": 100})
        self.assertEqual(st, 200)
        st, body = self.call("sam", "DELETE", f"/api/products/{made['id']}")
        self.assertEqual((st, "switched off" in body["error"]), (403, True))
        self.assertEqual(self.call("boss", "DELETE", f"/api/products/{made['id']}")[0], 200)
        # viewing a report still works; printing / exporting it does not
        self.assertEqual(self.call("sam", "GET", "/api/reports/income_statement")[0], 200)
        self.assertEqual(self.call("sam", "GET", "/api/reports/income_statement?format=pdf")[0], 403)
        self.assertEqual(self.call("sam", "GET", "/api/reports/income_statement?format=xlsx")[0], 403)
        # a tab with its own level
        self.assertEqual(self.call("sam", "GET", "/api/subscriptions")[0], 200)
        self.assertEqual(self.call("sam", "GET", "/api/service-plans")[0], 403)
        keys = [x["key"] for x in self.call("sam", "GET", "/api/forecast")[1]["sections"]]
        self.assertNotIn("loans", keys)
        self.assertIn("prices", keys)

    def test_demo_only_into_empty_books(self):
        self.assertEqual(self.call("boss", "POST", "/api/demo", {})[0], 400)                    # no confirmed warning
        self.assertEqual(self.call("boss", "POST", "/api/demo", {"replace": True, "confirm": "DEMO"})[0], 400)
        self.call("boss", "POST", "/api/products", {"code": "U-1", "name": "Unit 1", "kind": "unit", "price": 100})
        st, body = self.call("boss", "POST", "/api/demo", {"confirmed": True})
        self.assertEqual((st, "empty books" in body["error"]), (400, True))
        db2 = Database(":memory:")
        from erp import demo
        auth.create_user(db2, {"username": "boss", "password": "secret123"}, first=True)
        self.assertTrue(demo.is_empty(db2))
        self.app.db = db2
        self.tok["boss"] = auth.login(db2, "boss", "secret123")[0]
        self.assertEqual(self.call("boss", "POST", "/api/demo", {"confirmed": True})[0], 200)
        self.assertFalse(demo.is_empty(db2))
        self.assertEqual(self.call("boss", "POST", "/api/demo", {"confirmed": True})[0], 400)   # not twice


class PersonalLook(unittest.TestCase):
    """Colours, font, sizes and background pictures belong to each person (on their profile); only a super
    administrator sets the company look, and a large look is stored whole, never cut short."""

    def setUp(self):
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = Database(":memory:")
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "ada", "password": "secret123", "is_admin": True})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "ada")}

    def call(self, who, method, url, body=None):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json"}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        return r.status, json.loads(r.body)

    def test_look(self):
        import base64
        pic = "data:image/jpeg;base64," + base64.b64encode(b"x" * 300_000).decode()
        look = {"mode": "dark", "theme": {"button": "#e03a4b", "menu": "#FF7F11", "page": "javascript:alert(1)"}, "fontSize": "40",
                "density": "compact", "images": {"page": {"src": pic, "op": 3, "fit": "tile"}, "nowhere": {"src": pic}}}
        st, prof = self.call("ada", "PUT", "/api/me/profile", {"prefs": {"look": look}})
        self.assertEqual(st, 200)
        lk = prof["prefs"]["look"]
        self.assertEqual(lk["theme"], {"button": "#e03a4b", "menu": "#FF7F11"})     # a bad colour is dropped
        self.assertEqual((lk["fontSize"], lk["density"], lk["mode"]), ("22", "compact", "dark"))
        self.assertEqual(list(lk["images"]), ["page"])
        self.assertEqual((lk["images"]["page"]["op"], lk["images"]["page"]["fit"]), (1.0, "tile"))
        # only ada has it
        self.assertNotIn("look", self.call("boss", "GET", "/api/me/profile")[1]["prefs"])
        # not a picture → refused; clearing works
        self.assertEqual(self.call("ada", "PUT", "/api/me/profile", {"prefs": {"look": {"images": {"page": {"src": "data:text/html;base64,AAAA"}}}}})[0], 400)
        self.assertNotIn("look", self.call("ada", "PUT", "/api/me/profile", {"prefs": {"look": {}}})[1]["prefs"])
        # the company look: super administrators only — an administrator may still save the rest of the appearance
        self.assertEqual(self.call("ada", "PUT", "/api/settings", {"appearance": {"look": {"mode": "light"}}})[0], 403)
        st, sett = self.call("boss", "PUT", "/api/settings", {"appearance": {"look": {**look, "images": {"page": {"src": pic, "op": 0.2}}}}})
        self.assertEqual(st, 200)
        self.assertEqual(sett["appearance"]["look"]["images"]["page"]["src"], pic)        # stored whole (bigger than 20,000 characters)
        self.assertEqual(self.call("ada", "PUT", "/api/settings", {"appearance": {"look": sett["appearance"]["look"]}})[0], 200)


class Support(unittest.TestCase):
    """Help & support tickets: a reference built from the date and record number; people see their own, admins all."""

    def test_tickets(self):
        from erp import auth, support
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "Passw0rd!long1"}, first=True)
        auth.create_user(db, {"username": "sara", "password": "Passw0rd!long2"})
        boss, sara = (dict(db.one("SELECT * FROM users WHERE username=?", (n,))) for n in ("boss", "sara"))
        t = support.create(db, sara, {"kind": "bug", "subject": "Voucher will not save", "body": "steps", "channel": "whatsapp",
                                      "diag": {"version": "1.2.0", "screen": "#/payments"}})
        self.assertRegex(t["ref"], r"^BE-\d{6}-0001$")
        self.assertEqual(t["diag"]["screen"], "#/payments")
        support.create(db, boss, {"kind": "nonsense", "subject": "Question"})
        self.assertEqual([x["kind"] for x in support.tickets(db, boss)], ["bug", "bug"])   # unknown kind -> a problem report
        self.assertEqual(len(support.tickets(db, sara)), 1)
        support.set_status(db, sara, t["id"], "solved")
        self.assertEqual(support.tickets(db, sara)[0]["status"], "solved")
        with self.assertRaises(Exception):
            support.create(db, sara, {"subject": ""})


class PackageOffers(unittest.TestCase):
    """Offers priced as packages (no BOQ): tiers side by side; the chosen one is the offer's value."""

    def _offer(self, chosen=-1, style="table"):
        from erp import documents
        return dict(documents._content({"title": "DESIGN PACKAGES :{project}", "sections": [
            {"kind": "packages", "heading": "PACKAGES", "style": style, "chosen": chosen, "tiers": [
                {"name": "Essential", "price": 40000, "includes": "Layout\nMood board"},
                {"name": "Signature", "price": 70000, "includes": "Layout\nMood board\n3D renders"},
                {"name": "", "price": 5, "includes": "ignored: no name"}]},
            {"kind": "payment", "heading": "PAYMENT", "style": "table", "milestones": [{"label": "Advance", "pct": 50}, {"label": "Delivery", "pct": 50}]}]}),
            client="Client Name", project="Villa", date="2026-10-08")

    def test_not_chosen_yet(self):
        from erp import documents
        db = Database(":memory:")
        o = self._offer()
        self.assertIsNone(documents.pricing(o))
        html = documents.offer_html(db, o, preview=True)
        for s in ("Essential", "Signature", "3D renders", "✓", "—", "40,000", "70,000"):
            self.assertIn(s, html)
        self.assertNotIn("ignored", html)
        self.assertNotIn("35,000", html)          # no payment amounts before a package is chosen
        self.assertTrue(documents.offer_docx(db, o)[:2] == b"PK")

    def test_chosen_package_is_the_value(self):
        from erp import documents
        db = Database(":memory:")
        o = self._offer(chosen=1, style="lines")
        self.assertEqual(documents.offer_totals(documents.pricing(o))["net"], 70000)
        html = documents.offer_html(db, o, preview=True)
        self.assertIn("★ Signature", html)
        self.assertIn("35,000", html)             # 50% of the chosen package

    def test_settings_preview_has_no_real_bank(self):
        from erp import documents
        db = Database(":memory:")
        db.set_setting("banks", json.dumps([{"id": "m", "label": "Real", "bank": "Real Bank", "acct": "123456789"}]))
        sec = documents._content({"sections": [{"kind": "bank", "heading": "BANK"}]})["sections"]
        sec[0].update(bank_id="custom", bank_details=documents.SAMPLE_BANK)
        html = documents.offer_html(db, {"title": "X", "sections": sec, "client": "", "project": ""}, preview=True)
        self.assertNotIn("123456789", html)
        self.assertIn("0000 0000 0000", html)


class UniversalTemplates(unittest.TestCase):
    """Import templates carry the BeyondERP logo and 'Company Name' unless the company asks for its own identity."""

    def _company(self):
        from erp import branding
        db = Database(":memory:")
        db.set_setting("company_name", "Acme Interiors")
        db.set_setting("company_tagline", "Fit-out")
        db.set_setting("company_address", "1 Main St")
        branding.set_asset(db, "logo", {"mime": "image/png", "data": "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFBQIAX8jx0gAAAABJRU5ErkJggg=="})
        return db

    def _cover(self, xlsx):
        import io
        import zipfile
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        sets = {r[0]: r[1] for r in wb["Settings"].iter_rows(min_row=5, values_only=True) if r and r[0]}
        media = [n for n in zipfile.ZipFile(io.BytesIO(xlsx)).namelist() if n.startswith("xl/media/")]
        return wb["Cover"]["B2"].value, wb["Cover"]["B3"].value, sets, media

    def test_universal_by_default_and_safe_to_import(self):
        from erp import brand_assets, workbook
        db = self._company()
        x = workbook.build_workbook(db, template=True)
        name, tag, sets, media = self._cover(x)
        self.assertEqual((name, tag), ("COMPANY NAME", "COMPANY SLOGAN"))
        self.assertEqual(sets.get("company_name"), "Company Name")
        self.assertNotIn("company_address", sets)
        self.assertTrue(media)                      # the BeyondERP logo
        self.assertTrue(brand_assets.LOGO["data"])
        p = Path(tempfile.mkdtemp()) / "t.xlsx"
        p.write_bytes(x)
        interchange.import_data(db, interchange.read_xlsx(p))
        self.assertEqual((db.setting("company_name"), db.setting("company_tagline"), db.setting("company_address")),
                         ("Acme Interiors", "Fit-out", "1 Main St"))

    def test_own_identity_when_chosen(self):
        from erp import branding, workbook
        db = self._company()
        branding.save(db, {"apply": {"template": {"logo": True, "wm": False}}})
        name, tag, sets, media = self._cover(workbook.build_workbook(db, template=True))
        self.assertEqual((name, sets.get("company_name")), ("ACME INTERIORS", "Acme Interiors"))
        self.assertTrue(media)

    def test_new_books_are_not_branded_for_anyone(self):
        db = Database(":memory:")
        self.assertEqual(db.setting("company_name", ""), "")
        self.assertTrue(db.one("SELECT 1 FROM assets WHERE key='watermark'"))   # the faded BeyondERP mark


class WorkspaceCatchUp(unittest.TestCase):
    def test_heavy_changes_are_recognised(self):
        """After these, the device that made them saves a fresh copy at once, so phones load it instead of replaying."""
        from erp import cloud
        for path in ("/api/import?mode=replace&format=json", "/api/import", "/api/import-sheet", "/api/demo", "/api/erase"):
            self.assertTrue(cloud.is_heavy("POST", path), path)
        for m, path in (("POST", "/api/import-preview/check"), ("POST", "/api/docs/expense"), ("GET", "/api/import"), ("POST", "/api/bulk")):
            self.assertFalse(cloud.is_heavy(m, path), path)
        js = (Path(__file__).resolve().parents[1] / "web-src" / "bridge.js").read_text()
        self.assertIn("HEAVY = /^\\/api\\/(import(\\?|$)|import-sheet|demo|erase|setup)/", js)   # same rule in the browser
        self.assertIn("await breathe()", js)


class Forecasting(unittest.TestCase):
    """Driver-based price forecast: escalation from CPI/wages/FX/fuel/global, category mixes, overrides kept."""

    def test_forecast(self):
        import datetime as _dt
        from erp import demo, forecast
        db = Database(":memory:")
        demo.load(db)
        db.x("INSERT INTO stock_items(name, unit, category, std_cost, price) VALUES ('Price-only tile','m2','imported_materials',10000,15000)")
        self.assertEqual(db.one("SELECT COALESCE(SUM(qty),0) q FROM stock_moves WHERE item_id=(SELECT id FROM stock_items WHERE name='Price-only tile')")["q"], 0)
        forecast.save(db, {"horizon": 4, "drivers": {"cpi": {"rates": [0.03]}, "fx": {"rates": [0.02]}, "global": {"rates": [0.01]},
                                                     "wages": {"rates": [0.04]}, "fuel": {"rates": [0.05]}},
                           "overrides": {"item:x": {"1": 7}}})
        q = forecast.quarters(4, _dt.date(2026, 11, 5))
        self.assertEqual([p["key"] for p in q], ["2027-Q1", "2027-Q2", "2027-Q3", "2027-Q4"])
        a = forecast.assumptions(db)
        _dr, blocks = forecast.escalation(a, 4)
        self.assertAlmostEqual(blocks["imported"][0], 1.02 * 1.01 - 1)
        self.assertAlmostEqual(blocks["local"][3], 0.03)                     # the last quarter's rate carries on
        out = forecast.build(db, _dt.date(2026, 11, 5))
        items = next(s for s in out["sections"] if s["key"] == "items")["rows"]
        tile = next(r for r in items if r["label"] == "Price-only tile")   # zero stock still gets a price path
        self.assertEqual(tile["base"], 100.0)
        self.assertGreater(tile["values"][0], 100.0)
        self.assertGreater(tile["values"][3], tile["values"][0])
        # an override replaces that quarter only
        forecast.save(db, {**forecast.assumptions(db), "overrides": {tile["key"]: {"2": 150}}})
        tile2 = next(r for r in next(s for s in forecast.build(db, _dt.date(2026, 11, 5))["sections"] if s["key"] == "items")["rows"]
                     if r["label"] == "Price-only tile")
        self.assertEqual(tile2["values"][2], 150)
        self.assertEqual(tile2["edited"], [2])
        self.assertEqual(tile2["values"][0], tile["values"][0])


class BigBookSpeed(unittest.TestCase):
    """The speed-ups for large books return exactly what the plain queries did, and cached figures never go stale."""

    def test_party_lines_and_cache(self):
        from erp import demo
        db = Database(":memory:")
        demo.load(db)
        # the indexed party lookup equals "the line's party, else the entry's" for every party
        for (name,) in db.conn.execute("SELECT DISTINCT COALESCE(NULLIF(l.party,''), e.party) FROM journal_lines l "
                                       "JOIN journal_entries e ON e.id=l.entry_id WHERE COALESCE(NULLIF(l.party,''), e.party)<>''"):
            x = db.q(f"SELECT l.id FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id WHERE {a.PARTY}=? ORDER BY l.id", (name,))
            b = db.q(f"SELECT l.id FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id WHERE {a.PARTY_IS} ORDER BY l.id", (name,))
            self.assertEqual(x, b, name)
        # cached statements follow every write
        before = a.income_statement(db)["net_profit"]
        self.assertEqual(a.income_statement(db)["net_profit"], before)
        first = a.income_statement(db)
        first["net_profit"] = 1          # a caller changing its copy does not change the cache
        self.assertEqual(a.income_statement(db)["net_profit"], before)
        a.record_journal(db, {"date": a.today(), "kind": "Journal Voucher", "memo": "t",
                                "lines": [{"account": "6090", "debit": 100}, {"account": "1010", "credit": 100}]})
        self.assertAlmostEqual(a.income_statement(db)["net_profit"], before - 100)

    def test_ref_prefix(self):
        self.assertEqual(a.ref_prefix("LN:a*b[1]?:"), "LN:a[*]b[[]1][?]:*")
        db = Database(":memory:")
        for ref in ("LN:A[1]:pay", "LN:A11:pay", "LN:A[1]"):
            db.x("INSERT INTO journal_entries(date, kind, memo, ref, status) VALUES ('2026-01-01','Journal Voucher','',?,'posted')", (ref,))
        got = [r["ref"] for r in db.q("SELECT ref FROM journal_entries WHERE ref GLOB ?", (a.ref_prefix("LN:A[1]:"),))]
        self.assertEqual(got, ["LN:A[1]:pay"])

    def test_bankrec_report_needs_a_statement(self):
        from erp import reports
        db = Database(":memory:")
        with self.assertRaises(a.ValidationError):
            reports.build(db, "bank_reconciliation", {})


class AppUpdates(unittest.TestCase):
    """The installed desktop app learns about a newer version from the website; running from source never does."""

    def test_check(self):
        from unittest import mock
        from erp import appupdate
        self.assertGreater(appupdate.vt("1.10.0"), appupdate.vt("1.9.3"))
        self.assertFalse(appupdate.check()["newer"])                       # from source: updated with git
        site = {"version": "99.0.0", "files": [{"file": "BeyondERP-Offline-Setup.exe"}, {"file": "BeyondERP-Offline-mac.zip"}]}
        with mock.patch.object(appupdate.sys, "frozen", True, create=True), mock.patch.object(appupdate, "latest", return_value=site):
            r = appupdate.check()
            self.assertTrue(r["newer"])
            self.assertEqual(r["latest"], "99.0.0")
        with mock.patch.object(appupdate.sys, "frozen", True, create=True), \
                mock.patch.object(appupdate, "latest", return_value={**site, "version": "0.0.1"}):
            self.assertFalse(appupdate.check()["newer"])
        with mock.patch.object(appupdate.sys, "frozen", True, create=True), mock.patch.object(appupdate, "latest", side_effect=OSError):
            self.assertFalse(appupdate.check()["newer"])                   # offline: quietly no


class PriceOutlook(unittest.TestCase):
    def test_outlook(self):
        from erp import demo, forecast
        db = Database(":memory:")
        demo.load(db)
        self.assertFalse(forecast.outlook(db, accounts=["5010"])["ready"])   # nothing until the drivers are filled in
        forecast.save(db, {"drivers": {"cpi": {"rates": [0.03]}, "wages": {"rates": [0.04]}, "fx": {"rates": [0.02]},
                                       "global": {"rates": [0.01]}, "fuel": {"rates": [0.05]}}})
        o = forecast.outlook(db, accounts=["5010", "5030", "1010"])
        self.assertTrue(o["ready"])
        self.assertEqual(o["accounts"]["5030"]["cat"], "labour")             # direct labour moves with wages
        self.assertAlmostEqual(o["accounts"]["5030"]["total"], 1.04 ** 4 - 1, places=4)
        self.assertEqual(o["accounts"]["5010"]["cat"], "local_materials")
        self.assertNotIn("1010", o["accounts"])                             # cash is not a price
        # direct costs are part of the spending forecast now
        rows = next(s for s in forecast.build(db)["sections"] if s["key"] == "expenses")["rows"]
        self.assertTrue(any(r["key"] == "acct:5010" for r in rows))


class DefaultLogoAndMaterials(unittest.TestCase):
    def test_old_default_logo_is_replaced_uploaded_kept(self):
        import hashlib
        from erp import brand_assets
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "b.sqlite3"
            db = Database(path)
            fake_old = "T0xEREVGQVVMVA=="   # stands in for the old built-in mark
            db.x("INSERT OR REPLACE INTO assets(key,mime,w,h,data) VALUES ('logo','image/png',10,10,?)", (fake_old,))
            db.x("INSERT OR REPLACE INTO assets(key,mime,w,h,data) VALUES ('watermark','image/png',10,10,'VVBMT0FERUQ=')")
            db.close()
            Database.OLD_DEFAULT_ASSETS = {**Database.OLD_DEFAULT_ASSETS, hashlib.sha256(fake_old.encode()).hexdigest(): "LOGO"}
            db = Database(path)
            self.assertEqual(db.one("SELECT data FROM assets WHERE key='logo'")["data"], brand_assets.LOGO["data"])
            self.assertEqual(db.one("SELECT data FROM assets WHERE key='watermark'")["data"], "VVBMT0FERUQ=")   # an upload stays
            db.close()

    def test_journal_item_lines(self):
        from erp import demo
        db = Database(":memory:")
        demo.load(db)
        wh = db.one("SELECT id, name FROM warehouses LIMIT 1")
        bu = db.one("SELECT id, code FROM business_units LIMIT 1")
        eid = a.record_journal(db, {"date": "2026-10-05", "memo": "Tiles and paint", "party": "Bright Supplies", "bu_id": bu["id"],
                                    "items": [{"item": "Porcelain 60x120", "unit": "m2", "qty": 50, "rate": 420, "warehouse_id": wh["id"]},
                                              {"item": "Wall paint 18L", "unit": "drum", "qty": 4, "rate": 1500}],
                                    "lines": [{"account": "1010", "credit": 27000}]})
        e = a.get_entry(db, eid)
        self.assertEqual(e["amount"], 27000)
        self.assertEqual(db.one("SELECT l.qty FROM stock_levels l JOIN stock_items i ON i.id=l.item_id WHERE i.name='Porcelain 60x120'")["qty"], 50)
        self.assertTrue(all(c["ok"] for c in a.audit(db)))          # stock register = inventory accounts: counted once
        used = a.materials_used(db, q="paint")
        self.assertEqual(used["count"], 2)                           # the new line and the demo's wall paint issue/cost
        self.assertTrue(any(r["source"] == "bought" and r["bu_code"] == bu["code"] and r["amount"] == 6000 for r in used["rows"]))
        with self.assertRaises(a.ValidationError):                   # item lines are corrected by reversing, not edited
            a.edit_entry(db, eid, {"memo": "x"})
        # the books export and import back with the item lines intact
        data = interchange.export_data(db)
        db2 = Database(":memory:")
        interchange.import_data(db2, data)
        self.assertEqual(a.materials_used(db2)["total"], a.materials_used(db)["total"])
        self.assertEqual(db2.one("SELECT COUNT(*) n FROM material_lines")["n"], 2)
        # void: the stock and the history go with it
        a.void_entry(db, eid, "test")
        self.assertEqual(db.one("SELECT l.qty FROM stock_levels l JOIN stock_items i ON i.id=l.item_id WHERE i.name='Porcelain 60x120'")["qty"], 0)
        self.assertEqual(a.materials_used(db, q="Wall paint 18L")["count"], 1)
        with self.assertRaises(a.ValidationError):                   # a project line needs a project
            a.record_journal(db, {"date": "2026-10-05", "items": [{"item": "Sand", "qty": 1, "rate": 10}], "lines": [{"account": "1010", "credit": 10}]})

    def test_quantity_total_only_for_one_material(self):
        from erp import demo
        db = Database(":memory:")
        demo.load(db)
        self.assertIsNone(a.materials_used(db)["one_material"])                   # a project mixes units: no total
        one = a.materials_used(db, item="wall paint 18l")["one_material"]           # one material: its total, in its unit
        self.assertEqual((one["qty"], one["unit"]), (6, "drum"))


class Healthcare(unittest.TestCase):
    """Clinics and home healthcare: insurer share vs co-pay, provider shares, consumables from stock, prepaid
    packages, claims answered with rejections, the business types and their demo companies."""

    def setUp(self):
        from erp import care, industries
        self.care = care
        self.db = Database(":memory:")
        industries.apply(self.db, "healthcare")
        self.db.x("INSERT INTO business_units(code,name) VALUES ('CL1','Clinic 1')")
        self.bu = self.db.one("SELECT id FROM business_units WHERE code='CL1'")["id"]
        a.stock_move(self.db, {"type": "Receipt", "date": "2026-01-02", "warehouse": self.db.one("SELECT name FROM warehouses")["name"],
                               "item": "Gloves", "unit": "pair", "qty": 100, "rate": 5, "paid": 500})
        self.db.x("UPDATE stock_items SET sku='GLV' WHERE name='Gloves'")
        self.gp = care.save_service(self.db, {"name": "GP consultation", "category": "consultation", "price": 400, "insurer_price": 300,
                                              "provider_pct": 40, "consumables": [{"item": "GLV", "qty": 2}]})
        self.home = care.save_service(self.db, {"name": "Nursing shift", "category": "nursing", "price": 900, "covered": False, "provider_pct": 70})
        self.ins = care.save_patient(self.db, {"name": "Insured Ina", "payer": "insurer", "insurer": "Demo TPA", "coverage_pct": 80})
        self.cash = care.save_patient(self.db, {"name": "Cash Cal", "zone": "South District", "address": "1 Main St"})

    def bal(self, code, party=None):
        """Debit balance of an account (of one party on it)."""
        return self.db.one("SELECT COALESCE(SUM(l.debit-l.credit),0) b FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id "
                           "WHERE e.status='posted' AND l.account=? AND (? IS NULL OR l.party=?)", (code, party, party))["b"] / 100

    def test_type_template(self):
        from erp import industries
        self.assertTrue(self.db.one("SELECT 1 FROM accounts WHERE code='1105'"))
        self.assertNotIn("care", json.loads(self.db.setting("hidden_modules")))
        other = Database(":memory:")                       # every other business keeps the healthcare menus out of sight
        self.assertIn("care", json.loads(other.setting("hidden_modules")))
        industries.apply(other, "contracting")
        self.assertIn("care", json.loads(other.setting("hidden_modules")))

    def test_insured_visit_billing_claim(self):
        c = self.care
        vid = c.save_visit(self.db, {"date": "2026-02-01", "patient_id": self.ins, "provider": "Dr. A", "bu_id": self.bu,
                                     "lines": [{"service_id": self.gp, "qty": 1}]}, "boss")
        v = c.visit(self.db, vid)
        # the insurer's contract price (300): 80% insurer, 20% co-pay; the doctor gets 40%
        self.assertEqual((v["insurer_amt"], v["patient_amt"], v["provider_amt"]), (240, 60, 120))
        with self.assertRaises(a.ValidationError):
            c.bill_visit(self.db, vid, {}, "boss")         # not done yet
        c.set_status(self.db, vid, "arrived", {"time": "10:00"})
        c.set_status(self.db, vid, "done", {"time": "10:45", "vitals": {"bp": "120/80"}})
        self.assertEqual(c.visit(self.db, vid)["hours"], 0.75)
        r = c.bill_visit(self.db, vid, {"paid": 60, "method": "Cash"}, "boss")
        self.assertEqual(r["short"], [])
        self.assertEqual(self.bal("1105", "Demo TPA"), 240)
        self.assertEqual(self.bal("4061"), -300)
        self.assertEqual(self.bal("2115", "Dr. A"), -120)
        self.assertEqual(self.bal("5091"), 10)             # 2 pairs × 5 from stock
        with self.assertRaises(a.ValidationError):
            c.save_visit(self.db, {"date": "2026-02-01", "patient_id": self.ins, "lines": [{"service_id": self.gp}]}, "boss", vid)
        # the claim: paid 200, 40 rejected
        cid = c.make_claim(self.db, "Demo TPA", "2026-02-01", "2026-02-28", "boss")
        self.assertEqual(c.claims(self.db)[0]["claimed"], 240)
        self.assertEqual(c.unclaimed(self.db), [])
        with self.assertRaises(a.ValidationError):
            c.settle_claim(self.db, cid, {"paid": 200, "rejected": 0}, "boss")
        with self.assertRaises(a.ValidationError):
            c.unbill_visit(self.db, vid, "boss")           # in a claim
        c.settle_claim(self.db, cid, {"paid": 200, "rejected": 40, "date": "2026-03-20"}, "boss")
        self.assertEqual(self.bal("1105", "Demo TPA"), 0)
        self.assertEqual(self.bal("4056"), 40)

    def test_home_visit_partial_pay_and_unbill(self):
        c = self.care
        vid = c.save_visit(self.db, {"date": "2026-02-03", "patient_id": self.cash, "provider": "Nurse B", "place": "home", "travel_fee": 150,
                                     "bu_id": self.bu, "lines": [{"service_id": self.home, "qty": 1}]}, "boss")
        v = c.visit(self.db, vid)
        self.assertEqual((v["patient_amt"], v["zone"], v["address"]), (1050, "South District", "1 Main St"))
        c.set_status(self.db, vid, "done", {"check_in": "20:00", "time": "08:00"})
        self.assertEqual(c.visit(self.db, vid)["hours"], 12)          # an overnight shift
        c.bill_visit(self.db, vid, {"paid": 50}, "boss")
        self.assertEqual(self.bal("1100", "Cash Cal"), 1000)
        self.assertEqual((self.bal("4062"), self.bal("4063")), (-900, -150))
        self.assertEqual([p["owes"] for p in c.patients(self.db) if p["id"] == self.cash], [1000])
        c.unbill_visit(self.db, vid, "boss", "wrong patient")
        self.assertEqual((self.bal("1100", "Cash Cal"), self.bal("4062"), self.bal("2115")), (0, 0, 0))
        self.assertIsNone(c.visit(self.db, vid)["entry_id"])

    def test_package(self):
        c = self.care
        kid = c.sell_package(self.db, {"patient_id": self.cash, "service_id": self.home, "sessions": 4, "price": 3200, "paid": 2000,
                                       "date": "2026-02-01"}, "boss")
        self.assertEqual((self.bal("2108"), self.bal("1100", "Cash Cal")), (-3200, 1200))
        for day in ("2026-02-02", "2026-02-03"):
            vid = c.save_visit(self.db, {"date": day, "patient_id": self.cash, "place": "home", "package_id": kid, "provider": "Nurse B",
                                         "lines": [{"service_id": self.home}]}, "boss")
            self.assertEqual(c.visit(self.db, vid)["patient_amt"], 0)
            c.set_status(self.db, vid, "done")
            c.bill_visit(self.db, vid, {}, "boss")
        k = c.packages(self.db)[0]
        self.assertEqual((k["used"], k["left"], k["held"]), (2, 2, 1600))
        self.assertEqual((self.bal("2108"), self.bal("4062")), (-1600, -1600))
        self.assertEqual(self.bal("2115", "Nurse B"), -1120)            # 70% of 800 per session

    def test_api_permissions_and_report(self):
        from erp.server import handle
        app = type("App", (), {"static": False, "server": None})()
        app.db = self.db
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "desk", "password": "secret123", "perms": {"care": "add", "care.claims": "none", "care.reports": "none"}})
        tok = auth.login(self.db, "desk", "secret123")[0]
        call = lambda m, u, b=None: handle(app, m, u, {"x-requested-with": "BeyondERP", "x-session": tok, "content-type": "application/json"},  # noqa: E731
                                           json.dumps(b).encode() if b is not None else b"")
        r = call("POST", "/api/care/visits", {"date": "2026-02-05", "patient_id": self.ins, "lines": [{"service_id": self.gp}], "bu_id": self.bu})
        vid = json.loads(r.body)["id"]
        self.assertEqual(call("POST", f"/api/care/visits/{vid}/status", {"status": "done"}).status, 200)
        self.assertEqual(call("POST", f"/api/care/visits/{vid}/bill", {}).status, 200)
        self.assertEqual(call("POST", f"/api/care/visits/{vid}/void", {}).status, 403)    # undoing a bill needs edit
        self.assertEqual(call("GET", "/api/care/claims").status, 403)
        self.assertEqual(call("GET", "/api/care/report").status, 403)
        rep = self.care.report(self.db)
        self.assertEqual((rep["counts"], rep["insurers_due"][0]["balance"], rep["unclaimed"][0]["amount"]), ({"done": 1}, 240, 240))

    def test_demo_companies_and_backup(self):
        from erp import demo, industries, interchange
        for kind in ("healthcare", "homecare"):
            db = Database(":memory:")
            industries.apply(db, kind)
            demo.load(db, "boss", industry=kind)
            self.assertGreater(db.one("SELECT COUNT(*) n FROM care_visits WHERE status='done' AND entry_id IS NOT NULL")["n"], 200)
            self.assertTrue(db.one("SELECT 1 FROM care_claims WHERE status='settled'"))
            self.assertTrue(db.one("SELECT 1 FROM care_visits WHERE status='booked'"))
            tb = db.one("SELECT SUM(debit) d, SUM(credit) c FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id WHERE e.status='posted'")
            self.assertEqual(tb["d"], tb["c"])
            held = round(sum(k["held"] for k in self.care.packages(db)), 2)
            prepaid = db.one("SELECT COALESCE(SUM(credit-debit),0) b FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id "
                             "WHERE e.status='posted' AND l.account='2108'")["b"] / 100
            self.assertEqual(round(prepaid, 2), held)              # prepaid = the sessions still owed
            if kind == "homecare":
                self.assertTrue(all(v["place"] == "home" for v in self.care.visits(db)))
            # backup → restore into fresh books keeps every patient, visit and claim
            fresh = Database(":memory:")
            interchange.import_data(fresh, interchange.read_json(interchange.to_json(db).encode()), mode="replace")
            for t in self.care.TABLES:
                self.assertEqual(fresh.one(f"SELECT COUNT(*) n FROM {t}")["n"], db.one(f"SELECT COUNT(*) n FROM {t}")["n"], t)


class PublicDesk(unittest.TestCase):
    """Online booking, web forms, offers signed online and guest chat: a visitor without a sign-in uses the company's
    public pages (stand-in Supabase), the app claims what they sent and adds it to the books."""

    @classmethod
    def setUpClass(cls):
        from erp.server import ERPServer
        sys.path.insert(0, str(Path(__file__).parent))
        import mock_supabase
        cls.sb = mock_supabase.start()
        cls.sb_url = f"http://127.0.0.1:{cls.sb.server_address[1]}"
        cls.tmp = tempfile.TemporaryDirectory()
        cls.db = Database(Path(cls.tmp.name) / "d.sqlite3")
        cls.srv = ERPServer(("127.0.0.1", 0), cls.db, Path(cls.tmp.name) / "bk")
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()
        cls.base = f"http://127.0.0.1:{cls.srv.server_address[1]}"

    @classmethod
    def tearDownClass(cls):
        cls.srv.cloud.stop()
        cls.srv.shutdown()
        cls.srv.server_close()
        cls.db.close()
        cls.sb.shutdown()
        cls.tmp.cleanup()

    call = Api.call
    client = Api.client

    def visitor(self, fn, body):
        req = urllib.request.Request(f"{self.sb_url}/rest/v1/rpc/{fn}", data=json.dumps(body).encode(),
                                     headers={"Content-Type": "application/json", "apikey": "anon"}, method="POST")
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read() or b"null")
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"null")

    def cloud(self, op, name, params=None):
        st, r = self.call(op, "POST", "/api/cloud/call", {"kind": "rpc", "name": name, "params": params or {}})
        self.assertEqual(st, 200, r)
        return r["result"]

    def pull(self, op):
        rows = self.cloud(op, "desk_claim")
        st, r = self.call(op, "POST", "/api/desk/import", {"rows": rows})
        self.assertEqual(st, 200, r)
        return r["rows"]

    def test_full_cycle(self):
        from erp import care, industries
        op = self.client()
        self.assertEqual(self.call(op, "POST", "/api/setup", {"username": "admin", "password": "secret123"})[0], 200)
        industries.apply(self.db, "healthcare")
        sid = care.save_service(self.db, {"name": "GP consultation", "price": 400, "duration_min": 30})
        self.assertEqual(self.call(op, "PUT", "/api/cloud/config", {"url": self.sb_url, "key": "anon"})[0], 200)
        self.assertEqual(self.call(op, "POST", "/api/cloud/signin", {"email": "boss@x.com", "password": "pw123456", "create": True})[0], 200)
        self.assertEqual(self.call(op, "POST", "/api/cloud/claim", {})[0], 200)
        # the company publishes a booking page; a visitor books, the slot is held, a second visitor cannot take it
        st, d = self.call(op, "GET", "/api/desk")
        self.assertEqual([s["name"] for s in d["services"]], ["GP consultation"])
        self.assertEqual(self.cloud(op, "desk_my_publish"), "all")                        # administrators publish any page
        self.assertGreaterEqual(self.cloud(op, "desk_publishers", {"p_all": ["desk@x.com"], "p_sign": ["sales@x.com"]}), 1)
        page = self.cloud(op, "desk_page_save", {"p_id": None, "p_kind": "booking", "p_title": "Book a visit",
                                                  "p_config": {"services": d["services"], "per_slot": 1}})
        tok = page["token"]
        st, pg = self.visitor("public_page", {"p_token": tok})
        self.assertEqual((st, pg["kind"], pg["title"]), (200, "booking", "Book a visit"))
        st, b = self.visitor("public_send", {"p_token": tok, "p_kind": "booking", "p_payload": {
            "slot": "2031-03-04 10:30", "name": "Anna Bell", "phone": "0001234567", "service": "GP consultation", "service_id": sid}})
        self.assertEqual(st, 200, b)
        self.assertEqual(self.visitor("public_send", {"p_token": tok, "p_kind": "booking", "p_payload": {"slot": "2031-03-04 10:30", "name": "X"}})[0], 409)
        self.assertEqual(self.visitor("public_send", {"p_token": tok, "p_kind": "lead", "p_payload": {"name": "X"}})[0], 400)
        # the app brings it in: a booked visit for a new patient
        res = self.pull(op)
        self.assertEqual([r["state"] for r in res], ["done"])
        v = self.db.one("SELECT v.*, p.name, p.phone FROM care_visits v JOIN care_patients p ON p.id=v.patient_id")
        self.assertEqual((v["date"], v["time"], v["status"], v["name"], v["phone"]), ("2031-03-04", "10:30", "booked", "Anna Bell", "0001234567"))
        self.assertEqual(self.pull(op), [])                                  # each row once
        # the visitor moves it, then cancels; the visit follows
        self.assertEqual(self.visitor("public_manage", {"p_token": tok, "p_ref": b["ref"], "p_manage": "wrong", "p_action": "view"})[0], 403)
        self.assertEqual(self.visitor("public_manage", {"p_token": tok, "p_ref": b["ref"], "p_manage": b["manage"], "p_action": "move",
                                                        "p_slot": "2031-03-05 11:00"})[0], 200)
        self.pull(op)
        v = self.db.one("SELECT * FROM care_visits")
        self.assertEqual((v["date"], v["time"]), ("2031-03-05", "11:00"))
        # the team confirms; the visitor sees it
        st, dd = self.call(op, "GET", "/api/desk?from=2030-01-01")
        bk = dd["bookings"][0]
        self.assertEqual((bk["slot"], bk["name"]), ("2031-03-05 11:00", "Anna Bell"))
        self.cloud(op, "desk_reply", {"p_page": bk["page_id"], "p_ref": bk["ref"], "p_kind": "confirm", "p_body": "See you then"})
        self.assertEqual(self.call(op, "POST", f"/api/desk/rows/{bk['id']}/answer", {"how": "confirmed"})[0], 200)
        st, mine = self.visitor("public_manage", {"p_token": tok, "p_ref": b["ref"], "p_manage": b["manage"]})
        self.assertEqual(mine["messages"][-1]["body"], "See you then")
        self.visitor("public_manage", {"p_token": tok, "p_ref": b["ref"], "p_manage": b["manage"], "p_action": "cancel"})
        self.pull(op)
        self.assertEqual(self.db.one("SELECT status FROM care_visits")["status"], "cancelled")
        # a web form becomes a lead with a call-back
        form = self.cloud(op, "desk_page_save", {"p_id": None, "p_kind": "form", "p_title": "Contact us", "p_config": {}})
        self.visitor("public_send", {"p_token": form["token"], "p_kind": "lead", "p_payload": {"name": "Karim Nabil", "phone": "01112223334",
                                                                                              "fields": {"Full name": "Karim Nabil", "Message": "Need a quote"}}})
        self.pull(op)
        self.assertEqual(self.db.one("SELECT phone FROM parties WHERE name='Karim Nabil'")["phone"], "01112223334")
        self.assertIn("Need a quote", self.db.one("SELECT body FROM crm_activities WHERE party='Karim Nabil'")["body"])
        self.assertTrue(self.db.one("SELECT 1 FROM crm_followups WHERE party='Karim Nabil' AND done=0"))
        # a guest chat: the team answers, the guest reads it
        guest = self.cloud(op, "desk_page_save", {"p_id": None, "p_kind": "guest", "p_title": "Chat with us", "p_config": {}})
        st, g = self.visitor("public_send", {"p_token": guest["token"], "p_kind": "message", "p_payload": {"name": "Ali", "body": "Do you do home visits?"}})
        self.pull(op)
        st, dd = self.call(op, "GET", "/api/desk")
        th = dd["threads"][0]
        self.assertEqual((th["name"], th["unread"], dd["unread"]["messages"]), ("Ali", 1, 1))
        st, al = self.call(op, "GET", "/api/alerts")
        self.assertIn("desk", [r.get("kind") for r in al["rows"]])
        self.cloud(op, "desk_reply", {"p_page": th["page_id"], "p_ref": th["ref"], "p_kind": "reply", "p_body": "Yes, every day"})
        st, mine = self.visitor("public_manage", {"p_token": guest["token"], "p_ref": g["ref"], "p_manage": g["manage"]})
        self.assertEqual([m["body"] for m in mine["messages"]], ["Do you do home visits?", "Yes, every day"])
        # an offer signed online: accepted, signature kept with it
        st, o = self.call(op, "POST", "/api/offers", {"client": "Anna Bell", "project": "Home care plan", "status": "Sent",
                                                      "sections": [{"kind": "pricing", "heading": "Fees", "items": [{"desc": "Monthly nursing", "amount": 9000}]}]})
        self.assertEqual(st, 200, o)
        st, summ = self.call(op, "GET", f"/api/desk/offer/{o['id']}")
        self.assertEqual((summ["net"], summ["sections"][0]["items"][0]["desc"]), (9000, "Monthly nursing"))
        sign = self.cloud(op, "desk_page_save", {"p_id": None, "p_kind": "sign", "p_title": "Offer", "p_config": summ})
        png = "data:image/png;base64," + __import__("base64").b64encode(b"\x89PNG\r\n\x1a\n" + b"0" * 40).decode()
        self.visitor("public_send", {"p_token": sign["token"], "p_kind": "signature", "p_payload": {"offer_id": o["id"], "name": "Anna Bell", "image": png}})
        self.pull(op)
        self.assertEqual(self.db.one("SELECT status FROM offers WHERE id=?", (o["id"],))["status"], "Accepted")
        self.assertTrue(self.db.one("SELECT 1 FROM attachments WHERE owner='offer' AND owner_id=? AND title LIKE 'Signed by Anna Bell%'", (o["id"],)))
        # a switched-off page stops answering
        self.cloud(op, "desk_page_save", {"p_id": page["id"], "p_kind": "booking", "p_title": "Book a visit", "p_config": None, "p_active": False})
        self.assertEqual(self.visitor("public_page", {"p_token": tok})[0], 404)


class Automations(unittest.TestCase):
    """Rules that act by themselves: once per item, a follow-up for someone or a message in a team room."""

    def test_rules(self):
        from erp import automations, demo
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "secret123"}, first=True)
        demo.load(db, "boss")
        boss = db.one("SELECT * FROM users WHERE username='boss'")
        late = automations.preview(db, {"trigger": "receivable_overdue", "settings": {"days": 0}})
        self.assertTrue(late)
        with self.assertRaises(a.ValidationError):
            automations.save(db, {"name": "x", "trigger": "nope", "action": "followup"}, boss)
        with self.assertRaises(a.ValidationError):
            automations.save(db, {"name": "x", "trigger": "stock_low", "action": "chat", "params": {"room": "No such room"}}, boss)
        rid = automations.save(db, {"name": "Chase late payers", "trigger": "receivable_overdue", "settings": {"days": 0},
                                    "action": "followup", "params": {"assign": "boss", "title": "Collect"}}, boss)
        before = db.one("SELECT COUNT(*) n FROM crm_followups")["n"]
        done = automations.run(db, owner="boss")
        made = db.one("SELECT COUNT(*) n FROM crm_followups")["n"] - before
        self.assertEqual((done[0]["acted"], made), (len(late), len(late)))
        self.assertTrue(db.one("SELECT 1 FROM crm_followups WHERE title LIKE 'Collect: %' AND assigned_to='boss'"))
        self.assertEqual(automations.run(db, owner="boss"), [])            # once per item
        self.assertEqual(automations.run(db, owner="someone-else"), [])     # each person's app runs their own rules
        # a rule that posts in a team room
        room = db.one("SELECT name FROM chat_rooms WHERE kind<>'dm' ORDER BY id")["name"]
        automations.save(db, {"name": "Stalled deals", "trigger": "deal_stalled", "settings": {"days": 0}, "action": "chat",
                              "params": {"room": room, "message": "@boss have a look"}}, boss)
        automations.run(db, owner="boss")
        self.assertTrue(db.one("SELECT 1 FROM chat_messages WHERE body LIKE '%Stalled deals:%'") or
                        not automations.preview(db, {"trigger": "deal_stalled", "settings": {"days": 0}}))
        rows = automations.listing(db)
        self.assertEqual(rows[0]["fired"], len(late))
        cat = automations.catalog()
        self.assertIn("receivable_overdue", [x["key"] for x in cat["triggers"]])
        self.assertEqual(cat["example"]["action"], "followup")
        automations.delete(db, rid)
        self.assertFalse(db.one("SELECT 1 FROM automation_log WHERE automation_id=?", (rid,)))


class MyDay(unittest.TestCase):
    def test_start_end_report(self):
        from erp import crm, workdays
        db = Database(":memory:")
        auth.create_user(db, {"username": "nurse", "password": "secret123"}, first=True)
        u = db.one("SELECT * FROM users WHERE username='nurse'")
        self.assertIsNone(workdays.mine(db, u)["today"])
        with self.assertRaises(a.ValidationError):
            workdays.finish(db, u, {})                               # not started
        workdays.start(db, u, {"lat": 30.0444, "lng": 31.2357, "accuracy": 12})
        with self.assertRaises(a.ValidationError):
            workdays.start(db, u, {})                                # already started
        crm.add_activity(db, {"party": "Mona", "kind": "visit", "subject": "Wound dressing"}, "nurse")
        draft = workdays.mine(db, u)["draft"]
        self.assertIn("Wound dressing", draft)
        db.set_setting("workday_location", "required")
        with self.assertRaises(a.ValidationError):
            workdays.finish(db, u, {"report": "x"})                  # location required
        workdays.finish(db, u, {"report": draft, "lat": 30.05, "lng": 31.24})
        t = workdays.team(db)
        self.assertEqual((t["rows"][0]["username"], t["rows"][0]["start_lat"], t["rows"][0]["report"]), ("nurse", 30.0444, draft))
        self.assertIsNotNone(t["rows"][0]["hours"])


    def test_sessions_sites_and_requests(self):
        """Several check-ins a day add up; work sites restrict where people check in unless they may work anywhere;
        requests are approved into Time & Attendance and advances are paid into Employee advances."""
        from erp import hrrequests, workdays
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(db, {"username": "tech", "password": "secret123", "perms": {"team": "add"}})
        boss, tech = (db.one("SELECT * FROM users WHERE username=?", (u,)) for u in ("boss", "tech"))
        db.x("INSERT INTO employees(code, name, username, status) VALUES ('E1', 'Field Technician', 'tech', 'active')")
        emp = db.one("SELECT id FROM employees WHERE code='E1'")["id"]
        workdays.start(db, tech, {})
        workdays.finish(db, tech, {})
        workdays.start(db, tech, {})                                  # checking in again after checking out
        m = workdays.mine(db, tech)
        self.assertEqual((len(m["sessions"]), m["open"]), (2, True))
        hrrequests.save_policy(db, {"sites": [{"name": "Main office", "lat": 30.0444, "lng": 31.2357, "radius": 200}]})
        with self.assertRaises(a.ValidationError):
            workdays.finish(db, tech, {})                             # sites listed: location required
        with self.assertRaises(a.ValidationError):
            workdays.finish(db, tech, {"lat": 30.10, "lng": 31.30})   # several km away
        workdays.finish(db, tech, {"lat": 30.0450, "lng": 31.2360, "accuracy": 10})
        self.assertEqual(db.one("SELECT end_site FROM workdays ORDER BY id DESC LIMIT 1")["end_site"], "Main office")
        hrrequests.save_policy(db, {"anywhere": ["tech"]})
        workdays.start(db, tech, {})                                  # allowed anywhere
        # leave → attendance; rejecting removes it
        rid = hrrequests.submit(db, tech, {"kind": "annual", "date_from": "2030-03-03", "date_to": "2030-03-05", "reason": "Family"})
        with self.assertRaises(a.ValidationError):
            hrrequests.submit(db, tech, {"kind": "sick", "date_from": "2030-03-04"})   # overlaps
        with self.assertRaises(PermissionError):
            hrrequests.decide(db, tech, rid, True)                    # own request
        hrrequests.decide(db, boss, rid, True)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM attendance WHERE employee_id=? AND status='leave'", (emp,))["n"], 3)
        self.assertEqual(hrrequests.balance(db, "tech", "2030")["taken"], 3)
        hrrequests.decide(db, boss, rid, False, "Busy period")
        self.assertEqual(db.one("SELECT COUNT(*) n FROM attendance WHERE employee_id=? AND status='leave'", (emp,))["n"], 0)
        self.assertTrue(db.one("SELECT 1 FROM alert_feed WHERE username='tech'"))
        # late-arrival permission excuses lateness
        lid = hrrequests.submit(db, tech, {"kind": "late", "date_from": a.today(), "time_to": "11:00"})
        hrrequests.decide(db, boss, lid, True)
        self.assertEqual(hrrequests.excused(db, "tech", a.today()), {"late"})
        # salary advance
        aid = hrrequests.submit(db, tech, {"kind": "advance", "amount": 1500})
        with self.assertRaises(a.ValidationError):
            hrrequests.pay_advance(db, boss, aid, {})                 # not approved yet
        hrrequests.decide(db, boss, aid, True)
        eid = hrrequests.pay_advance(db, boss, aid, {"method": "Cash"})
        self.assertEqual(hrrequests.advance_balance(db, "Field Technician"), 1500)
        with self.assertRaises(a.ValidationError):
            hrrequests.decide(db, boss, aid, False)                   # paid: void the entry first
        self.assertTrue(eid)


class KnowledgeBaseAndApiKeys(unittest.TestCase):
    def test_kb(self):
        from erp import kb
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(db, {"username": "sam", "password": "secret123", "perms": {}})
        boss, sam = (auth.public_user(db.one("SELECT * FROM users WHERE username=?", (u,))) for u in ("boss", "sam"))
        kid = kb.save(db, {"title": "Return policy", "body": "# Returns\n- 14 days with receipt", "category": "Sales"}, sam)
        self.assertEqual([r["title"] for r in kb.listing(db, "receipt")["rows"]], ["Return policy"])
        self.assertEqual(kb.listing(db)["categories"], ["Sales"])
        other = kb.save(db, {"title": "Pricing", "body": "x"}, boss)
        with self.assertRaises(PermissionError):
            kb.save(db, {"title": "Hacked"}, sam, other)                 # someone else's article
        kb.save(db, {"title": "Return policy (2026)", "body": "30 days"}, boss, kid)   # administrators may
        self.assertEqual(kb.get(db, kid)["updated_by"], "boss")

    def test_new_permissions(self):
        """Workplace tools (knowledge base, my day) are open to people set up before they existed; super admins can
        narrow them. API keys got their own screen, starting from each person's "System settings" level."""
        from erp.server import handle
        app = type("App", (), {"static": False})()
        app.db = db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(db, {"username": "old", "password": "secret123", "perms": {"crm": "view", "settings": "edit", "settings.company": "edit", "settings.system": "none"}})
        auth.create_user(db, {"username": "sys", "password": "secret123", "perms": {"settings": "edit", "settings.system": "edit"}})
        auth.create_user(db, {"username": "reader", "password": "secret123", "perms": {"team": "view"}})
        auth.create_user(db, {"username": "office", "password": "secret123", "perms": {"team": "add", "team.myday": "none"}})
        pu = {u: auth.public_user(db.one("SELECT * FROM users WHERE username=?", (u,))) for u in ("old", "sys", "reader", "office")}
        self.assertEqual((auth.level_of(pu["old"], "team.kb"), auth.level_of(pu["old"], "team.myday")), ("add", "add"))
        self.assertEqual((auth.level_of(pu["old"], "settings.api"), auth.level_of(pu["sys"], "settings.api")), ("none", "edit"))
        tok = {u: auth.login(db, u, "secret123")[0] for u in ("boss", "old", "sys", "reader", "office")}

        def call(who, method, url, body=None):
            h = {"x-requested-with": "BeyondERP", "x-session": tok[who], "content-type": "application/json"}
            return self.handle_status(handle(app, method, url, h, json.dumps(body).encode() if body is not None else b""))
        self.handle_status = lambda r: r.status
        self.assertEqual(call("old", "GET", "/api/api-keys"), 403)       # had no System settings: still no API keys
        self.assertEqual(call("sys", "GET", "/api/api-keys"), 200)
        self.assertEqual(call("old", "POST", "/api/kb", {"title": "How we quote"}), 200)
        self.assertEqual(call("reader", "GET", "/api/kb"), 200)
        self.assertEqual(call("reader", "POST", "/api/kb", {"title": "x"}), 403)   # read only
        kid = db.one("SELECT id FROM kb_pages")["id"]
        self.assertEqual(call("office", "PUT", f"/api/kb/{kid}", {"title": "mine now"}), 403)   # someone else's article
        db.x("UPDATE users SET perms=? WHERE username='office'", (json.dumps({"team": "edit", "team.myday": "none"}),))
        self.assertEqual(call("office", "PUT", f"/api/kb/{kid}", {"title": "Tidied"}), 200)    # full rights on the KB
        self.assertEqual(call("office", "POST", "/api/my-day/start", {}), 403)               # my day switched off for them
        self.assertEqual(call("reader", "POST", "/api/my-day/start", {}), 403)               # view only
        self.assertEqual(call("old", "POST", "/api/my-day/start", {}), 200)

    def test_removed_picture_stays_removed(self):
        """Removing a background picture keeps a "none here" mark, so the company's picture does not show in its place."""
        from erp.looks import clean_look
        png = "data:image/png;base64,iVBORw0KGgo="
        out = clean_look({"images": {"page": {"off": 1}, "header": {"src": png, "op": 0.3}, "sidebar": {}, "bogus": {"off": 1}}})
        self.assertEqual(out["images"], {"page": {"off": 1}, "header": {"src": png, "op": 0.3, "fit": "cover"}})

    def test_admins_follow_the_grid(self):
        """Only super administrators skip permissions. An administrator whose grid was never saved keeps full access;
        once a super administrator saves it, they follow it (approvals and posting included)."""
        from erp import approvals
        db = Database(":memory:")
        auth.create_user(db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(db, {"username": "old_admin", "password": "secret123", "is_admin": True})
        auth.create_user(db, {"username": "new_admin", "password": "secret123", "is_admin": True,
                              "perms": {"crm": "edit", "accounting": "view", "_grid": 1}, "needs_approval": True})
        pu = {u: auth.public_user(db.one("SELECT * FROM users WHERE username=?", (u,))) for u in ("boss", "old_admin", "new_admin")}
        self.assertEqual([auth.level_of(pu["old_admin"], k) for k in ("accounting.journal", "settings.system")], ["edit", "edit"])
        self.assertEqual([auth.level_of(pu["new_admin"], k) for k in ("crm.deals", "accounting.journal", "settings.system")], ["edit", "view", "none"])
        self.assertTrue(pu["new_admin"]["perms"].get("_grid"))
        self.assertEqual(auth.level_of(pu["boss"], "settings.system"), "edit")
        self.assertTrue(approvals.needs_approval(db, pu["new_admin"]))            # "can post" applies to administrators too
        self.assertFalse(approvals.needs_approval(db, pu["boss"]))
        row = {"created_by": "sam", "module": "accounting.journal", "group_key": "site"}
        can = lambda u: (lambda m, lvl="view": auth.can(u, m, lvl))
        self.assertFalse(approvals.can_approve(pu["new_admin"], can(pu["new_admin"]), row))   # no approval permission
        self.assertTrue(approvals.can_approve(pu["old_admin"], can(pu["old_admin"]), row))    # full access, any group
        self.assertTrue(approvals.can_approve(pu["boss"], can(pu["boss"]), row))

    def test_api_keys(self):
        from erp.server import ERPServer
        tmp = tempfile.TemporaryDirectory()
        db = Database(Path(tmp.name) / "k.sqlite3")
        srv = ERPServer(("127.0.0.1", 0), db, Path(tmp.name) / "bk")
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        base = f"http://127.0.0.1:{srv.server_address[1]}"
        try:
            auth.create_user(db, {"username": "boss", "password": "secret123"}, first=True)
            auth.create_user(db, {"username": "viewer", "password": "secret123", "perms": {"crm": "view"}})
            ro = auth.create_api_key(db, db.one("SELECT id FROM users WHERE username='boss'")["id"], "reports", True)[1]
            rw = auth.create_api_key(db, db.one("SELECT id FROM users WHERE username='boss'")["id"], "writer", False)[1]
            lim = auth.create_api_key(db, db.one("SELECT id FROM users WHERE username='viewer'")["id"], "crm only", True)[1]

            def call(method, path, key, body=None):
                req = urllib.request.Request(base + path, data=json.dumps(body).encode() if body is not None else None, method=method,
                                             headers={"Authorization": "Bearer " + key, "Content-Type": "application/json"})
                try:
                    with urllib.request.urlopen(req) as r:
                        return r.status, json.loads(r.read())
                except urllib.error.HTTPError as e:
                    return e.code, json.loads(e.read())
            self.assertEqual(call("GET", "/api/kb", ro)[0], 200)
            self.assertEqual(call("POST", "/api/kb", ro, {"title": "x"})[0], 403)          # read-only
            self.assertEqual(call("POST", "/api/kb", rw, {"title": "From a script"})[0], 200)
            self.assertEqual(call("GET", "/api/users", rw)[0], 403)                         # never people / keys / passwords
            self.assertEqual(call("POST", "/api/api-keys", rw, {"name": "more"})[0], 403)
            self.assertEqual(call("GET", "/api/crm/contacts", lim)[0], 200)                # the person's own permissions
            self.assertEqual(call("GET", "/api/journal", lim)[0], 403)
            self.assertEqual(call("GET", "/api/kb", "berp_nope")[0], 401)
            docs = call("GET", "/api/api-docs?q=contacts", lim)[1]
            self.assertTrue(any("/api/crm/contacts" in r["path"] for r in docs["routes"]))
            auth.revoke_api_key(db, db.one("SELECT id FROM api_keys WHERE name='reports'")["id"])
            self.assertEqual(call("GET", "/api/kb", ro)[0], 401)
            self.assertTrue(db.one("SELECT last_used FROM api_keys WHERE name='writer'")["last_used"])
        finally:
            srv.shutdown()
            srv.server_close()
            db.close()
            tmp.cleanup()


class ImprovementsPack(unittest.TestCase):
    """Duplicate warnings, bank statement files, cash forecast, year-end close, project-level access, two-step sign-in
    and alerts that are checked without the app open."""

    def setUp(self):
        from erp import demo
        from erp.server import handle
        self.handle = handle
        self.app = type("App", (), {"static": False})()
        self.app.db = self.db = Database(":memory:")
        demo.load(self.db, "boss")
        auth.create_user(self.db, {"username": "boss", "password": "secret123"}, first=True)
        auth.create_user(self.db, {"username": "eng", "password": "secret123", "perms": {"purchasing": "add", "business_units": "view", "accounting": "view"}})
        self.tok = {u: auth.login(self.db, u, "secret123")[0] for u in ("boss", "eng")}

    def call(self, who, method, url, body=None):
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok[who], "content-type": "application/json"}
        r = self.handle(self.app, method, url, h, json.dumps(body).encode() if body is not None else b"")
        return r.status, json.loads(r.body)

    def test_duplicate_entries_are_flagged_until_confirmed(self):
        exp = {"description": "Office internet", "amount": 1234.5, "date": "2026-03-10", "vendor": "Internet provider", "account": "6045"}
        st, r = self.call("boss", "POST", "/api/docs/expense", exp)
        self.assertEqual(st, 200, r)
        n = self.db.one("SELECT COUNT(*) n FROM journal_entries")["n"]
        st, r = self.call("boss", "POST", "/api/docs/expense", {**exp, "date": "2026-03-12"})
        self.assertEqual((st, r["code"]), (409, "duplicate"))
        self.assertEqual(r["duplicates"][0]["amount"], 1234.5)
        self.assertEqual(self.db.one("SELECT COUNT(*) n FROM journal_entries")["n"], n)      # nothing saved
        st, r = self.call("boss", "POST", "/api/docs/expense", {**exp, "date": "2026-03-12", "allow_duplicate": 1})
        self.assertEqual(st, 200)
        self.assertNotIn("allow_duplicate", a.get_entry(self.db, r["id"])["meta"].get("form", {}))
        self.assertEqual(self.call("boss", "POST", "/api/docs/expense", {**exp, "date": "2026-03-20"})[0], 200)   # a week later: next month's bill
        jv = {"date": "2026-03-10", "memo": "Bank charges", "lines": [{"account": "6015", "debit": 50}, {"account": a.money_account(self.db, "bank"), "credit": 50}]}
        self.assertEqual(self.call("boss", "POST", "/api/journal", jv)[0], 200)
        self.assertEqual(self.call("boss", "POST", "/api/journal", jv)[0], 409)
        self.db.set_setting("duplicate_check", "off")
        self.assertEqual(self.call("boss", "POST", "/api/journal", jv)[0], 200)

    def test_bank_statement_files_mt940_camt_ofx_qif(self):
        from erp import bankrec
        mt = (":20:STMT1\n:25:EG123456\n:28C:1/1\n:60F:C260301EGP10000,00\n"
              ":61:2603020302C2500,00NTRFNONREF//B1\n:86:Client transfer ALFA\n"
              ":61:2603050305D120,50NCHGCHQ1001\n:86:Bank charges\nmonthly\n:62F:C260305EGP12379,50\n-")
        r = bankrec.parse(mt, "march.sta")
        self.assertEqual((r["format"][:5], r["opening"], r["closing"], [x["amount"] for x in r["rows"]]), ("MT940", 10000.0, 12379.5, [2500.0, -120.5]))
        self.assertEqual((r["rows"][1]["ref"], r["rows"][1]["description"], r["rows"][1]["balance"]), ("CHQ1001", "Bank charges monthly", 12379.5))
        camt = """<?xml version="1.0" encoding="UTF-8"?><Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.02"><BkToCstmrStmt><Stmt>
          <Bal><Tp><CdOrPrtry><Cd>OPBD</Cd></CdOrPrtry></Tp><Amt Ccy="EGP">500.00</Amt><CdtDbtInd>CRDT</CdtDbtInd></Bal>
          <Bal><Tp><CdOrPrtry><Cd>CLBD</Cd></CdOrPrtry></Tp><Amt Ccy="EGP">400.00</Amt><CdtDbtInd>CRDT</CdtDbtInd></Bal>
          <Ntry><Amt Ccy="EGP">100.00</Amt><CdtDbtInd>DBIT</CdtDbtInd><Sts>BOOK</Sts><BookgDt><Dt>2026-03-04</Dt></BookgDt><AcctSvcrRef>X9</AcctSvcrRef>
            <NtryDtls><TxDtls><RmtInf><Ustrd>Rent March</Ustrd></RmtInf><RltdPties><Cdtr><Nm>Landlord</Nm></Cdtr></RltdPties></TxDtls></NtryDtls></Ntry>
          </Stmt></BkToCstmrStmt></Document>"""
        r = bankrec.parse(camt, "s.xml")
        self.assertEqual((r["opening"], r["closing"], r["rows"][0]["amount"], r["rows"][0]["ref"], r["rows"][0]["description"]),
                         (500.0, 400.0, -100.0, "X9", "Rent March · Landlord"))
        with self.assertRaises(a.ValidationError):
            bankrec.parse('<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><BkToCstmrStmt/>', "bad.xml")
        ofx = ("OFXHEADER:100\nDATA:OFXSGML\n<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><BANKTRANLIST>"
               "<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260310120000<TRNAMT>-45.00<FITID>F1<NAME>FUEL STATION<MEMO>Card 1234"
               "<STMTTRN><TRNTYPE>CREDIT<DTPOSTED>20260311<TRNAMT>900<FITID>F2<NAME>CLIENT</BANKTRANLIST>"
               "<LEDGERBAL><BALAMT>1855.00<DTASOF>20260311</LEDGERBAL></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>")
        r = bankrec.parse(ofx, "x.ofx")
        self.assertEqual((r["format"], r["opening"], r["closing"], r["rows"][0]["description"]), ("OFX", 1000.0, 1855.0, "FUEL STATION · Card 1234"))
        r = bankrec.parse("!Type:Bank\nD03/15/2026\nT-75.00\nPElectricity\nN554\n^\nD16/03/2026\nT20\nPRefund\n^\n", "a.qif")
        self.assertEqual([(x["date"], x["amount"], x["ref"]) for x in r["rows"]], [("2026-03-15", -75.0, "554"), ("2026-03-16", 20.0, "")])
        csv_r = bankrec.parse("Date,Description,Amount\n01/03/2026,Fee,-5\n")     # the existing CSV reader still works
        self.assertEqual(csv_r["rows"][0]["amount"], -5.0)

    def test_cash_forecast_uses_only_dated_items(self):
        from erp import cashplan
        today = a.today()
        due = (datetime.date.fromisoformat(today) + datetime.timedelta(days=10)).isoformat()
        st, b = self.call("boss", "POST", "/api/bills", {"name": "Software subscription", "amount": 1000, "next_due": due, "cycle": "monthly"})
        self.assertEqual(st, 200, b)
        st, r = self.call("boss", "GET", "/api/cash-forecast?days=90")
        self.assertEqual(st, 200, r)
        bills = [x for x in r["items"] if x["source"] == "bills" and x["label"].startswith("Software subscription")]
        self.assertEqual([x["date"] for x in bills][:1], [due])
        self.assertEqual(len(bills), 3)                          # every monthly cycle inside 90 days
        self.assertEqual(r["cash_now"], cashplan.cash_now(self.db))
        w = {x["days"]: x for x in r["windows"]}
        in30 = sum(x["amount"] for x in r["items"] if x["amount"] > 0 and x["date"] <= w[30]["until"])
        self.assertAlmostEqual(w[30]["in"], in30, places=2)
        self.assertAlmostEqual(r["weeks"][-1]["balance"], round(r["opening"] + sum(x["amount"] for x in r["items"]), 2), places=2)
        self.assertAlmostEqual(r["opening"], round(r["cash_now"] + r["overdue"]["in"] - r["overdue"]["out"], 2), places=2)
        off = self.call("boss", "GET", "/api/cash-forecast?days=90&overdue=0")[1]
        self.assertEqual(off["opening"], off["cash_now"])
        self.assertEqual({x["key"] for x in r["undated"]}, {"receivables", "payables", "taxes", "purchase_orders"})
        self.assertEqual(self.call("eng", "GET", "/api/cash-forecast")[0], 403)

    def test_year_end_close_locks_and_moves_profit_to_retained_earnings(self):
        from erp import yearend
        db = self.db
        db.set_setting("lock_date", "")
        a.record_expense(db, {"description": "Old year cost", "amount": 500, "date": "2025-06-01", "account": "6045"})
        last_end = "2025-12-31"
        st, p = self.call("boss", "GET", f"/api/year-end/{last_end}")
        self.assertEqual((st, p["start"]), (200, "2025-01-01"), p)
        self.assertTrue({"trial_balance", "approvals", "cash"} <= {c["key"] for c in p["checks"]})
        bs_before = a.balance_sheet(db, "2026-06-30")
        st, r = self.call("boss", "POST", f"/api/year-end/{last_end}/close", {"notes": "Checked with the auditor"})
        self.assertEqual(st, 200, r)
        self.assertEqual(db.setting("lock_date"), last_end)
        with self.assertRaises(a.ValidationError):                               # the closed year is locked
            a.record_expense(db, {"description": "Late", "amount": 10, "date": "2025-12-30", "account": "6045"})
        bs = a.balance_sheet(db, "2026-06-30")
        eq = {x["subtype"]: x["amount"] for x in bs["equity"] if x["subtype"] in ("earnings", "earnings_closed")}
        prior = a.income_statement(db, None, last_end)["net_profit"]
        self.assertAlmostEqual(eq["earnings_closed"], prior, places=2)
        self.assertAlmostEqual(eq["earnings_closed"] + eq["earnings"], [x["amount"] for x in bs_before["equity"] if x["subtype"] == "earnings"][0], places=2)
        self.assertEqual((bs["total_equity"], bs["balanced"]), (bs_before["total_equity"], True))
        self.assertEqual(a.income_statement(db, "2025-01-01", last_end)["net_profit"], p["net_profit"])   # the year's own figures stay
        self.assertEqual(self.call("boss", "POST", f"/api/year-end/{last_end}/close", {})[0], 400)          # already closed
        self.assertEqual(self.call("boss", "POST", "/api/year-end/2025-12-30/close", {})[0], 400)          # not a year end
        self.assertEqual(self.call("boss", "PUT", "/api/settings", {"fiscal_year_start": "07-01"})[0], 400)
        self.assertEqual(self.call("eng", "POST", f"/api/year-end/{last_end}/reopen", {})[0], 403)
        st, r = self.call("boss", "POST", f"/api/year-end/{last_end}/reopen", {})
        self.assertEqual((st, db.setting("lock_date")), (200, ""))
        self.assertEqual(len(yearend.closed(db)), 0)
        self.assertNotIn("earnings_closed", {x["subtype"] for x in a.balance_sheet(db, "2026-06-30")["equity"]})

    def test_project_level_access(self):
        db = self.db
        codes = [r["code"] for r in db.q("SELECT code FROM business_units ORDER BY code")]
        mine, other = codes[0], codes[1]
        uid = auth.create_user(db, {"username": "site", "password": "secret123", "projects": [mine],
                                    "perms": {"business_units": "view", "accounting": "view", "purchasing": "add", "reports": "view", "assets": "view", "dashboard": "view"}})
        self.tok["site"] = auth.login(db, "site", "secret123")[0]
        self.assertEqual(auth.public_user(db.one("SELECT * FROM users WHERE id=?", (uid,)))["perms"]["_bus"], [mine])
        st, r = self.call("site", "GET", "/api/bus")
        self.assertEqual((st, [b["code"] for b in r["rows"]], "shared_overhead" in r), (200, [mine], False))
        st, j = self.call("site", "GET", "/api/journal?limit=100")
        mine_id = db.one("SELECT id FROM business_units WHERE code=?", (mine,))["id"]
        self.assertEqual(j["total"], db.one("SELECT COUNT(*) n FROM journal_entries WHERE bu_id=?", (mine_id,))["n"])
        self.assertTrue(j["rows"] and all(e["bu_code"] == mine for e in j["rows"]))
        theirs = db.one("SELECT e.id FROM journal_entries e JOIN business_units b ON b.id=e.bu_id WHERE b.code=? LIMIT 1", (other,))["id"]
        self.assertEqual(self.call("site", "GET", f"/api/journal/{theirs}")[0], 403)
        st, rep = self.call("site", "GET", "/api/reports/income_statement")
        self.assertEqual(st, 200, rep)
        self.assertEqual(self.call("site", "GET", "/api/reports/aging")[0], 403)
        self.assertEqual(self.call("site", "GET", f"/api/reports/income_statement?bu_id={db.one('SELECT id FROM business_units WHERE code=?', (other,))['id']}")[0], 403)
        self.assertEqual(self.call("site", "GET", "/api/cash-accounts")[0], 403)
        exp = {"description": "Site water", "amount": 120, "date": a.today(), "account": "6045"}
        self.assertEqual(self.call("site", "POST", "/api/docs/expense", exp)[0], 400)                      # no project = company overhead
        self.assertEqual(self.call("site", "POST", "/api/docs/expense", {**exp, "bu_code": other})[0], 403)
        self.assertEqual(self.call("site", "POST", "/api/docs/expense", {**exp, "bu_code": mine})[0], 200)
        self.tok["site"] = auth.login(db, "site", "secret123")[0]
        boot = self.call("site", "GET", "/api/bootstrap")[1]
        self.assertEqual([b["code"] for b in boot["bus"]], [mine])
        # only a super administrator changes it; the boss is never limited
        self.assertEqual(self.call("boss", "PUT", f"/api/users/{uid}", {"projects": []})[0], 200)
        self.assertNotIn("_bus", auth.public_user(db.one("SELECT * FROM users WHERE id=?", (uid,)))["perms"])
        self.assertEqual(len(self.call("site", "GET", "/api/bus")[1]["rows"]), len(codes))

    def test_two_step_sign_in(self):
        db = self.db
        st, r = self.call("eng", "GET", "/api/me/two-step")
        self.assertEqual((st, r["on"]), (200, False))
        self.assertTrue(r["uri"].startswith("otpauth://totp/") and r["secret"] in r["uri"])
        secret = r["secret"]
        self.assertEqual(self.call("eng", "POST", "/api/me/two-step", {"secret": secret, "code": "000000"})[0], 400)
        code = auth.totp(secret, auth._now_step())
        st, on = self.call("eng", "POST", "/api/me/two-step", {"secret": secret, "code": code})
        self.assertEqual((st, len(on["recovery_codes"])), (200, 10))
        auth._last_step.clear()
        with self.assertRaises(auth.SecondStepNeeded):                     # the password alone is no longer enough
            auth.login(db, "eng", "secret123")
        r = self.handle(self.app, "POST", "/api/login", {"x-requested-with": "BeyondERP", "content-type": "application/json"},
                        json.dumps({"username": "eng", "password": "secret123"}).encode())
        self.assertEqual((r.status, json.loads(r.body).get("two_step")), (400, True))
        with self.assertRaises(auth.SecondStepNeeded):
            auth.login(db, "eng", "secret123", "123456" if code != "123456" else "654321")
        tok, u = auth.login(db, "eng", "secret123", auth.totp(secret, auth._now_step()))
        self.assertTrue(u["two_step"])
        with self.assertRaises(auth.SecondStepNeeded):                     # the same code works once
            auth.login(db, "eng", "secret123", auth.totp(secret, auth._now_step()))
        rc = on["recovery_codes"][0]
        self.assertTrue(auth.login(db, "eng", "secret123", rc.lower())[0])  # a recovery code, once
        with self.assertRaises(auth.SecondStepNeeded):
            auth.login(db, "eng", "secret123", rc)
        # the company requires it for everyone: someone without it may only set it up
        auth.create_user(db, {"username": "clerk", "password": "secret123", "perms": {"purchasing": "view"}})
        self.tok["clerk"] = auth.login(db, "clerk", "secret123")[0]
        self.assertEqual(self.call("eng", "PUT", "/api/settings", {"two_step_required": "all"})[0], 403)   # super administrators decide
        self.assertEqual(self.call("boss", "PUT", "/api/settings", {"two_step_required": "all"})[0], 403)  # (boss has none either)
        db.set_setting("two_step_required", "all")
        st, r = self.call("clerk", "GET", "/api/purchase-orders")
        self.assertEqual((st, r.get("two_step_setup")), (403, True))
        self.assertEqual(self.call("clerk", "GET", "/api/me/two-step")[0], 200)
        self.assertEqual(self.call("eng", "POST", "/api/me/two-step/off", {"password": "secret123"})[0], 400)   # required: stays on
        db.set_setting("two_step_required", "")
        self.assertEqual(self.call("eng", "POST", "/api/me/two-step/off", {"password": "wrong"})[0], 400)
        uid = db.one("SELECT id FROM users WHERE username='eng'")["id"]
        self.assertEqual(self.call("clerk", "POST", f"/api/users/{uid}/two-step/off", {})[0], 403)
        self.assertEqual(self.call("boss", "POST", f"/api/users/{uid}/two-step/off", {})[0], 200)
        self.assertTrue(auth.login(db, "eng", "secret123")[0])
        self.assertNotIn("totp", json.dumps(interchange.export_extras(db, False, False)["people"]))      # never in backups

    def test_alerts_reach_people_with_their_app_closed(self):
        from erp import automations
        db = self.db
        # eng sets an alert about bills (he may see them); the boss's app finds it true and acts for him
        auth.update_user(db, db.one("SELECT id FROM users WHERE username='eng'")["id"], {"perms": {"purchasing": "add", "team": "add"}}, 1)
        self.tok["eng"] = auth.login(db, "eng", "secret123")[0]
        st, r = self.call("eng", "POST", "/api/my-alerts", {"name": "Bills soon", "trigger": "bill_due", "settings": {"days": 3, "at": "08:00"}})
        self.assertEqual(st, 200, r)
        today = datetime.date.fromisoformat(a.today())
        due = (today + datetime.timedelta(days=6)).isoformat()
        self.call("boss", "POST", "/api/bills", {"name": "Internet", "amount": 900, "next_due": due})
        # nothing yet today; the forecast says when it will come true (3 days before the due date, at 08:00)
        st, f = self.call("boss", "GET", f"/api/automations/forecast?local={today.isoformat()}T07:00")
        self.assertEqual(st, 200, f)
        mine = [x for x in f["items"] if x["title"] == "Bills soon"]
        self.assertEqual((mine[0]["day"], mine[0]["time"], mine[0]["to"]), ((today + datetime.timedelta(days=3)).isoformat(), "08:00", ["eng"]))
        self.assertEqual(len(mine), 1)                                 # once, on the first day it is true
        st, r = self.call("boss", "POST", "/api/automations/run", {"all": 1, "local": f"{today.isoformat()}T09:00"})
        self.assertTrue(r.get("_nochange"))                            # nothing true today: the shared log records nothing
        # the bill moves closer: the boss's change makes it true now — eng's bell gets it without his app
        sid = db.one("SELECT id FROM cost_subs WHERE name='Internet'")["id"]
        self.assertEqual(self.call("boss", "PUT", f"/api/bills/{sid}", {"name": "Internet", "amount": 900, "next_due": (today + datetime.timedelta(days=1)).isoformat()})[0], 200)
        st, r = self.call("boss", "POST", "/api/automations/run", {"all": 1, "local": f"{today.isoformat()}T09:00"})
        self.assertEqual(st, 200, r)
        self.assertTrue(any(x.get("acted") for x in r["done"]))
        self.assertEqual(db.one("SELECT COUNT(*) n FROM alert_feed WHERE username='eng'")["n"], 1)
        self.assertEqual([h["to"] for h in r["_push"]], [["eng"]])
        # someone who may not see bills does not run (or read) eng's alert
        auth.create_user(db, {"username": "clerk2", "password": "secret123", "perms": {"crm": "view"}})
        clerk = auth.public_user(db.one("SELECT * FROM users WHERE username='clerk2'"))
        self.assertEqual(automations.forecast(db, clerk)["covers"], [])

    def test_ai_batch_review_direct_mode_and_undo(self):
        from erp import ai
        db = self.db
        ai.save_policy(db, {"is_super": True}, {"enabled": True, "write": "batch"})
        self.assertIn("erp_write_batch", [t["name"] for t in ai.tools(ai.policy(db))])
        n0 = db.one("SELECT COUNT(*) n FROM journal_entries")["n"]
        lines = [{"method": "POST", "path": "/api/docs/expense", "summary": f"line {i}",
                  "body": {"description": f"Count line {i}", "amount": 100 + i, "date": "2026-10-05", "account": "6045", "allow_duplicate": 1}} for i in range(3)]
        st, rv = self.call("boss", "POST", "/api/ai/preview", {"changes": lines})
        self.assertEqual((st, rv["ok"], len(rv["items"]), rv["entries_total"]), (200, True, 3, 3), rv)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries")["n"], n0)          # nothing saved by the check
        h = {"x-requested-with": "BeyondERP", "x-session": self.tok["boss"], "content-type": "application/json", "x-ai-call": "1", "x-ai-review": rv["review"]}
        send = lambda i, path="/api/docs/expense?_ai=chg001": self.handle(self.app, "POST", path, {**h, "x-ai-batch": str(i)}, json.dumps(lines[i]["body"]).encode())  # noqa: E731
        self.assertEqual(send(1).status, 403)                                                # in the reviewed order only
        for i in range(3):
            self.assertEqual(send(i).status, 200)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries")["n"], n0 + 3)
        self.assertEqual(self.call("eng", "POST", "/api/ai-undo/chg001", {})[0], 403)         # only who asked (or a super admin)
        st, r = self.call("boss", "POST", "/api/ai-undo/chg001", {})
        self.assertEqual(st, 200, r)
        self.assertEqual(db.one("SELECT COUNT(*) n FROM journal_entries")["n"], n0)          # every record back as it was
        self.assertEqual(self.call("boss", "POST", "/api/ai-undo/chg001", {})[0], 400)       # once
        # direct mode: saved without approval, still recorded; an undo is refused once someone changed the record
        ai.save_policy(db, {"is_super": True}, {"write": "auto"})
        st, rv = self.call("boss", "POST", "/api/ai/preview", lines[0])
        h["x-ai-review"] = rv["review"]
        r = self.handle(self.app, "POST", "/api/docs/expense?_ai=chg002", {**h, "x-ai-batch": "0"}, json.dumps(lines[0]["body"]).encode())
        eid = json.loads(r.body)["id"]
        self.call("boss", "PUT", f"/api/journal/{eid}/tags", {"tags": ["checked"]})
        a.post_entry(db, date="2026-10-06", kind="Journal Voucher", memo="later work", lines=[{"account": "6045", "debit": 5}, {"account": a.money_account(db, "bank"), "credit": 5}])
        db.x("UPDATE journal_entries SET memo='edited by someone' WHERE id=?", (eid,))
        st, r = self.call("boss", "POST", "/api/ai-undo/chg002", {})
        self.assertEqual(st, 400)
        self.assertIn("were changed after", r["error"])
        self.assertTrue(db.one("SELECT 1 FROM journal_entries WHERE id=?", (eid,)))
