"""Cost allocation rules, ownership by stage, new rate capture, review flags and their import / export."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from erp import accounting as a  # noqa: E402
from erp import allocation, interchange, ratecap  # noqa: E402
from erp.db import Database  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAVE_XL = True
except ImportError:
    HAVE_XL = False


def books():
    db = Database(":memory:")
    for code, status, start, area in (("P-1", "In Progress", "2026-01-01", 100), ("P-2", "In Progress", "2026-01-01", 300),
                                      ("P-3", "Tendering", "", 0), ("P-4", "In Progress", "2026-06-01", 0)):
        db.x("INSERT INTO business_units(code,name,basis,status,start_date,area,overhead_pct) VALUES (?,?,?,?,?,?,0)",
             (code, code, "Fixed Price", status, start, area))
    return db


def alloc(db, **kw):
    return {db.one("SELECT code FROM business_units WHERE id=?", (b,))["code"]: v
            for b, v in allocation.allocate(db, **kw)["by_bu"].items()}


def total(slot):
    return slot["billable"] + slot["non_billable"]


class AllocationRules(unittest.TestCase):
    def test_no_rules_falls_back_to_project_overhead_pct(self):
        db = books()
        db.x("UPDATE business_units SET overhead_pct=0.5 WHERE code IN ('P-1','P-2')")
        a.record_expense(db, {"description": "Rent", "amount": 1000, "account": "6010", "date": "2026-02-01"})
        r = alloc(db)
        self.assertEqual((total(r["P-1"]), total(r["P-2"])), (50000, 50000))

    def test_fixed_even_single_none_and_tags(self):
        db = books()
        allocation.save(db, {"code": "EVEN", "method": "even", "is_default": 1})
        allocation.save(db, {"code": "FIX", "method": "fixed", "split": {"P-1": 60, "P-2": 40}, "billable": "Yes"})
        allocation.save(db, {"code": "ONE", "method": "single", "target": "P-2"})
        allocation.save(db, {"code": "OFF", "method": "none"})
        a.record_expense(db, {"description": "Rent", "amount": 900, "account": "6010", "date": "2026-02-01"})          # default
        a.record_expense(db, {"description": "Car", "amount": 1000, "account": "6010", "date": "2026-02-01", "alloc_rule": "FIX"})
        a.record_expense(db, {"description": "Ad", "amount": 300, "account": "6010", "date": "2026-02-01", "alloc_rule": "ONE"})
        a.record_expense(db, {"description": "Gift", "amount": 50, "account": "6010", "date": "2026-02-01", "alloc_rule": "OFF"})
        a.record_expense(db, {"description": "Fee", "amount": 70, "account": "6010", "date": "2026-02-01", "alloc_rule": "company"})
        r = alloc(db)
        # even: P-3 is tendering and P-4 has not started in February → P-1 and P-2 only
        self.assertEqual(r["P-1"], {"billable": 60000, "non_billable": 45000})
        self.assertEqual(r["P-2"], {"billable": 40000, "non_billable": 45000 + 30000})
        self.assertNotIn("P-3", r)
        out = allocation.allocate(db)
        self.assertEqual(out["company"], 12000)                   # the "none" rule and the company tag stay with the company
        with self.assertRaises(a.ValidationError):
            a.record_expense(db, {"description": "x", "amount": 1, "account": "6010", "alloc_rule": "NOPE"})

    def test_parts_add_up_exactly(self):
        db = books()
        allocation.save(db, {"code": "THIRDS", "method": "fixed", "split": {"P-1": 1 / 3, "P-2": 1 / 3, "P-4": 1 / 3}, "is_default": 1})
        a.record_expense(db, {"description": "Rent", "amount": 100, "account": "6010", "date": "2026-07-01"})
        r = allocation.allocate(db)
        self.assertEqual(sum(total(s) for s in r["by_bu"].values()), 10000)
        self.assertEqual(r["company"], 0)

    def test_driver_account_scope_dates_and_priority(self):
        db = books()
        allocation.save(db, {"code": "AREA", "method": "driver", "driver": "area", "is_default": 1})
        allocation.save(db, {"code": "UTIL", "method": "single", "target": "P-1", "accounts": "6040-6049", "priority": 10})
        allocation.save(db, {"code": "OLD", "method": "single", "target": "P-2", "accounts": "6040-6049", "priority": 5,
                             "date_to": "2026-01-31"})
        util = db.one("SELECT code FROM accounts WHERE code BETWEEN '6040' AND '6049' AND type='Expense' ORDER BY code")["code"]
        a.record_expense(db, {"description": "Rent", "amount": 400, "account": "6010", "date": "2026-02-01"})
        a.record_expense(db, {"description": "Power", "amount": 100, "account": util, "date": "2026-02-01"})
        r = alloc(db)
        self.assertEqual(total(r["P-1"]), 10000 + 10000)           # area 100 of 400 → 100, and the utility bill
        self.assertEqual(total(r["P-2"]), 30000)
        self.assertTrue(allocation.account_matches("6010, 602*, 6030-6060", "6025"))
        self.assertFalse(allocation.account_matches("6010, 602*", "6030"))

    def test_validation_and_rename_follow_tags(self):
        db = books()
        with self.assertRaises(a.ValidationError):
            allocation.save(db, {"code": "BAD", "method": "fixed", "split": {"P-1": 70, "P-2": 40}})
        with self.assertRaises(a.ValidationError):
            allocation.save(db, {"code": "BAD", "method": "fixed", "split": {"NOPE": 50}})
        rid = allocation.save(db, {"code": "A", "method": "single", "target": "P-1"})
        a.record_expense(db, {"description": "x", "amount": 10, "account": "6010", "date": "2026-02-01", "alloc_rule": "A"})
        allocation.save(db, {"code": "B", "method": "single", "target": "P-1"}, rid)
        self.assertEqual(db.one("SELECT json_extract(meta,'$.alloc_rule') r FROM journal_entries WHERE memo='x'")["r"], "B")
        with self.assertRaises(a.ValidationError):
            allocation.delete(db, rid)

    def test_billable_shared_cost_raises_cost_plus_billing(self):
        db = books()
        db.x("UPDATE business_units SET basis='Cost-Plus %', fee_pct=0.2 WHERE code='P-1'")
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Works", "qty": 1, "rate": 1000, "date": "2026-02-01"})
        before = next(b for b in a.bu_summaries(db) if b["code"] == "P-1")["expected_billing"]
        allocation.save(db, {"code": "SUP", "method": "single", "target": "P-1", "billable": "Yes", "is_default": 1})
        a.record_expense(db, {"description": "Supervision", "amount": 500, "account": "6010", "date": "2026-02-01"})
        b = next(b for b in a.bu_summaries(db) if b["code"] == "P-1")
        self.assertEqual(b["allocated_billable"], 500)
        self.assertAlmostEqual(b["expected_billing"] - before, 600, places=2)


class StageOwnership(unittest.TestCase):
    def test_profit_split_by_stage_cost(self):
        db = books()
        db.x("INSERT INTO partners(name,default_share) VALUES ('A',0.5),('B',0.5)")
        pid = {r["name"]: r["id"] for r in db.q("SELECT id,name FROM partners")}
        p1 = db.one("SELECT id FROM business_units WHERE code='P-1'")["id"]
        db.x("INSERT INTO partner_shares(partner_id,bu_id,share) VALUES (?,?,0.5),(?,?,0.5)", (pid["A"], p1, pid["B"], p1))
        db.x("INSERT INTO partner_stage_shares(partner_id,bu_id,stage,share) VALUES (?,?,'Finishing',1),(?,?,'Finishing',0)",
             (pid["A"], p1, pid["B"], p1))
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Shell", "qty": 1, "rate": 3000, "date": "2026-02-01"})
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Paint", "qty": 1, "rate": 1000, "date": "2026-02-01",
                                   "stage": "Finishing"})
        own = a.ownership(db, None, None)
        self.assertAlmostEqual(own(pid["A"], p1, 0.5), 0.75 * 0.5 + 0.25 * 1)
        self.assertAlmostEqual(own(pid["B"], p1, 0.5), 0.75 * 0.5)
        self.assertAlmostEqual(own(pid["A"], p1, 0.5) + own(pid["B"], p1, 0.5), 1)


class RateCaptureAndFlags(unittest.TestCase):
    def test_new_rates_and_ignore(self):
        db = books()
        db.x("INSERT INTO rates(category,description,unit,cost,markup,price) VALUES ('Paint','Wall paint','m2',8000,0.25,10000)")
        x = a.save_extract(db, {"bu_code": "P-1", "number": "INV-1", "date": "2026-03-01",
                                "lines": [{"description": "Wall  paint.", "unit": "M2", "qty": 10, "price": 120},
                                          {"description": "Ceiling board", "unit": "m2", "qty": 5, "price": 300}]})
        a.post_extract(db, x)
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Board sheets", "qty": 4, "unit": "pcs", "rate": 90,
                                   "date": "2026-03-02"})
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Board sheets", "qty": 2, "unit": "pcs", "rate": 95,
                                   "date": "2026-03-05"})
        rows = {r["description"]: r for r in ratecap.candidates(db)}
        self.assertEqual(set(rows), {"Ceiling board", "Board sheets"})     # wall paint is already listed
        self.assertEqual(rows["Ceiling board"]["price"]["last"], 300)
        self.assertEqual((rows["Board sheets"]["cost"]["last"], rows["Board sheets"]["cost"]["count"],
                          rows["Board sheets"]["cost"]["min"]), (95, 2, 90))
        ratecap.set_ignored(db, rows["Ceiling board"]["key"])
        self.assertEqual({r["description"] for r in ratecap.candidates(db)}, {"Board sheets"})
        self.assertEqual(len(ratecap.candidates(db, include_ignored=True)), 2)

    def test_review_flags(self):
        db = books()
        db.x("INSERT INTO partners(name,default_share) VALUES ('A',1)")
        a.record_partner_tx(db, {"partner": "A", "type": "Capital", "amount": 1000, "date": "2026-01-02"})
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Works", "qty": 1, "rate": 900, "date": "2026-02-01"})
        x = a.save_extract(db, {"bu_code": "P-1", "number": "INV-1", "date": "2026-03-01",
                                "lines": [{"description": "Works", "unit": "m2", "qty": 1, "price": 500}]})
        a.post_extract(db, x)
        a.record_partner_tx(db, {"partner": "A", "type": "Distribution", "amount": 100, "date": "2026-03-02"})
        flags = " | ".join(f["flag"] for f in a.review_flags(db))
        self.assertIn("P-1: costs are higher than income", flags)
        self.assertIn("A has taken", flags)
        self.assertIn("1 prices used in the books", flags)
        links = {f["flag"].split(":")[0].split(" has")[0]: f["href"] for f in a.review_flags(db)}
        p1 = db.one("SELECT id FROM business_units WHERE code='P-1'")["id"]
        self.assertEqual(links["P-1"], f"#/bu/{p1}")
        self.assertEqual(links["A"], "#/statement?party=A")
        self.assertTrue(any(h == "#/prices?ratecap=1" for h in links.values()))

    def test_funder_share_of_client_payments(self):
        """A funding partner repaid an agreed % of every client payment: only what is due to him now is held back from
        distributions (not all his money), whichever clients paid."""
        from erp import partner_pack
        db = books()
        db.set_setting("funding_partner", "F")
        for k, v in (("funding_repay_pct", "0.5"), ("reserve_pct", "0"), ("buffer_months", "0"), ("repay_funder_first", "0")):
            db.set_setting(k, v)
        db.x("INSERT INTO partners(name,default_share) VALUES ('F',0.5),('G',0.5)")
        a.record_journal(db, {"date": "2026-01-02", "memo": "lent", "lines": [{"account": "1020", "debit": 5000}, {"account": "3015", "credit": 5000}]})
        x = a.save_extract(db, {"bu_code": "P-1", "number": "INV-9", "date": "2026-01-05", "lines": [{"description": "Works", "qty": 1, "price": 3000}]})
        a.post_extract(db, x)
        a.record_collection(db, {"bu_code": "P-1", "amount": 2000, "date": "2026-01-10", "received_into": "1020", "credit_account": "1100"})
        due = partner_pack.compute(db)["dist"]["funder_due"]
        self.assertEqual((due["client_payments"], due["share"], due["repaid"], due["due"]), (2000.0, 1000.0, 0.0, 1000.0))
        before = partner_pack.compute(db)["dist"]["free"]
        a.record_journal(db, {"date": "2026-01-11", "memo": "his half", "lines": [{"account": "3015", "debit": 1000}, {"account": "1020", "credit": 1000}]})
        d = partner_pack.compute(db)["dist"]
        self.assertEqual((d["funder_due"]["due"], d["free"]), (0.0, before))      # paying him his half: cash and the hold both drop 1,000
        self.assertNotIn("repay the funding partner first", d["verdict"][0])
        db.set_setting("repay_funder_first", "1")                                 # ticked: all his money first, whatever the share
        d = partner_pack.compute(db)["dist"]
        self.assertEqual(d["max_safe"], 0.0)
        self.assertIn("repay the funding partner first", d["verdict"][0])

    def test_profit_not_in_cash_is_not_distributable_unless_allowed(self):
        """Profit billed but not collected is shown as not in cash; only cash is offered for distribution until the
        company allows more."""
        from erp import partner_pack
        db = books()
        for k, v in (("reserve_pct", "0"), ("buffer_months", "0")):
            db.set_setting(k, v)
        db.x("INSERT INTO partners(name,default_share) VALUES ('F',0.5),('G',0.5)")
        x = a.save_extract(db, {"bu_code": "P-1", "number": "INV-9", "date": "2026-01-05", "lines": [{"description": "Works", "qty": 1, "price": 3000}]})
        a.post_extract(db, x)
        a.record_collection(db, {"bu_code": "P-1", "amount": 1000, "date": "2026-01-10", "received_into": "1020", "credit_account": "1100"})
        d = partner_pack.compute(db)["dist"]
        self.assertEqual((d["cash_safe"], d["max_safe"], d["not_in_cash"], d["receivable"]), (1000.0, 1000.0, 2000.0, 2000.0))
        db.set_setting("distribute_noncash", "1")
        d = partner_pack.compute(db)["dist"]
        self.assertEqual((d["cash_safe"], d["max_safe"]), (1000.0, 3000.0))
        self.assertIn("1,000 from cash now, 2,000 not in cash yet", d["verdict"][0])

    def test_profit_by_category(self):
        """Each category's revenue less its costs; the rest stays outside, so the categories add up to the net profit."""
        import json
        from erp import categories
        db = books()
        x = a.save_extract(db, {"bu_code": "P-1", "number": "INV-9", "date": "2026-01-05", "lines": [
            {"description": "Rent", "qty": 1, "price": 5000, "account": "4045"},
            {"description": "Transport", "qty": 1, "price": 800, "account": "4047"}]})
        a.post_extract(db, x)
        a.record_expense(db, {"date": "2026-01-06", "account": "6075", "description": "Travel", "amount": 300, "paid": 300, "paid_from": "1010"})
        a.record_expense(db, {"date": "2026-01-06", "account": "6090", "description": "Truck", "amount": 500, "paid": 500, "paid_from": "1010"})
        db.set_setting("profit_categories", json.dumps([{"name": "Rent", "accounts": "4045"}, {"name": "Transport", "accounts": "4047, 6090"}]))
        d = categories.compute(db)
        got = {o["name"]: (o["revenue"], o["cost"], o["profit"]) for o in d["categories"]}
        self.assertEqual(got, {"Rent": (5000.0, 0.0, 5000.0), "Transport": (800.0, 500.0, 300.0)})
        self.assertEqual((d["rest"]["cost"], d["total"]["profit"]), (300.0, a.income_statement(db)["net_profit"]))
        with self.assertRaises(a.ValidationError):
            categories.clean([{"name": "Bad", "accounts": "40x5"}])
        self.assertEqual(categories.clean([{"name": "", "accounts": ""}, {"name": "A", "accounts": "60*, 6010-6040"}]),
                         [{"name": "A", "name_ar": "", "accounts": "60*, 6010-6040"}])

    def test_cash_position_follows_the_funding_facility(self):
        """Facility mode: the funding partner's account is the bank, so the cash position shows the room left in his
        facility instead of the company's own cash against a safety floor."""
        from erp import settlement
        db = books()
        db.x("INSERT INTO partners(name,default_share) VALUES ('F',1)")
        for k, v in (("funding_partner", "F"), ("funding_limit", "1000")):
            db.set_setting(k, v)
        self.assertIsNone(a.cash_position(db)["facility"])                 # capital mode: the floor as before
        db.set_setting("funding_mode", "facility"); settlement.apply_funding_mode(db, "facility")
        a.record_expense(db, {"date": "2026-01-05", "account": "6075", "description": "Travel", "amount": 300, "paid": 300,
                              "paid_from": "3015"})          # his account
        f = a.cash_position(db)["facility"]
        self.assertEqual((f["limit"], f["fronted"], f["room_after"]), (1000.0, 300.0, 700.0))
        from erp import automations, cashplan
        fc = cashplan.build(db)
        self.assertEqual((fc["basis"], fc["cash_now"]), ("facility", 700.0))      # the forecast starts from his facility
        hits = list(automations._t_cash_low(db, {"below": 800}))                 # the "cash below" rule watches it too
        self.assertEqual(len(hits), 1)
        self.assertIn("Funding room left is 700.00", hits[0]["text"])
        self.assertEqual(list(automations._t_cash_low(db, {"below": 500})), [])

    def test_cash_below_zero_flag(self):
        """Money paid out before any came in is flagged with the day and the lowest balance, even if cash is positive now."""
        db = books()
        a.record_expense(db, {"date": "2026-01-05", "account": "6075", "description": "Travel", "amount": 300, "paid": 300, "paid_from": "1010"})
        a.record_journal(db, {"date": "2026-01-20", "memo": "funds", "lines": [{"account": "1010", "debit": 1000}, {"account": "3010", "credit": 1000}]})
        flags = [f for f in a.review_flags(db) if f["href"] == "#/ledger?account=1010"]
        self.assertEqual(len(flags), 1)
        self.assertIn("below zero from 2026-01-05 (lowest -300.00 on 2026-01-05)", flags[0]["flag"])
        self.assertEqual([f for f in a.review_flags(db, date_to="2026-01-04") if f["href"].startswith("#/ledger")], [])


class DealPayables(unittest.TestCase):
    def test_two_deals_same_supplier_paid_in_stages(self):
        db = books()
        a.record_project_cost(db, {"bu_code": "P-1", "vendor": "Sup", "description": "Works A", "qty": 1, "rate": 1000,
                                   "paid": 300, "deal_ref": "SUP-001", "date": "2026-02-01"})
        a.record_project_cost(db, {"bu_code": "P-2", "vendor": "Sup", "description": "Works B", "qty": 1, "rate": 500,
                                   "paid": 0, "deal_ref": "SUP-002", "date": "2026-02-02"})
        a.record_payment(db, {"party": "Sup", "amount": 200, "deal_ref": "SUP-001", "date": "2026-02-10"})
        a.record_payment(db, {"party": "Sup", "amount": 100, "date": "2026-02-11"})          # no deal named
        d = a.deals(db)
        g = {x["ref"]: x for x in d["deals"]}
        self.assertEqual((g["SUP-001"]["booked"], g["SUP-001"]["paid"], g["SUP-001"]["owed"], g["SUP-001"]["state"]),
                         (1000, 500, 500, "stages"))
        self.assertEqual((g["SUP-002"]["owed"], g["SUP-002"]["state"], g["SUP-002"]["bu_code"]), (500, "not_paid", "P-2"))
        self.assertEqual((d["total_owed"], d["untagged"]), (900, -100))
        data = interchange.read_json(interchange.to_json(db).encode())
        dst = Database(":memory:")
        interchange.import_data(dst, data, mode="replace", user="t")
        self.assertEqual(a.deals(dst)["deals"], d["deals"])


class ImportExport(unittest.TestCase):
    def test_round_trip_of_rules_stages_mixes_and_tags(self):
        from erp import forecast
        src = books()
        src.x("INSERT INTO partners(name,default_share) VALUES ('A',0.5),('B',0.5)")
        allocation.save(src, {"code": "FIX", "name": "Shared site", "method": "fixed", "split": {"P-1": 60, "P-2": 40},
                              "billable": "Yes", "accounts": "60*", "is_default": 1, "notes": "n"})
        allocation.save(src, {"code": "DRV", "method": "driver", "driver": "revenue", "statuses": ["In Progress"],
                              "exclude": ["P-4"], "date_from": "2026-01-01"})
        p1 = src.one("SELECT id FROM business_units WHERE code='P-1'")["id"]
        for pid, share in ((1, 1), (2, 0)):
            src.x("INSERT INTO partner_stage_shares(partner_id,bu_id,stage,share) VALUES (?,?,'Finishing',?)", (pid, p1, share))
        forecast.save(src, {"mixes": {"Tiles": {"local": 0.5, "imported": 0.3, "labour": 0.2}}})
        a.record_expense(src, {"description": "Car", "amount": 100, "account": "6010", "date": "2026-02-01", "alloc_rule": "DRV"})
        a.record_expense(src, {"description": "Gift", "amount": 50, "account": "6010", "date": "2026-02-01", "alloc_rule": "company"})
        a.record_project_cost(src, {"bu_code": "P-1", "vendor": "Sup", "description": "Paint", "qty": 1, "rate": 10, "date": "2026-02-01",
                                    "stage": "Finishing"})
        a.record_journal(src, {"memo": "Accrual", "date": "2026-02-02", "alloc_rule": "DRV",
                               "lines": [{"account": "6010", "debit": 30}, {"account": "2010", "credit": 30}]})
        data = interchange.read_json(interchange.to_json(src).encode())
        for k in ("Allocation Rules", "Partner Stage Shares", "Forecast Mixes"):
            self.assertTrue(data.get(k), k)
        dst = Database(":memory:")
        interchange.import_data(dst, data, mode="replace", user="t")
        self.assertEqual(allocation.export_rows(src), allocation.export_rows(dst))
        q = "SELECT p.name, b.code, s.stage, s.share FROM partner_stage_shares s JOIN partners p ON p.id=s.partner_id " \
            "JOIN business_units b ON b.id=s.bu_id ORDER BY 1"
        self.assertEqual(src.q(q), dst.q(q))
        self.assertEqual(forecast.assumptions(dst)["mixes"], forecast.assumptions(src)["mixes"])
        meta = "SELECT memo, json_extract(meta,'$.alloc_rule') r, json_extract(meta,'$.no_alloc') n, " \
               "json_extract(meta,'$.stage') s FROM journal_entries ORDER BY memo"
        self.assertEqual(src.q(meta), dst.q(meta))
        self.assertEqual(allocation.allocate(src)["by_bu"].__len__(), allocation.allocate(dst)["by_bu"].__len__())

    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_workbook_sheets_have_the_new_columns(self):
        from erp import workbook
        db = books()
        allocation.save(db, {"code": "FIX", "method": "fixed", "split": {"P-1": 100}, "is_default": 1})
        path = str(Path(tempfile.mkdtemp()) / "wb.xlsx")
        with open(path, "wb") as f:
            f.write(workbook.build_workbook(db))
        data = interchange.read_xlsx(path)
        self.assertEqual([r["code"] for r in data["Allocation Rules"]], ["FIX"])
        cols = {s["name"]: [c["key"] for c in s["cols"]] for s in interchange.SHEETS}
        self.assertIn("alloc_rule", cols["Operating Expenses"])
        self.assertIn("stage", cols["Direct Costs"])
        self.assertTrue({"alloc_rule", "stage"} <= set(cols["Journal Vouchers"]))


class LegacyWorkbookSheets(unittest.TestCase):
    @unittest.skipUnless(HAVE_XL, "openpyxl not installed")
    def test_rules_tags_stage_estimator_and_forecast_mixes(self):
        import openpyxl
        from erp import forecast, legacy
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Project Register"
        ws.append(["Project ID", "Client Name", "Project Type", "Location", "Total Area (Sqm)", "Start Date", "Status",
                   "Total Contract Value", "Billing Basis / Fee Type", "Supervision Fee %"])
        ws.append(["PRJ-1", "Villa A", "Villa", "X", 100, "2026-01-01", "In Progress", 0, "Cost-Plus %", 0.25])
        ws.append(["PRJ-2", "Villa B", "Villa", "X", 100, "2026-01-01", "In Progress", 0, "Cost-Plus %", 0.25])
        ca = wb.create_sheet("Cost Allocation")
        ca.append(["COST ALLOCATION RULES"])
        ca.append(["Allocation rule", "Billable to client?", "Pooled amount", "PRJ-1", "PRJ-2", "Total %"])
        ca.append(["Even split", "No", 0, 0.5, 0.5, 1])
        ca.append(["Site supervision", "Yes", 0, 1, 0, 1])
        ca.append(["Broken", "No", 0, 0.5, 0.2, 0.7])
        ca.append(["Unallocated (company only)", "No", 0, 0, 0, 0])
        gj = wb.create_sheet("General Journal")
        gj.append(["GENERAL JOURNAL"])
        gj.append(["JE #", "Date", "Source / Type", "Project", "Stage", "Account", "Description", "Vendor / Party",
                   "Deal Ref", "Qty", "Unit", "Rate", "Debit", "Credit", "Note", "Alloc Rule"])
        for r in (["JE-1", "2026-01-05", "Opex", None, None, "6010", "Supervisor", None, None, None, None, None, 300, None, None, "Site supervision"],
                  ["JE-1", "2026-01-05", "Opex", None, None, "1020", "Supervisor", None, None, None, None, None, None, 300, None, "Site supervision"],
                  ["JE-2", "2026-01-06", "Project cost", "PRJ-1", "Finishing", "5020", "Paint", "Sub", None, None, None, None, 100, None, None, None],
                  ["JE-2", "2026-01-06", "Project cost", "PRJ-1", "Finishing", "1020", "Paint", "Sub", None, None, None, None, None, 100, None, None],
                  ["JE-3", "2026-01-07", "Opex", None, None, "6010", "Rent", None, None, None, None, None, 200, None, None, None],
                  ["JE-3", "2026-01-07", "Opex", None, None, "1020", "Rent", None, None, None, None, None, None, 200, None, None]):
            gj.append(r)
        qe = wb.create_sheet("Quick Estimator")
        qe.append(["QUICK ESTIMATOR"])
        qe.append([])
        qe.append([])
        qe.append(["Client Name:", "Client X"])
        qe.append(["Total Area (Sqm):", 120])
        qe.append([])
        qe.append(["Category", "Subcategory / Specification", "Unit", "Quantity", "Unit Rate (Client)", "Unit Cost (Historical)",
                   "Total Line Price", "Line Cost"])
        qe.append(["Paint", "Wall paint two coats", "m2", 200, 120, 90, 24000, 18000])
        qe.append([None, None, None, None, 0, 0, 0, 0])
        qe.append(["Total Project Price:", 24000])
        pf = wb.create_sheet("2027 Price Forecast")
        pf.append(["PRICE FORECAST"])
        pf.append(["Driver", "Q1", "Q2"])
        pf.append(["Domestic CPI", 0.5, 0.5])                       # the sheet's own rates must not be copied
        pf.append([])
        pf.append(["Cost category", "Local mat %", "Imported %", "Labour %", "Energy/Fuel %", "Σ check"])
        pf.append(["Paint", 0.25, 0.45, 0.05, 0.25, 1])
        pf.append(["Transport", 0.15, 0, 0.35, 0.5, 1])
        path = str(Path(tempfile.mkdtemp()) / "legacy.xlsx")
        wb.save(path)
        data = legacy.read_legacy_workbook(path)
        notes = " ".join(data["_notes"])
        self.assertEqual([(r["code"], r["method"], r["billable"], r["is_default"]) for r in data["Allocation Rules"]],
                         [("EVEN-SPLIT", "fixed", "No", "Yes"), ("SITE-SUPERVISION", "fixed", "Yes", "No"),
                          ("UNALLOCATED-COMPANY-ONLY", "none", "No", "No")])
        self.assertIn("Broken (70.00%)", notes)
        db = Database(":memory:")
        interchange.import_data(db, data, mode="replace", user="t")
        tags = {r["ref"]: (r["rule"], r["stage"]) for r in db.q(
            "SELECT ref, json_extract(meta,'$.alloc_rule') rule, json_extract(meta,'$.stage') stage FROM journal_entries "
            "WHERE ref LIKE 'JE-%'")}
        self.assertEqual(tags, {"JE-1": ("SITE-SUPERVISION", None), "JE-2": (None, "Finishing"), "JE-3": (None, None)})
        r = alloc(db)
        self.assertEqual(r["PRJ-1"], {"billable": 30000, "non_billable": 10000})
        self.assertEqual(r["PRJ-2"], {"billable": 0, "non_billable": 10000})
        boq = db.one("SELECT b.title, l.description, l.qty, l.price, l.cost FROM boqs b JOIN boq_lines l ON l.boq_id=b.id "
                     "WHERE b.title LIKE 'Quick estimate%'")
        self.assertEqual(dict(boq), {"title": "Quick estimate — Client X", "description": "Wall paint two coats", "qty": 200,
                                     "price": 12000, "cost": 9000})
        fc = forecast.assumptions(db)
        self.assertEqual(fc["mixes"]["Paint"], {"local": 0.25, "imported": 0.45, "labour": 0.05, "energy": 0.25})
        self.assertFalse(any((d or {}).get("rates") for d in fc["drivers"].values()))


if __name__ == "__main__":
    unittest.main()
